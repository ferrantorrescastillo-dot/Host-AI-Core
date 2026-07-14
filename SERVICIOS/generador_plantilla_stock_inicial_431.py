from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List
import json


@dataclass
class ResultadoPlantillaStock:
    total_articulos: int
    ruta_destino: str
    estado: str
    mensaje: str


class GeneradorPlantillaStockInicial431:
    """
    Genera una plantilla Excel para rellenar stock inicial real.

    Fuente:
    DATOS/db/articulos.json

    Salida:
    DATOS/stock_inicial_4_3_1.xlsx
    """

    COLUMNAS = [
        "Codigo",
        "Articulo",
        "Proveedor",
        "Familia",
        "Unidad",
        "Stock actual",
        "Stock mínimo",
        "Ubicación",
        "Observaciones stock",
    ]

    UNIDADES_SUGERIDAS = ["kg", "g", "l", "ml", "ud", "caja", "bolsa", "bote", "paquete"]

    UBICACIONES_SUGERIDAS = [
        "Cámara",
        "Congelador",
        "Almacén seco",
        "Barra",
        "Cocina",
        "Producción",
    ]

    def __init__(self, ruta_articulos: str = "DATOS/db/articulos.json") -> None:
        self.ruta_articulos = Path(ruta_articulos)

    def generar_excel(self, ruta_destino: str = "DATOS/stock_inicial_4_3_1.xlsx") -> ResultadoPlantillaStock:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.worksheet.datavalidation import DataValidation
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para generar la plantilla Excel.") from exc

        articulos = self._leer_articulos()
        ruta = Path(ruta_destino)
        ruta.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Stock inicial"

        ws.append(self.COLUMNAS)

        for articulo in articulos:
            ws.append([
                articulo.get("codigo"),
                articulo.get("nombre"),
                articulo.get("proveedor"),
                articulo.get("familia"),
                "",
                "",
                "",
                "",
                "",
            ])

        header_fill = PatternFill("solid", fgColor="166534")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        widths = {
            "A": 14,
            "B": 48,
            "C": 24,
            "D": 24,
            "E": 14,
            "F": 16,
            "G": 16,
            "H": 22,
            "I": 36,
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Hoja de ayudas
        ayudas = wb.create_sheet("Ayudas")
        ayudas.append(["Unidades sugeridas"])
        for unidad in self.UNIDADES_SUGERIDAS:
            ayudas.append([unidad])

        ayudas["C1"] = "Ubicaciones sugeridas"
        for i, ubicacion in enumerate(self.UBICACIONES_SUGERIDAS, start=2):
            ayudas[f"C{i}"] = ubicacion

        ayudas.column_dimensions["A"].width = 22
        ayudas.column_dimensions["C"].width = 28

        # Validaciones
        if len(articulos) > 0:
            dv_unidad = DataValidation(type="list", formula1='"kg,g,l,ml,ud,caja,bolsa,bote,paquete"', allow_blank=True)
            dv_ubicacion = DataValidation(type="list", formula1='"Cámara,Congelador,Almacén seco,Barra,Cocina,Producción"', allow_blank=True)

            ws.add_data_validation(dv_unidad)
            ws.add_data_validation(dv_ubicacion)

            dv_unidad.add(f"E2:E{len(articulos)+1}")
            dv_ubicacion.add(f"H2:H{len(articulos)+1}")

        resumen = wb.create_sheet("Resumen")
        resumen.append(["HOST AI 4.3.1 - Plantilla Stock Inicial", ""])
        resumen.append(["Artículos incluidos", len(articulos)])
        resumen.append(["Estado", "Plantilla generada"])
        resumen.append(["Instrucción", "Rellenar Unidad, Stock actual, Stock mínimo y Ubicación."])

        for cell in resumen[1]:
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
            cell.font = Font(color="FFFFFF", bold=True)

        resumen.column_dimensions["A"].width = 34
        resumen.column_dimensions["B"].width = 60

        wb.save(ruta)

        return ResultadoPlantillaStock(
            total_articulos=len(articulos),
            ruta_destino=str(ruta),
            estado="ok" if articulos else "sin_articulos",
            mensaje="Plantilla generada correctamente." if articulos else "No hay artículos para generar plantilla.",
        )

    def _leer_articulos(self) -> List[Dict[str, Any]]:
        if not self.ruta_articulos.exists():
            return []

        try:
            contenido = json.loads(self.ruta_articulos.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return []

        return contenido if isinstance(contenido, list) else []


__all__ = ["GeneradorPlantillaStockInicial431", "ResultadoPlantillaStock"]
