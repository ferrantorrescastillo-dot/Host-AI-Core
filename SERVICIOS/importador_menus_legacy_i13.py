from __future__ import annotations

from dataclasses import dataclass, asdict, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable
import re
import unicodedata

from openpyxl import load_workbook


def _norm(valor: Any) -> str:
    texto = str(valor or "").strip().lower()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-z0-9]+", " ", texto).strip()


def _numero(valor: Any) -> float | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("€", "").replace("%", "").replace(" ", "")
    if texto.count(",") == 1 and texto.count(".") == 0:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except (TypeError, ValueError):
        return None


def _codigo_menu(nombre: str, usados: set[str]) -> str:
    base = re.sub(r"[^A-Z0-9]+", "-", _norm(nombre).upper()).strip("-")[:26] or "MENU"
    codigo = f"MENU-{base}"
    contador = 2
    while codigo in usados:
        codigo = f"MENU-{base[:21]}-{contador}"
        contador += 1
    usados.add(codigo)
    return codigo


SECCIONES = {
    "coctel": "Cóctel",
    "cocktail": "Cóctel",
    "aperitivo": "Cóctel",
    "aperitivos": "Cóctel",
    "entrante": "Entrante",
    "entrantes": "Entrante",
    "primero": "Primero",
    "primer plato": "Primero",
    "segundo": "Segundo",
    "segundo plato": "Segundo",
    "principal": "Principal",
    "plato principal": "Principal",
    "postre": "Postre",
    "postres": "Postre",
    "bodega": "Bodega",
    "bebida": "Bebidas",
    "bebidas": "Bebidas",
    "recena": "Recena",
    "infantil": "Infantil",
    "menu infantil": "Infantil",
    "pan": "Pan",
    "cafes": "Cafés",
    "cafe": "Cafés",
}

RESUMEN_CLAVES = {
    "precio venta": "precio_venta",
    "pvp": "precio_venta",
    "precio menu": "precio_venta",
    "coste total": "coste_total",
    "coste menu": "coste_total",
    "coste por persona": "coste_por_persona",
    "coste pax": "coste_por_persona",
    "food cost": "food_cost_pct",
    "beneficio": "beneficio",
    "margen": "margen_pct",
}

PALABRAS_NO_PLATO = {
    "menu", "menú", "plato", "cantidad", "precio", "coste", "total", "unidad",
    "racion", "raciones", "pax", "personas", "observaciones", "descripcion",
}


@dataclass
class PlatoMenuI13:
    seccion: str
    nombre_excel: str
    coste_declarado: float | None = None
    receta_id: str = ""
    receta_nombre: str = ""
    confianza: float = 0.0
    estado: str = "sin_resolver"  # exacto | probable | sin_resolver
    candidatos: list[dict[str, Any]] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


@dataclass
class MenuPreviaI13:
    hoja: str
    nombre: str
    menu_id: str
    platos: list[PlatoMenuI13]
    precio_venta: float | None = None
    coste_total_declarado: float | None = None
    coste_por_persona: float | None = None
    food_cost_pct: float | None = None
    beneficio: float | None = None
    margen_pct: float | None = None
    duplicado: bool = False
    duplicado_id: str = ""
    bloqueado: bool = False
    avisos: list[str] = field(default_factory=list)

    def a_dict(self) -> dict[str, Any]:
        return asdict(self)


class ImportadorMenusLegacyI13:
    """I1.3: detecta, vincula e importa hojas de menú legacy.

    La vista previa nunca escribe. La importación requiere confirmación explícita.
    """

    def __init__(self, umbral_probable: float = 0.72):
        self.umbral_probable = float(umbral_probable)

    @staticmethod
    def _recetas(recetas: Iterable[dict[str, Any]] | None) -> list[dict[str, Any]]:
        salida = []
        for receta in recetas or []:
            nombre = str(receta.get("nombre") or "").strip()
            if not nombre:
                continue
            salida.append({
                "receta_id": str(receta.get("receta_id") or receta.get("id") or ""),
                "nombre": nombre,
                "grupo": str(receta.get("grupo") or ""),
                "coste_total": _numero(receta.get("coste_total")),
                "coste_por_racion": _numero(receta.get("coste_por_racion")),
            })
        return salida

    @staticmethod
    def _menus_existentes(menus: Iterable[dict[str, Any]] | None) -> list[dict[str, Any]]:
        return [m for m in (menus or []) if str(m.get("nombre") or "").strip()]

    def _vincular_plato(self, nombre: str, seccion: str, coste: float | None, recetas: list[dict[str, Any]]) -> PlatoMenuI13:
        buscado = _norm(nombre)
        puntuados: list[tuple[float, dict[str, Any]]] = []
        for receta in recetas:
            candidato = _norm(receta["nombre"])
            score = 1.0 if buscado == candidato else SequenceMatcher(None, buscado, candidato).ratio()
            if buscado and (buscado in candidato or candidato in buscado):
                score = max(score, 0.88)
            if score >= 0.40:
                puntuados.append((score, receta))
        puntuados.sort(key=lambda x: x[0], reverse=True)
        candidatos = [
            {"receta_id": r["receta_id"], "nombre": r["nombre"], "confianza": round(score, 3)}
            for score, r in puntuados[:5]
        ]
        if puntuados and puntuados[0][0] >= 0.98:
            estado = "exacto"
        elif puntuados and puntuados[0][0] >= self.umbral_probable:
            estado = "probable"
        else:
            estado = "sin_resolver"
        mejor = puntuados[0][1] if puntuados and estado != "sin_resolver" else {}
        avisos = []
        if estado == "probable":
            avisos.append("Coincidencia probable: revisar antes de importar.")
        elif estado == "sin_resolver":
            avisos.append("Plato sin receta vinculada.")
        return PlatoMenuI13(
            seccion=seccion,
            nombre_excel=nombre,
            coste_declarado=coste,
            receta_id=str(mejor.get("receta_id") or ""),
            receta_nombre=str(mejor.get("nombre") or ""),
            confianza=round(puntuados[0][0], 3) if puntuados else 0.0,
            estado=estado,
            candidatos=candidatos,
            avisos=avisos,
        )

    @staticmethod
    def _nombre_menu(ws) -> str:
        for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
            textos = [str(v).strip() for v in fila if isinstance(v, str) and str(v).strip()]
            for texto in textos:
                normal = _norm(texto)
                if "menu" in normal and len(normal) > 4 and normal not in {"menu", "menu boda"}:
                    return texto
        return ws.title.strip()

    @staticmethod
    def _valor_derecha(valores: list[Any], indice: int) -> float | None:
        for valor in valores[indice + 1:]:
            numero = _numero(valor)
            if numero is not None:
                return numero
        return None

    def _detectar_hoja(self, ws, recetas: list[dict[str, Any]], usados: set[str], por_nombre: dict[str, dict[str, Any]]) -> dict[str, Any] | None:
        nombre_menu = self._nombre_menu(ws)
        seccion_actual = "Sin sección"
        platos: list[PlatoMenuI13] = []
        resumen: dict[str, float | None] = {
            "precio_venta": None,
            "coste_total": None,
            "coste_por_persona": None,
            "food_cost_pct": None,
            "beneficio": None,
            "margen_pct": None,
        }
        vistos: set[tuple[str, str]] = set()

        for fila in ws.iter_rows(values_only=True):
            valores = list(fila)
            textos = [(idx, str(v).strip()) for idx, v in enumerate(valores) if isinstance(v, str) and str(v).strip()]
            if not textos:
                continue

            # Primero secciones y campos resumen.
            seccion_detectada = None
            for _, texto in textos:
                normal = _norm(texto)
                if normal in SECCIONES:
                    seccion_detectada = SECCIONES[normal]
                    break
            if seccion_detectada:
                seccion_actual = seccion_detectada
                # Una fila de sección puede contener platos en columnas posteriores.

            fila_es_resumen = False
            for idx, texto in textos:
                normal = _norm(texto)
                clave = next((campo for etiqueta, campo in RESUMEN_CLAVES.items() if etiqueta in normal), None)
                if clave:
                    valor = self._valor_derecha(valores, idx)
                    if valor is not None:
                        resumen[clave] = valor
                    fila_es_resumen = True
            if fila_es_resumen:
                continue

            for idx, texto in textos:
                normal = _norm(texto)
                if not normal or normal in SECCIONES or normal in PALABRAS_NO_PLATO:
                    continue
                if any(etiqueta in normal for etiqueta in RESUMEN_CLAVES):
                    continue
                if "menu" in normal and len(textos) == 1:
                    continue
                if len(normal) < 3 or normal.isdigit():
                    continue
                # Descarta encabezados evidentes.
                if normal.startswith(("ficha tecnica", "materia prima", "articulo", "elaboracion")):
                    continue
                clave_visto = (seccion_actual, normal)
                if clave_visto in vistos:
                    continue
                vistos.add(clave_visto)
                coste = self._valor_derecha(valores, idx)
                platos.append(self._vincular_plato(texto, seccion_actual, coste, recetas))

        # Exige al menos dos platos para no confundir una hoja auxiliar con un menú.
        if len(platos) < 2:
            return None

        duplicado = por_nombre.get(_norm(nombre_menu))
        menu_id = str(duplicado.get("menu_id") or duplicado.get("id") or "") if duplicado else _codigo_menu(nombre_menu, usados)
        avisos: list[str] = []
        if not any(p.estado == "exacto" for p in platos):
            avisos.append("No se encontró ninguna receta vinculada de forma exacta.")
        if resumen["coste_total"] is not None:
            suma = sum(p.coste_declarado or 0 for p in platos)
            if suma and abs(suma - float(resumen["coste_total"])) > max(1.0, float(resumen["coste_total"]) * 0.08):
                avisos.append("El coste total declarado no coincide con la suma de platos.")

        return MenuPreviaI13(
            hoja=ws.title,
            nombre=nombre_menu,
            menu_id=menu_id,
            platos=platos,
            precio_venta=resumen["precio_venta"],
            coste_total_declarado=resumen["coste_total"],
            coste_por_persona=resumen["coste_por_persona"],
            food_cost_pct=resumen["food_cost_pct"],
            beneficio=resumen["beneficio"],
            margen_pct=resumen["margen_pct"],
            duplicado=bool(duplicado),
            duplicado_id=str(duplicado.get("menu_id") or duplicado.get("id") or "") if duplicado else "",
            bloqueado=not platos,
            avisos=avisos,
        ).a_dict()

    def preparar(self, ruta_excel: str | Path, hojas=None, recetas_existentes=None, menus_existentes=None) -> dict[str, Any]:
        ruta = Path(ruta_excel)
        if not ruta.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta}")
        wb = load_workbook(ruta, data_only=True, read_only=True)
        seleccion = list(hojas or [])
        if seleccion:
            inexistentes = [h for h in seleccion if h not in wb.sheetnames]
            if inexistentes:
                raise ValueError(f"No existen las hojas: {', '.join(inexistentes)}")
        else:
            seleccion = [
                h for h in wb.sheetnames
                if "menu" in _norm(h) and not _norm(h).startswith("mp ") and "materia prima" not in _norm(h)
            ]
        recetas = self._recetas(recetas_existentes)
        existentes = self._menus_existentes(menus_existentes)
        usados = {str(m.get("menu_id") or m.get("id") or "").upper() for m in existentes}
        por_nombre = {_norm(m.get("nombre")): m for m in existentes}
        menus = []
        for hoja in seleccion:
            detectado = self._detectar_hoja(wb[hoja], recetas, usados, por_nombre)
            if detectado:
                menus.append(detectado)
        wb.close()
        resumen = {
            "hojas_analizadas": len(seleccion),
            "menus_detectados": len(menus),
            "duplicados": sum(1 for m in menus if m["duplicado"]),
            "bloqueados": sum(1 for m in menus if m["bloqueado"]),
            "platos": sum(len(m["platos"]) for m in menus),
            "vinculos_exactos": sum(1 for m in menus for p in m["platos"] if p["estado"] == "exacto"),
            "vinculos_probables": sum(1 for m in menus for p in m["platos"] if p["estado"] == "probable"),
            "sin_resolver": sum(1 for m in menus for p in m["platos"] if p["estado"] == "sin_resolver"),
        }
        return {
            "ok": True,
            "ruta": str(ruta),
            "menus": menus,
            "resumen": resumen,
            "modo": "vista_previa_sin_escritura",
        }

    @staticmethod
    def _guardar_en_repositorio(repositorio, menu: dict[str, Any]) -> Any:
        # Motor específico de menús, si existe.
        for metodo in ("registrar_menu", "crear_menu", "guardar_menu"):
            funcion = getattr(repositorio, metodo, None)
            if callable(funcion):
                try:
                    return funcion(**menu)
                except TypeError:
                    return funcion(menu)
        # Base local Host AI: guardar colección completa.
        cargar = getattr(repositorio, "cargar", None)
        guardar = getattr(repositorio, "guardar", None)
        if callable(cargar) and callable(guardar):
            actuales = list(cargar("menus") or [])
            actuales.append(menu)
            guardar("menus", actuales)
            return menu
        raise TypeError("No se encontró un repositorio compatible para guardar menús.")

    def importar(self, previa: dict[str, Any], repositorio_menus, seleccion=None,
                 confirmar: bool = False, permitir_probables: bool = False,
                 permitir_sin_resolver: bool = False) -> dict[str, Any]:
        if not confirmar:
            raise ValueError("La importación requiere confirmación explícita.")
        seleccionados = set(seleccion or [m["menu_id"] for m in previa.get("menus", [])])
        importados = []
        omitidos = []
        errores = []
        for menu in previa.get("menus", []):
            if menu["menu_id"] not in seleccionados:
                continue
            motivos = []
            if menu.get("duplicado"):
                motivos.append("duplicado")
            if menu.get("bloqueado"):
                motivos.append("bloqueado")
            if not permitir_probables and any(p["estado"] == "probable" for p in menu["platos"]):
                motivos.append("vínculos probables")
            if not permitir_sin_resolver and any(p["estado"] == "sin_resolver" for p in menu["platos"]):
                motivos.append("platos sin receta vinculada")
            if motivos:
                omitidos.append({"menu_id": menu["menu_id"], "nombre": menu["nombre"], "motivo": ", ".join(motivos)})
                continue
            secciones: dict[str, list[dict[str, Any]]] = {}
            for plato in menu["platos"]:
                secciones.setdefault(plato["seccion"], []).append({
                    "receta_id": plato["receta_id"],
                    "nombre": plato["receta_nombre"] or plato["nombre_excel"],
                    "nombre_original": plato["nombre_excel"],
                    "coste_declarado": plato.get("coste_declarado"),
                })
            dato = {
                "menu_id": menu["menu_id"],
                "nombre": menu["nombre"],
                "secciones": secciones,
                "precio_venta": menu.get("precio_venta"),
                "coste_total_declarado": menu.get("coste_total_declarado"),
                "coste_por_persona": menu.get("coste_por_persona"),
                "food_cost_pct": menu.get("food_cost_pct"),
                "beneficio": menu.get("beneficio"),
                "margen_pct": menu.get("margen_pct"),
                "origen": {"tipo": "excel_legacy_i13", "hoja": menu["hoja"]},
                "activo": True,
            }
            try:
                importados.append(self._guardar_en_repositorio(repositorio_menus, dato))
            except Exception as exc:  # pragma: no cover - detalle del repositorio real
                errores.append({"menu_id": menu["menu_id"], "nombre": menu["nombre"], "error": str(exc)})
        return {
            "ok": not errores,
            "importados": importados,
            "omitidos": omitidos,
            "errores": errores,
            "resumen": {"importados": len(importados), "omitidos": len(omitidos), "errores": len(errores)},
        }


def formatear_previa_i13(previa: dict[str, Any]) -> str:
    r = previa["resumen"]
    lineas = [
        "I1.3 — VISTA PREVIA DE MENÚS",
        "=" * 72,
        f"Hojas: {r['hojas_analizadas']} | Menús: {r['menus_detectados']} | Duplicados: {r['duplicados']} | Bloqueados: {r['bloqueados']}",
        f"Platos: {r['platos']} | Exactos: {r['vinculos_exactos']} | Probables: {r['vinculos_probables']} | Sin resolver: {r['sin_resolver']}",
    ]
    for indice, menu in enumerate(previa.get("menus", []), 1):
        estados = []
        if menu.get("duplicado"):
            estados.append("DUPLICADO")
        if menu.get("bloqueado"):
            estados.append("BLOQUEADO")
        extra = f" | {'/'.join(estados)}" if estados else ""
        lineas.append(f"{indice}. {menu['menu_id']} | {menu['nombre']} | {len(menu['platos'])} platos{extra}")
        seccion_anterior = None
        for plato in menu["platos"]:
            if plato["seccion"] != seccion_anterior:
                lineas.append(f"   [{plato['seccion']}]")
                seccion_anterior = plato["seccion"]
            destino = plato["receta_nombre"] or "SIN VINCULAR"
            coste = f" | coste {plato['coste_declarado']:.2f} €" if plato.get("coste_declarado") is not None else ""
            lineas.append(f"   - {plato['nombre_excel']} -> {destino} [{plato['estado']}, {plato['confianza']:.0%}]{coste}")
        lineas.append(
            "   Económico: "
            f"venta={menu.get('precio_venta')} | coste={menu.get('coste_total_declarado')} | "
            f"food cost={menu.get('food_cost_pct')}"
        )
        for aviso in menu.get("avisos", []):
            lineas.append(f"   AVISO: {aviso}")
    lineas += ["=" * 72, "VISTA PREVIA: no se ha modificado la base de datos."]
    return "\n".join(lineas)
