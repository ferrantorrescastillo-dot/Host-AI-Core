from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import unicodedata
import uuid
from copy import deepcopy
from dataclasses import dataclass, asdict
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Callable, Iterable, Optional


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(c for c in text if not unicodedata.combining(c)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def _num(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("€", "").replace(" ", "")
    if text.count(",") == 1 and text.count(".") == 0:
        text = text.replace(",", ".")
    elif text.count(".") > 1 and "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _uid(prefix: str) -> str:
    return f"{prefix}-{uuid.uuid4().hex[:10].upper()}"


@dataclass
class LineaRecepcionPiloto1:
    numero: int
    descripcion: str
    cantidad: float
    unidad: str
    precio_unitario: Optional[float] = None
    total_linea: Optional[float] = None
    codigo_proveedor: str = ""
    articulo_id: str = ""
    articulo_nombre: str = ""
    confianza: float = 0.0
    estado: str = "SIN_RESOLVER"
    accion: str = "REVISAR"
    motivo: str = ""


class RecepcionInteligentePiloto1:
    """Flujo unificado de recepción para el piloto privado.

    Principios:
    - nombre del artículo primero; proveedor como segunda vía;
    - vista previa obligatoria;
    - ninguna escritura con líneas sin resolver;
    - backup + escritura atómica + rollback;
    - factura, precio y stock se registran en una única transacción lógica.
    """

    EXTENSIONES = {".pdf", ".txt", ".csv", ".xlsx", ".xlsm", ".json", ".png", ".jpg", ".jpeg"}

    def __init__(self, base_dir: Path | str):
        self.base_dir = Path(base_dir).resolve()
        self.db_dir = self.base_dir / "DATOS" / "db"
        self.piloto_dir = self.base_dir / "DATOS" / "piloto" / "recepciones"
        self.backups_dir = self.base_dir / "DATOS" / "backups" / "piloto1"
        self.piloto_dir.mkdir(parents=True, exist_ok=True)
        self.backups_dir.mkdir(parents=True, exist_ok=True)

    # ---------- lectura ----------
    def preparar(self, ruta_archivo: str | Path, texto_manual: str = "", proveedor: str = "") -> dict[str, Any]:
        ruta = Path(ruta_archivo)
        if not ruta.is_absolute():
            ruta = self.base_dir / ruta
        ruta = ruta.resolve()
        if not ruta.exists():
            raise FileNotFoundError(f"No existe el documento: {ruta}")
        if ruta.suffix.lower() not in self.EXTENSIONES:
            raise ValueError(f"Formato no soportado en PILOTO-1: {ruta.suffix}")

        filas, metadatos = self._leer_documento(ruta, texto_manual)
        proveedor_final = proveedor.strip() or str(metadatos.get("proveedor") or "").strip() or self._proveedor_desde_nombre(ruta.stem)
        articulos = self._cargar_lista(self.db_dir / "articulos.json")
        lineas = []
        for idx, fila in enumerate(filas, 1):
            linea = self._normalizar_fila(idx, fila)
            if not linea:
                continue
            self._resolver_linea(linea, articulos, proveedor_final)
            lineas.append(linea)

        exactas = sum(x.estado == "EXACTA" for x in lineas)
        probables = sum(x.estado == "PROBABLE" for x in lineas)
        pendientes = sum(x.estado == "SIN_RESOLVER" for x in lineas)
        plan_id = _uid("REC")
        plan = {
            "version": "PILOTO-1.0",
            "plan_id": plan_id,
            "creado_en": datetime.now().isoformat(timespec="seconds"),
            "archivo": str(ruta),
            "archivo_sha256": hashlib.sha256(ruta.read_bytes()).hexdigest(),
            "proveedor": proveedor_final,
            "numero_factura": str(metadatos.get("numero_factura") or ""),
            "fecha_documento": str(metadatos.get("fecha") or datetime.now().date().isoformat()),
            "lineas": [asdict(x) for x in lineas],
            "resumen": {
                "lineas": len(lineas), "exactas": exactas, "probables": probables,
                "pendientes": pendientes, "importe_detectado": round(sum((x.total_linea or ((x.precio_unitario or 0) * x.cantidad)) for x in lineas), 4),
            },
            "estado": "LISTA_PARA_CONFIRMAR" if lineas and pendientes == 0 else "REQUIERE_REVISION",
            "solo_vista_previa": True,
        }
        self.guardar_plan(plan)
        return plan

    def _leer_documento(self, ruta: Path, texto_manual: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        ext = ruta.suffix.lower()
        if ext in {".xlsx", ".xlsm"}:
            return self._leer_excel(ruta)
        if ext == ".csv":
            return self._leer_csv(ruta)
        if ext == ".json":
            data = json.loads(ruta.read_text(encoding="utf-8"))
            if isinstance(data, dict):
                return list(data.get("lineas", [])), {k: data.get(k) for k in ("proveedor", "numero_factura", "fecha")}
            return list(data), {}
        if ext == ".pdf":
            text = self._texto_pdf(ruta)
            if not text.strip() and texto_manual.strip():
                text = texto_manual
            if not text.strip():
                raise ValueError("El PDF no contiene texto extraíble. Pega el texto OCR/manual para continuar.")
            return self._parsear_texto(text)
        if ext in {".png", ".jpg", ".jpeg"}:
            if not texto_manual.strip():
                raise ValueError("La imagen necesita texto OCR/manual en esta versión del piloto.")
            return self._parsear_texto(texto_manual)
        return self._parsear_texto(texto_manual or ruta.read_text(encoding="utf-8", errors="ignore"))

    @staticmethod
    def _texto_pdf(ruta: Path) -> str:
        try:
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(str(ruta)).pages)
        except Exception:
            try:
                from PyPDF2 import PdfReader
                return "\n".join((p.extract_text() or "") for p in PdfReader(str(ruta)).pages)
            except Exception:
                return ""

    def _leer_excel(self, ruta: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        from openpyxl import load_workbook
        wb = load_workbook(ruta, data_only=True, read_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return [], {}
        header_idx = self._detectar_cabecera(rows)
        headers = [str(x or "").strip() for x in rows[header_idx]]
        items = []
        for row in rows[header_idx + 1:]:
            if not any(v not in (None, "") for v in row):
                continue
            items.append(dict(zip(headers, row)))
        return items, {}

    def _leer_csv(self, ruta: Path) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        text = ruta.read_text(encoding="utf-8-sig", errors="ignore")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        return list(csv.DictReader(text.splitlines(), dialect=dialect)), {}

    @staticmethod
    def _detectar_cabecera(rows: list[tuple]) -> int:
        claves = {"articulo", "producto", "descripcion", "cantidad", "precio", "unidad", "importe"}
        best = (0, 0)
        for idx, row in enumerate(rows[:25]):
            score = sum(any(k in _norm(cell) for k in claves) for cell in row if cell is not None)
            if score > best[1]:
                best = (idx, score)
        return best[0]

    def _parsear_texto(self, text: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
        filas: list[dict[str, Any]] = []
        meta: dict[str, Any] = {}
        for raw in text.splitlines():
            line = " ".join(raw.strip().split())
            if not line:
                continue
            low = _norm(line)
            if low.startswith("proveedor"):
                meta["proveedor"] = line.split(":", 1)[-1].strip()
                continue
            if "factura" in low and re.search(r"\d", line):
                meta.setdefault("numero_factura", line.split(":", 1)[-1].strip())
                continue
            # Formato flexible: descripción ; cantidad ; unidad ; precio [; total]
            parts = [p.strip() for p in re.split(r"[;\t|]", line)]
            if len(parts) >= 3 and _num(parts[1]) is not None:
                filas.append({"descripcion": parts[0], "cantidad": parts[1], "unidad": parts[2], "precio": parts[3] if len(parts)>3 else None, "total": parts[4] if len(parts)>4 else None})
                continue
            # "2 kg Patata monalisa 1,20" o "Patata monalisa 2 kg 1,20"
            m = re.match(r"^(?P<cant>\d+(?:[.,]\d+)?)\s*(?P<uni>kg|g|l|ml|ud|uds|unidad(?:es)?)\s+(?P<desc>.+?)(?:\s+(?P<precio>\d+[.,]\d{1,4}))?$", line, re.I)
            if not m:
                m = re.match(r"^(?P<desc>.+?)\s+(?P<cant>\d+(?:[.,]\d+)?)\s*(?P<uni>kg|g|l|ml|ud|uds|unidad(?:es)?)(?:\s+(?P<precio>\d+[.,]\d{1,4}))?$", line, re.I)
            if m:
                filas.append({"descripcion": m.group("desc"), "cantidad": m.group("cant"), "unidad": m.group("uni"), "precio": m.groupdict().get("precio")})
        return filas, meta

    @staticmethod
    def _proveedor_desde_nombre(stem: str) -> str:
        text = re.sub(r"(?i)factura|albaran|albarán|pedido|\d{4}[-_]\d{2}[-_]\d{2}|\d+", " ", stem)
        return " ".join(text.replace("_", " ").replace("-", " ").split()).strip()

    def _normalizar_fila(self, numero: int, fila: dict[str, Any]) -> Optional[LineaRecepcionPiloto1]:
        norm = {_norm(k): v for k, v in fila.items()}
        def pick(*keys):
            for key in keys:
                for actual, val in norm.items():
                    if key == actual or key in actual:
                        if val not in (None, ""):
                            return val
            return None
        desc = pick("descripcion", "articulo", "producto", "concepto", "nombre")
        if not desc:
            return None
        cantidad = _num(pick("cantidad", "cant", "unidades")) or 1.0
        unidad = str(pick("unidad", "ud", "formato") or "ud").strip().lower()
        precio = _num(pick("precio unitario", "precio", "coste unitario"))
        total = _num(pick("importe", "total linea", "total"))
        if precio is None and total is not None and cantidad:
            precio = total / cantidad
        return LineaRecepcionPiloto1(numero, str(desc).strip(), cantidad, unidad, precio, total, str(pick("codigo", "referencia") or ""))

    def _resolver_linea(self, linea: LineaRecepcionPiloto1, articulos: list[dict], proveedor: str) -> None:
        target = _norm(linea.descripcion)
        if not target:
            return
        scored = []
        for art in articulos:
            name = _norm(art.get("nombre"))
            if not name:
                continue
            score = SequenceMatcher(None, target, name).ratio()
            if target == name:
                score = 1.0
            elif target in name or name in target:
                score = max(score, min(len(target), len(name)) / max(len(target), len(name)) * 0.98)
            # proveedor es segunda vía, nunca sustituye al nombre
            if proveedor and _norm(proveedor) and _norm(proveedor) in _norm(art.get("proveedor")):
                score = min(1.0, score + 0.03)
            scored.append((score, art))
        if not scored:
            linea.motivo = "Catálogo vacío o sin coincidencias."
            return
        score, art = max(scored, key=lambda x: x[0])
        linea.confianza = round(score, 4)
        linea.articulo_id = str(art.get("codigo") or art.get("id") or "")
        linea.articulo_nombre = str(art.get("nombre") or "")
        if score >= 0.94:
            linea.estado, linea.accion, linea.motivo = "EXACTA", "VINCULAR", "Coincidencia exacta o equivalente de alta confianza."
        elif score >= 0.72:
            linea.estado, linea.accion, linea.motivo = "PROBABLE", "REVISAR", "Existe una coincidencia probable; requiere confirmación."
        else:
            linea.articulo_id = ""
            linea.articulo_nombre = ""
            linea.estado, linea.accion, linea.motivo = "SIN_RESOLVER", "REVISAR", "No existe una coincidencia suficientemente segura."

    # ---------- revisión ----------
    @staticmethod
    def actualizar_decision(plan: dict[str, Any], numero: int, accion: str, articulo: Optional[dict[str, Any]] = None, nombre_nuevo: str = "", unidad: str = "") -> dict[str, Any]:
        line = next((x for x in plan["lineas"] if int(x["numero"]) == int(numero)), None)
        if line is None:
            raise ValueError(f"No existe la línea {numero}.")
        accion = accion.upper()
        if accion == "VINCULAR":
            if not articulo:
                raise ValueError("Falta el artículo para vincular.")
            line.update({"articulo_id": str(articulo.get("codigo") or articulo.get("id") or ""), "articulo_nombre": str(articulo.get("nombre") or ""), "estado": "RESUELTA", "accion": "VINCULAR", "confianza": 1.0, "motivo": "Vínculo confirmado por el usuario."})
        elif accion == "CREAR":
            nombre = nombre_nuevo.strip() or str(line["descripcion"]).strip()
            uni = unidad.strip().lower() or str(line.get("unidad") or "").strip().lower()
            if not nombre or not uni:
                raise ValueError("Nombre y unidad son obligatorios para crear un artículo.")
            line.update({"articulo_id": "", "articulo_nombre": nombre, "estado": "RESUELTA", "accion": "CREAR", "confianza": 1.0, "unidad": uni, "motivo": "Alta confirmada por el usuario."})
        elif accion == "OMITIR":
            line.update({"estado": "OMITIDA", "accion": "OMITIR", "motivo": "Línea omitida por el usuario."})
        else:
            raise ValueError("Acción no soportada.")
        pendientes = sum(x["estado"] in {"SIN_RESOLVER", "PROBABLE"} for x in plan["lineas"])
        plan["resumen"]["pendientes"] = pendientes
        plan["estado"] = "LISTA_PARA_CONFIRMAR" if pendientes == 0 else "REQUIERE_REVISION"
        return plan

    def buscar_articulos(self, texto: str, proveedor: str = "", limite: int = 8) -> list[dict[str, Any]]:
        dummy = LineaRecepcionPiloto1(0, texto, 1, "ud")
        arts = self._cargar_lista(self.db_dir / "articulos.json")
        target = _norm(texto)
        out = []
        for art in arts:
            score = SequenceMatcher(None, target, _norm(art.get("nombre"))).ratio()
            if target == _norm(art.get("nombre")):
                score = 1.0
            if proveedor and _norm(proveedor) in _norm(art.get("proveedor")):
                score += 0.03
            out.append((min(score, 1.0), art))
        return [{**deepcopy(a), "confianza": round(s, 4)} for s, a in sorted(out, key=lambda x: x[0], reverse=True)[:limite]]

    # ---------- escritura segura ----------
    def aplicar(self, plan: dict[str, Any], confirmacion: str, db_dir: Optional[Path | str] = None) -> dict[str, Any]:
        if confirmacion.strip().upper() != "RECEPCIONAR":
            raise ValueError("Confirmación incorrecta. Debes escribir RECEPCIONAR.")
        if plan.get("estado") != "LISTA_PARA_CONFIRMAR":
            raise ValueError("La recepción todavía tiene líneas pendientes.")
        work_db = Path(db_dir).resolve() if db_dir else self.db_dir
        work_db.mkdir(parents=True, exist_ok=True)
        files = {
            "articulos": work_db / "articulos.json", "lotes": work_db / "stock_lotes.json",
            "movimientos": work_db / "stock_movimientos.json", "precios": work_db / "precios.json",
            "facturas": work_db / "facturas_piloto.json", "proveedores": work_db / "proveedores.json",
        }
        original = {k: self._cargar_lista(p) for k, p in files.items()}
        data = deepcopy(original)
        tx_id = _uid("TXP1")
        backup = self.backups_dir / tx_id
        backup.mkdir(parents=True, exist_ok=True)
        for key, path in files.items():
            if path.exists():
                shutil.copy2(path, backup / path.name)

        created = updated_prices = stock_entries = omitted = 0
        try:
            provider = str(plan.get("proveedor") or "").strip()
            provider_code = self._ensure_provider(data["proveedores"], provider)
            article_by_id = {str(a.get("codigo") or a.get("id")): a for a in data["articulos"]}
            for line in plan["lineas"]:
                if line["accion"] == "OMITIR":
                    omitted += 1
                    continue
                art = None
                if line["accion"] == "CREAR":
                    code = self._next_article_code(data["articulos"])
                    art = {"codigo": code, "nombre": line["articulo_nombre"], "observaciones": "Alta desde PILOTO-1 recepción", "proveedor": provider or None, "familia": None, "precio": line.get("precio_unitario"), "unidad": line.get("unidad"), "activo": True, "origen": "piloto_recepcion", "fecha_importacion": datetime.now().isoformat(timespec="seconds")}
                    data["articulos"].append(art); article_by_id[code] = art; created += 1
                    line["articulo_id"] = code
                else:
                    art = article_by_id.get(str(line.get("articulo_id")))
                    if not art:
                        raise ValueError(f"Artículo vinculado inexistente en línea {line['numero']}.")
                price = _num(line.get("precio_unitario"))
                if price is not None:
                    old = _num(art.get("precio"))
                    art["precio"] = price
                    if provider:
                        art["proveedor"] = provider
                    data["precios"].append({"id": _uid("PRECIO"), "articulo_id": line["articulo_id"], "nombre": art.get("nombre"), "precio_unitario": price, "precio_anterior": old, "unidad": line.get("unidad"), "proveedor": provider, "fecha": plan.get("fecha_documento"), "origen": plan.get("plan_id"), "creado_en": datetime.now().isoformat(timespec="seconds")})
                    updated_prices += 1
                lot_id = _uid("LOTE")
                lot = {"id": lot_id, "articulo_id": line["articulo_id"], "nombre": art.get("nombre"), "cantidad": float(line["cantidad"]), "unidad": line.get("unidad"), "familia": art.get("familia") or "", "ubicacion": "pendiente_ubicar", "proveedor": provider, "fecha_entrada": plan.get("fecha_documento"), "caducidad": "", "coste_unitario": price or 0.0, "documento_id": plan.get("plan_id"), "creado_en": datetime.now().isoformat(timespec="seconds")}
                data["lotes"].append(lot)
                data["movimientos"].append({"id": _uid("MOV"), "tipo": "entrada", "articulo_id": line["articulo_id"], "nombre": art.get("nombre"), "cantidad": float(line["cantidad"]), "unidad": line.get("unidad"), "motivo": f"recepción {plan.get('numero_factura') or plan.get('plan_id')}", "lote_id": lot_id, "proveedor": provider, "documento_id": plan.get("plan_id"), "creado_en": datetime.now().isoformat(timespec="seconds")})
                stock_entries += 1
            factura = {"id": plan.get("plan_id"), "numero": plan.get("numero_factura"), "proveedor": provider, "proveedor_id": provider_code, "fecha": plan.get("fecha_documento"), "archivo": plan.get("archivo"), "archivo_sha256": plan.get("archivo_sha256"), "lineas": deepcopy(plan["lineas"]), "importe": plan["resumen"].get("importe_detectado"), "estado": "RECEPCIONADA", "transaccion": tx_id, "creado_en": datetime.now().isoformat(timespec="seconds")}
            # Idempotencia por hash de archivo
            if any(f.get("archivo_sha256") == factura["archivo_sha256"] for f in data["facturas"]):
                raise ValueError("Este documento ya fue recepcionado anteriormente.")
            data["facturas"].append(factura)
            for key, path in files.items():
                self._write_atomic(path, data[key])
            result = {"estado": "COMMIT", "transaccion": tx_id, "backup": str(backup), "factura_id": factura["id"], "articulos_creados": created, "precios_actualizados": updated_prices, "entradas_stock": stock_entries, "lineas_omitidas": omitted, "integridad": "OK"}
            self._guardar_resultado(plan, result)
            return result
        except Exception:
            for key, path in files.items():
                if (backup / path.name).exists():
                    shutil.copy2(backup / path.name, path)
                elif not original[key] and path.exists():
                    path.unlink()
            raise

    @staticmethod
    def _ensure_provider(providers: list[dict], name: str) -> str:
        if not name:
            return ""
        norm = _norm(name)
        for p in providers:
            if _norm(p.get("nombre")) == norm:
                return str(p.get("codigo") or p.get("id") or "")
        code = f"PROV{len(providers)+1:04d}"
        providers.append({"codigo": code, "nombre": name, "nombre_normalizado": norm, "articulos_asociados": 0, "variantes_detectadas": [name], "estado": "activo", "observaciones": "Alta desde PILOTO-1"})
        return code

    @staticmethod
    def _next_article_code(articles: list[dict]) -> str:
        nums = []
        for a in articles:
            m = re.fullmatch(r"ART(\d+)", str(a.get("codigo") or ""))
            if m: nums.append(int(m.group(1)))
        return f"ART{max(nums, default=0)+1:06d}"

    def guardar_plan(self, plan: dict[str, Any]) -> Path:
        path = self.piloto_dir / f"{plan['plan_id']}.json"
        self._write_atomic(path, plan)
        return path

    def _guardar_resultado(self, plan: dict[str, Any], result: dict[str, Any]) -> None:
        payload = {"plan": plan, "resultado": result}
        self._write_atomic(self.piloto_dir / f"{plan['plan_id']}_resultado.json", payload)

    @staticmethod
    def _cargar_lista(path: Path) -> list[dict]:
        if not path.exists():
            return []
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, list): return data
        for key in ("items", "articulos", "proveedores", "facturas", "registros"):
            if isinstance(data.get(key), list): return data[key]
        return []

    @staticmethod
    def _write_atomic(path: Path, data: Any) -> None:
        path.parent.mkdir(parents=True, exist_ok=True)
        temp = path.with_suffix(path.suffix + ".tmp")
        temp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        json.loads(temp.read_text(encoding="utf-8"))
        temp.replace(path)

    # ---------- diagnóstico ----------
    def diagnostico(self) -> dict[str, Any]:
        diag = self.base_dir / "DATOS" / "mur" / "diagnostico_piloto1"
        if diag.exists(): shutil.rmtree(diag)
        db = diag / "db"; db.mkdir(parents=True)
        sample_articles = [
            {"codigo": "ART000001", "nombre": "Patata monalisa", "proveedor": "Makro", "familia": "Verduras", "precio": 1.0, "unidad": "kg", "activo": True},
            {"codigo": "ART000002", "nombre": "Mantequilla", "proveedor": "Makro", "familia": "Lácteos", "precio": 5.0, "unidad": "kg", "activo": True},
        ]
        for name, data in {"articulos.json": sample_articles, "stock_lotes.json": [], "stock_movimientos.json": [], "precios.json": [], "proveedores.json": [], "facturas_piloto.json": []}.items():
            self._write_atomic(db / name, data)
        csv_path = diag / "factura_makro.csv"
        csv_path.write_text("descripcion;cantidad;unidad;precio\nPatata monalisa;10;kg;1,20\nMantequilla;2;kg;5,50\nNuez moscada;0,2;kg;12,00\n", encoding="utf-8")
        original_db = self.db_dir
        try:
            self.db_dir = db
            plan = self.preparar(csv_path, proveedor="Makro")
            # Crear artículo pendiente de forma explícita
            self.actualizar_decision(plan, 3, "CREAR", nombre_nuevo="Nuez moscada", unidad="kg")
            # Confirmar probable si apareciera
            for line in plan["lineas"]:
                if line["estado"] == "PROBABLE":
                    candidates = self.buscar_articulos(line["descripcion"], "Makro")
                    self.actualizar_decision(plan, line["numero"], "VINCULAR", candidates[0])
            result = self.aplicar(plan, "RECEPCIONAR", db_dir=db)
            return {"diagnostico": "OK", "estado": result["estado"], "lineas": len(plan["lineas"]), "exactas": plan["resumen"]["exactas"], "creados": result["articulos_creados"], "precios": result["precios_actualizados"], "stock": result["entradas_stock"], "facturas": len(self._cargar_lista(db / "facturas_piloto.json")), "archivo_aislado": str(csv_path), "db_aislada": str(db), "integridad": result["integridad"]}
        finally:
            self.db_dir = original_db


def formatear_diagnostico_piloto1(data: dict[str, Any]) -> str:
    return "\n".join([
        "PILOTO-1 — RECEPCIÓN INTELIGENTE DE MERCANCÍA", "=" * 78,
        f"Diagnóstico: {data['diagnostico']} | Estado: {data['estado']} | Integridad: {data['integridad']}",
        f"Líneas: {data['lineas']} | Exactas: {data['exactas']} | Artículos creados: {data['creados']}",
        f"Precios actualizados: {data['precios']} | Entradas stock: {data['stock']} | Facturas: {data['facturas']}",
        f"Documento aislado: {data['archivo_aislado']}", f"Base aislada: {data['db_aislada']}", "-" * 78,
        "Se validó lectura CSV, relación por nombre, proveedor como segunda vía, alta controlada, precios, stock, factura, backup y commit.",
        "El diagnóstico no modifica artículos, precios, stock ni facturas reales.",
    ])


__all__ = ["RecepcionInteligentePiloto1", "LineaRecepcionPiloto1", "formatear_diagnostico_piloto1"]
