from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from MODELOS.analisis_estructural_excel_555b import (
    HojaClasificada555B,
    ResultadoAnalisisEstructural555B,
)
from SERVICIOS.detector_bloques_recetas_555b import DetectorBloquesRecetas555B


def _normalizar(valor: object) -> str:
    texto = str(valor or "").strip().lower()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    return " ".join(re.sub(r"[^a-z0-9]+", " ", texto).split())


class ClasificadorHojasExcel555B:
    TIPOS = (
        "ARTICULOS",
        "FICHAS_TECNICAS",
        "MENUS",
        "RESUMENES",
        "AUXILIARES",
        "DESCONOCIDAS",
    )

    def analizar(self, ruta_archivo: str | Path) -> dict[str, Any]:
        ruta = Path(ruta_archivo)
        resultado = ResultadoAnalisisEstructural555B(archivo=str(ruta))
        if not ruta.exists():
            resultado.errores.append(f"No se encuentra el archivo: {ruta}")
            return resultado.to_dict()
        if ruta.suffix.lower() not in {".xlsx", ".xlsm"}:
            resultado.errores.append("Esta fase admite únicamente archivos .xlsx o .xlsm.")
            return resultado.to_dict()

        try:
            libro = load_workbook(ruta, read_only=True, data_only=True)
        except Exception as exc:
            resultado.errores.append(f"No se pudo abrir el Excel: {exc}")
            return resultado.to_dict()

        try:
            nombres = [hoja.title for hoja in libro.worksheets]
            for hoja in libro.worksheets:
                clasificada = self._clasificar_hoja(hoja)
                clasificada.hoja_relacionada = self._buscar_relacion(hoja.title, nombres, clasificada.tipo)
                resultado.hojas.append(clasificada)
        finally:
            libro.close()
        return resultado.to_dict()

    def _clasificar_hoja(self, hoja) -> HojaClasificada555B:
        nombre = _normalizar(hoja.title)
        filas = list(hoja.iter_rows(min_row=1, max_row=min(hoja.max_row or 1, 80), values_only=True))
        textos = [_normalizar(v) for fila in filas for v in fila if v not in (None, "")]
        corpus = " | ".join(textos)
        puntuaciones = {tipo: 0 for tipo in self.TIPOS}
        motivos: dict[str, list[str]] = {tipo: [] for tipo in self.TIPOS}

        def sumar(tipo: str, puntos: int, motivo: str) -> None:
            puntuaciones[tipo] += puntos
            motivos[tipo].append(motivo)

        if "listado de articulos" in nombre or nombre.startswith("articulos"):
            sumar("ARTICULOS", 100, "nombre de hoja de catálogo de artículos")
        if any(token in corpus for token in ("proveedor", "familia", "precio")) and any(token in corpus for token in ("articulo", "codigo")):
            sumar("ARTICULOS", 45, "campos propios de catálogo de artículos")

        if nombre.startswith("m p ") or "ficha tecnica plato" in corpus:
            sumar("FICHAS_TECNICAS", 95, "patrón de matriz/ficha técnica")
        if any(token in corpus for token in ("precio kg", "euros racion", "gr o kg", "coste racion")):
            sumar("FICHAS_TECNICAS", 40, "campos de coste y dosificación")

        if "menu" in nombre and not nombre.startswith("m p "):
            sumar("MENUS", 85, "nombre de hoja de menú")
        if any(token in corpus for token in ("aperitivo", "plato principal", "postre")) and "ficha tecnica plato" not in corpus:
            sumar("MENUS", 25, "estructura de composición de menú")

        if "resumen" in nombre or "total articulos" in corpus:
            sumar("RESUMENES", 100, "hoja de resumen")

        if nombre in {"hoja 1", "sheet1"} or "codigos 4 1 3" in nombre or "plantilla" in nombre:
            sumar("AUXILIARES", 85, "hoja auxiliar o plantilla")
        if (hoja.max_row or 0) <= 1 and (hoja.max_column or 0) <= 1:
            sumar("AUXILIARES", 40, "hoja vacía o casi vacía")

        tipo = max(puntuaciones, key=puntuaciones.get)
        maximo = puntuaciones[tipo]
        if maximo <= 0:
            tipo = "DESCONOCIDAS"
            confianza = 0.0
            razones = ["sin señales suficientes"]
        else:
            segundo = sorted(puntuaciones.values(), reverse=True)[1]
            confianza = round(min(100.0, 55 + maximo * 0.35 + max(0, maximo - segundo) * 0.15), 2)
            razones = motivos[tipo]

        bloques = []
        if tipo == "FICHAS_TECNICAS":
            bloques = DetectorBloquesRecetas555B().detectar_en_filas(filas)

        accion = {
            "ARTICULOS": "REUTILIZAR_CATALOGO",
            "FICHAS_TECNICAS": "ANALIZAR_BLOQUES",
            "MENUS": "RELACIONAR_CON_FICHAS",
            "RESUMENES": "IGNORAR_IMPORTACION",
            "AUXILIARES": "IGNORAR_IMPORTACION",
            "DESCONOCIDAS": "REVISION_MANUAL",
        }[tipo]
        return HojaClasificada555B(
            nombre=hoja.title,
            tipo=tipo,
            confianza=confianza,
            motivos=razones,
            bloques=bloques,
            accion_recomendada=accion,
        )

    @staticmethod
    def _buscar_relacion(nombre: str, nombres: list[str], tipo: str) -> str | None:
        if tipo not in {"MENUS", "FICHAS_TECNICAS"}:
            return None
        base = _normalizar(nombre)
        base = re.sub(r"^m p ", "", base)
        base = re.sub(r"^menu ", "", base)
        palabras = {p for p in base.split() if len(p) > 2 and p not in {"finde", "precio"}}
        mejor: tuple[int, str] | None = None
        for candidato in nombres:
            if candidato == nombre:
                continue
            c_norm = _normalizar(candidato)
            if tipo == "MENUS" and not c_norm.startswith("m p "):
                continue
            if tipo == "FICHAS_TECNICAS" and "menu" not in c_norm:
                continue
            c_base = re.sub(r"^m p ", "", c_norm)
            c_base = re.sub(r"^menu ", "", c_base)
            coincidencias = len(palabras & set(c_base.split()))
            if coincidencias and (mejor is None or coincidencias > mejor[0]):
                mejor = (coincidencias, candidato)
        return mejor[1] if mejor else None
