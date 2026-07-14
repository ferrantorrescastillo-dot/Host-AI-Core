from __future__ import annotations

import re
import unicodedata
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _norm(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).lower().strip()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


class DetectorMultihojaMenusI13411:
    """Detecta hojas de menú combinando pistas de nombre y validación estructural."""

    VERSION = "I1.3.4.1.1"
    SECCIONES = {
        "aperitivo", "aperitivos", "coctel", "frio", "frios", "caliente", "calientes",
        "entrante", "entrantes", "primero", "primeros", "segundo", "segundos", "postre",
        "postres", "a compartir", "en mesa o coctel", "bodega", "bebidas",
    }
    CABECERAS_ECONOMICAS = {
        "p v p", "pvp", "neto", "costo", "coste", "food cost", "euros racion",
        "precio kg", "gr o kg", "benefici", "beneficio", "escandallo",
    }
    EXCLUSION_NOMBRE = (
        "listado de articulos", "plantilla", "ficha tecnica", "hoja 1",
    )

    @staticmethod
    def _es_mp(nombre_normalizado: str) -> bool:
        tokens = nombre_normalizado.split()
        return (
            nombre_normalizado.startswith("m p ")
            or nombre_normalizado.startswith("mp ")
            or (len(tokens) >= 2 and tokens[0] == "m" and tokens[1] == "p")
        )

    def _evaluar_hoja(self, ws) -> dict[str, Any]:
        nombre = ws.title
        nn = _norm(nombre)
        razones: list[str] = []
        exclusiones: list[str] = []
        score = 0

        if self._es_mp(nn):
            exclusiones.append("La hoja es M.P/MP (materias primas o fichas), no un menú final.")
            score -= 12
        if any(x in nn for x in self.EXCLUSION_NOMBRE):
            exclusiones.append("El nombre identifica una hoja técnica, plantilla o vacía.")
            score -= 12

        if "menu" in nn:
            score += 4
            razones.append("El nombre de la hoja contiene MENU/MENÚ.")
        if any(x in nn for x in ("boda", "comunion", "calçotada", "calcotada", "bbq", "fin de ano", "agencia")):
            score += 1
            razones.append("El nombre contiene una pista de evento o formato gastronómico.")

        textos: list[str] = []
        filas_con_coste = 0
        max_row = min(int(ws.max_row or 0), 120)
        max_col = min(int(ws.max_column or 0), 12)
        for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col, values_only=True):
            vals = list(row)
            norm_vals = [_norm(v) for v in vals if v not in (None, "")]
            textos.extend(norm_vals)
            if vals and isinstance(vals[0], str) and str(vals[0]).strip():
                numericos = [v for v in vals[1:5] if isinstance(v, (int, float))]
                if numericos:
                    filas_con_coste += 1

        primeros = textos[:12]
        if any(t.startswith("menu ") or t == "menu" for t in primeros):
            score += 4
            razones.append("El encabezado interno identifica un menú.")
        cabeceras = sorted({t for t in textos[:50] if t in self.CABECERAS_ECONOMICAS})
        if len(cabeceras) >= 3:
            score += 4
            razones.append("Contiene cabeceras económicas de menú (PVP/neto/coste/food cost).")
        elif cabeceras:
            score += 1
            razones.append("Contiene alguna cabecera económica compatible.")

        secciones = sorted({t for t in textos[:100] if t.rstrip(" s") in self.SECCIONES or t in self.SECCIONES})
        if secciones:
            score += 2
            razones.append("Contiene secciones gastronómicas reconocibles.")
        if filas_con_coste >= 3:
            score += 2
            razones.append("Contiene varias líneas de platos con valores económicos.")

        if any("ficha tecnica plato" in t for t in textos[:15]):
            score -= 7
            exclusiones.append("El contenido corresponde a una ficha técnica de plato.")

        if exclusiones and score < 7:
            clasificacion = "NO_MENU"
        elif score >= 8:
            clasificacion = "MENU_CONFIRMADO"
        elif score >= 5:
            clasificacion = "MENU_PROBABLE"
        else:
            clasificacion = "NO_MENU"

        return {
            "hoja": nombre,
            "clasificacion": clasificacion,
            "puntuacion": score,
            "razones": razones,
            "exclusiones": exclusiones,
            "secciones_detectadas": secciones,
            "cabeceras_economicas": cabeceras,
            "filas_con_coste": filas_con_coste,
        }

    def detectar(self, ruta_excel: str | Path) -> dict[str, Any]:
        ruta = Path(ruta_excel)
        if not ruta.exists():
            raise FileNotFoundError(ruta)
        wb = load_workbook(ruta, read_only=True, data_only=True)
        try:
            evaluadas = [self._evaluar_hoja(ws) for ws in wb.worksheets]
        finally:
            wb.close()
        confirmadas = [x for x in evaluadas if x["clasificacion"] == "MENU_CONFIRMADO"]
        probables = [x for x in evaluadas if x["clasificacion"] == "MENU_PROBABLE"]
        descartadas = [x for x in evaluadas if x["clasificacion"] == "NO_MENU"]
        return {
            "version": self.VERSION,
            "archivo": str(ruta.resolve()),
            "hojas_totales": len(evaluadas),
            "confirmadas": confirmadas,
            "probables": probables,
            "descartadas": descartadas,
            "hojas_confirmadas": [x["hoja"] for x in confirmadas],
            "hojas_probables": [x["hoja"] for x in probables],
        }


def formatear_deteccion_multihoja_i13411(resultado: dict[str, Any]) -> str:
    lineas = [
        "I1.3.4.1.1 — DETECCIÓN MULTIMENÚ",
        "=" * 78,
        f"Hojas del libro: {resultado.get('hojas_totales', 0)} | Menús confirmados: {len(resultado.get('confirmadas', []))} | Probables: {len(resultado.get('probables', []))}",
    ]
    candidatos = resultado.get("confirmadas", []) + resultado.get("probables", [])
    for i, item in enumerate(candidatos, 1):
        etiqueta = "CONFIRMADO" if item.get("clasificacion") == "MENU_CONFIRMADO" else "PROBABLE"
        lineas.append(f"{i}. [{etiqueta}] {item.get('hoja')} | puntuación {item.get('puntuacion')}")
    if not candidatos:
        lineas.append("No se detectaron hojas candidatas a menú.")
    lineas.append("Las hojas M.P/MP, plantillas, fichas técnicas, listados y hojas vacías se excluyen.")
    return "\n".join(lineas)
