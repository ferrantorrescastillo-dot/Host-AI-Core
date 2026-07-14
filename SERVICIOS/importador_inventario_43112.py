from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional
from datetime import datetime
import json


@dataclass
class RegistroInventarioContado:
    codigo: str
    articulo: str
    unidad: str
    stock_sistema: float
    stock_contado: float
    diferencia: float
    ubicacion: Optional[str]
    proveedor: Optional[str]
    familia: Optional[str]
    observaciones_inventario: Optional[str]
    fecha_importacion: str


@dataclass
class ResultadoImportacionInventario:
    total_filas: int
    importados: int
    ignorados: int
    errores: int
    ruta_destino: str
    mensajes: List[str]


class ImportadorInventario43112:
    """
    Importa el inventario físico contado desde la plantilla 4.3.11.1.

    Entrada:
    DATOS/inventario_4_3_11_1.xlsx

    Salida:
    DATOS/db/inventario_contado_4_3_11_2.json

    Importante:
    - No modifica stock_inicial.json.
    - Solo guarda lo contado para comparar en el siguiente paso.
    """

    HOJA = "Inventario"

    def importar_desde_excel(
        self,
        ruta_excel: str,
        ruta_destino: str = "DATOS/db/inventario_contado_4_3_11_2.json",
    ) -> ResultadoImportacionInventario:
        filas = self._leer_excel(ruta_excel)
        return self.importar_filas(filas, ruta_destino)

    def importar_filas(
        self,
        filas: List[Dict[str, Any]],
        ruta_destino: str = "DATOS/db/inventario_contado_4_3_11_2.json",
    ) -> ResultadoImportacionInventario:
        registros: List[Dict[str, Any]] = []
        ignorados = 0
        errores = 0
        mensajes: List[str] = []

        for numero, fila in enumerate(filas, start=2):
            codigo = self._texto(fila.get("Codigo"))
            articulo = self._texto(fila.get("Articulo"))
            unidad = self._texto(fila.get("Unidad"))
            stock_sistema = self._numero(fila.get("Stock sistema"))
            stock_contado = self._numero(fila.get("Stock contado"))

            # Si no hay stock contado, la fila no se ha contado todavía.
            if stock_contado is None:
                ignorados += 1
                continue

            if not codigo:
                errores += 1
                mensajes.append(f"Fila {numero}: falta código.")
                continue

            if not articulo:
                errores += 1
                mensajes.append(f"Fila {numero}: falta artículo.")
                continue

            if not unidad:
                errores += 1
                mensajes.append(f"Fila {numero}: falta unidad.")
                continue

            if stock_sistema is None:
                errores += 1
                mensajes.append(f"Fila {numero}: stock sistema inválido.")
                continue

            if stock_contado < 0:
                errores += 1
                mensajes.append(f"Fila {numero}: stock contado negativo.")
                continue

            diferencia = round(stock_contado - stock_sistema, 4)

            registro = RegistroInventarioContado(
                codigo=codigo,
                articulo=articulo,
                unidad=unidad,
                stock_sistema=stock_sistema,
                stock_contado=stock_contado,
                diferencia=diferencia,
                ubicacion=self._texto(fila.get("Ubicación")),
                proveedor=self._texto(fila.get("Proveedor")),
                familia=self._texto(fila.get("Familia")),
                observaciones_inventario=self._texto(fila.get("Observaciones inventario")),
                fecha_importacion=datetime.now().isoformat(timespec="seconds"),
            )

            registros.append(asdict(registro))

        ruta = Path(ruta_destino)
        ruta.parent.mkdir(parents=True, exist_ok=True)
        ruta.write_text(json.dumps(registros, ensure_ascii=False, indent=2), encoding="utf-8")

        mensajes.append(
            f"Inventario importado. Importados: {len(registros)}. Ignorados: {ignorados}. Errores: {errores}."
        )

        return ResultadoImportacionInventario(
            total_filas=len(filas),
            importados=len(registros),
            ignorados=ignorados,
            errores=errores,
            ruta_destino=str(ruta),
            mensajes=mensajes,
        )

    def _leer_excel(self, ruta_excel: str) -> List[Dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para leer el inventario.") from exc

        ruta = Path(ruta_excel)
        if not ruta.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta_excel}")

        wb = load_workbook(ruta, data_only=False)

        if self.HOJA not in wb.sheetnames:
            raise ValueError(f"No existe la hoja '{self.HOJA}'. Hojas disponibles: {wb.sheetnames}")

        ws = wb[self.HOJA]
        headers = [cell.value for cell in ws[1]]
        filas: List[Dict[str, Any]] = []

        for row in ws.iter_rows(min_row=2, values_only=False):
            values = []
            for cell in row:
                # Si es fórmula de diferencia, no dependemos de ella; calculamos diferencia nosotros.
                values.append(cell.value)

            if not any(values):
                continue

            filas.append({str(headers[i]).strip(): values[i] if i < len(values) else None for i in range(len(headers))})

        return filas

    def _texto(self, valor: Any) -> Optional[str]:
        if valor is None:
            return None
        texto = str(valor).strip()
        return texto if texto else None

    def _numero(self, valor: Any) -> Optional[float]:
        if valor is None or str(valor).strip() == "":
            return None
        if isinstance(valor, (int, float)):
            return float(valor)
        limpio = str(valor).replace(" ", "").replace(",", ".")
        # Si accidentalmente llega una fórmula de Excel, no intentamos evaluarla.
        if limpio.startswith("="):
            return None
        try:
            return float(limpio)
        except ValueError:
            return None


__all__ = [
    "ImportadorInventario43112",
    "RegistroInventarioContado",
    "ResultadoImportacionInventario",
]
