from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional
import unicodedata


@dataclass
class ClasificacionArticulo:
    codigo: Optional[str]
    articulo: str
    observaciones: Optional[str]
    proveedor: Optional[str]
    familia_original: Optional[str]
    precio: Optional[float]
    tipo_sugerido: str
    familia_sugerida: str
    confianza: float
    revisar: str
    motivo: str


@dataclass
class InformeClasificacionArticulos:
    total_articulos: int
    clasificaciones: List[ClasificacionArticulo] = field(default_factory=list)

    @property
    def total_revisar(self) -> int:
        return sum(1 for item in self.clasificaciones if item.revisar in {"SI", "REVISAR"})

    @property
    def resumen_tipos(self) -> Dict[str, int]:
        resumen: Dict[str, int] = {}
        for item in self.clasificaciones:
            resumen[item.tipo_sugerido] = resumen.get(item.tipo_sugerido, 0) + 1
        return resumen


class ClasificadorInteligenteArticulos412:
    COLUMNAS_SALIDA = [
        "Codigo", "Articulo", "Observaciones", "Proveedor", "Familia", "Precio",
        "Tipo sugerido", "Familia sugerida", "Confianza", "Revisar", "Motivo",
    ]

    def clasificar_filas(self, filas: Iterable[Dict[str, Any]]) -> InformeClasificacionArticulos:
        clasificaciones = [self.clasificar_articulo(self._normalizar_fila(fila)) for fila in filas]
        return InformeClasificacionArticulos(len(clasificaciones), clasificaciones)

    def clasificar_articulo(self, fila: Dict[str, Any]) -> ClasificacionArticulo:
        codigo = self._texto(fila.get("codigo"))
        articulo = self._texto(fila.get("articulo")) or ""
        observaciones = self._texto(fila.get("observaciones"))
        proveedor = self._texto(fila.get("proveedor"))
        familia_original = self._texto(fila.get("familia"))
        precio = self._precio(fila.get("precio"))

        nombre = self._normalizar(articulo)
        proveedor_norm = self._normalizar(proveedor)

        tipo = "OTRO"
        familia = "Sin clasificar"
        confianza = 0.45
        revisar = "SI"
        motivo = "No clasificado por reglas claras."

        if not articulo:
            confianza = 0.0
            motivo = "Artículo vacío."
        elif nombre.startswith("a.p") or proveedor_norm.startswith("m.p aperitivos"):
            tipo, familia, confianza, revisar, motivo = "APERITIVO", "Aperitivos", 0.98, "NO", "Prefijo A.P o proveedor de aperitivos."
        elif "menu" in nombre:
            tipo, familia, confianza, revisar, motivo = "MENU", "Menús", 0.95, "NO", "Contiene la palabra menú."
        elif nombre.startswith("m.p") or proveedor_norm == "m.p" or proveedor_norm.startswith("m.p "):
            tipo, familia, confianza, revisar, motivo = "MATERIA_PRIMA", "Materia prima", 0.88, "REVISAR", "Prefijo o proveedor M.P."

        tipo, familia, confianza, revisar, motivo = self._reglas_especificas(nombre, tipo, familia, confianza, revisar, motivo)

        return ClasificacionArticulo(
            codigo, articulo, observaciones, proveedor, familia_original, precio,
            tipo, familia, round(confianza, 2), revisar, motivo
        )

    def leer_excel(self, ruta_excel: str, hoja: str = "Listado de Artículos") -> List[Dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para leer archivos .xlsx.") from exc

        ruta = Path(ruta_excel)
        if not ruta.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta_excel}")

        wb = load_workbook(ruta, data_only=True)
        if hoja not in wb.sheetnames:
            raise ValueError(f"No existe la hoja '{hoja}'. Hojas disponibles: {wb.sheetnames}")

        ws = wb[hoja]
        headers = [cell.value for cell in ws[1]]
        filas = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            filas.append({str(headers[i]): row[i] if i < len(row) else None for i in range(len(headers))})
        return filas

    def exportar_excel_clasificado(self, ruta_origen: str, ruta_destino: str, hoja_origen: str = "Listado de Artículos") -> InformeClasificacionArticulos:
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para exportar archivos .xlsx.") from exc

        filas = self.leer_excel(ruta_origen, hoja_origen)
        informe = self.clasificar_filas(filas)

        wb = load_workbook(ruta_origen)
        for nombre_hoja in ["Clasificacion 4.1.2", "Resumen 4.1.2"]:
            if nombre_hoja in wb.sheetnames:
                del wb[nombre_hoja]

        ws = wb.create_sheet("Clasificacion 4.1.2")
        ws.append(self.COLUMNAS_SALIDA)

        for item in informe.clasificaciones:
            ws.append([
                item.codigo, item.articulo, item.observaciones, item.proveedor,
                item.familia_original, item.precio, item.tipo_sugerido,
                item.familia_sugerida, item.confianza, item.revisar, item.motivo,
            ])

        header_fill = PatternFill("solid", fgColor="0F766E")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        widths = {"A": 14, "B": 48, "C": 30, "D": 24, "E": 22, "F": 12, "G": 20, "H": 24, "I": 12, "J": 12, "K": 42}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for row in ws.iter_rows(min_row=2):
            revisar = row[9].value
            if revisar == "SI":
                fill = PatternFill("solid", fgColor="FECACA")
            elif revisar == "REVISAR":
                fill = PatternFill("solid", fgColor="FEF3C7")
            else:
                fill = None
            if fill:
                for cell in row:
                    cell.fill = fill

        resumen = wb.create_sheet("Resumen 4.1.2")
        resumen.append(["HOST AI 4.1.2 - Resumen clasificación", ""])
        resumen.append(["Total artículos", informe.total_articulos])
        resumen.append(["Artículos a revisar", informe.total_revisar])
        resumen.append(["Estado", "BORRADOR PARA REVISIÓN"])
        resumen.append(["", ""])
        resumen.append(["Tipo sugerido", "Cantidad"])
        for tipo, cantidad in sorted(informe.resumen_tipos.items()):
            resumen.append([tipo, cantidad])

        for cell in resumen[1]:
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
            cell.font = Font(color="FFFFFF", bold=True)
        resumen.column_dimensions["A"].width = 38
        resumen.column_dimensions["B"].width = 18

        wb.save(ruta_destino)
        return informe

    def _normalizar_fila(self, fila: Dict[str, Any]) -> Dict[str, Any]:
        mapa = {"codigo": "", "articulo": "", "observaciones": "", "proveedor": "", "familia": "", "precio": None}
        for clave, valor in fila.items():
            clave_norm = self._normalizar(clave)
            if clave_norm in {"codigo", "cod", "id"}:
                mapa["codigo"] = valor
            elif clave_norm in {"articulo", "producto", "nombre"}:
                mapa["articulo"] = valor
            elif clave_norm in {"observaciones", "observacion", "obs", "notas"}:
                mapa["observaciones"] = valor
            elif clave_norm in {"proveedor", "supplier"}:
                mapa["proveedor"] = valor
            elif clave_norm in {"familia", "categoria"}:
                mapa["familia"] = valor
            elif clave_norm in {"precio", "coste", "precio_compra"}:
                mapa["precio"] = valor
        return mapa

    def _reglas_especificas(self, nombre: str, tipo: str, familia: str, confianza: float, revisar: str, motivo: str):
        if any(p in nombre for p in ["coca", "fanta", "sprite", "agua", "cerveza", "vino", "vermut", "refresco", "nestea", "aquarius", "cava"]):
            return "BEBIDA", "Bebidas", max(confianza, 0.94), "NO", "Palabra de bebida detectada."

        if tipo != "APERITIVO" and any(p in nombre for p in ["postre", "tarta", "brownie", "flan", "helado", "sorbete", "mousse", "pastel"]):
            return "POSTRE", "Postres", max(confianza, 0.90), "NO", "Palabra de postre detectada."

        if tipo not in {"APERITIVO", "BEBIDA", "MENU", "POSTRE"} and any(p in nombre for p in ["salsa", "fondo", "caldo", "fumet", "base", "crema", "pure", "sofrito", "romesco", "alioli", "mayonesa", "bechamel"]):
            tipo, familia, confianza, revisar, motivo = "ELABORACION", "Elaboraciones", max(confianza, 0.82), "REVISAR", "Palabra de elaboración detectada."

        if any(p in nombre for p in ["servilleta", "caja", "bolsa", "film", "aluminio", "guante", "bandeja", "vaso", "tapa", "envase"]):
            return "CONSUMIBLE", "Consumibles", 0.95, "NO", "Consumible detectado."

        reglas_familia = [
            ("Aceites", ["aceite", "aove"]),
            ("Harinas", ["harina", "maicena", "panko"]),
            ("Arroces y cereales", ["arroz", "fideo", "pasta", "quinoa"]),
            ("Carnes", ["pollo", "ternera", "cerdo", "secreto", "carrillera", "solomillo", "chuleta", "costilla", "pato", "butifarra"]),
            ("Pescados", ["bacalao", "merluza", "sardina", "salmon", "corvina", "atun"]),
            ("Mariscos", ["gamba", "langostino", "sepia", "calamar", "pulpo", "mejillon", "almeja", "bogavante"]),
            ("Verduras", ["tomate", "cebolla", "ajo", "puerro", "zanahoria", "pimiento", "alcachofa", "calcot", "patata", "berenjena", "calabacin"]),
            ("Frutas", ["limon", "naranja", "manzana", "uva", "pera", "maracuya"]),
            ("Lácteos", ["leche", "nata", "queso", "mantequilla", "yogur"]),
            ("Huevos", ["huevo"]),
            ("Especias y condimentos", ["sal", "pimienta", "pimenton", "curry", "comino", "azafran", "jengibre", "canela"]),
            ("Panadería", ["pan", "brioche", "briox", "coca", "tostad", "mollete"]),
            ("Salsas", ["salsa", "romesco", "mayonesa", "alioli"]),
            ("Conservas", ["lata", "conserva", "anchoa", "oliva"]),
            ("Congelados", ["congelad"]),
        ]

        if tipo in {"MATERIA_PRIMA", "OTRO", "ELABORACION"}:
            for familia_detectada, palabras in reglas_familia:
                if any(palabra in nombre for palabra in palabras):
                    if tipo == "OTRO":
                        tipo, confianza, revisar, motivo = "MATERIA_PRIMA", max(confianza, 0.82), "REVISAR", "Familia de materia prima detectada."
                    elif tipo == "MATERIA_PRIMA":
                        confianza = max(confianza, 0.90)
                        revisar = "NO" if confianza >= 0.90 else revisar
                    familia = familia_detectada
                    break
        return tipo, familia, confianza, revisar, motivo

    def _normalizar(self, texto: Any) -> str:
        texto = "" if texto is None else str(texto)
        texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
        return texto.strip().lower()

    def _texto(self, valor: Any) -> Optional[str]:
        if valor is None:
            return None
        texto = str(valor).strip()
        return texto if texto else None

    def _precio(self, valor: Any) -> Optional[float]:
        if valor is None or str(valor).strip() == "":
            return None
        if isinstance(valor, (int, float)):
            return float(valor)
        limpio = str(valor).replace("€", "").replace(" ", "").strip()
        if "," in limpio and "." in limpio:
            limpio = limpio.replace(".", "").replace(",", ".")
        else:
            limpio = limpio.replace(",", ".")
        try:
            return float(limpio)
        except ValueError:
            return None


__all__ = ["ClasificadorInteligenteArticulos412", "ClasificacionArticulo", "InformeClasificacionArticulos"]
