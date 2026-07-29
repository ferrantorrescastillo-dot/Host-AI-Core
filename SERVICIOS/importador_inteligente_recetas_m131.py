from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any
import json, re, unicodedata


def _norm(v: Any) -> str:
    s = str(v or '').strip().lower()
    s = ''.join(c for c in unicodedata.normalize('NFD', s) if unicodedata.category(c) != 'Mn')
    return re.sub(r'\s+', ' ', s).strip(' .,:;-_')


def clasificar_prefijo_articulo(articulo: dict[str, Any]) -> str:
    codigo = str(articulo.get('codigo') or articulo.get('id') or '').strip().upper()
    nombre = str(articulo.get('nombre') or articulo.get('articulo') or '').strip().upper()
    texto = f'{codigo} {nombre}'.strip()
    if re.match(r'^A[. ]?P\b', texto):
        return 'APERITIVO'
    if re.match(r'^M[. ]?P\b', texto):
        return 'MATERIA_PRIMA'
    return 'SIN_CLASIFICAR'


@dataclass
class IngredienteDetectado:
    texto_original: str
    nombre: str
    cantidad: float | None
    unidad: str
    articulo_id: str = ''
    articulo_nombre: str = ''
    confianza: float = 0.0
    estado: str = 'SIN_RESOLVER'
    clase_articulo: str = ''


@dataclass
class BorradorRecetaM131:
    nombre: str
    rendimiento: float | None
    ingredientes: list[IngredienteDetectado]
    elaboracion: str
    origen_tipo: str
    origen_ruta: str = ''
    bloqueos: list[str] | None = None

    def to_dict(self) -> dict[str, Any]:
        d = asdict(self)
        d['bloqueos'] = list(self.bloqueos or [])
        return d


class ImportadorInteligenteRecetasM131:
    """Ingesta segura de recetas desde texto, Word, PDF y Excel.

    M1.3.1 crea un borrador y lo entrega al ResolutorRecetasM13. No resuelve
    todavía artículos inexistentes de forma anidada (M1.3.2).

    Dependencias opcionales:
    - Word (.docx): requiere python-docx.
    - PDF (.pdf): requiere pypdf.
    """

    UNIDADES = {
        'kg':'kg','kilo':'kg','kilos':'kg','g':'g','gr':'g','gramo':'g','gramos':'g',
        'l':'l','litro':'l','litros':'l','ml':'ml','cl':'cl','ud':'ud','u':'ud',
        'unidad':'ud','unidades':'ud','pieza':'ud','piezas':'ud'
    }

    def __init__(self, ruta_articulos: str | Path):
        self.ruta_articulos = Path(ruta_articulos)

    def desde_texto(self, texto: str, origen_tipo: str='TEXTO', origen_ruta: str='') -> BorradorRecetaM131:
        lineas = [x.strip() for x in str(texto or '').replace('\r','').split('\n')]
        lineas_utiles = [x for x in lineas if x]
        if not lineas_utiles:
            raise ValueError('La receta está vacía.')
        nombre = self._detectar_nombre(lineas_utiles)
        rendimiento = self._detectar_rendimiento(lineas_utiles)
        ingredientes, inicio_elab = self._detectar_ingredientes(lineas)
        elaboracion = '\n'.join(x for x in lineas[inicio_elab:] if x.strip()).strip() if inicio_elab is not None else ''
        borrador = BorradorRecetaM131(nombre, rendimiento, ingredientes, elaboracion, origen_tipo, origen_ruta, [])
        self._resolver_articulos(borrador)
        self._validar(borrador)
        return borrador

    def desde_archivo(self, ruta: str | Path, hoja: str | None=None) -> BorradorRecetaM131:
        p = Path(ruta)
        if not p.exists(): raise FileNotFoundError(p)
        ext = p.suffix.lower()
        if ext == '.docx': texto = self._leer_docx(p); tipo='WORD'
        elif ext == '.pdf': texto = self._leer_pdf(p); tipo='PDF'
        elif ext in {'.xlsx','.xlsm'}: texto = self._leer_excel(p, hoja); tipo='EXCEL'
        elif ext in {'.txt','.md'}: texto = p.read_text(encoding='utf-8'); tipo='TEXTO'
        else: raise ValueError(f'Formato no soportado en M1.3.1: {ext}')
        return self.desde_texto(texto, tipo, str(p))

    def payload_creacion(self, borrador: BorradorRecetaM131) -> dict[str, Any]:
        if borrador.bloqueos:
            raise ValueError('El borrador contiene bloqueos: ' + ' | '.join(borrador.bloqueos))
        return {
            'nombre': borrador.nombre,
            'raciones_base': borrador.rendimiento,
            'ingredientes': [
                {'articulo_id': i.articulo_id, 'cantidad': i.cantidad, 'unidad': i.unidad}
                for i in borrador.ingredientes
            ],
            'observaciones': f'Importada desde {borrador.origen_tipo}.\n{borrador.elaboracion}'.strip(),
            'origen_importacion': borrador.origen_tipo,
        }

    def _detectar_nombre(self, lineas: list[str]) -> str:
        for x in lineas:
            n=_norm(x)
            if n in {'ingredientes','elaboracion','elaboración','preparacion','preparación'}: continue
            if re.search(r'\b(raciones?|rendimiento|para)\b\s*[:=]?\s*\d+', n): continue
            return x.strip(' :-')
        return 'Receta sin nombre'

    def _detectar_rendimiento(self, lineas: list[str]) -> float | None:
        for x in lineas:
            m=re.search(r'(?:rendimiento|raciones?|para)\s*[:=]?\s*(\d+(?:[.,]\d+)?)', _norm(x))
            if m: return float(m.group(1).replace(',','.'))
        return None

    def _detectar_ingredientes(self, lineas: list[str]) -> tuple[list[IngredienteDetectado], int | None]:
        out=[]; en_ing=False; inicio_elab=None
        for idx,x in enumerate(lineas):
            n=_norm(x)
            if n in {'ingredientes','ingredientes:'}: en_ing=True; continue
            if n in {'elaboracion','preparacion','procedimiento','pasos','elaboración','preparación'}:
                inicio_elab=idx+1; break
            m=re.match(r'^\s*(\d+(?:[.,]\d+)?)\s*(kg|kilos?|g|gr|gramos?|l|litros?|ml|cl|ud|u|unidades?|piezas?)\s+(.*)$', x, re.I)
            if m:
                cantidad=float(m.group(1).replace(',','.')); unidad=self.UNIDADES.get(_norm(m.group(2)),_norm(m.group(2)))
                out.append(IngredienteDetectado(x,m.group(3).strip(' .,-'),cantidad,unidad)); en_ing=True
            elif en_ing and x.strip() and not re.search(r'\b(raciones?|rendimiento)\b', n):
                # ingrediente sin cantidad: se conserva como bloqueo, nunca se inventa
                out.append(IngredienteDetectado(x,x.strip(' .,-'),None,''))
        return out,inicio_elab

    def _resolver_articulos(self, borrador: BorradorRecetaM131) -> None:
        articulos=json.loads(self.ruta_articulos.read_text(encoding='utf-8')) if self.ruta_articulos.exists() else []
        for ing in borrador.ingredientes:
            q=_norm(ing.nombre); candidatos=[]
            for a in articulos:
                clase=clasificar_prefijo_articulo(a)
                if clase == 'APERITIVO':
                    continue
                nombre=str(a.get('nombre') or a.get('articulo') or '')
                n=_norm(nombre)
                if not n: continue
                if n==q: score=1.0
                elif q in n or n in q: score=.92
                else:
                    tq,tn=set(q.split()),set(n.split()); score=len(tq&tn)/max(len(tq|tn),1)
                if score>=.55: candidatos.append((score,a,clase))
            candidatos.sort(key=lambda z:-z[0])
            if candidatos:
                score,a,clase=candidatos[0]
                ing.articulo_id=str(a.get('codigo') or a.get('id') or a.get('articulo_id') or '')
                ing.articulo_nombre=str(a.get('nombre') or a.get('articulo') or '')
                ing.confianza=round(score,4);ing.clase_articulo=clase
                ing.estado='VINCULADO' if score>=.92 else 'PROBABLE'

    def _validar(self, b: BorradorRecetaM131) -> None:
        bloqueos=[]
        if not b.nombre: bloqueos.append('Falta el nombre de la receta.')
        if b.rendimiento is None or b.rendimiento<=0: bloqueos.append('Falta un rendimiento válido.')
        if not b.ingredientes: bloqueos.append('No se detectaron ingredientes.')
        for i in b.ingredientes:
            if i.cantidad is None or not i.unidad: bloqueos.append(f'Ingrediente sin cantidad/unidad: {i.nombre}.')
            if not i.articulo_id: bloqueos.append(f'Artículo no resuelto: {i.nombre}.')
            if i.clase_articulo == 'APERITIVO': bloqueos.append(f'Un A.P no puede usarse como materia prima automática: {i.articulo_nombre}.')
        b.bloqueos=list(dict.fromkeys(bloqueos))

    @staticmethod
    def _leer_docx(p: Path) -> str:
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError('Falta python-docx para importar archivos Word (.docx).') from exc
        d=Document(p); partes=[x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows: partes.append(' '.join(c.text.strip() for c in row.cells if c.text.strip()))
        return '\n'.join(partes)

    @staticmethod
    def _leer_pdf(p: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RuntimeError('Falta pypdf para importar PDF.') from exc
        return '\n'.join((page.extract_text() or '') for page in PdfReader(str(p)).pages)

    @staticmethod
    def _leer_excel(p: Path, hoja: str | None) -> str:
        from openpyxl import load_workbook
        wb=load_workbook(p,data_only=True,read_only=True)
        try:
            ws=wb[hoja] if hoja else wb[wb.sheetnames[0]]
            filas=[]
            for row in ws.iter_rows(values_only=True):
                vals=[str(v).strip() for v in row if v not in (None,'')]
                if vals: filas.append(' '.join(vals))
            return '\n'.join(filas)
        finally:
            wb.close()

__all__=['ImportadorInteligenteRecetasM131','BorradorRecetaM131','IngredienteDetectado','clasificar_prefijo_articulo']
