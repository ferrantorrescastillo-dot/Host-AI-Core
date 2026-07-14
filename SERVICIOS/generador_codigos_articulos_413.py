from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional


@dataclass
class CodigoArticuloGenerado:
    codigo_original: Optional[str]
    codigo_final: str
    articulo: str
    accion: str
    revisar: str
    motivo: str


@dataclass
class InformeCodigosArticulos:
    total_articulos: int
    codigos_generados: int
    codigos_respetados: int
    codigos_duplicados: int
    resultados: List[CodigoArticuloGenerado] = field(default_factory=list)

    @property
    def total_revisar(self) -> int:
        return sum(1 for item in self.resultados if item.revisar == "SI")


class GeneradorCodigosArticulos413:
    """
    Genera códigos internos estables para artículos.

    Regla principal:
    - Si Codigo está vacío: genera ART000001, ART000002...
    - Si Codigo existe y no está duplicado: lo respeta.
    - Si Codigo existe duplicado: lo respeta, pero marca revisión.
    """

    def __init__(self, prefijo: str = "ART", digitos: int = 6) -> None:
        self.prefijo = prefijo
        self.digitos = digitos

    def generar_para_filas(self, filas: Iterable[Dict[str, Any]]) -> InformeCodigosArticulos:
        filas_lista = list(filas)
        codigos_existentes = []
        for fila in filas_lista:
            codigo = self._texto(fila.get("Codigo", fila.get("codigo", "")))
            if codigo:
                codigos_existentes.append(codigo)

        duplicados = self._detectar_duplicados(codigos_existentes)
        usados = set(codigos_existentes)
        contador = 1

        resultados: List[CodigoArticuloGenerado] = []
        generados = 0
        respetados = 0
        duplicados_count = 0

        for fila in filas_lista:
            codigo_original = self._texto(fila.get("Codigo", fila.get("codigo", "")))
            articulo = self._texto(fila.get("Articulo", fila.get("articulo", ""))) or ""

            if codigo_original:
                if codigo_original in duplicados:
                    duplicados_count += 1
                    resultados.append(CodigoArticuloGenerado(
                        codigo_original=codigo_original,
                        codigo_final=codigo_original,
                        articulo=articulo,
                        accion="respetado_duplicado",
                        revisar="SI",
                        motivo="Código existente duplicado. Revisar antes de importar.",
                    ))
                else:
                    respetados += 1
                    resultados.append(CodigoArticuloGenerado(
                        codigo_original=codigo_original,
                        codigo_final=codigo_original,
                        articulo=articulo,
                        accion="respetado",
                        revisar="NO",
                        motivo="Código existente respetado.",
                    ))
                continue

            nuevo_codigo = self._siguiente_codigo(usados, contador)
            while nuevo_codigo in usados:
                contador += 1
                nuevo_codigo = self._siguiente_codigo(usados, contador)

            usados.add(nuevo_codigo)
            contador += 1
            generados += 1

            resultados.append(CodigoArticuloGenerado(
                codigo_original=None,
                codigo_final=nuevo_codigo,
                articulo=articulo,
                accion="generado",
                revisar="NO",
                motivo="Código generado automáticamente por Host AI.",
            ))

        return InformeCodigosArticulos(
            total_articulos=len(filas_lista),
            codigos_generados=generados,
            codigos_respetados=respetados,
            codigos_duplicados=duplicados_count,
            resultados=resultados,
        )

    def leer_excel(self, ruta_excel: str, hoja: str = "Listado de Artículos") -> List[Dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para leer archivos .xlsx.") from exc

        ruta = Path(ruta_excel)
        if not ruta.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta_excel}")

        wb = load_workbook(ruta, data_only=True)
        if hoja not in wb.sheetnames:
            raise ValueError(f"No existe la hoja '{hoja}'. Hojas disponibles: {wb.sheetnames}")

        ws = wb[hoja]
        headers = [cell.value for cell in ws[1]]
        filas = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            filas.append({str(headers[i]): row[i] if i < len(row) else None for i in range(len(headers))})
        return filas

    def exportar_excel_con_codigos(
        self,
        ruta_origen: str,
        ruta_destino: str,
        hoja_origen: str = "Listado de Artículos",
        hoja_destino: str = "Codigos 4.1.3",
    ) -> InformeCodigosArticulos:
        try:
            from openpyxl import load_workbook
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para exportar archivos .xlsx.") from exc

        filas = self.leer_excel(ruta_origen, hoja_origen)
        informe = self.generar_para_filas(filas)

        wb = load_workbook(ruta_origen)
        if hoja_destino in wb.sheetnames:
            del wb[hoja_destino]
        if "Resumen 4.1.3" in wb.sheetnames:
            del wb["Resumen 4.1.3"]

        ws = wb.create_sheet(hoja_destino)
        ws.append([
            "Codigo original",
            "Codigo Host AI",
            "Articulo",
            "Accion",
            "Revisar",
            "Motivo",
        ])

        for item in informe.resultados:
            ws.append([
                item.codigo_original,
                item.codigo_final,
                item.articulo,
                item.accion,
                item.revisar,
                item.motivo,
            ])

        header_fill = PatternFill("solid", fgColor="7C3AED")
        header_font = Font(color="FFFFFF", bold=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        widths = {"A": 18, "B": 18, "C": 48, "D": 22, "E": 12, "F": 48}
        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for row in ws.iter_rows(min_row=2):
            if row[4].value == "SI":
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor="FECACA")
            elif row[3].value == "generado":
                for cell in row:
                    cell.fill = PatternFill("solid", fgColor="DCFCE7")

        resumen = wb.create_sheet("Resumen 4.1.3")
        resumen.append(["HOST AI 4.1.3 - Generación códigos internos", ""])
        resumen.append(["Total artículos", informe.total_articulos])
        resumen.append(["Códigos generados", informe.codigos_generados])
        resumen.append(["Códigos respetados", informe.codigos_respetados])
        resumen.append(["Códigos duplicados", informe.codigos_duplicados])
        resumen.append(["A revisar", informe.total_revisar])
        resumen.append(["Estado", "REVISAR DUPLICADOS" if informe.codigos_duplicados else "APTO"])

        for cell in resumen[1]:
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
            cell.font = Font(color="FFFFFF", bold=True)

        resumen.column_dimensions["A"].width = 32
        resumen.column_dimensions["B"].width = 20

        wb.save(ruta_destino)
        return informe

    def _siguiente_codigo(self, usados: set[str], contador: int) -> str:
        return f"{self.prefijo}{contador:0{self.digitos}d}"

    def _detectar_duplicados(self, codigos: List[str]) -> set[str]:
        vistos = set()
        duplicados = set()
        for codigo in codigos:
            if codigo in vistos:
                duplicados.add(codigo)
            vistos.add(codigo)
        return duplicados

    def _texto(self, valor: Any) -> Optional[str]:
        if valor is None:
            return None
        texto = str(valor).strip()
        return texto if texto else None


__all__ = [
    "GeneradorCodigosArticulos413",
    "CodigoArticuloGenerado",
    "InformeCodigosArticulos",
]
