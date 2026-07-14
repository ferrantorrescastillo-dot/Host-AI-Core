from __future__ import annotations

from pathlib import Path
from typing import Dict, List, Any
from datetime import datetime, date
import json

try:
    from openpyxl import load_workbook, Workbook
    from openpyxl.utils import get_column_letter
except Exception as exc:
    raise ImportError(
        "Falta openpyxl. Instala con: pip install openpyxl"
    ) from exc

from MODELOS.excel_importacion import AnalisisExcel, HojaExcel, ColumnaExcel


class LectorUniversalExcel:
    """
    Lector Universal de Excel v3.0.2.1.

    Lee cualquier .xlsx y devuelve:
    - hojas
    - filas/columnas
    - columnas detectadas
    - tipos de datos
    - porcentaje de vacío
    - vista previa
    - informe JSON exportable
    """

    def __init__(self, base_dir: Path):
        self.base_dir = Path(base_dir)
        self.importaciones_dir = self.base_dir / "DATOS" / "importaciones"
        self.importaciones_dir.mkdir(parents=True, exist_ok=True)

    def analizar_archivo(
        self,
        ruta_archivo: str,
        filas_preview: int = 10,
        exportar_json: bool = True,
    ) -> Dict[str, Any]:
        ruta = Path(ruta_archivo)
        if not ruta.is_absolute():
            ruta = self.base_dir / ruta

        if not ruta.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta}")

        if ruta.suffix.lower() not in [".xlsx", ".xlsm"]:
            raise ValueError("Solo se soportan archivos .xlsx o .xlsm en esta versión.")

        wb = load_workbook(ruta, data_only=True, read_only=False)
        hojas = []

        for ws in wb.worksheets:
            hoja = self._analizar_hoja(ws, filas_preview)
            hojas.append(hoja)

        analisis = AnalisisExcel(
            archivo=str(ruta),
            nombre_archivo=ruta.name,
            hojas=hojas,
            total_hojas=len(hojas),
            total_filas=sum(h.filas for h in hojas),
            total_columnas=sum(h.columnas for h in hojas),
            total_celdas_utilizadas=sum(h.celdas_utilizadas for h in hojas),
        )

        datos = analisis.to_dict()
        if exportar_json:
            exportado = self.exportar_analisis_json(datos)
            datos["json_exportado"] = exportado["archivo"]

        datos["lectura_host_ai"] = (
            f"Excel analizado: {ruta.name}. "
            f"{datos['total_hojas']} hojas, {datos['total_filas']} filas, "
            f"{datos['total_columnas']} columnas."
        )
        return datos

    def exportar_analisis_json(self, analisis: Dict[str, Any], nombre: str = "") -> Dict[str, Any]:
        nombre = nombre or f"analisis_excel_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        destino = self.importaciones_dir / nombre
        destino.write_text(json.dumps(analisis, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        return {
            "archivo": str(destino),
            "lectura_host_ai": f"Análisis Excel exportado: {destino.name}.",
        }

    def _analizar_hoja(self, ws, filas_preview: int) -> HojaExcel:
        max_row = ws.max_row or 0
        max_col = ws.max_column or 0
        celdas_utilizadas = 0

        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in [None, ""]:
                    celdas_utilizadas += 1

        headers = self._detectar_headers(ws, max_col)
        columnas = []
        for col_idx in range(1, max_col + 1):
            col = self._analizar_columna(ws, col_idx, headers.get(col_idx, ""), max_row)
            columnas.append(col)

        preview = self._vista_previa(ws, headers, max_row, max_col, filas_preview)

        return HojaExcel(
            nombre=ws.title,
            filas=max_row,
            columnas=max_col,
            celdas_utilizadas=celdas_utilizadas,
            columnas_detectadas=columnas,
            vista_previa=preview,
        )

    def _detectar_headers(self, ws, max_col: int) -> Dict[int, str]:
        """
        Versión inicial: considera la primera fila como cabecera.
        En próximos sprints detectaremos cabeceras desplazadas.
        """
        headers = {}
        for col_idx in range(1, max_col + 1):
            value = ws.cell(row=1, column=col_idx).value
            headers[col_idx] = str(value).strip() if value not in [None, ""] else f"Columna {get_column_letter(col_idx)}"
        return headers

    def _analizar_columna(self, ws, col_idx: int, nombre: str, max_row: int) -> ColumnaExcel:
        valores = []
        total = max(0, max_row - 1)

        for row_idx in range(2, max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value not in [None, ""]:
                valores.append(value)

        celdas_con_dato = len(valores)
        porcentaje_vacio = round((1 - (celdas_con_dato / total)) * 100, 2) if total else 100.0

        return ColumnaExcel(
            indice=col_idx,
            letra=get_column_letter(col_idx),
            nombre_detectado=nombre,
            tipo_detectado=self._detectar_tipo(valores),
            total_celdas=total,
            celdas_con_dato=celdas_con_dato,
            porcentaje_vacio=porcentaje_vacio,
            ejemplos=[self._serializar(v) for v in valores[:5]],
        )

    def _vista_previa(self, ws, headers: Dict[int, str], max_row: int, max_col: int, filas_preview: int) -> List[Dict[str, Any]]:
        preview = []
        limite = min(max_row, filas_preview + 1)
        for row_idx in range(2, limite + 1):
            fila = {}
            for col_idx in range(1, max_col + 1):
                nombre = headers.get(col_idx, f"Columna {get_column_letter(col_idx)}")
                fila[nombre] = self._serializar(ws.cell(row=row_idx, column=col_idx).value)
            preview.append(fila)
        return preview

    def _detectar_tipo(self, valores: List[Any]) -> str:
        if not valores:
            return "vacio"

        total = len(valores)
        nums = sum(1 for v in valores if isinstance(v, (int, float)) and not isinstance(v, bool))
        fechas = sum(1 for v in valores if isinstance(v, (datetime, date)))
        bools = sum(1 for v in valores if isinstance(v, bool))

        textos_moneda = 0
        textos = 0
        for v in valores:
            if isinstance(v, str):
                textos += 1
                limpio = v.replace("€", "").replace(",", ".").strip()
                try:
                    float(limpio)
                    if "€" in v:
                        textos_moneda += 1
                except Exception:
                    pass

        if fechas / total >= 0.6:
            return "fecha"
        if bools / total >= 0.6:
            return "booleano"
        if nums / total >= 0.6:
            return "numero"
        if (nums + textos_moneda) / total >= 0.6:
            return "moneda"
        if textos / total >= 0.6:
            return "texto"
        return "mixto"

    def _serializar(self, value: Any) -> Any:
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        return value
