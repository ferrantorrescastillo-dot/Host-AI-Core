from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable
import json
import math
import re
import unicodedata

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


@dataclass
class IngredienteDetectado:
    nombre: str
    cantidad: float | None = None
    unidad: str = ""
    precio_unitario: float | None = None
    coste: float | None = None
    fila_excel: int = 0
    datos_originales: list[Any] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)


@dataclass
class EscandalloDetectado:
    hoja: str
    nombre: str
    fila_inicio: int
    fila_fin: int
    tipo: str = ""
    rendimiento_cantidad: float | None = None
    rendimiento_unidad: str = ""
    coste_total_declarado: float | None = None
    coste_unitario_declarado: float | None = None
    coste_ingredientes_calculado: float | None = None
    ingredientes: list[IngredienteDetectado] = field(default_factory=list)
    confianza: float = 0.0
    avisos: list[str] = field(default_factory=list)
    marcadores: list[str] = field(default_factory=list)

    def a_dict(self) -> dict[str, Any]:
        return asdict(self)


class DetectorEscandallosAntiguosI11:
    """Detector no destructivo para hojas legacy con múltiples fichas técnicas.

    No importa ni modifica datos. Lee el libro con ``data_only=False`` para
    conservar valores y fórmulas, identifica bloques de ficha técnica y devuelve
    una vista estructurada que I1.2 podrá revisar e importar.
    """

    MARCADORES_FICHA = ("FICHA TECNICA", "FICHA TÉCNICA")
    MARCADORES_NOMBRE = ("ARTICULO", "ARTÍCULO", "ELABORACION", "ELABORACIÓN", "RECETA")
    CABECERAS_INGREDIENTES = ("€/UN", "€ / UN", "EUROS/RACION", "€ / RACION", "KG", "BRUTO", "NETO")
    TIPOS = {
        "TAPA", "PLATO", "ELABORADOS", "ELABORADO", "SALSA", "FONDO", "GUARNICION",
        "GUARNICIÓN", "POSTRE", "APERITIVO", "ENTRANTE", "PRIMERO", "SEGUNDO", "BEBIDA",
    }
    UNIDADES = {
        "KG": "kg", "KGS": "kg", "G": "g", "GR": "g", "GRAMOS": "g", "L": "L",
        "LT": "L", "LTS": "L", "LITRO": "L", "LITROS": "L", "UNID": "u", "UNID.": "u",
        "UNIDAD": "u", "UNIDADES": "u", "UD": "u", "U": "u", "RACION": "ración",
        "RACIÓN": "ración", "RACIONES": "ración", "PAX": "pax",
    }

    def __init__(self, precio_anomalo: float = 500.0, tolerancia_coste_pct: float = 5.0):
        self.precio_anomalo = float(precio_anomalo)
        self.tolerancia_coste_pct = float(tolerancia_coste_pct)

    @staticmethod
    def _texto(valor: Any) -> str:
        if valor is None:
            return ""
        return str(valor).strip()

    @classmethod
    def _normalizar(cls, valor: Any) -> str:
        texto = cls._texto(valor).upper()
        texto = "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")
        return re.sub(r"\s+", " ", texto).strip()

    @staticmethod
    def _numero(valor: Any) -> float | None:
        if valor is None or isinstance(valor, bool):
            return None
        if isinstance(valor, (int, float)):
            if isinstance(valor, float) and (math.isnan(valor) or math.isinf(valor)):
                return None
            return float(valor)
        texto = str(valor).strip().replace("€", "").replace(" ", "")
        if not texto or texto.startswith("="):
            return None
        if texto.count(",") == 1 and texto.count(".") == 0:
            texto = texto.replace(",", ".")
        elif texto.count(".") > 1 and texto.count(",") == 1:
            texto = texto.replace(".", "").replace(",", ".")
        try:
            return float(texto)
        except ValueError:
            return None

    @staticmethod
    def _fila_valores(ws, fila: int) -> list[Any]:
        return [ws.cell(fila, col).value for col in range(1, ws.max_column + 1)]

    @classmethod
    def _fila_texto(cls, valores: Iterable[Any]) -> str:
        return " | ".join(cls._texto(v) for v in valores if cls._texto(v))

    @classmethod
    def _es_marcador_ficha(cls, valores: list[Any]) -> bool:
        texto = cls._normalizar(cls._fila_texto(valores))
        return any(cls._normalizar(m) in texto for m in cls.MARCADORES_FICHA)

    @classmethod
    def _es_inicio_receta(cls, valores: list[Any]) -> bool:
        """Detecta filas que presentan una receta, no la cabecera de ingredientes.

        En los libros Boronat suele existir un único marcador ``FICHA TÉCNICA``
        al principio de la hoja y cada ficha posterior empieza directamente con
        ``ARTÍCULO | <nombre>``. Por eso el inicio real se obtiene de estas filas.
        """
        tokens = [cls._normalizar(v) for v in valores]
        no_vacios = [(i, t) for i, t in enumerate(tokens) if t]
        marcadores = {"ARTICULO", "ELABORACION", "RECETA"}
        cabecera_tokens = {"KG", "€/UN", "€ / UN", "BRUTO", "NETO", "%NETO", "%DESP", "SUCIO", "€ / RACION"}
        if sum(1 for _, t in no_vacios if t in cabecera_tokens) >= 2:
            return False
        for idx, token in no_vacios:
            if token not in marcadores:
                continue
            candidatos = [cls._texto(v) for v in valores[idx + 1:] if cls._texto(v)]
            if not candidatos:
                continue
            candidato = cls._normalizar(candidatos[0])
            if candidato and candidato not in cabecera_tokens and candidato not in marcadores:
                return True
        return False

    @classmethod
    def _nombre_desde_fila_inicio(cls, valores: list[Any]) -> str | None:
        norm = [cls._normalizar(v) for v in valores]
        for idx, token in enumerate(norm):
            if token in {"ARTICULO", "ELABORACION", "RECETA"}:
                candidatos = [cls._texto(v) for v in valores[idx + 1:] if cls._texto(v)]
                if candidatos:
                    return candidatos[0]
        return None

    @classmethod
    def _es_cabecera_ingredientes(cls, valores: list[Any]) -> bool:
        tokens = [cls._normalizar(v) for v in valores if cls._texto(v)]
        unido = " | ".join(tokens)
        tiene_articulo = any(t in {"ARTICULO", "INGREDIENTE", "PRODUCTO"} for t in tokens)
        puntos = sum(1 for marcador in cls.CABECERAS_INGREDIENTES if cls._normalizar(marcador) in unido)
        return tiene_articulo and puntos >= 2

    @classmethod
    def _buscar_nombre(cls, ws, fila_ficha: int, limite: int) -> tuple[str, int, list[str]]:
        avisos: list[str] = []
        for fila in range(fila_ficha + 1, min(limite, fila_ficha + 12) + 1):
            valores = cls._fila_valores(ws, fila)
            norm = [cls._normalizar(v) for v in valores]
            for idx, token in enumerate(norm):
                if token in {"ARTICULO", "ELABORACION", "RECETA"}:
                    candidatos = [cls._texto(v) for v in valores[idx + 1:] if cls._texto(v)]
                    if candidatos:
                        nombre = candidatos[0]
                        if len(nombre) >= 2:
                            return nombre, fila, avisos
        # Respaldo: primera fila de texto razonable tras el marcador.
        for fila in range(fila_ficha + 1, min(limite, fila_ficha + 10) + 1):
            textos = [cls._texto(v) for v in cls._fila_valores(ws, fila) if cls._texto(v)]
            textos = [t for t in textos if cls._normalizar(t) not in cls.TIPOS]
            if len(textos) == 1 and not any(x in cls._normalizar(textos[0]) for x in ("COSTE", "FICHA", "UNID")):
                avisos.append("Nombre inferido sin marcador ARTÍCULO/RECETA.")
                return textos[0], fila, avisos
        return f"Escandallo sin nombre (fila {fila_ficha})", fila_ficha, ["No se pudo detectar el nombre del escandallo."]

    @classmethod
    def _buscar_tipo(cls, ws, fila_nombre: int, limite: int) -> str:
        for fila in range(fila_nombre, min(limite, fila_nombre + 4) + 1):
            for valor in cls._fila_valores(ws, fila):
                norm = cls._normalizar(valor)
                if norm in cls.TIPOS:
                    return cls._texto(valor)
        return ""

    @classmethod
    def _buscar_rendimiento_costes(
        cls, ws, fila_nombre: int, limite: int, ws_valores=None
    ) -> tuple[float | None, str, float | None, float | None, list[str]]:
        avisos: list[str] = []
        cantidad = None
        unidad = ""
        coste_total = None
        coste_unitario = None
        for fila in range(fila_nombre, min(limite, fila_nombre + 12) + 1):
            valores = cls._fila_valores(ws, fila)
            valores_num = cls._fila_valores(ws_valores, fila) if ws_valores is not None else valores
            norm = [cls._normalizar(v) for v in valores]
            for idx, token in enumerate(norm):
                if token in cls.UNIDADES and idx + 1 < len(valores_num):
                    posible = cls._numero(valores_num[idx + 1])
                    if posible is not None and posible > 0:
                        unidad = cls.UNIDADES[token]
                        cantidad = posible
                if "COSTE TOTAL" in token:
                    nums = [cls._numero(v) for v in valores_num[idx + 1:] if cls._numero(v) is not None]
                    if nums:
                        coste_total = nums[0]
                if token in {"COSTE PAX", "COSTE RACION", "COSTE POR RACION", "COSTE UNIDAD", "COSTE UNID", "COSTE KG"}:
                    nums = [cls._numero(v) for v in valores_num[idx + 1:] if cls._numero(v) is not None]
                    if nums:
                        coste_unitario = nums[0]
        if cantidad is None:
            avisos.append("No se detectó rendimiento/cantidad base.")
        if coste_total is None:
            avisos.append("No se detectó coste total declarado.")
        return cantidad, unidad, coste_total, coste_unitario, avisos

    @classmethod
    def _encontrar_cabecera(cls, ws, desde: int, hasta: int) -> int | None:
        for fila in range(desde, hasta + 1):
            if cls._es_cabecera_ingredientes(cls._fila_valores(ws, fila)):
                return fila
        return None

    @classmethod
    def _mapear_columnas(cls, valores: list[Any]) -> dict[str, int]:
        mapa: dict[str, int] = {}
        for idx, valor in enumerate(valores):
            token = cls._normalizar(valor)
            if token in {"ARTICULO", "INGREDIENTE", "PRODUCTO"} and "nombre" not in mapa:
                mapa["nombre"] = idx
            elif token in {"KG", "CANTIDAD", "CANT.", "UNIDADES", "UNID", "GR O KG"} and "cantidad" not in mapa:
                mapa["cantidad"] = idx
            elif token in {"€/UN", "€ / UN", "PRECIO KG", "PRECIO", "COSTE UNITARIO"} and "precio" not in mapa:
                mapa["precio"] = idx
            elif token in {"€ / RACION", "EUROS/RACION", "COSTE", "COSTE TOTAL"} and "coste" not in mapa:
                mapa["coste"] = idx
            elif token in {"UNIDAD", "UD", "U.M.", "UM"} and "unidad" not in mapa:
                mapa["unidad"] = idx
        return mapa

    def _leer_ingredientes(self, ws, cabecera: int, fin_bloque: int, ws_valores=None) -> tuple[list[IngredienteDetectado], int, list[str]]:
        valores_cab = self._fila_valores(ws, cabecera)
        mapa = self._mapear_columnas(valores_cab)
        avisos: list[str] = []
        if "nombre" not in mapa:
            return [], cabecera, ["Cabecera detectada sin columna de artículo."]
        ingredientes: list[IngredienteDetectado] = []
        vacias_consecutivas = 0
        ultima_fila = cabecera
        for fila in range(cabecera + 1, fin_bloque + 1):
            valores = self._fila_valores(ws, fila)
            valores_num = self._fila_valores(ws_valores, fila) if ws_valores is not None else valores
            if self._es_marcador_ficha(valores) or self._es_inicio_receta(valores):
                break
            nombre = self._texto(valores[mapa["nombre"]] if mapa["nombre"] < len(valores) else None)
            if not nombre:
                if not any(self._texto(v) for v in valores):
                    vacias_consecutivas += 1
                    if vacias_consecutivas >= 2 and ingredientes:
                        break
                continue
            vacias_consecutivas = 0
            norm_nombre = self._normalizar(nombre)
            if norm_nombre in {"ESCANDALLO", "TOTAL", "COSTE TOTAL"} or "FICHA TECNICA" in norm_nombre:
                break
            cantidad = self._numero(valores_num[mapa["cantidad"]]) if "cantidad" in mapa and mapa["cantidad"] < len(valores_num) else None
            precio = self._numero(valores_num[mapa["precio"]]) if "precio" in mapa and mapa["precio"] < len(valores_num) else None
            coste = self._numero(valores_num[mapa["coste"]]) if "coste" in mapa and mapa["coste"] < len(valores_num) else None
            unidad = self._texto(valores[mapa["unidad"]]) if "unidad" in mapa and mapa["unidad"] < len(valores) else ""
            linea_avisos: list[str] = []
            if cantidad is None:
                linea_avisos.append("Cantidad no detectada.")
            if precio is not None and abs(precio) >= self.precio_anomalo:
                linea_avisos.append(f"Precio unitario anómalo: {precio:g}.")
            if coste is None and cantidad is not None and precio is not None:
                coste = cantidad * precio
            ingredientes.append(IngredienteDetectado(
                nombre=nombre,
                cantidad=cantidad,
                unidad=unidad,
                precio_unitario=precio,
                coste=coste,
                fila_excel=fila,
                datos_originales=valores,
                avisos=linea_avisos,
            ))
            ultima_fila = fila
        if not ingredientes:
            avisos.append("No se detectaron ingredientes dentro del bloque.")
        return ingredientes, ultima_fila, avisos

    def detectar_hoja(self, ws, ws_valores=None) -> list[EscandalloDetectado]:
        inicios = [
            fila for fila in range(1, ws.max_row + 1)
            if self._es_inicio_receta(self._fila_valores(ws, fila))
        ]
        # Compatibilidad con hojas donde cada ficha sí repite FICHA TÉCNICA.
        if not inicios:
            inicios = [
                fila for fila in range(1, ws.max_row + 1)
                if self._es_marcador_ficha(self._fila_valores(ws, fila))
            ]
        resultados: list[EscandalloDetectado] = []
        for pos, fila_inicio in enumerate(inicios):
            limite = (inicios[pos + 1] - 1) if pos + 1 < len(inicios) else ws.max_row
            valores_inicio = self._fila_valores(ws, fila_inicio)
            nombre_directo = self._nombre_desde_fila_inicio(valores_inicio)
            if nombre_directo:
                nombre, fila_nombre, avisos_nombre = nombre_directo, fila_inicio, []
            else:
                nombre, fila_nombre, avisos_nombre = self._buscar_nombre(ws, fila_inicio, limite)
            tipo = self._buscar_tipo(ws, fila_nombre, limite)
            rendimiento, unidad, coste_total, coste_unitario, avisos_meta = self._buscar_rendimiento_costes(
                ws, fila_nombre, limite, ws_valores=ws_valores
            )
            cabecera = self._encontrar_cabecera(ws, fila_nombre, limite)
            avisos = avisos_nombre + avisos_meta
            if cabecera is None:
                ingredientes: list[IngredienteDetectado] = []
                fila_fin = limite
                avisos.append("No se detectó la cabecera de ingredientes.")
            else:
                ingredientes, fila_fin, avisos_ing = self._leer_ingredientes(
                    ws, cabecera, limite, ws_valores=ws_valores
                )
                avisos.extend(avisos_ing)
            costes = [i.coste for i in ingredientes if i.coste is not None]
            coste_calc = round(sum(costes), 6) if costes else None
            if coste_total is None and coste_calc is not None:
                avisos = [a for a in avisos if a != "No se detectó coste total declarado."]
                avisos.append("Coste total no declarado; se usa la suma de ingredientes como referencia.")
            if coste_total is not None and coste_calc is not None and coste_total != 0:
                diferencia_pct = abs(coste_calc - coste_total) / abs(coste_total) * 100
                if diferencia_pct > self.tolerancia_coste_pct:
                    avisos.append(
                        f"El coste calculado ({coste_calc:.4f}) difiere del declarado "
                        f"({coste_total:.4f}) en {diferencia_pct:.1f}%."
                    )
            if coste_unitario is None and coste_total is not None and rendimiento:
                coste_unitario = coste_total / rendimiento
            confianza = 25.0
            confianza += 20.0 if nombre and "sin nombre" not in nombre.lower() else 0.0
            confianza += 15.0 if rendimiento is not None else 0.0
            confianza += 15.0 if (coste_total is not None or coste_calc is not None) else 0.0
            confianza += 25.0 if ingredientes else 0.0
            confianza = min(100.0, confianza)
            resultados.append(EscandalloDetectado(
                hoja=ws.title,
                nombre=nombre,
                fila_inicio=fila_inicio,
                fila_fin=fila_fin,
                tipo=tipo,
                rendimiento_cantidad=rendimiento,
                rendimiento_unidad=unidad,
                coste_total_declarado=coste_total,
                coste_unitario_declarado=coste_unitario,
                coste_ingredientes_calculado=coste_calc,
                ingredientes=ingredientes,
                confianza=confianza,
                avisos=avisos,
                marcadores=[f"INICIO_RECETA:{fila_inicio}", f"NOMBRE:{fila_nombre}"] + ([f"CABECERA:{cabecera}"] if cabecera else []),
            ))
        return resultados

    def analizar(self, ruta_archivo: str | Path, hojas: list[str] | None = None, exportar_json: bool = True) -> dict[str, Any]:
        ruta = Path(str(ruta_archivo).strip().strip('"')).expanduser()
        if not ruta.exists() or not ruta.is_file():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta}")
        if ruta.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError("I1.1 admite archivos .xlsx o .xlsm.")
        wb = load_workbook(ruta, data_only=False, read_only=False)
        # Segundo libro para leer el valor cacheado de fórmulas de costes.
        wb_valores = load_workbook(ruta, data_only=True, read_only=False)
        seleccion = hojas or list(wb.sheetnames)
        inexistentes = [h for h in seleccion if h not in wb.sheetnames]
        if inexistentes:
            raise ValueError(f"No existen estas hojas: {', '.join(inexistentes)}")
        escandallos: list[EscandalloDetectado] = []
        hojas_analizadas: list[dict[str, Any]] = []
        for nombre_hoja in seleccion:
            ws = wb[nombre_hoja]
            ws_valores = wb_valores[nombre_hoja]
            detectados = self.detectar_hoja(ws, ws_valores=ws_valores)
            escandallos.extend(detectados)
            hojas_analizadas.append({
                "hoja": nombre_hoja,
                "filas": ws.max_row,
                "columnas": ws.max_column,
                "escandallos_detectados": len(detectados),
            })
        total_ingredientes = sum(len(e.ingredientes) for e in escandallos)
        total_avisos = sum(len(e.avisos) + sum(len(i.avisos) for i in e.ingredientes) for e in escandallos)
        resultado = {
            "ok": True,
            "modo": "solo_deteccion_no_importa",
            "archivo": str(ruta.resolve()),
            "hojas": hojas_analizadas,
            "resumen": {
                "hojas_analizadas": len(hojas_analizadas),
                "escandallos_detectados": len(escandallos),
                "ingredientes_detectados": total_ingredientes,
                "avisos": total_avisos,
                "confianza_media": round(sum(e.confianza for e in escandallos) / len(escandallos), 2) if escandallos else 0.0,
            },
            "escandallos": [e.a_dict() for e in escandallos],
        }
        if exportar_json:
            salida = ruta.with_name(f"{ruta.stem}_I1.1_DETECCION.json")
            salida.write_text(json.dumps(resultado, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
            resultado["archivo_json"] = str(salida.resolve())
        return resultado


def formatear_resumen_i11(resultado: dict[str, Any], max_recetas: int = 50) -> str:
    r = resultado.get("resumen", {})
    lineas = [
        "=" * 72,
        "I1.1.1 — DETECTOR MULTI-ESCANDALLOS (SOLO ANÁLISIS)",
        "=" * 72,
        f"Hojas analizadas: {r.get('hojas_analizadas', 0)}",
        f"Escandallos detectados: {r.get('escandallos_detectados', 0)}",
        f"Ingredientes detectados: {r.get('ingredientes_detectados', 0)}",
        f"Avisos: {r.get('avisos', 0)}",
        f"Confianza media: {r.get('confianza_media', 0)}%",
        "",
    ]
    for idx, e in enumerate(resultado.get("escandallos", [])[:max_recetas], start=1):
        rendimiento = "-"
        if e.get("rendimiento_cantidad") is not None:
            rendimiento = f"{e['rendimiento_cantidad']:g} {e.get('rendimiento_unidad') or ''}".strip()
        coste = e.get("coste_total_declarado")
        coste_txt = f"{coste:.4f} €" if isinstance(coste, (int, float)) else "-"
        lineas.append(
            f"{idx}. [{e.get('hoja')}] {e.get('nombre')} | fila {e.get('fila_inicio')} | "
            f"{len(e.get('ingredientes', []))} ingrediente(s) | rendimiento {rendimiento} | "
            f"coste {coste_txt} | confianza {e.get('confianza', 0):.0f}%"
        )
        for aviso in e.get("avisos", []):
            lineas.append(f"   AVISO: {aviso}")
        for ing in e.get("ingredientes", []):
            for aviso in ing.get("avisos", []):
                lineas.append(f"   AVISO ingrediente '{ing.get('nombre')}': {aviso}")
    if len(resultado.get("escandallos", [])) > max_recetas:
        lineas.append(f"... y {len(resultado['escandallos']) - max_recetas} escandallos más en el JSON.")
    if resultado.get("archivo_json"):
        lineas.extend(["", f"Informe JSON: {resultado['archivo_json']}"])
    lineas.append("No se ha importado ni modificado ningún dato de Host AI.")
    lineas.append("=" * 72)
    return "\n".join(lineas)
