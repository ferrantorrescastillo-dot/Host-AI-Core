from __future__ import annotations

import json
import re
import tempfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from SERVICIOS.asistente_resolucion_incidencias_601 import AsistenteResolucionIncidencias601
from SERVICIOS.biblioteca_recetas_601 import RepositorioBibliotecaRecetas601
from SERVICIOS.repositorio_productos_maestro_601 import RepositorioProductosMaestro601


ESTADO_INCIDENCIA_PENDIENTE = "PENDIENTE"
ESTADO_INCIDENCIA_RESUELTA = "RESUELTA"
ESTADO_IMPORTACION_COMPLETADA = "COMPLETADA"
ESTADO_IMPORTACION_CANCELADA = "CANCELADA"
ESTADO_IMPORTACION_CON_INCIDENCIAS = "CON_INCIDENCIAS"

TIPO_PRODUCTO_INEXISTENTE = "PRODUCTO_INEXISTENTE"
TIPO_PRODUCTO_SIN_PRECIO = "PRODUCTO_SIN_PRECIO"
TIPO_PRODUCTO_DUPLICADO = "PRODUCTO_DUPLICADO"
TIPO_INGREDIENTE_SIN_CANTIDAD = "INGREDIENTE_SIN_CANTIDAD"
TIPO_PROVEEDOR_INEXISTENTE = "PROVEEDOR_INEXISTENTE"
TIPO_RECETA_INCOMPLETA = "RECETA_INCOMPLETA"
TIPO_UNIDAD_DESCONOCIDA = "UNIDAD_DESCONOCIDA"
TIPO_RECETA_INEXISTENTE = "RECETA_INEXISTENTE"
TIPO_CANTIDAD_AUSENTE = "CANTIDAD_AUSENTE"
TIPO_MERMA_INVALIDA = "MERMA_INVALIDA"
TIPO_COSTE_IMPORTADO_DISCREPANTE = "COSTE_IMPORTADO_DISCREPANTE"
TIPO_MENU_INEXISTENTE = "MENU_INEXISTENTE"
TIPO_PLATO_INEXISTENTE = "PLATO_INEXISTENTE"
TIPO_ESCANDALLO_INEXISTENTE = "ESCANDALLO_INEXISTENTE"
TIPO_PRODUCTO_SIN_PRECIO_MENU = "PRODUCTO_SIN_PRECIO_MENU"
TIPO_REFERENCIA_DUPLICADA_MENU = "REFERENCIA_DUPLICADA_MENU"


@dataclass
class DocumentoImportacion601:
    origen: str
    etiqueta_origen: str
    texto: str = ""
    archivos: list[str] = field(default_factory=list)
    advertencias: list[str] = field(default_factory=list)
    tipo_contenido: str = "RECETAS"
    payload: dict[str, Any] = field(default_factory=dict)


class RepositorioCentroImportacion601:
    """Persistencia del historial e incidencias del centro de importación 6.0.1."""

    VERSION = "6.0.1"

    def __init__(self, base_dir: Path):
        self.base_dir = Path(base_dir).resolve()
        self.path = self.base_dir / "DATOS" / "db" / "centro_importacion_601.json"
        self.repo_catalogo = RepositorioProductosMaestro601(self.base_dir)
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._asegurar_archivo()

    def _asegurar_archivo(self) -> None:
        if self.path.exists():
            return
        self._guardar_payload({"historial": [], "incidencias": []})

    def _leer_payload(self) -> dict[str, Any]:
        try:
            data = json.loads(self.path.read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError, OSError):
            data = {}
        historial = data.get("historial") if isinstance(data, dict) else None
        incidencias = data.get("incidencias") if isinstance(data, dict) else None
        return {
            "historial": historial if isinstance(historial, list) else [],
            "incidencias": incidencias if isinstance(incidencias, list) else [],
        }

    def _guardar_payload(self, payload: dict[str, Any]) -> None:
        data = {
            "version_modelo": self.VERSION,
            "actualizado_en": datetime.now().isoformat(timespec="seconds"),
            "historial": payload.get("historial", []),
            "incidencias": payload.get("incidencias", []),
        }
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", delete=False, dir=self.path.parent, suffix=".tmp") as tmp:
            json.dump(data, tmp, ensure_ascii=False, indent=2)
            tmp.flush()
            temp_path = Path(tmp.name)
        temp_path.replace(self.path)

    def _siguiente_id(self, prefijo: str, registros: list[dict[str, Any]], clave: str = "id") -> str:
        ultimo = 0
        for item in registros:
            rid = str(item.get(clave) or "")
            if rid.startswith(prefijo):
                try:
                    ultimo = max(ultimo, int(rid.split("-")[-1]))
                except ValueError:
                    continue
        return f"{prefijo}-{ultimo + 1:06d}"

    def registrar_importacion(self, registro: dict[str, Any]) -> dict[str, Any]:
        payload = self._leer_payload()
        historial = list(payload["historial"])
        rid = self._siguiente_id("IMP601", historial)
        fila = {
            "id": rid,
            "fecha": datetime.now().isoformat(timespec="seconds"),
            "origen": str(registro.get("origen") or ""),
            "numero_recetas": int(registro.get("numero_recetas") or 0),
            "numero_incidencias": int(registro.get("numero_incidencias") or 0),
            "estado": str(registro.get("estado") or ESTADO_IMPORTACION_CON_INCIDENCIAS),
        }
        historial.append(fila)
        payload["historial"] = historial
        self._guardar_payload(payload)
        return fila

    def registrar_incidencias(self, importacion_id: str, incidencias: list[dict[str, Any]]) -> list[dict[str, Any]]:
        if not incidencias:
            return []
        payload = self._leer_payload()
        todas = list(payload["incidencias"])
        nuevas: list[dict[str, Any]] = []
        for incidencia in incidencias:
            iid = self._siguiente_id("INC601", todas + nuevas)
            fila = {
                "id": iid,
                "importacion_id": importacion_id,
                "fecha": datetime.now().isoformat(timespec="seconds"),
                "tipo": str(incidencia.get("tipo") or "INCIDENCIA"),
                "detalle": str(incidencia.get("detalle") or ""),
                "estado": ESTADO_INCIDENCIA_PENDIENTE,
            }
            nuevas.append(fila)
        payload["incidencias"] = todas + nuevas
        self._guardar_payload(payload)
        return nuevas

    def listar_historial(self) -> list[dict[str, Any]]:
        return list(self._leer_payload()["historial"])

    def incidencias_pendientes(self) -> list[dict[str, Any]]:
        incidencias = list(self._leer_payload()["incidencias"])
        return [i for i in incidencias if str(i.get("estado") or "") == ESTADO_INCIDENCIA_PENDIENTE]

    def listar_productos(self) -> list[dict[str, Any]]:
        productos = self.repo_catalogo.listar_productos(incluir_archivados=True)
        salida: list[dict[str, Any]] = []
        for p in productos:
            item = dict(p)
            item.setdefault("id", item.get("codigo"))
            salida.append(item)
        return salida

    def listar_proveedores(self) -> list[dict[str, Any]]:
        proveedores = self.repo_catalogo.listar_proveedores()
        salida: list[dict[str, Any]] = []
        for p in proveedores:
            item = dict(p)
            item.setdefault("id", item.get("codigo"))
            salida.append(item)
        return salida

    def _siguiente_id_simple(self, prefijo: str, registros: list[dict[str, Any]]) -> str:
        return self._siguiente_id(prefijo, registros, clave="id")

    def crear_producto(self, producto: dict[str, Any]) -> dict[str, Any]:
        data = dict(producto)
        if "cantidad_por_envase" in data and "cantidad_formato" not in data:
            data["cantidad_formato"] = data.get("cantidad_por_envase")
        if "unidad_receta" in data and "unidad_recetas" not in data:
            data["unidad_recetas"] = data.get("unidad_receta")
        if "unidad_compra" in data and "unidad_base" not in data:
            data["unidad_base"] = data.get("unidad_compra")
        creado = self.repo_catalogo.crear_producto(data)
        creado = dict(creado)
        creado.setdefault("id", creado.get("codigo"))
        return creado

    def actualizar_producto(self, producto_id: str, cambios: dict[str, Any]) -> dict[str, Any] | None:
        data = dict(cambios)
        if "cantidad_por_envase" in data and "cantidad_formato" not in data:
            data["cantidad_formato"] = data.get("cantidad_por_envase")
        if "unidad_receta" in data and "unidad_recetas" not in data:
            data["unidad_recetas"] = data.get("unidad_receta")
        if "unidad_compra" in data and "unidad_base" not in data:
            data["unidad_base"] = data.get("unidad_compra")
        codigo = str(producto_id or "").strip()
        try:
            actualizado = self.repo_catalogo.editar_producto(codigo, data)
            actualizado = dict(actualizado)
            actualizado.setdefault("id", actualizado.get("codigo"))
            return actualizado
        except ValueError:
            return None

    def asegurar_producto_basico(self, nombre: str) -> dict[str, Any]:
        encontrados = self.repo_catalogo.buscar_productos({"nombre": str(nombre or "")})
        if encontrados:
            out = dict(encontrados[0])
            out.setdefault("id", out.get("codigo"))
            return out
        creado = self.repo_catalogo.crear_producto({"nombre": nombre})
        creado = dict(creado)
        creado.setdefault("id", creado.get("codigo"))
        return creado

    def crear_proveedor(self, proveedor: dict[str, Any]) -> dict[str, Any]:
        nuevo = self.repo_catalogo.crear_proveedor(proveedor)
        out = dict(nuevo)
        out.setdefault("id", out.get("codigo"))
        return out


class LectorBase601:
    origen = "DESCONOCIDO"

    def leer(self, entrada: str) -> DocumentoImportacion601:
        raise NotImplementedError

    @staticmethod
    def _normalizar_ruta(entrada: str) -> str:
        return str(entrada or "").strip().strip('"')


class LectorWord601(LectorBase601):
    origen = "WORD"

    def leer(self, entrada: str) -> DocumentoImportacion601:
        ruta = self._normalizar_ruta(entrada)
        advertencias: list[str] = []
        if not ruta.lower().endswith(".docx"):
            advertencias.append("El archivo no tiene extensión .docx")
        return DocumentoImportacion601(
            origen=self.origen,
            etiqueta_origen=ruta,
            archivos=[ruta] if ruta else [],
            advertencias=advertencias + ["Lector Word pendiente de implementación."],
        )


class LectorExcel601(LectorBase601):
    origen = "EXCEL"

    def leer(self, entrada: str) -> DocumentoImportacion601:
        ruta = self._normalizar_ruta(entrada)
        advertencias: list[str] = []
        if not ruta.lower().endswith(".xlsx"):
            advertencias.append("El archivo no tiene extensión .xlsx")
        texto = ""
        try:
            from openpyxl import load_workbook

            wb = load_workbook(filename=ruta, data_only=True)
            ws = wb.active
            lineas: list[str] = []
            for fila in ws.iter_rows(values_only=True):
                celdas = list(fila or [])
                if not celdas or all(v is None or str(v).strip() == "" for v in celdas):
                    continue
                primero = str(celdas[0]).strip() if len(celdas) > 0 and celdas[0] is not None else ""
                segundo = str(celdas[1]).strip() if len(celdas) > 1 and celdas[1] is not None else ""
                if segundo != "" or (len(celdas) > 1 and primero):
                    lineas.append(f"{primero}: {segundo}")
                else:
                    extras = [str(v).strip() for v in celdas if v is not None and str(v).strip()]
                    if not extras:
                        continue
                    lineas.append(extras[0])
            texto = "\n".join(lineas)
            if not texto:
                advertencias.append("Excel sin contenido legible para importación de recetas.")
        except Exception as exc:
            advertencias.append(f"No se pudo leer Excel: {exc}")
        return DocumentoImportacion601(
            origen=self.origen,
            etiqueta_origen=ruta,
            texto=texto,
            archivos=[ruta] if ruta else [],
            advertencias=advertencias,
        )


class LectorPdf601(LectorBase601):
    origen = "PDF"

    def leer(self, entrada: str) -> DocumentoImportacion601:
        ruta = self._normalizar_ruta(entrada)
        advertencias: list[str] = []
        if not ruta.lower().endswith(".pdf"):
            advertencias.append("El archivo no tiene extensión .pdf")
        return DocumentoImportacion601(
            origen=self.origen,
            etiqueta_origen=ruta,
            archivos=[ruta] if ruta else [],
            advertencias=advertencias + ["Lector PDF pendiente de implementación."],
        )


class LectorImagen601(LectorBase601):
    origen = "IMAGEN"
    EXTENSIONES = {".jpg", ".jpeg", ".png", ".heic"}

    def leer(self, entrada: str) -> DocumentoImportacion601:
        texto = self._normalizar_ruta(entrada)
        rutas = [r.strip().strip('"') for r in texto.split(",") if r.strip()]
        advertencias: list[str] = []
        for ruta in rutas:
            if Path(ruta).suffix.lower() not in self.EXTENSIONES:
                advertencias.append(f"Formato no previsto para imagen: {ruta}")
        return DocumentoImportacion601(
            origen=self.origen,
            etiqueta_origen=", ".join(rutas),
            archivos=rutas,
            advertencias=advertencias + ["Extracción desde imágenes pendiente de implementación."],
        )


class LectorTexto601(LectorBase601):
    origen = "TEXTO"

    def leer(self, entrada: str) -> DocumentoImportacion601:
        return DocumentoImportacion601(
            origen=self.origen,
            etiqueta_origen="texto_pegado",
            texto=str(entrada or ""),
            advertencias=[],
        )


class FlujoImportacionUnificado601:
    """Motor único de flujo de importación para todas las fuentes."""

    _UNIDADES_VALIDAS = {"kg", "g", "l", "ml", "u", "ud", "uds", "unidad", "unidades"}
    _HEADERS_EQUIV = {
        "nombre_escandallo": ["escandallo", "plato", "elaboracion", "elaboración", "receta"],
        "receta": ["receta", "elaboracion", "elaboración", "plato"],
        "raciones": ["raciones", "rendimiento", "comensales"],
        "ingrediente": ["ingrediente", "producto", "articulo", "artículo"],
        "producto": ["producto", "articulo", "artículo", "ingrediente"],
        "cantidad": ["cantidad", "peso"],
        "unidad": ["unidad", "um"],
        "merma": ["merma"],
        "precio": ["precio", "precio compra", "coste unitario", "precio_unitario"],
        "proveedor": ["proveedor"],
        "coste_linea": ["coste linea", "coste línea", "importe", "coste_linea"],
        "precio_venta": ["precio venta", "pvp", "venta"],
        "observaciones": ["observaciones", "observacion", "notas", "nota"],
    }
    _CAMPOS_BASE_ESC = [
        "nombre_escandallo",
        "receta",
        "raciones",
        "ingrediente",
        "producto",
        "cantidad",
        "unidad",
        "merma",
        "precio",
        "proveedor",
        "coste_linea",
        "precio_venta",
        "observaciones",
        "origen",
        "confianza",
        "campos_dudosos",
        "fila_o_posicion_origen",
    ]

    def __init__(self, base_dir: Path):
        self.base_dir = Path(base_dir).resolve()
        self.repo_importaciones = RepositorioCentroImportacion601(self.base_dir)
        self.repo_catalogo = RepositorioProductosMaestro601(self.base_dir)
        self.repo_recetas = RepositorioBibliotecaRecetas601(self.base_dir)
        from SERVICIOS.biblioteca_menus_601 import BibliotecaMenus601

        self.servicio_menus = BibliotecaMenus601(self.base_dir)
        from SERVICIOS.biblioteca_escandallos_601 import BibliotecaEscandallos601

        self.servicio_escandallos = BibliotecaEscandallos601(self.base_dir, repo_incidencias=self.repo_importaciones)
        self.asistente_incidencias = AsistenteResolucionIncidencias601(self.repo_importaciones, self.base_dir)

    def ejecutar(self, documento: DocumentoImportacion601) -> dict[str, Any]:
        if str(documento.tipo_contenido or "").upper() == "ESCANDALLOS":
            return self._ejecutar_importacion_escandallos(documento)
        if str(documento.tipo_contenido or "").upper() == "MENUS":
            return self._ejecutar_importacion_menus(documento)

        analisis = self._analizar_documento(documento)
        recetas_detectadas = self._detectar_recetas(analisis)
        ingredientes_detectados = self._detectar_ingredientes(recetas_detectadas)
        existentes, nuevos = self._buscar_productos(recetas_detectadas)
        incidencias = self._detectar_incidencias(
            documento=documento,
            recetas=recetas_detectadas,
            ingredientes=ingredientes_detectados,
            productos_nuevos=nuevos,
        )
        resumen = self._resumen(
            recetas=recetas_detectadas,
            ingredientes=ingredientes_detectados,
            nuevos=nuevos,
            incidencias=incidencias,
        )
        return {
            "documento": documento,
            "analisis": analisis,
            "recetas": recetas_detectadas,
            "ingredientes": ingredientes_detectados,
            "productos_existentes": existentes,
            "productos_nuevos": nuevos,
            "incidencias": incidencias,
            "resumen": resumen,
        }

    @staticmethod
    def _norm(txt: Any) -> str:
        return " ".join(str(txt or "").strip().lower().split())

    @staticmethod
    def _to_float(value: Any) -> float | None:
        if value in (None, ""):
            return None
        if isinstance(value, (int, float)):
            return float(value)
        txt = str(value).strip().replace("€", "").replace(" ", "")
        if "," in txt and "." in txt:
            txt = txt.replace(".", "").replace(",", ".")
        else:
            txt = txt.replace(",", ".")
        try:
            return float(txt)
        except ValueError:
            return None

    @classmethod
    def _split_tabular_line(cls, line: str) -> list[str]:
        text = str(line or "").strip()
        if not text:
            return []
        if "\t" in text:
            return [c.strip() for c in text.split("\t")]
        if ";" in text:
            return [c.strip() for c in text.split(";")]
        if "|" in text:
            return [c.strip() for c in text.split("|")]
        return [text]

    @classmethod
    def _detectar_mapeo_headers(cls, headers: list[str]) -> dict[str, int]:
        mapping: dict[str, int] = {}
        normalized = [cls._norm(h) for h in headers]
        for canon, variantes in cls._HEADERS_EQUIV.items():
            for idx, h in enumerate(normalized):
                if h in {cls._norm(v) for v in variantes}:
                    mapping[canon] = idx
                    break
        return mapping

    def inspeccionar_excel_escandallos(self, ruta_excel: str) -> dict[str, Any]:
        ruta = str(ruta_excel or "").strip().strip('"')
        hojas: list[str] = []
        mapeos: dict[str, dict[str, int]] = {}
        advertencias: list[str] = []
        try:
            from openpyxl import load_workbook

            wb = load_workbook(filename=ruta, data_only=True)
            hojas = list(wb.sheetnames)
            for hoja in hojas:
                ws = wb[hoja]
                encabezado: list[str] = []
                for fila in ws.iter_rows(values_only=True):
                    celdas = [str(v).strip() for v in list(fila or [])]
                    if any(celdas):
                        encabezado = celdas
                        break
                mapeos[hoja] = self._detectar_mapeo_headers(encabezado)
        except Exception as exc:
            advertencias.append(f"No se pudo inspeccionar Excel: {exc}")
        return {"ruta": ruta, "hojas": hojas, "mapeos_sugeridos": mapeos, "advertencias": advertencias}

    def crear_documento_escandallos_desde_texto(self, texto: str, etiqueta: str = "texto_pegado") -> DocumentoImportacion601:
        rows = [l for l in str(texto or "").splitlines() if l.strip()]
        if not rows:
            return DocumentoImportacion601(origen="TEXTO", etiqueta_origen=etiqueta, texto=texto, tipo_contenido="ESCANDALLOS", payload={"lineas": []})

        parsed = [self._split_tabular_line(r) for r in rows]
        header_idx = 0
        mapping = self._detectar_mapeo_headers(parsed[0])
        if not mapping and len(parsed) > 1:
            mapping = self._detectar_mapeo_headers(parsed[1])
            if mapping:
                header_idx = 1

        if not mapping:
            mapping = {
                "nombre_escandallo": 0,
                "ingrediente": 1,
                "cantidad": 2,
                "unidad": 3,
                "merma": 4,
                "precio": 5,
                "coste_linea": 6,
                "raciones": 7,
                "precio_venta": 8,
                "observaciones": 9,
            }

        lineas_raw: list[dict[str, Any]] = []
        for idx, cols in enumerate(parsed):
            if idx == header_idx:
                continue
            if not cols or all(not str(x).strip() for x in cols):
                continue
            lineas_raw.append({
                "hoja": "texto_pegado",
                "fila": idx + 1,
                "columnas": cols,
            })

        return DocumentoImportacion601(
            origen="TEXTO",
            etiqueta_origen=etiqueta,
            texto=str(texto or ""),
            tipo_contenido="ESCANDALLOS",
            payload={"lineas": lineas_raw, "mapping": mapping, "origen": "texto", "hojas": ["texto_pegado"]},
        )

    def crear_documento_escandallos_desde_excel(
        self,
        ruta_excel: str,
        *,
        hojas: list[str] | None = None,
        mapping_forzado: dict[str, str | int] | None = None,
    ) -> DocumentoImportacion601:
        ruta = str(ruta_excel or "").strip().strip('"')
        advertencias: list[str] = []
        lineas_raw: list[dict[str, Any]] = []
        mapping_global: dict[str, int] = {}
        hojas_usadas: list[str] = []
        try:
            from openpyxl import load_workbook

            wb = load_workbook(filename=ruta, data_only=True)
            disponibles = list(wb.sheetnames)
            seleccion = list(hojas or disponibles)
            for hoja in seleccion:
                if hoja not in disponibles:
                    advertencias.append(f"Hoja no encontrada: {hoja}")
                    continue
                ws = wb[hoja]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue

                header_row_idx = -1
                header_cells: list[str] = []
                for ridx, fila in enumerate(rows):
                    celdas = [str(v).strip() if v is not None else "" for v in list(fila or [])]
                    if any(celdas):
                        header_row_idx = ridx
                        header_cells = celdas
                        break
                if header_row_idx < 0:
                    continue

                mapping = self._detectar_mapeo_headers(header_cells)
                if mapping_forzado:
                    for canon, pos in mapping_forzado.items():
                        if isinstance(pos, int):
                            mapping[canon] = pos
                        else:
                            txt = self._norm(pos)
                            for i, h in enumerate([self._norm(hh) for hh in header_cells]):
                                if h == txt:
                                    mapping[canon] = i
                                    break

                if not mapping_global:
                    mapping_global = dict(mapping)

                for ridx, fila in enumerate(rows[header_row_idx + 1 :], start=header_row_idx + 2):
                    celdas = [str(v).strip() if v is not None else "" for v in list(fila or [])]
                    if not celdas or all(not c for c in celdas):
                        continue
                    lineas_raw.append({"hoja": hoja, "fila": ridx, "columnas": celdas})
                hojas_usadas.append(hoja)
        except Exception as exc:
            advertencias.append(f"No se pudo leer Excel: {exc}")

        return DocumentoImportacion601(
            origen="EXCEL",
            etiqueta_origen=ruta,
            archivos=[ruta] if ruta else [],
            advertencias=advertencias,
            tipo_contenido="ESCANDALLOS",
            payload={"lineas": lineas_raw, "mapping": mapping_global, "origen": "excel", "hojas": hojas_usadas},
        )

    def crear_documento_menus_desde_texto(self, texto: str, etiqueta: str = "texto_pegado") -> DocumentoImportacion601:
        contenido = str(texto or "")
        bloques = [b.strip() for b in re.split(r"\n\s*\n", contenido) if b.strip()]
        payload_menus: list[dict[str, Any]] = []
        actual: dict[str, Any] | None = None
        for linea in contenido.splitlines():
            l = str(linea or "").strip()
            if not l:
                continue
            if l.lower().startswith("menu:"):
                if actual:
                    payload_menus.append(actual)
                nombre = l.split(":", 1)[1].strip() or f"Menu importado {len(payload_menus) + 1}"
                actual = {
                    "nombre": nombre,
                    "tipo": "",
                    "comensales": 0,
                    "precio_venta_comensal": 0,
                    "composicion": {
                        "Aperitivos": [],
                        "Entrantes": [],
                        "Principales": [],
                        "Postres": [],
                        "Bodega": [],
                        "Extras": [],
                    },
                    "fuente": etiqueta,
                }
                continue
            if actual is None:
                actual = {
                    "nombre": f"Menu importado {len(payload_menus) + 1}",
                    "tipo": "",
                    "comensales": 0,
                    "precio_venta_comensal": 0,
                    "composicion": {
                        "Aperitivos": [],
                        "Entrantes": [],
                        "Principales": [],
                        "Postres": [],
                        "Bodega": [],
                        "Extras": [],
                    },
                    "fuente": etiqueta,
                }

            if l.lower().startswith("tipo:"):
                actual["tipo"] = l.split(":", 1)[1].strip()
            elif l.lower().startswith("comensales:"):
                actual["comensales"] = self._to_float(l.split(":", 1)[1].strip())
            elif l.lower().startswith("precio:"):
                actual["precio_venta_comensal"] = self._to_float(l.split(":", 1)[1].strip())
            else:
                m = re.match(r"^(aperitivos|entrantes|principales|postres|bodega|extras)\s*[:|-]\s*(.+)$", l, re.IGNORECASE)
                if m:
                    seccion = m.group(1).capitalize()
                    ref = m.group(2).strip()
                    actual["composicion"].setdefault(seccion, []).append(
                        {"tipo_referencia": "ESCANDALLO", "referencia": ref, "cantidad": 1}
                    )
                else:
                    actual["composicion"]["Principales"].append(
                        {"tipo_referencia": "ESCANDALLO", "referencia": l, "cantidad": 1}
                    )
        if actual:
            payload_menus.append(actual)

        if not payload_menus and bloques:
            for i, bloque in enumerate(bloques, start=1):
                payload_menus.append(
                    {
                        "nombre": f"Menu importado {i}",
                        "tipo": "",
                        "comensales": 0,
                        "precio_venta_comensal": 0,
                        "composicion": {
                            "Aperitivos": [],
                            "Entrantes": [],
                            "Principales": [
                                {"tipo_referencia": "ESCANDALLO", "referencia": ln.strip(), "cantidad": 1}
                                for ln in bloque.splitlines()
                                if ln.strip()
                            ],
                            "Postres": [],
                            "Bodega": [],
                            "Extras": [],
                        },
                        "fuente": etiqueta,
                    }
                )

        return DocumentoImportacion601(
            origen="TEXTO",
            etiqueta_origen=etiqueta,
            texto=contenido,
            tipo_contenido="MENUS",
            payload={"menus": payload_menus, "origen": "texto"},
        )

    def crear_documento_menus_desde_excel(self, ruta_excel: str) -> DocumentoImportacion601:
        ruta = str(ruta_excel or "").strip().strip('"')
        advertencias: list[str] = []
        menus: list[dict[str, Any]] = []
        try:
            from openpyxl import load_workbook

            wb = load_workbook(filename=ruta, data_only=True)
            ws = wb.active
            headers: list[str] = []
            for fila in ws.iter_rows(values_only=True):
                celdas = [str(v).strip() if v is not None else "" for v in list(fila or [])]
                if any(celdas):
                    headers = celdas
                    break

            idx: dict[str, int] = {}
            for i, h in enumerate(headers):
                n = self._norm(h)
                if n in {"menu", "nombre", "nombre_menu"}:
                    idx["nombre"] = i
                elif n in {"tipo"}:
                    idx["tipo"] = i
                elif n in {"comensales", "raciones"}:
                    idx["comensales"] = i
                elif n in {"precio", "precio_comensal", "pvp"}:
                    idx["precio"] = i
                elif n in {"aperitivos", "entrantes", "principales", "postres", "bodega", "extras"}:
                    idx[n.capitalize()] = i

            for fila in ws.iter_rows(min_row=2, values_only=True):
                celdas = [str(v).strip() if v is not None else "" for v in list(fila or [])]
                if not any(celdas):
                    continue
                nombre = celdas[idx["nombre"]] if "nombre" in idx and idx["nombre"] < len(celdas) else ""
                if not nombre:
                    continue
                menu = {
                    "nombre": nombre,
                    "tipo": celdas[idx["tipo"]] if "tipo" in idx and idx["tipo"] < len(celdas) else "",
                    "comensales": self._to_float(celdas[idx["comensales"]]) if "comensales" in idx and idx["comensales"] < len(celdas) else 0,
                    "precio_venta_comensal": self._to_float(celdas[idx["precio"]]) if "precio" in idx and idx["precio"] < len(celdas) else 0,
                    "composicion": {
                        "Aperitivos": [],
                        "Entrantes": [],
                        "Principales": [],
                        "Postres": [],
                        "Bodega": [],
                        "Extras": [],
                    },
                    "fuente": ruta,
                }
                for sec in ["Aperitivos", "Entrantes", "Principales", "Postres", "Bodega", "Extras"]:
                    pos = idx.get(sec)
                    if isinstance(pos, int) and pos < len(celdas):
                        valor = celdas[pos]
                        if valor:
                            for token in [x.strip() for x in valor.split("|") if x.strip()]:
                                menu["composicion"][sec].append(
                                    {"tipo_referencia": "ESCANDALLO", "referencia": token, "cantidad": 1}
                                )
                menus.append(menu)
        except Exception as exc:
            advertencias.append(f"No se pudo leer Excel: {exc}")

        return DocumentoImportacion601(
            origen="EXCEL",
            etiqueta_origen=ruta,
            archivos=[ruta] if ruta else [],
            advertencias=advertencias,
            tipo_contenido="MENUS",
            payload={"menus": menus, "origen": "excel"},
        )

    def _crear_modelo_intermedio_menus(self, documento: DocumentoImportacion601) -> list[dict[str, Any]]:
        return list((documento.payload or {}).get("menus") or [])

    def _detectar_incidencias_menus(self, menus: list[dict[str, Any]]) -> list[dict[str, Any]]:
        incidencias: list[dict[str, Any]] = []
        nombres: dict[str, int] = {}
        for i, menu in enumerate(menus):
            nombre = str(menu.get("nombre") or "").strip()
            if not nombre:
                incidencias.append({
                    "tipo": TIPO_MENU_INEXISTENTE,
                    "detalle": f"Menú sin nombre en bloque {i + 1}",
                    "menu_index": i,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })
                continue

            clave = self._norm(nombre)
            nombres[clave] = nombres.get(clave, 0) + 1
            comp = dict(menu.get("composicion") or {})
            total_refs = 0
            refs_vistas: set[tuple[str, str]] = set()
            for seccion, items in comp.items():
                for item in list(items or []):
                    tipo_ref = str(item.get("tipo_referencia") or "ESCANDALLO").upper()
                    referencia = str(item.get("referencia") or "").strip()
                    if not referencia:
                        continue
                    total_refs += 1
                    k = (self._norm(seccion), self._norm(referencia))
                    if k in refs_vistas:
                        incidencias.append({
                            "tipo": TIPO_REFERENCIA_DUPLICADA_MENU,
                            "detalle": f"Referencia duplicada en {nombre}: {referencia}",
                            "menu_index": i,
                            "estado": ESTADO_INCIDENCIA_PENDIENTE,
                        })
                    refs_vistas.add(k)

                    if tipo_ref == "ESCANDALLO":
                        esc = self.servicio_menus.repo_esc.obtener(referencia)
                        if not esc:
                            receta_match = self.repo_recetas.obtener(referencia)
                            if not receta_match:
                                clave_rec = self._norm(referencia)
                                for cand in self.repo_recetas.listar(incluir_archivadas=True):
                                    if clave_rec == self._norm(cand.get("nombre")):
                                        receta_match = cand
                                        break
                            if receta_match:
                                item["tipo_referencia"] = "RECETA"
                                continue
                            incidencias.append({
                                "tipo": TIPO_ESCANDALLO_INEXISTENTE,
                                "detalle": f"Escandallo no encontrado para {nombre}: {referencia}",
                                "menu_index": i,
                                "estado": ESTADO_INCIDENCIA_PENDIENTE,
                            })
                    elif tipo_ref == "RECETA":
                        receta = self.repo_recetas.obtener(referencia)
                        if not receta:
                            clave_rec = self._norm(referencia)
                            for cand in self.repo_recetas.listar(incluir_archivadas=True):
                                if clave_rec == self._norm(cand.get("nombre")):
                                    receta = cand
                                    break
                        if not receta:
                            incidencias.append({
                                "tipo": TIPO_PLATO_INEXISTENTE,
                                "detalle": f"Receta no encontrada para {nombre}: {referencia}",
                                "menu_index": i,
                                "estado": ESTADO_INCIDENCIA_PENDIENTE,
                            })
                    elif tipo_ref == "PRODUCTO":
                        prod = self.repo_catalogo.obtener_producto(referencia)
                        if not prod:
                            encontrados = self.repo_catalogo.buscar_productos({"nombre": referencia})
                            prod = encontrados[0] if encontrados else None
                        if not prod:
                            incidencias.append({
                                "tipo": TIPO_PLATO_INEXISTENTE,
                                "detalle": f"Producto no encontrado para {nombre}: {referencia}",
                                "menu_index": i,
                                "estado": ESTADO_INCIDENCIA_PENDIENTE,
                            })
                        elif self._to_float(prod.get("precio")) <= 0:
                            incidencias.append({
                                "tipo": TIPO_PRODUCTO_SIN_PRECIO_MENU,
                                "detalle": f"Producto sin precio para {nombre}: {referencia}",
                                "menu_index": i,
                                "estado": ESTADO_INCIDENCIA_PENDIENTE,
                            })

            if total_refs == 0:
                incidencias.append({
                    "tipo": TIPO_PLATO_INEXISTENTE,
                    "detalle": f"Menú sin composición válida: {nombre}",
                    "menu_index": i,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })

        for nombre_norm, count in nombres.items():
            if count > 1:
                incidencias.append({
                    "tipo": TIPO_REFERENCIA_DUPLICADA_MENU,
                    "detalle": f"Nombre de menú duplicado: {nombre_norm}",
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })
        return incidencias

    def _resumen_menus(self, menus: list[dict[str, Any]], incidencias: list[dict[str, Any]]) -> dict[str, Any]:
        return {
            "menus_detectados": len(menus),
            "menus_validos": max(0, len(menus) - len({i.get('menu_index') for i in incidencias if isinstance(i.get('menu_index'), int)})),
            "incidencias": len(incidencias),
            "escandallos_inexistentes": sum(1 for i in incidencias if i.get("tipo") == TIPO_ESCANDALLO_INEXISTENTE),
            "platos_inexistentes": sum(1 for i in incidencias if i.get("tipo") == TIPO_PLATO_INEXISTENTE),
            "productos_sin_precio": sum(1 for i in incidencias if i.get("tipo") == TIPO_PRODUCTO_SIN_PRECIO_MENU),
            "duplicados": sum(1 for i in incidencias if i.get("tipo") == TIPO_REFERENCIA_DUPLICADA_MENU),
        }

    def _ejecutar_importacion_menus(self, documento: DocumentoImportacion601) -> dict[str, Any]:
        modelo = self._crear_modelo_intermedio_menus(documento)
        incidencias = self._detectar_incidencias_menus(modelo)
        resumen = self._resumen_menus(modelo, incidencias)
        return {
            "tipo": "MENUS",
            "documento": documento,
            "modelo_intermedio": modelo,
            "incidencias": incidencias,
            "resumen": resumen,
            "recetas": [],
            "ingredientes": [],
            "productos_existentes": [],
            "productos_nuevos": [],
        }

    def _confirmar_importacion_menus(self, contexto: dict[str, Any], guardar_pendientes: bool = False) -> dict[str, Any]:
        modelo = list(contexto.get("modelo_intermedio") or [])
        incidencias = list(contexto.get("incidencias") or [])
        pendientes_por_menu = {i.get("menu_index") for i in incidencias if isinstance(i.get("menu_index"), int)}

        guardados = 0
        for i, menu in enumerate(modelo):
            if i in pendientes_por_menu and not guardar_pendientes:
                continue

            res = self.servicio_menus.nuevo_menu(
                {
                    "nombre": menu.get("nombre"),
                    "tipo": menu.get("tipo"),
                    "comensales_recomendado": menu.get("comensales") or 0,
                    "precio_venta_comensal": menu.get("precio_venta_comensal") or 0,
                    "composicion": menu.get("composicion") or {},
                    "observaciones": f"Importado desde Centro de Importación 601 ({contexto.get('documento').origen})",
                }
            )
            if res.get("ok"):
                guardados += 1
                if guardar_pendientes:
                    creado = res.get("menu") or {}
                    self.servicio_menus.repo.actualizar(
                        str(creado.get("menu_id") or creado.get("codigo") or ""),
                        {**creado, "estado": ESTADO_CON_INCIDENCIAS},
                        motivo_historial="importacion_pendiente",
                    )

        estado = ESTADO_IMPORTACION_COMPLETADA
        if incidencias:
            estado = ESTADO_IMPORTACION_CON_INCIDENCIAS
        registro = self.repo_importaciones.registrar_importacion(
            {
                "origen": f"MENUS_{contexto.get('documento').origen}",
                "numero_recetas": len(modelo),
                "numero_incidencias": len(incidencias),
                "estado": estado,
            }
        )
        incidencias_guardadas = self.repo_importaciones.registrar_incidencias(registro["id"], incidencias)
        return {
            "importacion": registro,
            "menus_guardados": guardados,
            "incidencias": incidencias_guardadas,
            "resumen_final": {
                "importacion_finalizada": True,
                "menus_guardados": guardados,
                "incidencias_pendientes": len(incidencias_guardadas),
            },
        }

    def _map_linea_intermedia(self, raw: dict[str, Any], mapping: dict[str, int], origen_etiqueta: str) -> dict[str, Any]:
        cols = list(raw.get("columnas") or [])

        def val(campo: str) -> str:
            idx = mapping.get(campo)
            if isinstance(idx, int) and 0 <= idx < len(cols):
                return str(cols[idx] or "").strip()
            return ""

        item = {
            "nombre_escandallo": val("nombre_escandallo") or val("receta"),
            "receta": val("receta") or val("nombre_escandallo"),
            "raciones": val("raciones"),
            "ingrediente": val("ingrediente") or val("producto"),
            "producto": val("producto") or val("ingrediente"),
            "cantidad": val("cantidad"),
            "unidad": val("unidad"),
            "merma": val("merma"),
            "precio": val("precio"),
            "proveedor": val("proveedor"),
            "coste_linea": val("coste_linea"),
            "precio_venta": val("precio_venta"),
            "observaciones": val("observaciones"),
            "origen": origen_etiqueta,
            "confianza": 0.0,
            "campos_dudosos": [],
            "fila_o_posicion_origen": f"{raw.get('hoja', 'hoja')}:{raw.get('fila', '?')}",
        }

        filled = sum(1 for k in self._CAMPOS_BASE_ESC if k not in {"confianza", "campos_dudosos"} and str(item.get(k) or "").strip())
        item["confianza"] = round(filled / max(1, len(self._CAMPOS_BASE_ESC) - 2), 4)
        if not item.get("nombre_escandallo"):
            item["campos_dudosos"].append("nombre_escandallo")
        if not item.get("producto"):
            item["campos_dudosos"].append("producto")
        if not item.get("cantidad"):
            item["campos_dudosos"].append("cantidad")
        return item

    def _crear_modelo_intermedio_escandallos(self, documento: DocumentoImportacion601) -> list[dict[str, Any]]:
        payload = dict(documento.payload or {})
        mapping = dict(payload.get("mapping") or {})
        raws = list(payload.get("lineas") or [])
        out: list[dict[str, Any]] = []
        for idx, raw in enumerate(raws):
            item = self._map_linea_intermedia(raw, mapping, documento.origen)
            item["linea_index"] = idx
            out.append(item)
        return out

    def _ejecutar_importacion_escandallos(self, documento: DocumentoImportacion601) -> dict[str, Any]:
        modelo = self._crear_modelo_intermedio_escandallos(documento)
        incidencias = self._detectar_incidencias_escandallos(modelo)
        resumen = self._resumen_escandallos(modelo, incidencias)
        return {
            "tipo": "ESCANDALLOS",
            "documento": documento,
            "modelo_intermedio": modelo,
            "incidencias": incidencias,
            "resumen": resumen,
            "recetas": [],
            "ingredientes": [],
            "productos_existentes": [],
            "productos_nuevos": [],
        }

    def _detectar_incidencias_escandallos(self, modelo: list[dict[str, Any]]) -> list[dict[str, Any]]:
        incidencias: list[dict[str, Any]] = []
        proveedores = {self._norm(p.get("nombre")) for p in self.repo_catalogo.listar_proveedores()}
        recetas = self.repo_recetas.listar(incluir_archivadas=True)
        recetas_nombre = {self._norm(r.get("nombre")) for r in recetas}
        recetas_codigo = {self._norm(r.get("codigo")) for r in recetas}

        for idx, linea in enumerate(modelo):
            producto_txt = str(linea.get("producto") or "").strip()
            receta_txt = str(linea.get("receta") or "").strip()
            cantidad = self._to_float(linea.get("cantidad"))
            merma = self._to_float(linea.get("merma"))
            unidad = self._norm(linea.get("unidad"))
            proveedor = self._norm(linea.get("proveedor"))

            if receta_txt and self._norm(receta_txt) not in recetas_nombre and self._norm(receta_txt) not in recetas_codigo:
                incidencias.append({
                    "tipo": TIPO_RECETA_INEXISTENTE,
                    "detalle": f"Receta inexistente: {receta_txt}",
                    "linea_index": idx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })

            if not cantidad and self._norm(linea.get("cantidad")) not in {"al gusto", "c/s", "cs"}:
                incidencias.append({
                    "tipo": TIPO_CANTIDAD_AUSENTE,
                    "detalle": f"Cantidad ausente en {linea.get('fila_o_posicion_origen')}",
                    "linea_index": idx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })

            if unidad and unidad not in self._UNIDADES_VALIDAS:
                incidencias.append({
                    "tipo": TIPO_UNIDAD_DESCONOCIDA,
                    "detalle": f"Unidad desconocida '{linea.get('unidad')}' en {linea.get('fila_o_posicion_origen')}",
                    "linea_index": idx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })

            if merma is not None and (merma < 0 or merma >= 100):
                incidencias.append({
                    "tipo": TIPO_MERMA_INVALIDA,
                    "detalle": f"Merma inválida '{linea.get('merma')}' en {linea.get('fila_o_posicion_origen')}",
                    "linea_index": idx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })

            encontrados = self.repo_catalogo.buscar_productos({"nombre": producto_txt}) if producto_txt else []
            if not encontrados:
                incidencias.append({
                    "tipo": TIPO_PRODUCTO_INEXISTENTE,
                    "detalle": f"Producto inexistente: {producto_txt or '[vacío]'}",
                    "producto": producto_txt,
                    "linea_index": idx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })
            elif len(encontrados) > 1:
                incidencias.append({
                    "tipo": TIPO_PRODUCTO_DUPLICADO,
                    "detalle": f"Producto ambiguo: {producto_txt}",
                    "producto": producto_txt,
                    "candidatos": [str(e.get("nombre") or "") for e in encontrados[:10]],
                    "linea_index": idx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })
            else:
                prod = encontrados[0]
                linea["producto_codigo_resuelto"] = str(prod.get("codigo") or "")
                if self._to_float(prod.get("precio")) in (None, 0.0):
                    incidencias.append({
                        "tipo": TIPO_PRODUCTO_SIN_PRECIO,
                        "detalle": f"Producto sin precio: {producto_txt}",
                        "producto": producto_txt,
                        "linea_index": idx,
                        "estado": ESTADO_INCIDENCIA_PENDIENTE,
                    })

            if proveedor and proveedor not in proveedores:
                incidencias.append({
                    "tipo": TIPO_PROVEEDOR_INEXISTENTE,
                    "detalle": f"Proveedor inexistente: {linea.get('proveedor')}",
                    "proveedor": linea.get("proveedor"),
                    "linea_index": idx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })

            coste_importado = self._to_float(linea.get("coste_linea"))
            if coste_importado is not None and coste_importado >= 0 and linea.get("producto_codigo_resuelto") and cantidad:
                calculo = self.servicio_escandallos.motor.calcular(
                    nombre_escandallo="verificacion_importacion",
                    numero_raciones=1,
                    lineas_entrada=[
                        {
                            "producto_codigo": linea.get("producto_codigo_resuelto"),
                            "producto": linea.get("producto"),
                            "cantidad_neta": cantidad,
                            "unidad_receta": linea.get("unidad"),
                            "merma_especifica": linea.get("merma"),
                        }
                    ],
                    precio_venta_total=0.0,
                )
                lineas_calc = list(calculo.get("lineas") or [])
                if lineas_calc:
                    coste_calc = self._to_float(lineas_calc[0].get("coste_linea")) or 0.0
                    linea["coste_calculado_sistema"] = coste_calc
                    if abs(coste_importado - coste_calc) > 0.01:
                        incidencias.append({
                            "tipo": TIPO_COSTE_IMPORTADO_DISCREPANTE,
                            "detalle": f"Coste discrepante en {linea.get('fila_o_posicion_origen')}",
                            "linea_index": idx,
                            "coste_importado": round(coste_importado, 6),
                            "coste_calculado": round(coste_calc, 6),
                            "estado": ESTADO_INCIDENCIA_PENDIENTE,
                        })

        return incidencias

    def _resumen_escandallos(self, modelo: list[dict[str, Any]], incidencias: list[dict[str, Any]]) -> dict[str, Any]:
        escandallos_detectados = len({self._norm(x.get("nombre_escandallo") or x.get("receta")) for x in modelo if self._norm(x.get("nombre_escandallo") or x.get("receta"))})
        lineas_analizadas = len(modelo)
        por_linea: dict[int, list[dict[str, Any]]] = {}
        for inc in incidencias:
            idx = inc.get("linea_index")
            if isinstance(idx, int):
                por_linea.setdefault(idx, []).append(inc)

        criticos = {
            TIPO_PRODUCTO_INEXISTENTE,
            TIPO_PRODUCTO_DUPLICADO,
            TIPO_CANTIDAD_AUSENTE,
            TIPO_UNIDAD_DESCONOCIDA,
            TIPO_MERMA_INVALIDA,
        }
        lineas_validas = 0
        for idx in range(lineas_analizadas):
            tipos = {i.get("tipo") for i in por_linea.get(idx, [])}
            if not (tipos & criticos):
                lineas_validas += 1

        return {
            "archivo": "; ".join(list((modelo[0].get("origen") for _ in [0])) if modelo else []),
            "escandallos_detectados": escandallos_detectados,
            "lineas_analizadas": lineas_analizadas,
            "lineas_validas": lineas_validas,
            "posibles_duplicados": sum(1 for i in incidencias if i.get("tipo") == TIPO_PRODUCTO_DUPLICADO),
            "productos_inexistentes": sum(1 for i in incidencias if i.get("tipo") == TIPO_PRODUCTO_INEXISTENTE),
            "discrepancias_coste": sum(1 for i in incidencias if i.get("tipo") == TIPO_COSTE_IMPORTADO_DISCREPANTE),
            "registros_incompletos": sum(1 for i in incidencias if i.get("tipo") in {TIPO_CANTIDAD_AUSENTE, TIPO_RECETA_INEXISTENTE, TIPO_MERMA_INVALIDA}),
            "incidencias": len(incidencias),
        }

    def _agrupar_modelo_por_escandallo(self, modelo: list[dict[str, Any]]) -> dict[str, list[dict[str, Any]]]:
        grupos: dict[str, list[dict[str, Any]]] = {}
        for fila in modelo:
            clave = str(fila.get("nombre_escandallo") or fila.get("receta") or "Escandallo importado").strip()
            grupos.setdefault(clave, []).append(fila)
        return grupos

    def _aplicar_costes_manuales(self, escandallo_calc: dict[str, Any], lineas_origen: list[dict[str, Any]]) -> dict[str, Any]:
        lineas_calc = list(escandallo_calc.get("lineas") or [])
        for idx, origen in enumerate(lineas_origen):
            if idx >= len(lineas_calc):
                continue
            if str(origen.get("decision_coste") or "") == "MANTENER_IMPORTADO":
                manual = self._to_float(origen.get("coste_linea_manual") or origen.get("coste_linea"))
                if manual is None:
                    continue
                lineas_calc[idx]["coste_linea"] = round(manual, 6)
                lineas_calc[idx]["dato_manual"] = True
                lineas_calc[idx]["observaciones"] = (str(lineas_calc[idx].get("observaciones") or "") + " | Coste importado conservado manualmente").strip(" |")

        coste_total = round(sum(float(l.get("coste_linea") or 0.0) for l in lineas_calc), 6)
        raciones = float(escandallo_calc.get("numero_raciones") or 0.0)
        escandallo_calc["lineas"] = lineas_calc
        escandallo_calc["coste_total"] = coste_total
        escandallo_calc["coste_por_racion"] = round(coste_total / raciones, 6) if raciones > 0 else 0.0
        venta = float(escandallo_calc.get("precio_venta_total") or 0.0)
        if venta > 0:
            beneficio = round(venta - coste_total, 6)
            escandallo_calc["beneficio_total"] = beneficio
            escandallo_calc["beneficio_por_racion"] = round(beneficio / raciones, 6) if raciones > 0 else 0.0
            escandallo_calc["margen_monetario"] = beneficio
            escandallo_calc["margen_porcentual"] = round((beneficio / venta) * 100.0, 6) if venta else 0.0
            escandallo_calc["coste_sobre_venta_porcentual"] = round((coste_total / venta) * 100.0, 6) if venta else 0.0
        return escandallo_calc

    def _confirmar_importacion_escandallos(self, contexto: dict[str, Any], guardar_pendientes: bool = False) -> dict[str, Any]:
        resolucion = self.asistente_incidencias.resolver(contexto)
        contexto["incidencias"] = resolucion["incidencias"]
        acciones = resolucion["acciones"]

        incidencias_pendientes = [
            i for i in list(contexto.get("incidencias") or [])
            if str(i.get("estado") or ESTADO_INCIDENCIA_PENDIENTE) != ESTADO_INCIDENCIA_RESUELTA
        ]
        pendientes_por_linea = {
            i.get("linea_index") for i in incidencias_pendientes if isinstance(i.get("linea_index"), int)
        }

        modelo = list(contexto.get("modelo_intermedio") or [])
        grupos = self._agrupar_modelo_por_escandallo(modelo)

        guardados = 0
        for nombre, filas in grupos.items():
            lineas_entrada: list[dict[str, Any]] = []
            filas_origen: list[dict[str, Any]] = []
            for fila in filas:
                idx = int(fila.get("linea_index") or -1)
                if idx in pendientes_por_linea and not guardar_pendientes:
                    continue
                lineas_entrada.append(
                    {
                        "producto_codigo": fila.get("producto_codigo_resuelto") or "",
                        "producto": fila.get("producto") or fila.get("ingrediente") or "",
                        "cantidad_texto": fila.get("cantidad") or "",
                        "unidad_receta": fila.get("unidad") or "",
                        "merma_especifica": fila.get("merma") or "",
                        "observaciones": fila.get("observaciones") or "",
                    }
                )
                filas_origen.append(fila)

            if not lineas_entrada:
                continue

            raciones = self._to_float(filas[0].get("raciones")) or 1.0
            precio_venta_total = self._to_float(filas[0].get("precio_venta")) or 0.0
            receta_txt = str(filas[0].get("receta") or "").strip()
            receta = self.repo_recetas.obtener(receta_txt) if receta_txt else None
            calc = self.servicio_escandallos.motor.calcular(
                nombre_escandallo=nombre,
                numero_raciones=raciones,
                lineas_entrada=lineas_entrada,
                precio_venta_total=precio_venta_total,
                receta_asociada={
                    "id": str((receta or {}).get("id") or ""),
                    "codigo": str((receta or {}).get("codigo") or ""),
                    "nombre": str((receta or {}).get("nombre") or receta_txt),
                    "version": int((receta or {}).get("version") or 0),
                },
            )
            calc = self._aplicar_costes_manuales(calc, filas_origen)
            calc["observaciones_economicas"] = "Importado desde Centro de Importación 601"
            calc["familia"] = str((receta or {}).get("familia") or "")
            if guardar_pendientes and str(calc.get("estado") or "") == "OPERATIVO":
                calc["estado"] = "CON_INCIDENCIAS"
            self.servicio_escandallos.repo_esc.guardar_nuevo(calc)
            guardados += 1

        estado = ESTADO_IMPORTACION_COMPLETADA if not incidencias_pendientes else ESTADO_IMPORTACION_CON_INCIDENCIAS
        registro = self.repo_importaciones.registrar_importacion(
            {
                "origen": f"ESCANDALLOS_{contexto.get('documento').origen}",
                "numero_recetas": len(grupos),
                "numero_incidencias": len(incidencias_pendientes),
                "estado": estado,
            }
        )
        incidencias_guardadas = self.repo_importaciones.registrar_incidencias(registro["id"], incidencias_pendientes)
        return {
            "importacion": registro,
            "escandallos_guardados": guardados,
            "incidencias": incidencias_guardadas,
            "acciones": acciones,
            "resumen_final": {
                "importacion_finalizada": True,
                "escandallos_guardados": guardados,
                "productos_nuevos": int(acciones.get("productos_nuevos") or 0),
                "precios_añadidos": int(acciones.get("precios_añadidos") or 0),
                "proveedores_nuevos": int(acciones.get("proveedores_nuevos") or 0),
                "incidencias_pendientes": len(incidencias_guardadas),
            },
        }

    def confirmar_importacion(self, contexto: dict[str, Any], *, guardar_pendientes: bool = False) -> dict[str, Any]:
        if str((contexto.get("tipo") or contexto.get("documento").tipo_contenido or "")).upper() == "ESCANDALLOS":
            return self._confirmar_importacion_escandallos(contexto, guardar_pendientes=guardar_pendientes)
        if str((contexto.get("tipo") or contexto.get("documento").tipo_contenido or "")).upper() == "MENUS":
            return self._confirmar_importacion_menus(contexto, guardar_pendientes=guardar_pendientes)

        resolucion = self.asistente_incidencias.resolver(contexto)
        contexto["incidencias"] = resolucion["incidencias"]
        acciones = resolucion["acciones"]

        recetas = list(contexto.get("recetas") or [])
        incidencias_todas = list(contexto.get("incidencias") or [])
        incidencias_pendientes = [
            i for i in incidencias_todas
            if str(i.get("estado") or ESTADO_INCIDENCIA_PENDIENTE) != ESTADO_INCIDENCIA_RESUELTA
        ]

        creadas = 0
        for receta in recetas:
            resultado = self.repo_recetas.crear(receta)
            if resultado.get("ok"):
                creadas += 1
            else:
                for error in resultado.get("errores", []):
                    incidencias_pendientes.append({
                        "tipo": "RECETA_NO_IMPORTADA",
                        "detalle": f"{receta.get('nombre', 'Receta')}: {error}",
                        "estado": ESTADO_INCIDENCIA_PENDIENTE,
                    })

        estado = ESTADO_IMPORTACION_COMPLETADA if not incidencias_pendientes else ESTADO_IMPORTACION_CON_INCIDENCIAS
        registro = self.repo_importaciones.registrar_importacion({
            "origen": contexto.get("documento").origen,
            "numero_recetas": len(recetas),
            "numero_incidencias": len(incidencias_pendientes),
            "estado": estado,
        })
        incidencias_guardadas = self.repo_importaciones.registrar_incidencias(registro["id"], incidencias_pendientes)
        resumen_final = {
            "importacion_finalizada": True,
            "recetas_creadas": creadas,
            "productos_nuevos": int(acciones.get("productos_nuevos") or 0),
            "precios_añadidos": int(acciones.get("precios_añadidos") or 0),
            "proveedores_nuevos": int(acciones.get("proveedores_nuevos") or 0),
            "incidencias_pendientes": len(incidencias_guardadas),
        }
        return {
            "importacion": registro,
            "recetas_creadas": creadas,
            "incidencias": incidencias_guardadas,
            "acciones": acciones,
            "resumen_final": resumen_final,
        }

    def cancelar_importacion(self, contexto: dict[str, Any]) -> dict[str, Any]:
        numero_registros = len(contexto.get("recetas") or [])
        if str((contexto.get("tipo") or contexto.get("documento").tipo_contenido or "")).upper() == "ESCANDALLOS":
            numero_registros = len(contexto.get("modelo_intermedio") or [])
        if str((contexto.get("tipo") or contexto.get("documento").tipo_contenido or "")).upper() == "MENUS":
            numero_registros = len(contexto.get("modelo_intermedio") or [])
        registro = self.repo_importaciones.registrar_importacion({
            "origen": contexto.get("documento").origen,
            "numero_recetas": numero_registros,
            "numero_incidencias": len(contexto.get("incidencias") or []),
            "estado": ESTADO_IMPORTACION_CANCELADA,
        })
        return {"importacion": registro}

    @staticmethod
    def _analizar_documento(documento: DocumentoImportacion601) -> dict[str, Any]:
        return {
            "origen": documento.origen,
            "etiqueta_origen": documento.etiqueta_origen,
            "texto": documento.texto,
            "texto_disponible": bool(documento.texto.strip()),
            "archivos": list(documento.archivos),
            "advertencias": list(documento.advertencias),
        }

    def _detectar_recetas(self, analisis: dict[str, Any]) -> list[dict[str, Any]]:
        if not analisis.get("texto_disponible"):
            return []
        texto = str(analisis.get("texto") or "")
        if not texto:
            texto = ""
        return self._extraer_recetas_texto(texto)

    def _extraer_recetas_texto(self, texto: str) -> list[dict[str, Any]]:
        lineas = [l.strip() for l in str(texto or "").splitlines() if l.strip()]
        recetas: list[dict[str, Any]] = []
        actual: Optional[dict[str, Any]] = None
        for linea in lineas:
            if linea.lower().startswith("receta:"):
                if actual:
                    recetas.append(actual)
                nombre = linea.split(":", 1)[1].strip() or f"Receta importada {len(recetas) + 1}"
                actual = {
                    "nombre": nombre,
                    "codigo": "",
                    "familia": "",
                    "tipo": "",
                    "numero_raciones": 0,
                    "ingredientes": [],
                    "cantidades": [],
                    "elaboracion": "",
                    "tiempo_elaboracion": "",
                    "conservacion": "",
                    "alergenos": [],
                    "observaciones": "",
                    "fotografia": "",
                }
                continue
            if actual is None:
                continue
            if linea.lower().startswith("raciones:"):
                actual["numero_raciones"] = linea.split(":", 1)[1].strip()
            elif linea.lower().startswith("elaboracion:"):
                actual["elaboracion"] = linea.split(":", 1)[1].strip()
            else:
                m = re.match(r"^-?\s*([^:|-]+)\s*[:|-]\s*(.*)$", linea)
                if m:
                    actual["ingredientes"].append(m.group(1).strip())
                    actual["cantidades"].append(m.group(2).strip())
                elif actual.get("elaboracion"):
                    actual["elaboracion"] = f"{actual['elaboracion']} {linea}".strip()
                else:
                    actual["elaboracion"] = linea
        if actual:
            recetas.append(actual)

        if not recetas and lineas:
            receta_unica = {
                "nombre": "Receta importada 1",
                "codigo": "",
                "familia": "",
                "tipo": "",
                "numero_raciones": 0,
                "ingredientes": [],
                "cantidades": [],
                "elaboracion": " ".join(lineas),
                "tiempo_elaboracion": "",
                "conservacion": "",
                "alergenos": [],
                "observaciones": "Importación desde texto libre.",
                "fotografia": "",
            }
            recetas.append(receta_unica)
        return recetas

    @staticmethod
    def _detectar_ingredientes(recetas: list[dict[str, Any]]) -> list[dict[str, str]]:
        ingredientes: list[dict[str, str]] = []
        for receta in recetas:
            for nombre, cantidad in zip(receta.get("ingredientes", []), receta.get("cantidades", [])):
                ingredientes.append({"nombre": str(nombre), "cantidad": str(cantidad)})
        return ingredientes

    def _buscar_productos(self, recetas: list[dict[str, Any]]) -> tuple[list[str], list[str]]:
        vistos: set[str] = set()
        for receta in recetas:
            for nombre in receta.get("ingredientes", []):
                n = str(nombre).strip().lower()
                if n:
                    vistos.add(n)
        existentes: list[str] = []
        nuevos: list[str] = []
        for nombre in sorted(vistos):
            encontrados = self.repo_catalogo.buscar_productos({"nombre": nombre})
            if encontrados:
                existentes.append(nombre)
            else:
                nuevos.append(nombre)
        return existentes, nuevos

    @staticmethod
    def _detectar_incidencias(
        *,
        documento: DocumentoImportacion601,
        recetas: list[dict[str, Any]],
        ingredientes: list[dict[str, str]],
        productos_nuevos: list[str],
    ) -> list[dict[str, Any]]:
        incidencias: list[dict[str, Any]] = []

        for advertencia in documento.advertencias:
            incidencias.append({"tipo": "LECTOR_PENDIENTE", "detalle": advertencia, "estado": ESTADO_INCIDENCIA_PENDIENTE})

        if not recetas:
            incidencias.append({"tipo": "RECETAS_NO_DETECTADAS", "detalle": "No se detectaron recetas en el documento.", "estado": ESTADO_INCIDENCIA_PENDIENTE})

        vistos_nombres: dict[str, int] = {}
        for ridx, receta in enumerate(recetas):
            nombre = str(receta.get("nombre") or "").strip().lower()
            if nombre:
                vistos_nombres[nombre] = vistos_nombres.get(nombre, 0) + 1

            if not receta.get("elaboracion"):
                incidencias.append({
                    "tipo": TIPO_RECETA_INCOMPLETA,
                    "detalle": f"{receta.get('nombre', 'Receta sin nombre')} sin elaboración.",
                    "receta": receta.get("nombre", "Receta sin nombre"),
                    "receta_index": ridx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })
            if not receta.get("ingredientes"):
                incidencias.append({
                    "tipo": TIPO_RECETA_INCOMPLETA,
                    "detalle": f"{receta.get('nombre', 'Receta sin nombre')} sin ingredientes.",
                    "receta": receta.get("nombre", "Receta sin nombre"),
                    "receta_index": ridx,
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })

            ingredientes_receta = list(receta.get("ingredientes") or [])
            cantidades_receta = list(receta.get("cantidades") or [])
            while len(cantidades_receta) < len(ingredientes_receta):
                cantidades_receta.append("")
            receta["cantidades"] = cantidades_receta

            for iidx, ing in enumerate(ingredientes_receta):
                cantidad = str(cantidades_receta[iidx] or "").strip()
                if not cantidad:
                    incidencias.append({
                        "tipo": TIPO_INGREDIENTE_SIN_CANTIDAD,
                        "detalle": f"Ingrediente sin cantidad en {receta.get('nombre', 'Receta sin nombre')}: {ing}",
                        "receta": receta.get("nombre", "Receta sin nombre"),
                        "receta_index": ridx,
                        "ingrediente": ing,
                        "ingrediente_index": iidx,
                        "estado": ESTADO_INCIDENCIA_PENDIENTE,
                    })

        for nombre, repeticiones in vistos_nombres.items():
            if repeticiones > 1:
                incidencias.append({
                    "tipo": TIPO_PRODUCTO_DUPLICADO,
                    "detalle": f"Posible duplicado detectado: {nombre}",
                    "producto": nombre,
                    "candidatos": [nombre, f"{nombre} 5 kg", f"{nombre} natural triturado"],
                    "estado": ESTADO_INCIDENCIA_PENDIENTE,
                })

        for producto in productos_nuevos:
            incidencias.append({
                "tipo": TIPO_PRODUCTO_INEXISTENTE,
                "detalle": f"Producto inexistente: {producto}",
                "producto": producto,
                "estado": ESTADO_INCIDENCIA_PENDIENTE,
            })
            incidencias.append({
                "tipo": TIPO_PRODUCTO_SIN_PRECIO,
                "detalle": f"Producto sin precio detectado: {producto}",
                "producto": producto,
                "estado": ESTADO_INCIDENCIA_PENDIENTE,
            })

        if productos_nuevos:
            incidencias.append({
                "tipo": TIPO_PROVEEDOR_INEXISTENTE,
                "detalle": "Proveedor no identificado para productos nuevos.",
                "proveedor": "Proveedor por definir",
                "estado": ESTADO_INCIDENCIA_PENDIENTE,
            })

        return incidencias

    @staticmethod
    def _resumen(*, recetas: list[dict[str, Any]], ingredientes: list[dict[str, str]], nuevos: list[str], incidencias: list[dict[str, str]]) -> dict[str, int]:
        sin_precio = sum(1 for i in incidencias if i.get("tipo") == TIPO_PRODUCTO_SIN_PRECIO)
        duplicados = sum(1 for i in incidencias if i.get("tipo") == TIPO_PRODUCTO_DUPLICADO)
        incompletas = sum(1 for i in incidencias if i.get("tipo") == TIPO_RECETA_INCOMPLETA)
        total_incidencias = sum(1 for i in incidencias if str(i.get("estado") or ESTADO_INCIDENCIA_PENDIENTE) == ESTADO_INCIDENCIA_PENDIENTE)
        return {
            "recetas_detectadas": len(recetas),
            "ingredientes": len(ingredientes),
            "incidencias": total_incidencias,
            "productos_nuevos": len(nuevos),
            "productos_sin_precio": sin_precio,
            "posibles_duplicados": duplicados,
            "recetas_incompletas": incompletas,
        }


class CentroImportacionUI601:
    def __init__(self, base_dir: Path):
        self.base_dir = Path(base_dir).resolve()
        self.flujo = FlujoImportacionUnificado601(self.base_dir)
        self.repo = RepositorioCentroImportacion601(self.base_dir)

    def ejecutar(self) -> None:
        while True:
            self._imprimir_menu()
            op = input("Elige una opción: ").strip()
            if op == "1":
                self.importar_word()
            elif op == "2":
                self.importar_excel()
            elif op == "3":
                self.importar_pdf()
            elif op == "4":
                self.importar_imagenes()
            elif op == "5":
                self.pegar_texto()
            elif op == "6":
                self.ver_historial()
            elif op == "7":
                self.ver_incidencias_pendientes()
            elif op == "8":
                self.importar_escandallos()
            elif op == "9":
                self.importar_menus()
            elif op == "0":
                return
            else:
                print("Opción no válida.")

    @staticmethod
    def _imprimir_menu() -> None:
        print("\nCENTRO DE IMPORTACIÓN")
        print("1. Importar Word")
        print("2. Importar Excel")
        print("3. Importar PDF")
        print("4. Importar imágenes")
        print("5. Pegar texto")
        print("6. Historial de importaciones")
        print("7. Incidencias pendientes")
        print("8. Importar escandallos")
        print("9. Importar menús")
        print("0. Volver")

    def importar_word(self) -> None:
        ruta = input("Ruta del Word (.docx): ").strip()
        self._ejecutar_importacion(LectorWord601(), ruta)

    def importar_excel(self) -> None:
        ruta = input("Ruta del Excel (.xlsx): ").strip()
        self._ejecutar_importacion(LectorExcel601(), ruta)

    def importar_pdf(self) -> None:
        ruta = input("Ruta del PDF: ").strip()
        self._ejecutar_importacion(LectorPdf601(), ruta)

    def importar_imagenes(self) -> None:
        rutas = input("Rutas de imágenes separadas por coma (JPG, PNG, HEIC): ").strip()
        self._ejecutar_importacion(LectorImagen601(), rutas)

    def pegar_texto(self) -> None:
        print("Pega el texto. Finaliza con una línea única que contenga FIN")
        lineas: list[str] = []
        while True:
            linea = input()
            if linea.strip().upper() == "FIN":
                break
            lineas.append(linea)
        self._ejecutar_importacion(LectorTexto601(), "\n".join(lineas))

    def importar_escandallos(self) -> None:
        print("\nIMPORTAR ESCANDALLOS")
        print("1. Texto pegado")
        print("2. Excel (.xlsx)")
        op = input("Origen: ").strip()

        if op == "1":
            print("Pega el texto tabulado/separado por ; o |. Finaliza con FIN")
            lineas: list[str] = []
            while True:
                linea = input()
                if linea.strip().upper() == "FIN":
                    break
                lineas.append(linea)
            doc = self.flujo.crear_documento_escandallos_desde_texto("\n".join(lineas))
        elif op == "2":
            ruta = input("Ruta del Excel (.xlsx): ").strip()
            inspeccion = self.flujo.inspeccionar_excel_escandallos(ruta)
            if inspeccion.get("advertencias"):
                for a in inspeccion.get("advertencias", []):
                    print(f"- {a}")
            hojas = list(inspeccion.get("hojas") or [])
            if not hojas:
                print("No se detectaron hojas válidas.")
                return

            print("Hojas detectadas:")
            for i, h in enumerate(hojas, 1):
                print(f"{i}. {h}")
            seleccion = input("Elige hojas (ej. 1,2 o *): ").strip()
            if seleccion == "*" or not seleccion:
                hojas_sel = hojas
            else:
                hojas_sel = []
                for tok in [t.strip() for t in seleccion.split(",") if t.strip()]:
                    if tok.isdigit():
                        pos = int(tok)
                        if 1 <= pos <= len(hojas):
                            hojas_sel.append(hojas[pos - 1])
                if not hojas_sel:
                    hojas_sel = hojas

            mapeo_sugerido = {}
            if hojas_sel:
                mapeo_sugerido = dict((inspeccion.get("mapeos_sugeridos") or {}).get(hojas_sel[0], {}))
            print("Mapeo detectado (campo -> columna):")
            for campo, col in sorted(mapeo_sugerido.items()):
                print(f"- {campo}: {col}")

            mapping_forzado: dict[str, int] = {}
            if input("¿Corregir mapeo manualmente? (s/N): ").strip().lower() == "s":
                campos = [
                    "nombre_escandallo", "receta", "raciones", "ingrediente", "producto", "cantidad",
                    "unidad", "merma", "precio", "proveedor", "coste_linea", "precio_venta", "observaciones",
                ]
                for campo in campos:
                    actual = mapeo_sugerido.get(campo, "")
                    val = input(f"Columna para {campo} [{actual}]: ").strip()
                    if val.isdigit():
                        mapping_forzado[campo] = int(val)

            doc = self.flujo.crear_documento_escandallos_desde_excel(ruta, hojas=hojas_sel, mapping_forzado=mapping_forzado or None)
        else:
            print("Opción no válida.")
            return

        contexto = self.flujo.ejecutar(doc)
        self._mostrar_resumen_escandallos(contexto)

        while True:
            print("Opciones:")
            print("1. Revisar incidencias")
            print("2. Continuar con válidos")
            print("3. Guardar como pendientes")
            print("4. Cancelar")
            opf = input("Elige una opción: ").strip()
            if opf == "1":
                self._revisar(contexto)
            elif opf == "2":
                res = self.flujo.confirmar_importacion(contexto, guardar_pendientes=False)
                print(f"Importación registrada: {res['importacion']['id']}")
                print(f"Escandallos guardados: {res.get('escandallos_guardados', 0)}")
                print(f"Incidencias registradas: {len(res.get('incidencias', []))}")
                return
            elif opf == "3":
                res = self.flujo.confirmar_importacion(contexto, guardar_pendientes=True)
                print(f"Importación registrada: {res['importacion']['id']}")
                print(f"Escandallos guardados: {res.get('escandallos_guardados', 0)}")
                print(f"Incidencias registradas: {len(res.get('incidencias', []))}")
                return
            elif opf == "4":
                cancelada = self.flujo.cancelar_importacion(contexto)
                print(f"Importación cancelada: {cancelada['importacion']['id']}")
                return
            else:
                print("Opción no válida.")

    def importar_menus(self) -> None:
        print("\nIMPORTAR MENUS")
        print("1. Texto pegado")
        print("2. Excel (.xlsx)")
        print("3. Word (.docx)")
        print("4. PDF")
        print("5. Imagen")
        op = input("Origen: ").strip()

        if op == "1":
            print("Pega el texto. Finaliza con FIN")
            lineas: list[str] = []
            while True:
                linea = input()
                if linea.strip().upper() == "FIN":
                    break
                lineas.append(linea)
            doc = self.flujo.crear_documento_menus_desde_texto("\n".join(lineas))
        elif op == "2":
            ruta = input("Ruta del Excel (.xlsx): ").strip()
            doc = self.flujo.crear_documento_menus_desde_excel(ruta)
        elif op == "3":
            ruta = input("Ruta del Word (.docx): ").strip()
            docw = LectorWord601().leer(ruta)
            doc = self.flujo.crear_documento_menus_desde_texto(docw.texto or "", etiqueta=docw.etiqueta_origen or ruta)
            doc.origen = "WORD"
            doc.archivos = [ruta]
            doc.advertencias = list(docw.advertencias or [])
        elif op == "4":
            ruta = input("Ruta del PDF: ").strip()
            docp = LectorPdf601().leer(ruta)
            doc = self.flujo.crear_documento_menus_desde_texto(docp.texto or "", etiqueta=docp.etiqueta_origen or ruta)
            doc.origen = "PDF"
            doc.archivos = [ruta]
            doc.advertencias = list(docp.advertencias or [])
        elif op == "5":
            rutas = input("Rutas imágenes separadas por coma: ").strip()
            doci = LectorImagen601().leer(rutas)
            doc = self.flujo.crear_documento_menus_desde_texto(doci.texto or "", etiqueta=doci.etiqueta_origen or "imagenes")
            doc.origen = "IMAGEN"
            doc.archivos = list(doci.archivos or [])
            doc.advertencias = list(doci.advertencias or [])
        else:
            print("Opción no válida.")
            return

        contexto = self.flujo.ejecutar(doc)
        self._mostrar_resumen_menus(contexto)

        while True:
            print("Opciones:")
            print("1. Revisar incidencias")
            print("2. Continuar con válidos")
            print("3. Guardar como pendientes")
            print("4. Cancelar")
            opf = input("Elige una opción: ").strip()
            if opf == "1":
                self._revisar_menus(contexto)
            elif opf == "2":
                res = self.flujo.confirmar_importacion(contexto, guardar_pendientes=False)
                print(f"Importación registrada: {res['importacion']['id']}")
                print(f"Menús guardados: {res.get('menus_guardados', 0)}")
                print(f"Incidencias registradas: {len(res.get('incidencias', []))}")
                return
            elif opf == "3":
                res = self.flujo.confirmar_importacion(contexto, guardar_pendientes=True)
                print(f"Importación registrada: {res['importacion']['id']}")
                print(f"Menús guardados: {res.get('menus_guardados', 0)}")
                print(f"Incidencias registradas: {len(res.get('incidencias', []))}")
                return
            elif opf == "4":
                cancelada = self.flujo.cancelar_importacion(contexto)
                print(f"Importación cancelada: {cancelada['importacion']['id']}")
                return
            else:
                print("Opción no válida.")

    def _ejecutar_importacion(self, lector: LectorBase601, entrada: str) -> None:
        documento = lector.leer(entrada)
        contexto = self.flujo.ejecutar(documento)
        self._mostrar_resumen(contexto["resumen"])

        while True:
            print("Opciones:")
            print("1. Revisar")
            print("2. Continuar")
            print("3. Cancelar")
            op = input("Elige una opción: ").strip()

            if op == "1":
                self._revisar(contexto)
            elif op == "2":
                resultado = self.flujo.confirmar_importacion(contexto)
                print(f"Importación registrada: {resultado['importacion']['id']}")
                print(f"Recetas creadas: {resultado['recetas_creadas']}")
                print(f"Incidencias registradas: {len(resultado['incidencias'])}")
                self._mostrar_resumen_final(resultado.get("resumen_final") or {})
                return
            elif op == "3":
                cancelada = self.flujo.cancelar_importacion(contexto)
                print(f"Importación cancelada: {cancelada['importacion']['id']}")
                return
            else:
                print("Opción no válida.")

    @staticmethod
    def _mostrar_resumen(resumen: dict[str, int]) -> None:
        print("-" * 36)
        print(f"{resumen['recetas_detectadas']} recetas detectadas")
        print(f"{resumen['ingredientes']} ingredientes")
        print(f"{resumen['incidencias']} incidencias")
        print(f"{resumen['productos_nuevos']} productos nuevos")
        print(f"{resumen['productos_sin_precio']} productos sin precio")
        print(f"{resumen['posibles_duplicados']} posibles duplicados")
        print(f"{resumen['recetas_incompletas']} receta incompleta")
        print("-" * 36)

    @staticmethod
    def _mostrar_resumen_escandallos(contexto: dict[str, Any]) -> None:
        resumen = dict(contexto.get("resumen") or {})
        doc = contexto.get("documento")
        archivo = getattr(doc, "etiqueta_origen", "") or "texto_pegado"
        print(f"\nArchivo: {archivo}")
        print(f"Escandallos detectados: {resumen.get('escandallos_detectados', 0)}")
        print(f"Líneas analizadas: {resumen.get('lineas_analizadas', 0)}")
        print(f"Líneas válidas: {resumen.get('lineas_validas', 0)}")
        print(f"Posibles duplicados: {resumen.get('posibles_duplicados', 0)}")
        print(f"Productos inexistentes: {resumen.get('productos_inexistentes', 0)}")
        print(f"Discrepancias de coste: {resumen.get('discrepancias_coste', 0)}")
        print(f"Registros incompletos: {resumen.get('registros_incompletos', 0)}")

    @staticmethod
    def _mostrar_resumen_menus(contexto: dict[str, Any]) -> None:
        resumen = dict(contexto.get("resumen") or {})
        doc = contexto.get("documento")
        archivo = getattr(doc, "etiqueta_origen", "") or "texto_pegado"
        print(f"\nArchivo: {archivo}")
        print(f"Menús detectados: {resumen.get('menus_detectados', 0)}")
        print(f"Menús válidos: {resumen.get('menus_validos', 0)}")
        print(f"Incidencias: {resumen.get('incidencias', 0)}")
        print(f"Escandallos inexistentes: {resumen.get('escandallos_inexistentes', 0)}")
        print(f"Platos inexistentes: {resumen.get('platos_inexistentes', 0)}")
        print(f"Productos sin precio: {resumen.get('productos_sin_precio', 0)}")
        print(f"Duplicados: {resumen.get('duplicados', 0)}")

    @staticmethod
    def _mostrar_resumen_final(resumen: dict[str, Any]) -> None:
        if not resumen:
            return
        print("-" * 36)
        print("Importación finalizada")
        print(f"{resumen.get('recetas_creadas', 0)} recetas creadas")
        print(f"{resumen.get('productos_nuevos', 0)} productos nuevos")
        print(f"{resumen.get('precios_añadidos', 0)} precios añadidos")
        print(f"{resumen.get('proveedores_nuevos', 0)} proveedores nuevos")
        print(f"{resumen.get('incidencias_pendientes', 0)} incidencias pendientes")
        print("-" * 36)

    @staticmethod
    def _revisar(contexto: dict[str, Any]) -> None:
        print("\nREVISIÓN")
        print("Recetas detectadas:")
        for receta in contexto.get("recetas", []):
            print(f"- {receta.get('nombre')} | ingredientes: {len(receta.get('ingredientes', []))}")

        print("\nIncidencias detectadas:")
        incidencias = contexto.get("incidencias", [])
        if not incidencias:
            print("- Sin incidencias")
            return
        for incidencia in incidencias:
            print(f"- {incidencia.get('tipo')}: {incidencia.get('detalle')}")

    @staticmethod
    def _revisar_menus(contexto: dict[str, Any]) -> None:
        print("\nREVISIÓN MENÚS")
        for menu in contexto.get("modelo_intermedio", []):
            nombre = str(menu.get("nombre") or "Menú")
            total_items = sum(len(list(v or [])) for v in (menu.get("composicion") or {}).values())
            print(f"- {nombre} | items: {total_items}")
        print("\nIncidencias detectadas:")
        incidencias = contexto.get("incidencias", [])
        if not incidencias:
            print("- Sin incidencias")
            return
        for incidencia in incidencias:
            print(f"- {incidencia.get('tipo')}: {incidencia.get('detalle')}")

    def ver_historial(self) -> None:
        print("\nHISTORIAL DE IMPORTACIONES")
        historial = self.repo.listar_historial()
        if not historial:
            print("No hay importaciones registradas.")
            return
        for fila in historial:
            print(
                f"- {fila.get('fecha')} | {fila.get('id')} | {fila.get('origen')} | "
                f"recetas={fila.get('numero_recetas')} | incidencias={fila.get('numero_incidencias')} | {fila.get('estado')}"
            )

    def ver_incidencias_pendientes(self) -> None:
        print("\nINCIDENCIAS PENDIENTES")
        incidencias = self.repo.incidencias_pendientes()
        if not incidencias:
            print("No hay incidencias pendientes.")
            return
        for incidencia in incidencias:
            print(
                f"- {incidencia.get('fecha')} | {incidencia.get('id')} | {incidencia.get('tipo')} | "
                f"{incidencia.get('detalle')}"
            )


__all__ = [
    "CentroImportacionUI601",
    "FlujoImportacionUnificado601",
    "RepositorioCentroImportacion601",
    "DocumentoImportacion601",
]
