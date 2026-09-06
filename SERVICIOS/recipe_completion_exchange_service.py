from __future__ import annotations

import base64
import binascii
import hashlib
import json
import math
import re
import unicodedata
from collections import Counter
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
from itertools import islice
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook

from SERVICIOS.biblioteca_recetas_601 import RepositorioBibliotecaRecetas601
from SERVICIOS.receta_documentacion_batch_service import RecetaDocumentacionBatchService
from SERVICIOS.recipe_completion_contract import (
    BATCH_MASS_SAFE_FIELDS,
    CONTRACT_FORMAT,
    CONTRACT_VERSION,
    ERROR_MESSAGES,
    FIELD_CONTRACTS,
    NO_APLICA_FIELDS,
    PRICE_BASIS_VALUES,
    PRICE_FIELD_CONTRACTS,
    PRICE_NORMALIZED_UNITS,
    PRICE_REFERENCE_COLUMNS,
    RECIPE_BOOLEAN_FIELDS,
    RECIPE_DOCUMENTATION_FIELDS,
    RECIPE_ENUM_FIELDS,
    RECIPE_INGREDIENT_UNITS,
    RECIPE_JSON_FIELDS,
    RECIPE_NUMBER_FIELDS,
    RECIPE_OPERATIONAL_UNITS,
    RECIPE_TIME_FIELDS,
    RECIPE_UNIT_FIELDS,
    SUPPORTED_CONTRACT_VERSIONS,
    WORKBOOK_INSTRUCTIONS,
    master_ai_prompt,
)
from SERVICIOS.receta_documentacion_write_service import (
    RecetaDocumentacionError,
    RecetaDocumentacionWriteService,
    classify_recipe_proposals_for_review,
    is_no_aplica,
    recipe_completion_fingerprint,
)


class RecipeCompletionExchangeError(RecetaDocumentacionError):
    pass


class RecipeCompletionExchangeService:
    """Contrato XLSX de propuestas externas; nunca modifica recetas."""

    FORMAT = CONTRACT_FORMAT
    VERSION = CONTRACT_VERSION
    SUPPORTED_VERSIONS = SUPPORTED_CONTRACT_VERSIONS
    MAX_FILE_BYTES = 10 * 1024 * 1024
    MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
    MAX_ZIP_MEMBERS = 200
    MAX_ROWS = 5000
    MAX_COLUMNS = 200
    MAX_TEXT_LENGTH = 8000
    # Un mapa de procedencia para todos los campos puede superar 8 KiB sin que
    # ninguno de sus valores sea grande. Se mantiene un límite específico y
    # acotado para no confundir ese caso legítimo con texto libre ilimitado.
    MAX_METADATA_LENGTH = 32_767
    REQUIRED_HEADERS = ("recipe_id", "recipe_fingerprint", "nombre", "origen_propuesta")
    CONTEXT_HEADERS = (
        "import_id", "contexto_documental", "ingredientes_contexto",
        "recetas_relacionadas", "datos_documentales", "datos_calculados",
        "estados_campos_operativos",
        "ingredientes", "ingredientes_estructurados", "cantidades", "unidad",
        "categoria_contexto", "texto_original", "procedencia_documental",
        "articulos_relacionados", "menus_contexto", "contexto_servicio",
        "rendimiento_contexto", "unidad_rendimiento_contexto", "raciones_contexto",
        "campos_confirmados", "propuestas_previas", "campos_pendientes", "campos_pendientes_claves",
        "campos_proponibles", "campos_requieren_revision_humana",
    )
    BOOLEAN_FIELDS = RECIPE_BOOLEAN_FIELDS
    NUMBER_FIELDS = RECIPE_NUMBER_FIELDS
    TIME_FIELDS = RECIPE_TIME_FIELDS
    JSON_FIELDS = RECIPE_JSON_FIELDS
    ENUM_FIELDS = RECIPE_ENUM_FIELDS
    UNIT_FIELDS = RECIPE_UNIT_FIELDS
    EXTERNAL_ORIGINS = frozenset({"CHATGPT", "OPENAI_API", "OTRO_PROVEEDOR_IA", "HUMANO", "ARCHIVO_EXTERNO"})

    def __init__(
        self,
        base_dir: Path,
        *,
        repository: RepositorioBibliotecaRecetas601 | None = None,
        recipe_service: RecetaDocumentacionWriteService | None = None,
        batch_service: RecetaDocumentacionBatchService | None = None,
    ) -> None:
        self.base_dir = Path(base_dir).resolve()
        self.repository = repository or RepositorioBibliotecaRecetas601(base_dir)
        self.recipe_service = recipe_service or RecetaDocumentacionWriteService(base_dir, repository=self.repository)
        self.batch_service = batch_service or RecetaDocumentacionBatchService(
            base_dir, repository=self.repository, recipe_service=self.recipe_service,
        )

    def export(
        self,
        *,
        recipe_ids: list[str],
        scope: str = "BIBLIOTECA",
        import_id: str = "",
    ) -> dict[str, Any]:
        requested = list(dict.fromkeys(str(value).strip() for value in recipe_ids if str(value).strip()))
        recipes = {self._identity(item): item for item in self.repository.listar() if self._identity(item)}
        included: list[dict[str, Any]] = []
        excluded: list[dict[str, str]] = []
        for recipe_id in requested:
            recipe = recipes.get(recipe_id)
            if not recipe:
                excluded.append({"recipe_id": recipe_id, "motivo": "RECETA_NO_EXISTE"})
            elif not self._missing(recipe):
                excluded.append({"recipe_id": recipe_id, "motivo": "YA_COMPLETA"})
            else:
                included.append(recipe)
        contexts = {
            self._identity(recipe): self.recipe_service.completion_context(self._identity(recipe))
            for recipe in included
        }
        pending_price_articles = self._pending_price_articles(included, contexts)

        workbook = Workbook()
        metadata = workbook.active
        metadata.title = "METADATA"
        for key, value in (
            ("format", self.FORMAT), ("version", self.VERSION), ("scope", str(scope or "BIBLIOTECA")),
            ("exported_at", datetime.now(timezone.utc).isoformat(timespec="seconds")),
            ("import_id", str(import_id or "")), ("recipe_count", len(included)),
            ("prompt_id", "HOSTAI_COMPLETADO_MASIVO"),
            ("schema_source", "RECIPE_COMPLETION_FIELD_CONTRACT"),
            ("pending_price_article_count", len(pending_price_articles)),
        ):
            metadata.append([key, value])
        prompt_sheet = workbook.create_sheet("PROMPT_IA")
        prompt_sheet.append(["prompt_id", "version", "contenido"])
        prompt_sheet.append(["HOSTAI_COMPLETADO_MASIVO", self.VERSION, master_ai_prompt()])
        instructions = workbook.create_sheet("INSTRUCCIONES")
        instructions.append(["regla", "detalle"])
        for rule, detail in WORKBOOK_INSTRUCTIONS:
            instructions.append([rule, detail])
        schema = workbook.create_sheet("SCHEMA")
        schema.append([
            "campo", "etiqueta", "que_es", "por_que_host_ai_lo_necesita", "como_estimarlo",
            "tipo_esperado", "enum_permitidos_json", "unidades_admitidas_json", "shape_json",
            "validaciones", "relaciones_coherencia", "admite_no_aplica", "cuando_no_aplica",
            "politica_revision", "safe", "critical", "participa_readiness", "solo_si_pendiente",
            "ejemplo_valido", "ejemplo_no_aplica",
        ])
        for field in sorted(RECIPE_DOCUMENTATION_FIELDS):
            spec = FIELD_CONTRACTS[field]
            schema.append([
                field, spec["label"], spec["what"], spec["why"], spec["how"],
                self._schema_type(field), json.dumps(list(spec["enum"]), ensure_ascii=False),
                json.dumps(list(spec["units"]), ensure_ascii=False),
                json.dumps(spec["shape"], ensure_ascii=False, separators=(",", ":")) if spec["shape"] else "",
                spec["validations"], spec["relations"], spec["allows_no_aplica"], spec["no_aplica_when"],
                spec["policy"], spec["safe"], spec["critical"], spec["readiness"], True,
                self._cell_value(spec["example"]),
                '{"estado":"NO_APLICA","origen":"IA_PROPUESTA","confianza":0.85,"motivo":"No corresponde a esta receta","fuente":"Contexto del XLSX","modelo":"modelo usado"}' if field in NO_APLICA_FIELDS else "",
            ])
        self._append_price_sheets(workbook, pending_price_articles)
        sheet = workbook.create_sheet("RECETAS")
        fields = sorted(RECIPE_DOCUMENTATION_FIELDS)
        headers = [
            *self.REQUIRED_HEADERS[:3], "origen", *self.CONTEXT_HEADERS,
            *[f"{field}_actual" for field in fields],
            *[f"{field}_propuesto" for field in fields],
            "metadatos_propuestas", "origen_propuesta",
        ]
        sheet.append(headers)
        for recipe in included:
            pending = self._missing(recipe)
            recipe_id = self._identity(recipe)
            context = contexts[recipe_id]
            values: dict[str, Any] = {
                "recipe_id": recipe_id,
                "recipe_fingerprint": recipe_completion_fingerprint(recipe),
                "nombre": recipe.get("nombre") or "",
                "origen": recipe.get("origen") or "",
                "import_id": str(import_id or ""),
                "contexto_documental": self._cell_value(context.get("contexto_documental") or {}),
                "ingredientes_contexto": self._cell_value(context.get("ingredientes_contexto") or []),
                "recetas_relacionadas": self._cell_value(context.get("recetas_relacionadas") or []),
                "datos_documentales": self._cell_value(context.get("datos_documentales") or {}),
                "datos_calculados": self._cell_value(context.get("datos_calculados") or {}),
                "estados_campos_operativos": self._cell_value(context.get("estados_campos_operativos") or {}),
                "ingredientes": self._cell_value(recipe.get("ingredientes") or []),
                "ingredientes_estructurados": self._cell_value(recipe.get("ingredientes_estructurados") or []),
                "cantidades": self._cell_value(recipe.get("cantidades") or []),
                "unidad": self._cell_value(recipe.get("unidad") or ""),
                "categoria_contexto": self._cell_value(
                    recipe.get("categoria") or recipe.get("familia") or recipe.get("tipo") or ""
                ),
                "texto_original": self._cell_value(context.get("texto_original")),
                "procedencia_documental": self._cell_value(context.get("procedencia_documental") or {}),
                "articulos_relacionados": self._cell_value(context.get("articulos_relacionados") or []),
                "menus_contexto": self._cell_value(context.get("menus") or []),
                "contexto_servicio": self._cell_value(context.get("contexto_servicio") or []),
                "rendimiento_contexto": self._cell_value(recipe.get("rendimiento")),
                "unidad_rendimiento_contexto": self._cell_value(recipe.get("unidad_rendimiento")),
                "raciones_contexto": self._cell_value(recipe.get("numero_raciones")),
                "campos_confirmados": self._cell_value(context.get("campos_confirmados") or {}),
                "propuestas_previas": self._cell_value(context.get("propuestas_previas") or {}),
                "campos_pendientes": self._cell_value(pending),
                "campos_pendientes_claves": self._cell_value(self._pending_keys(recipe)),
                "campos_proponibles": self._cell_value(sorted(RECIPE_DOCUMENTATION_FIELDS)),
                "campos_requieren_revision_humana": self._cell_value(
                    (recipe.get("completitud") or {}).get("campos_requieren_revision_humana") or []
                ),
                "metadatos_propuestas": "",
                "origen_propuesta": "",
            }
            for field in fields:
                values[f"{field}_actual"] = self._cell_value(recipe.get(field))
                values[f"{field}_propuesto"] = ""
            sheet.append([values.get(header, "") for header in headers])
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
        buffer = BytesIO()
        workbook.save(buffer)
        content = buffer.getvalue()
        suffix = re.sub(r"[^A-Za-z0-9_-]+", "-", str(import_id or scope or "biblioteca")).strip("-")[:50]
        return {
            "ok": True,
            "contrato": {"format": self.FORMAT, "version": self.VERSION},
            "filename": f"hostai-completado-recetas-{suffix or 'biblioteca'}.xlsx",
            "tipo_mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "contenido_base64": base64.b64encode(content).decode("ascii"),
            "recetas_solicitadas": len(requested), "recetas_exportadas": len(included),
            "recipe_ids": [self._identity(item) for item in included], "excluidas": excluded,
            "datos_reales_modificados": False,
        }

    def import_package(
        self,
        *,
        filename: str,
        content_base64: str,
        expected_recipe_ids: list[str] | None = None,
        scope: str = "BIBLIOTECA",
        import_id: str = "",
        source: str = "ARCHIVO_EXTERNO",
    ) -> dict[str, Any]:
        raw = self._decode_file(filename, content_base64)
        workbook = self._load_safe_workbook(raw)
        metadata = self._metadata(workbook)
        if metadata.get("format") != self.FORMAT:
            raise RecipeCompletionExchangeError("invalid_completion_format", "El archivo no es un paquete de completado de recetas de Host AI.")
        if metadata.get("version") not in self.SUPPORTED_VERSIONS:
            raise RecipeCompletionExchangeError("unsupported_completion_version", "La version del paquete de completado no es compatible.")
        if import_id and metadata.get("import_id") != str(import_id):
            raise RecipeCompletionExchangeError("completion_context_mismatch", "El archivo pertenece a otra importacion de origen.")
        if scope and metadata.get("scope") != str(scope):
            raise RecipeCompletionExchangeError("completion_scope_mismatch", "El ambito del archivo no coincide con el flujo actual.")
        if "RECETAS" not in workbook.sheetnames:
            raise RecipeCompletionExchangeError("missing_recipes_sheet", "El archivo no contiene la hoja RECETAS.")
        sheet = workbook["RECETAS"]
        # Algunos editores externos eliminan la dimension calculada de la hoja.
        # En modo read_only, openpyxl expone entonces max_row/max_column=None.
        # La iteracion acotada mantiene el limite sin depender de ese metadato.
        rows = list(islice(sheet.iter_rows(), self.MAX_ROWS + 2))
        if len(rows) > self.MAX_ROWS + 1 or (rows and max(len(row) for row in rows) > self.MAX_COLUMNS):
            raise RecipeCompletionExchangeError("completion_sheet_too_large", "La hoja RECETAS supera los limites permitidos.")
        if not rows:
            raise RecipeCompletionExchangeError("empty_completion_sheet", "La hoja RECETAS esta vacia.")
        headers = [str(cell.value or "").strip() for cell in rows[0]]
        if len(headers) != len(set(headers)):
            raise RecipeCompletionExchangeError("duplicate_completion_headers", "La hoja RECETAS contiene cabeceras duplicadas.")
        missing_headers = [header for header in self.REQUIRED_HEADERS if header not in headers]
        if metadata.get("version") == "0.3" and "import_id" not in headers:
            missing_headers.append("import_id")
        if missing_headers:
            raise RecipeCompletionExchangeError("invalid_completion_schema", "Faltan columnas obligatorias: " + ", ".join(missing_headers))
        formula_rows = sorted({cell.row for row in rows[1:] for cell in row if cell.data_type == "f"})
        if formula_rows:
            raise RecipeCompletionExchangeError("unsafe_formula", "El archivo contiene formulas y no puede importarse de forma segura.")

        records = [dict(zip(headers, [cell.value for cell in row])) for row in rows[1:] if any(cell.value not in (None, "") for cell in row)]
        identities = [str(record.get("recipe_id") or "").strip() for record in records]
        duplicates = {recipe_id for recipe_id, count in Counter(identities).items() if recipe_id and count > 1}
        expected = None if expected_recipe_ids is None else {
            str(value).strip() for value in expected_recipe_ids if str(value).strip()
        }
        repository_recipes = {self._identity(item): item for item in self.repository.listar() if self._identity(item)}
        proposal_columns = {f"{field}_propuesto": field for field in RECIPE_DOCUMENTATION_FIELDS}
        unknown_proposal_columns = sorted(
            header for header in headers if header.endswith("_propuesto") and header not in proposal_columns
        )
        source_default = self._source(source)
        validation_rows: list[dict[str, Any]] = []
        proposal_results: list[dict[str, Any]] = []
        field_totals = {
            "recibidos": 0, "utiles": 0, "operativos_agrupables": 0,
            "requieren_revision": 0,
            "rechazados": 0, "bloqueados_criticos": 0,
        }
        fields_summary = {
            field: {
                "politica": str(FIELD_CONTRACTS[field]["policy"]),
                "recibidos": 0, "utiles": 0, "operativos_agrupables": 0,
                "requieren_revision": 0,
                "rechazados": 0, "bloqueados_criticos": 0,
            }
            for field in sorted(RECIPE_DOCUMENTATION_FIELDS)
        }

        for index, record in enumerate(records, start=2):
            recipe_id = str(record.get("recipe_id") or "").strip()
            row_result: dict[str, Any] = {"fila": index, "recipe_id": recipe_id, "estado": "RECHAZADA", "errores": [], "avisos": []}
            if not recipe_id:
                row_result["errores"].append("ID_VACIO")
            elif recipe_id in duplicates:
                row_result["errores"].append("ID_DUPLICADO")
            elif expected is not None and recipe_id not in expected:
                row_result["errores"].append("FUERA_DE_LA_SELECCION_EXPORTADA")
            row_import_id = str(record.get("import_id") or "").strip()
            if metadata.get("version") == "0.3" and row_import_id != str(metadata.get("import_id") or ""):
                row_result["errores"].append("IMPORTACION_FILA_NO_COINCIDE")
            recipe = repository_recipes.get(recipe_id)
            if recipe_id and not recipe:
                row_result["errores"].append("RECETA_NO_EXISTE")
            exported_fingerprint = str(record.get("recipe_fingerprint") or "").strip()
            if recipe and exported_fingerprint != recipe_completion_fingerprint(recipe):
                row_result["errores"].append("RECETA_CAMBIO_DESDE_EXPORTACION")

            proposed: dict[str, Any] = {}
            impossible_fields: dict[str, str] = {}
            trace_by_field: dict[str, dict[str, Any]] = {}
            supplied_metadata: dict[str, dict[str, Any]] = {}
            try:
                supplied_metadata = self._proposal_metadata(record.get("metadatos_propuestas"))
            except RecipeCompletionExchangeError as exc:
                row_result["errores"].append(f"metadatos_propuestas:{exc.code}")
            for column, field in proposal_columns.items():
                value = record.get(column)
                if value in (None, ""):
                    continue
                field_totals["recibidos"] += 1
                fields_summary[field]["recibidos"] += 1
                trace_by_field[field] = {
                    "columna": column, "campo": field, "valor_bruto": value,
                    "tipo_bruto": type(value).__name__, "estado": "RECIBIDO",
                    "politica": fields_summary[field]["politica"], "motivos": [],
                }
                impossible_reason = self._impossible_reason(value)
                if impossible_reason is not None:
                    impossible_fields[field] = impossible_reason
                    field_totals["requieren_revision"] += 1
                    fields_summary[field]["requieren_revision"] += 1
                    trace_by_field[field]["estado"] = "IMPOSIBLE_ESTIMAR_RAZONABLEMENTE"
                    trace_by_field[field]["motivos"] = [impossible_reason]
                    continue
                try:
                    no_aplica_metadata = self._no_aplica_metadata(value)
                    if no_aplica_metadata is not None:
                        if field not in NO_APLICA_FIELDS:
                            raise RecipeCompletionExchangeError(
                                "no_aplica_no_permitido", "El campo no admite NO_APLICA.",
                            )
                        field_metadata = {
                            **no_aplica_metadata,
                            **dict(supplied_metadata.get(field) or {}),
                            "estado_campo": "NO_APLICA",
                        }
                        reason = str(
                            field_metadata.get("motivo") or field_metadata.get("evidencia") or ""
                        ).strip()
                        if not reason:
                            raise RecipeCompletionExchangeError(
                                "missing_no_aplica_reason",
                                "NO_APLICA requiere un motivo o evidencia por campo.",
                            )
                        supplied_metadata[field] = field_metadata
                    proposed[field] = self._typed_value(field, value)
                    if field == "alergenos":
                        allergen_metadata = supplied_metadata.get(field) or {}
                        if not str(allergen_metadata.get("motivo") or allergen_metadata.get("evidencia") or "").strip() \
                                or not str(allergen_metadata.get("fuente") or "").strip():
                            proposed.pop(field, None)
                            raise RecipeCompletionExchangeError(
                                "insufficient_allergen_provenance",
                                "Los alérgenos necesitan motivo y fuente explícitos.",
                            )
                    trace_by_field[field]["valor_normalizado"] = proposed[field]
                    trace_by_field[field]["estado"] = "NORMALIZADO"
                except RecipeCompletionExchangeError as exc:
                    field_totals["rechazados"] += 1
                    fields_summary[field]["rechazados"] += 1
                    row_result["errores"].append(f"{field}:{exc.code}")
                    trace_by_field[field]["estado"] = "RECHAZADO"
                    trace_by_field[field]["motivos"] = [exc.code]
            if not trace_by_field:
                row_result["avisos"].append("SIN_PROPUESTAS_RECIBIDAS")

            fatal_row_error = any(":" not in error for error in row_result["errores"])
            if not fatal_row_error and recipe and proposed:
                proposal = self.recipe_service.proposal(recipe_id=recipe_id, proposed=proposed)
                coverage = dict(proposal.get("completitud") or {})
                resolution = dict(coverage.get("resolucion_campos") or {})
                for field, reason in impossible_fields.items():
                    resolution[field] = {
                        "estado": "IMPOSIBLE_ESTIMAR_RAZONABLEMENTE",
                        "origen": None, "motivo": reason,
                    }
                if impossible_fields:
                    coverage["resolucion_campos"] = resolution
                    coverage["pendientes"] = sorted(set(coverage.get("pendientes") or []) | set(impossible_fields))
                    coverage["production_ready_provisional"] = False
                    readiness = dict(coverage.get("production_ready") or {})
                    readiness["provisional"] = False
                    readiness["bloqueos_provisionales"] = sorted(
                        set(readiness.get("bloqueos_provisionales") or []) | set(impossible_fields)
                    )
                    coverage["production_ready"] = readiness
                    proposal["completitud"] = coverage
                clean = dict(proposal.get("datos_propuestos_ia") or {})
                blocked = list(proposal.get("propuestas_bloqueadas_revision") or [])
                safe, grouped, individual = classify_recipe_proposals_for_review(clean)
                field_totals["utiles"] += len(safe)
                field_totals["operativos_agrupables"] += len(grouped)
                field_totals["requieren_revision"] += len(individual)
                field_totals["bloqueados_criticos"] += len(blocked)
                field_totals["rechazados"] += max(0, len(proposed) - len(clean) - len(blocked))
                for field in safe:
                    fields_summary[field]["utiles"] += 1
                    trace_by_field[field]["estado"] = "ACEPTADO_SELECCION_MASIVA"
                for field in grouped:
                    fields_summary[field]["operativos_agrupables"] += 1
                    trace_by_field[field]["estado"] = "ACEPTADO_REVISION_AGRUPADA"
                for field in individual:
                    fields_summary[field]["requieren_revision"] += 1
                    trace_by_field[field]["estado"] = "ACEPTADO_REVISION_INDIVIDUAL"
                for entry in blocked:
                    blocked_field = str(entry.get("campo") or "")
                    if blocked_field in fields_summary:
                        fields_summary[blocked_field]["bloqueados_criticos"] += 1
                        trace_by_field[blocked_field]["estado"] = "BLOQUEADO_CONTENIDO_CRITICO"
                        trace_by_field[blocked_field]["motivos"] = list(entry.get("motivos") or [])
                for field in set(proposed) - set(clean) - {
                    str(entry.get("campo") or "") for entry in blocked
                }:
                    fields_summary[field]["rechazados"] += 1
                    reasons = [
                        str(entry.get("reason") or "RECHAZADO_POR_POLITICA")
                        for entry in proposal.get("campos_descartados") or []
                        if str(entry.get("key") or "") == field
                    ]
                    trace_by_field[field]["estado"] = "RECHAZADO"
                    trace_by_field[field]["motivos"] = reasons or ["RECHAZADO_POR_POLITICA"]
                if blocked:
                    row_result["avisos"].append("CONTENIDO_CRITICO_REQUIERE_REVISION")
                row_result["campos_utiles"] = sorted(safe)
                row_result["campos_operativos_agrupables"] = sorted(grouped)
                row_result["campos_requieren_revision_individual"] = sorted(individual)
                row_result["campos_aceptados"] = sorted(clean)
                row_result["campos_bloqueados"] = blocked
                row_result["tiene_propuestas_utiles"] = bool(safe)
                row_result["tiene_propuestas"] = bool(clean)
                row_result["requiere_revision"] = bool(individual or blocked or row_result["errores"])
                row_result["campos_imposibles_estimar"] = [
                    {"campo": field, "motivo": reason}
                    for field, reason in impossible_fields.items()
                ]
                row_result["requiere_revision"] = bool(row_result["requiere_revision"] or impossible_fields)
                if impossible_fields and row_result["estado"] == "UTIL":
                    row_result["estado"] = "UTIL_Y_REQUIERE_REVISION"
                if clean:
                    if safe and row_result["requiere_revision"]:
                        row_result["estado"] = "UTIL_Y_REQUIERE_REVISION"
                    elif safe:
                        row_result["estado"] = "UTIL"
                    elif grouped and not row_result["requiere_revision"]:
                        row_result["estado"] = "OPERATIVA_AGRUPABLE"
                    else:
                        row_result["estado"] = "REQUIERE_REVISION"
                    row_source = self._source(record.get("origen_propuesta") or source_default)
                    proposal["metadatos_propuestas"] = {
                        field: self._normalized_metadata(
                            supplied_metadata.get(field), row_source=row_source,
                            fallback=(proposal.get("metadatos_propuestas") or {}).get(field),
                        )
                        for field in clean
                    }
                    row_result["origen_propuesta"] = row_source
                    proposal_results.append({
                        "recipe_id": recipe_id, "recipe_fingerprint": exported_fingerprint,
                        "proposal_result": proposal, "proposal_source": {
                            "fuente": "ARCHIVO_EXTERNO", "origen_externo": row_source,
                        },
                    })
                else:
                    row_result["avisos"].append("SIN_PROPUESTAS_UTILIZABLES")
                    row_result["estado"] = "REQUIERE_REVISION" if blocked else "RECHAZADA"
            elif not fatal_row_error:
                row_result["estado"] = "REQUIERE_REVISION"
                row_result["tiene_propuestas_utiles"] = False
                row_result["tiene_propuestas"] = False
                row_result["requiere_revision"] = True
            row_result["traza_campos"] = list(trace_by_field.values())
            row_result["rechazos_detallados"] = [
                {
                    "campo": trace.get("campo"), "estado": trace.get("estado"),
                    "motivos": list(trace.get("motivos") or []),
                    "mensajes": [
                        self._actionable_error(f"{trace.get('campo')}:{reason}")
                        for reason in trace.get("motivos") or []
                    ],
                }
                for trace in trace_by_field.values()
                if trace.get("estado") in {"RECHAZADO", "BLOQUEADO_CONTENIDO_CRITICO"}
            ]
            row_result["errores_detallados"] = [
                self._actionable_error(error) for error in row_result["errores"]
            ]
            row_result["avisos_detallados"] = [
                self._actionable_error(warning) for warning in row_result["avisos"]
            ]
            validation_rows.append(row_result)

        dynamic_candidates = self._proposed_price_candidates(proposal_results)
        price_validation = self._validate_price_references(
            workbook, dynamic_candidates=dynamic_candidates,
        )
        validation = {
            "scope": str(scope or metadata.get("scope") or "BIBLIOTECA"),
            "import_id": str(import_id or metadata.get("import_id") or ""),
            "filas_recibidas": len(records),
            "filas_utiles": sum(bool(row.get("tiene_propuestas_utiles")) for row in validation_rows),
            "filas_con_propuestas": sum(bool(row.get("tiene_propuestas")) for row in validation_rows),
            "filas_requieren_revision": sum(bool(row.get("requiere_revision")) for row in validation_rows),
            "filas_rechazadas": sum(row["estado"] == "RECHAZADA" for row in validation_rows),
            "campos": field_totals,
            "campos_por_nombre": {field: counts for field, counts in fields_summary.items() if counts["recibidos"]},
            "columnas_propuesta_desconocidas": unknown_proposal_columns,
            "filas": validation_rows, "referencias_precio": price_validation,
            "datos_reales_modificados": False,
        }
        file_receipt = {
            "nombre": filename, "tamano": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest(),
        }
        batch = self.batch_service.start_external(
            proposal_results=proposal_results,
            validation=validation,
            price_references=self._consolidated_price_references(price_validation),
            file_receipt=file_receipt,
            import_id=str(import_id or metadata.get("import_id") or ""),
            scope=str(scope or metadata.get("scope") or "BIBLIOTECA"),
        )
        return {
            "ok": True, "contrato": {
                "format": self.FORMAT, "version": self.VERSION,
                "version_importada": metadata.get("version"),
            },
            "archivo": file_receipt,
            "validacion": validation, "batch": batch, "datos_reales_modificados": False,
        }

    def _decode_file(self, filename: str, content_base64: str) -> bytes:
        safe_name = Path(str(filename or "")).name
        if not safe_name.lower().endswith(".xlsx") or safe_name != str(filename or ""):
            raise RecipeCompletionExchangeError("invalid_completion_filename", "Solo se admite un nombre de archivo XLSX seguro.")
        try:
            raw = base64.b64decode(str(content_base64 or ""), validate=True)
        except (ValueError, binascii.Error):
            raise RecipeCompletionExchangeError("invalid_completion_base64", "El contenido del archivo no es base64 valido.")
        if not raw or len(raw) > self.MAX_FILE_BYTES:
            raise RecipeCompletionExchangeError("invalid_completion_size", "El archivo esta vacio o supera el limite permitido.")
        return raw

    def _load_safe_workbook(self, raw: bytes):
        try:
            with ZipFile(BytesIO(raw)) as archive:
                members = archive.infolist()
                if len(members) > self.MAX_ZIP_MEMBERS or sum(item.file_size for item in members) > self.MAX_UNCOMPRESSED_BYTES:
                    raise RecipeCompletionExchangeError("unsafe_completion_archive", "El contenedor XLSX supera los limites de seguridad.")
            return load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        except RecipeCompletionExchangeError:
            raise
        except (BadZipFile, OSError, ValueError, KeyError):
            raise RecipeCompletionExchangeError("invalid_completion_xlsx", "El archivo XLSX esta corrupto o no es valido.")

    def _metadata(self, workbook) -> dict[str, str]:
        if "METADATA" not in workbook.sheetnames:
            raise RecipeCompletionExchangeError("missing_completion_metadata", "El archivo no contiene metadatos de contrato.")
        return {
            str(row[0].value or "").strip(): str(row[1].value or "").strip()
            for row in workbook["METADATA"].iter_rows(min_col=1, max_col=2)
            if row[0].value not in (None, "")
        }

    def _typed_value(self, field: str, value: Any) -> Any:
        if isinstance(value, str) and len(value) > self.MAX_TEXT_LENGTH:
            raise RecipeCompletionExchangeError("value_too_long", "El valor supera la longitud maxima.")
        no_aplica = self._no_aplica_metadata(value)
        if no_aplica is not None:
            value = {"estado": "NO_APLICA"}
        try:
            return self.recipe_service.normalize_proposed_value(field, value)
        except RecetaDocumentacionError as exc:
            raise RecipeCompletionExchangeError(exc.code, str(exc)) from exc

    @staticmethod
    def _impossible_reason(value: Any) -> str | None:
        if not isinstance(value, str):
            return None
        prefix = "PENDIENTE_IMPOSIBLE_DE_ESTIMAR"
        text = value.strip()
        if not text.upper().startswith(prefix):
            return None
        reason = text[len(prefix):].lstrip(" :-")
        return reason or "No existe base razonable para estimar este campo."

    def _proposal_metadata(self, value: Any) -> dict[str, dict[str, Any]]:
        if value in (None, ""):
            return {}
        if isinstance(value, str):
            if len(value) > self.MAX_METADATA_LENGTH:
                raise RecipeCompletionExchangeError("value_too_long", "Los metadatos superan la longitud maxima.")
            try:
                value = json.loads(value)
            except ValueError:
                raise RecipeCompletionExchangeError("invalid_proposal_metadata", "Los metadatos de propuestas no son JSON valido.")
        if not isinstance(value, dict):
            raise RecipeCompletionExchangeError("invalid_proposal_metadata", "Los metadatos de propuestas deben ser un objeto.")
        output: dict[str, dict[str, Any]] = {}
        for field, metadata in value.items():
            if field not in RECIPE_DOCUMENTATION_FIELDS or not isinstance(metadata, dict):
                continue
            output[field] = dict(metadata)
        return output

    @staticmethod
    def _no_aplica_metadata(value: Any) -> dict[str, Any] | None:
        parsed = value
        if isinstance(value, str):
            text = value.strip()
            if text.upper() == "NO_APLICA":
                return {}
            if not text.startswith("{"):
                return None
            try:
                parsed = json.loads(text)
            except ValueError:
                return None
        if isinstance(parsed, dict) and str(parsed.get("estado") or "").strip().upper() == "NO_APLICA":
            return dict(parsed)
        return None

    def _pending_price_articles(
        self,
        recipes: list[dict[str, Any]],
        contexts: dict[str, dict[str, Any]],
    ) -> list[dict[str, Any]]:
        """Consolida una sola búsqueda por artículo/candidato sin inventar identidad."""
        del recipes  # La trazabilidad se obtiene del contexto canónico por recipe_id.
        consolidated: dict[str, dict[str, Any]] = {}
        linked_identities = {
            (
                self._normalized_identity(article.get("nombre") or article.get("nombre_original")),
                self._unit_family(article.get("unidad_base") or article.get("unidad")),
            )
            for context in contexts.values()
            for article in (context.get("articulos_relacionados") or [])
            if isinstance(article, dict)
        }
        linked_names = {identity for identity, _ in linked_identities if identity}
        for recipe_id, context in contexts.items():
            for article in context.get("articulos_relacionados") or []:
                if not isinstance(article, dict):
                    continue
                article_id = str(article.get("article_id") or article.get("articulo_id") or "").strip()
                name = str(article.get("nombre") or article.get("nombre_original") or "").strip()
                if not article_id or not name:
                    continue
                if article.get("precio") not in (None, "") or article.get("precio_referencia") not in (None, ""):
                    continue
                row = consolidated.setdefault(article_id, {
                    "article_key": article_id, "article_id": article_id,
                    "nombre_canonico": name, "unidad_base": article.get("unidad_base") or "",
                    "recetas": set(), "estado_precio": "PRECIO_REFERENCIA_PENDIENTE",
                    "identidad_protegida": True,
                })
                row["recetas"].add(recipe_id)
            for ingredient in context.get("ingredientes_contexto") or []:
                if not isinstance(ingredient, dict):
                    continue
                name = str(
                    ingredient.get("nombre_original") or ingredient.get("name_raw")
                    or ingredient.get("nombre") or ingredient.get("ingrediente") or ""
                ).strip()
                normalized = self._normalized_identity(name)
                unit = str(
                    ingredient.get("unidad_normalizada") or ingredient.get("unidad")
                    or ingredient.get("unidad_original") or ""
                ).strip().casefold()
                family = self._unit_family(unit)
                if not normalized or normalized in linked_names or (normalized, family) in linked_identities:
                    continue
                key_seed = f"{normalized}|{family}"
                key = f"CANDIDATO-{hashlib.sha256(key_seed.encode('utf-8')).hexdigest()[:12].upper()}"
                row = consolidated.setdefault(key, {
                    "article_key": key, "article_id": "", "nombre_canonico": name,
                    "unidad_base": family, "recetas": set(),
                    "estado_precio": "PRECIO_REFERENCIA_PENDIENTE",
                    "identidad_protegida": False,
                })
                row["recetas"].add(recipe_id)
        result = []
        for row in consolidated.values():
            result.append({**row, "recetas": sorted(row["recetas"])})
        return sorted(result, key=lambda item: (str(item["nombre_canonico"]).casefold(), item["article_key"]))

    @staticmethod
    def _normalized_identity(value: Any) -> str:
        text = unicodedata.normalize("NFKD", str(value or ""))
        text = "".join(char for char in text if not unicodedata.combining(char)).casefold()
        return " ".join(re.findall(r"[a-z0-9]+", text))

    @staticmethod
    def _unit_family(value: Any) -> str:
        unit = str(value or "").strip().casefold()
        return "kg" if unit in {"kg", "g"} else "l" if unit in {"l", "ml", "cl"} else "u" if unit == "u" else ""

    @staticmethod
    def _append_price_sheets(workbook: Workbook, rows: list[dict[str, Any]]) -> None:
        pending = workbook.create_sheet("ARTICULOS_PENDIENTES")
        pending.append([
            "article_key", "article_id", "nombre_canonico", "unidad_base", "recetas_json",
            "estado_precio", "identidad_protegida", "instruccion",
        ])
        references = workbook.create_sheet("PRECIOS_REFERENCIA")
        references.append(list(PRICE_REFERENCE_COLUMNS))
        price_schema = workbook.create_sheet("SCHEMA_PRECIOS")
        price_schema.append(["campo", "descripcion", "tipo_esperado", "protegido"])
        for field in PRICE_REFERENCE_COLUMNS:
            description, expected_type, protected = PRICE_FIELD_CONTRACTS[field]
            price_schema.append([field, description, expected_type, protected])
        for row in rows:
            recipes_json = json.dumps(row["recetas"], ensure_ascii=False, separators=(",", ":"))
            pending.append([
                row["article_key"], row["article_id"], row["nombre_canonico"], row["unidad_base"],
                recipes_json, row["estado_precio"], row["identidad_protegida"],
                "No modificar identidad; buscar una sola referencia trazable si existe acceso web.",
            ])
            values = {column: "" for column in PRICE_REFERENCE_COLUMNS}
            values.update({
                "article_key": row["article_key"], "article_id": row["article_id"],
                "nombre_canonico": row["nombre_canonico"], "unidad_base": row["unidad_base"],
                "recetas_json": recipes_json, "procedencia": "REFERENCIA_EXTERNA",
                "estado_referencia": "PRECIO_REFERENCIA_PENDIENTE",
            })
            references.append([values[column] for column in PRICE_REFERENCE_COLUMNS])

    def _validate_price_references(
        self, workbook, *, dynamic_candidates: dict[str, dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        """Valida propuestas externas de precio sin persistirlas ni elevar su autoridad."""
        empty = {
            "filas_recibidas": 0, "referencias_utiles": 0, "pendientes": 0,
            "rechazadas": 0, "duplicadas": 0, "filas": [],
            "datos_reales_modificados": False,
        }
        has_pending = "ARTICULOS_PENDIENTES" in workbook.sheetnames
        has_references = "PRECIOS_REFERENCIA" in workbook.sheetnames
        if not has_pending and not has_references:
            return empty
        if has_pending != has_references:
            raise RecipeCompletionExchangeError(
                "invalid_price_reference_schema",
                "ARTICULOS_PENDIENTES y PRECIOS_REFERENCIA deben conservarse juntas.",
            )

        pending_rows = self._sheet_records(
            workbook["ARTICULOS_PENDIENTES"], "ARTICULOS_PENDIENTES",
            expected_headers=(
                "article_key", "article_id", "nombre_canonico", "unidad_base", "recetas_json",
                "estado_precio", "identidad_protegida", "instruccion",
            ),
        )
        pending: dict[str, dict[str, Any]] = {}
        for _, record in pending_rows:
            key = str(record.get("article_key") or "").strip()
            if not key or key in pending:
                raise RecipeCompletionExchangeError(
                    "invalid_price_reference_schema", "article_key debe ser único y no vacío.",
                )
            pending[key] = {
                "article_id": str(record.get("article_id") or "").strip(),
                "nombre_canonico": str(record.get("nombre_canonico") or "").strip(),
                "unidad_base": str(record.get("unidad_base") or "").strip().casefold(),
                "recetas": self._json_string_list(record.get("recetas_json")),
            }
        pending.update({
            key: deepcopy(value) for key, value in dict(dynamic_candidates or {}).items()
            if key not in pending
        })

        reference_rows = self._sheet_records(
            workbook["PRECIOS_REFERENCIA"], "PRECIOS_REFERENCIA",
            expected_headers=PRICE_REFERENCE_COLUMNS,
        )
        results: list[dict[str, Any]] = []
        seen: set[str] = set()
        useful = rejected = duplicates = awaiting = 0
        proposal_fields = {
            "producto_encontrado", "comercio_fuente", "marca", "precio_observado",
            "moneda", "formato_envase", "cantidad_envase", "unidad_envase",
            "precio_normalizado", "unidad_precio_normalizado", "url_fuente",
            "fecha_consulta", "observacion_equivalencia", "confianza", "price_basis",
        }
        for row_number, record in reference_rows:
            key = str(record.get("article_key") or "").strip()
            state = str(record.get("estado_referencia") or "").strip().upper()
            has_proposal = any(record.get(field) not in (None, "") for field in proposal_fields)
            row_result: dict[str, Any] = {
                "fila": row_number, "article_key": key, "estado": "RECHAZADA",
                "errores": [], "referencia": None,
            }
            published = pending.get(key)
            if not key or published is None:
                row_result["errores"].append("ARTICLE_KEY_DESCONOCIDA")
            elif key in seen:
                row_result["errores"].append("REFERENCIA_DUPLICADA")
                duplicates += 1
            else:
                seen.add(key)
            if published:
                if str(record.get("article_id") or "").strip() != published["article_id"]:
                    row_result["errores"].append("ARTICLE_ID_MODIFICADO")
                if self._normalized_identity(record.get("nombre_canonico")) != self._normalized_identity(published["nombre_canonico"]):
                    row_result["errores"].append("NOMBRE_CANONICO_MODIFICADO")
                supplied_unit = str(record.get("unidad_base") or "").strip().casefold()
                if published["unidad_base"] and supplied_unit != published["unidad_base"]:
                    row_result["errores"].append("UNIDAD_BASE_MODIFICADA")
                elif not published["unidad_base"] and supplied_unit not in {"", *PRICE_NORMALIZED_UNITS}:
                    row_result["errores"].append("UNIDAD_BASE_MODIFICADA")
            if not has_proposal and not row_result["errores"] and state in {"", "PRECIO_REFERENCIA_PENDIENTE"}:
                row_result["estado"] = "PRECIO_REFERENCIA_PENDIENTE"
                awaiting += 1
            elif not row_result["errores"]:
                try:
                    row_result["referencia"] = {
                        **self._normalize_price_reference(record),
                        "recetas": list(published.get("recetas") or []),
                    }
                    row_result["estado"] = "REFERENCIA_PROPUESTA"
                    useful += 1
                except RecipeCompletionExchangeError as exc:
                    row_result["errores"].append(exc.code)
            if row_result["errores"]:
                rejected += 1
                row_result["errores_detallados"] = [self._actionable_error(code) for code in row_result["errores"]]
            results.append(row_result)
        response = {
            "filas_recibidas": len(reference_rows), "referencias_utiles": useful,
            "pendientes": awaiting, "rechazadas": rejected, "duplicadas": duplicates,
            "filas": results, "datos_reales_modificados": False,
        }
        response["referencias_consolidadas"] = len(self._consolidated_price_references(response))
        return response

    def _proposed_price_candidates(
        self, proposal_results: list[dict[str, Any]],
    ) -> dict[str, dict[str, Any]]:
        """Autoriza solo claves legibles que correspondan a candidatos realmente propuestos."""
        consolidated: dict[str, dict[str, Any]] = {}
        for proposal_row in proposal_results:
            recipe_id = str(proposal_row.get("recipe_id") or "")
            proposal = dict(proposal_row.get("proposal_result") or {})
            ingredients = (proposal.get("datos_propuestos_ia") or {}).get("ingredientes_estructurados") or []
            for ingredient in ingredients:
                if not isinstance(ingredient, dict) or str(ingredient.get("estado_relacion") or "").upper() != "CANDIDATO_NUEVO":
                    continue
                name = str(ingredient.get("nombre_original") or ingredient.get("name_raw") or "").strip()
                unit = self._unit_family(
                    ingredient.get("unidad_normalizada") or ingredient.get("unidad")
                )
                if not name or not unit:
                    continue
                key = self._dynamic_candidate_key(name, unit)
                row = consolidated.setdefault(key, {
                    "article_id": "", "nombre_canonico": name,
                    "unidad_base": unit, "recetas": [],
                })
                if recipe_id and recipe_id not in row["recetas"]:
                    row["recetas"].append(recipe_id)
        return consolidated

    def _consolidated_price_references(self, validation: dict[str, Any]) -> list[dict[str, Any]]:
        """Una referencia por nombre exacto y familia, priorizando identidad canÃ³nica."""
        selected: dict[tuple[str, str], dict[str, Any]] = {}
        for row in validation.get("filas") or []:
            reference = row.get("referencia") if isinstance(row, dict) else None
            if not isinstance(reference, dict):
                continue
            identity = (
                self._normalized_identity(reference.get("nombre_canonico")),
                self._unit_family(reference.get("unidad_normalizada")),
            )
            current = selected.get(identity)
            if current is None or (reference.get("article_id") and not current.get("article_id")):
                selected[identity] = deepcopy(reference)
            elif current:
                current["recetas"] = sorted(set(current.get("recetas") or []) | set(reference.get("recetas") or []))
        return sorted(selected.values(), key=lambda item: (
            self._normalized_identity(item.get("nombre_canonico")),
            str(item.get("article_key") or ""),
        ))

    @classmethod
    def _dynamic_candidate_key(cls, name: Any, unit: Any) -> str:
        return f"CANDIDATO_NUEVO|{cls._normalized_identity(name)}|{cls._unit_family(unit)}"

    @staticmethod
    def _json_string_list(value: Any) -> list[str]:
        try:
            parsed = json.loads(str(value or "[]"))
        except (TypeError, ValueError):
            return []
        return [str(item) for item in parsed if str(item).strip()] if isinstance(parsed, list) else []

    def _sheet_records(
        self, sheet, name: str, *, expected_headers: tuple[str, ...] | None = None,
    ) -> list[tuple[int, dict[str, Any]]]:
        rows = list(islice(sheet.iter_rows(), self.MAX_ROWS + 2))
        if len(rows) > self.MAX_ROWS + 1 or (rows and max(len(row) for row in rows) > self.MAX_COLUMNS):
            raise RecipeCompletionExchangeError("completion_sheet_too_large", f"La hoja {name} supera los límites permitidos.")
        if not rows:
            raise RecipeCompletionExchangeError("invalid_price_reference_schema", f"La hoja {name} está vacía.")
        headers = [str(cell.value or "").strip() for cell in rows[0]]
        if len(headers) != len(set(headers)):
            raise RecipeCompletionExchangeError("duplicate_completion_headers", f"La hoja {name} contiene cabeceras duplicadas.")
        if expected_headers is not None and tuple(headers) != tuple(expected_headers):
            raise RecipeCompletionExchangeError("invalid_price_reference_schema", f"La hoja {name} no conserva el contrato público.")
        if any(cell.data_type == "f" for row in rows[1:] for cell in row):
            raise RecipeCompletionExchangeError("unsafe_formula", f"La hoja {name} contiene fórmulas.")
        return [
            (row_number, dict(zip(headers, [cell.value for cell in row])))
            for row_number, row in enumerate(rows[1:], start=2)
            if any(cell.value not in (None, "") for cell in row)
        ]

    def _normalize_price_reference(self, record: dict[str, Any]) -> dict[str, Any]:
        required_text = ("producto_encontrado", "comercio_fuente", "url_fuente", "fecha_consulta")
        if any(not str(record.get(field) or "").strip() for field in required_text):
            raise RecipeCompletionExchangeError("invalid_price_reference", "Faltan producto, comercio, URL o fecha.")
        if str(record.get("procedencia") or "").strip().upper() != "REFERENCIA_EXTERNA":
            raise RecipeCompletionExchangeError("invalid_price_authority", "La procedencia debe ser REFERENCIA_EXTERNA.")
        if str(record.get("estado_referencia") or "").strip().upper() != "REFERENCIA_PROPUESTA":
            raise RecipeCompletionExchangeError("invalid_price_reference_state", "El estado debe ser REFERENCIA_PROPUESTA.")
        basis = str(record.get("price_basis") or "").strip().upper()
        if basis not in PRICE_BASIS_VALUES:
            raise RecipeCompletionExchangeError("invalid_price_basis", "price_basis no está admitido.")
        unit = str(record.get("unidad_envase") or "").strip().casefold()
        normalized_unit = str(record.get("unidad_precio_normalizado") or "").strip().casefold()
        unit_map = {"g": ("kg", 1000.0), "kg": ("kg", 1.0), "ml": ("l", 1000.0), "l": ("l", 1.0), "u": ("u", 1.0)}
        if unit not in unit_map or normalized_unit not in PRICE_NORMALIZED_UNITS or unit_map[unit][0] != normalized_unit:
            raise RecipeCompletionExchangeError("invalid_price_unit", "La unidad comercial no normaliza a la unidad publicada.")
        observed = self._finite_positive(record.get("precio_observado"), "precio_observado")
        amount = self._finite_positive(record.get("cantidad_envase"), "cantidad_envase")
        normalized = self._finite_positive(record.get("precio_normalizado"), "precio_normalizado")
        expected = observed / (amount / unit_map[unit][1])
        if abs(normalized - expected) > max(0.01, expected * 0.02):
            raise RecipeCompletionExchangeError("inconsistent_normalized_price", "El precio normalizado no coincide con precio y formato.")
        try:
            confidence = float(record.get("confianza"))
        except (TypeError, ValueError):
            raise RecipeCompletionExchangeError("invalid_confidence", "La confianza debe ser numérica.")
        if not math.isfinite(confidence) or confidence < 0 or confidence > 1:
            raise RecipeCompletionExchangeError("invalid_confidence", "La confianza debe estar entre 0 y 1.")
        parsed_url = urlparse(str(record.get("url_fuente") or "").strip())
        if parsed_url.scheme not in {"http", "https"} or not parsed_url.netloc:
            raise RecipeCompletionExchangeError("invalid_source", "La URL de referencia no es HTTP(S) válida.")
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(record.get("fecha_consulta") or "").strip()):
            raise RecipeCompletionExchangeError("invalid_reference_date", "Use fecha ISO YYYY-MM-DD.")
        return {
            "article_key": str(record.get("article_key") or "").strip(),
            "article_id": str(record.get("article_id") or "").strip(),
            "nombre_canonico": str(record.get("nombre_canonico") or "").strip(),
            "producto": str(record.get("producto_encontrado") or "").strip(),
            "tienda_referencia": str(record.get("comercio_fuente") or "").strip(),
            "marca": str(record.get("marca") or "").strip(),
            "precio_comercial": observed, "moneda": str(record.get("moneda") or "EUR").strip().upper(),
            "formato_comercial": str(record.get("formato_envase") or "").strip(),
            "cantidad_formato": amount, "unidad_formato": unit,
            "precio_normalizado": normalized, "unidad_normalizada": normalized_unit,
            "url": str(record.get("url_fuente") or "").strip(),
            "consultado_en": str(record.get("fecha_consulta") or "").strip(),
            "evidencia": str(record.get("observacion_equivalencia") or "").strip(),
            "confianza": confidence, "price_basis": basis,
            "origen": "REFERENCIA_EXTERNA", "autoridad": "REFERENCIA_NO_REAL",
            "estado_revision": "REQUIERE_REVISION_HUMANA",
        }

    @staticmethod
    def _finite_positive(value: Any, field: str) -> float:
        try:
            number = float(str(value).replace(",", "."))
        except (TypeError, ValueError):
            raise RecipeCompletionExchangeError("invalid_price", f"{field} debe ser numérico.")
        if not math.isfinite(number) or number <= 0:
            raise RecipeCompletionExchangeError("invalid_price", f"{field} debe ser finito y mayor que cero.")
        return number

    @classmethod
    def _schema_type(cls, field: str) -> str:
        kind = str(FIELD_CONTRACTS[field]["type"])
        return {
            "boolean": "boolean", "number": "number>0", "percentage": "number[0,1)",
            "duration": "duration<string>", "enum": "enum<string>", "json": "json",
        }.get(kind, "string")

    @staticmethod
    def _schema_units(field: str) -> list[str]:
        return list(FIELD_CONTRACTS[field]["units"])

    @staticmethod
    def _schema_restrictions(field: str) -> str:
        return str(FIELD_CONTRACTS[field]["validations"])

    @staticmethod
    def _schema_shape(field: str) -> Any:
        return deepcopy(FIELD_CONTRACTS[field]["shape"])

    @staticmethod
    def _actionable_error(error: str) -> str:
        field, separator, code = str(error or "").partition(":")
        if not separator:
            field, code = "", field
        detail = ERROR_MESSAGES.get(code, code.replace("_", " ").lower())
        return f"{field}: {detail}" if field else detail

    def _normalized_metadata(
        self, value: dict[str, Any] | None, *, row_source: str,
        fallback: dict[str, Any] | None,
    ) -> dict[str, Any]:
        metadata = {**dict(fallback or {}), **dict(value or {})}
        origin = str(metadata.get("origen") or (
            "CONFIRMADO" if row_source == "HUMANO" else "IA_PROPUESTA"
        )).strip().upper()
        if origin not in {
            "DOCUMENTO", "REAL", "CALCULADO", "CONTEXTO_INTERNO", "IA_PROPUESTA",
            "REFERENCIA_EXTERNA", "CONFIRMADO",
        }:
            origin = "IA_PROPUESTA"
        confidence = metadata.get("confianza")
        try:
            confidence = max(0.0, min(1.0, float(confidence))) if confidence not in (None, "") else None
        except (TypeError, ValueError):
            confidence = None
        normalized = {
            "origen": origin,
            "confianza": confidence,
            "motivo": str(metadata.get("motivo") or metadata.get("evidencia") or "Propuesta recibida en paquete externo."),
            "estado_revision": "REQUIERE_REVISION_HUMANA",
            "fuente": str(metadata.get("fuente") or "ARCHIVO_EXTERNO"),
            "url": str(metadata.get("url") or "") or None,
            "fecha": str(metadata.get("fecha") or datetime.now(timezone.utc).isoformat(timespec="seconds")),
            "modelo": str(metadata.get("modelo") or metadata.get("proveedor") or row_source),
        }
        if str(metadata.get("estado_campo") or "").upper() == "NO_APLICA":
            normalized["estado_campo"] = "NO_APLICA"
        return normalized

    def _source(self, value: Any) -> str:
        source = str(value or "ARCHIVO_EXTERNO").strip().upper()
        return source if source in self.EXTERNAL_ORIGINS else "ARCHIVO_EXTERNO"

    @staticmethod
    def _cell_value(value: Any) -> str:
        if value in (None, ""):
            return ""
        if isinstance(value, (dict, list, tuple, bool, int, float)):
            return json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
        return str(value)

    @staticmethod
    def _missing(recipe: dict[str, Any]) -> list[str]:
        return list((recipe.get("completitud") or {}).get("campos_obligatorios_pendientes") or [])

    def _pending_keys(self, recipe: dict[str, Any]) -> list[str]:
        return self.recipe_service.pending_field_keys(recipe)

    @staticmethod
    def _identity(recipe: dict[str, Any]) -> str:
        return str(recipe.get("id") or recipe.get("codigo") or "")


__all__ = ["RecipeCompletionExchangeService", "RecipeCompletionExchangeError"]
