from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, List, Optional
from datetime import datetime
import json


@dataclass
class RegistroStockInicial:
    codigo: str
    articulo: str
    proveedor: Optional[str]
    familia: Optional[str]
    unidad: str
    stock_actual: float
    stock_minimo: float
    ubicacion: Optional[str]
    observaciones_stock: Optional[str]
    fecha_importacion: str


@dataclass
class ResultadoImportacionStockInicial:
    total_filas: int
    importados: int
    actualizados: int
    ignorados: int
    errores: int
    ruta_destino: str
    mensajes: List[str]


class ImportadorStockInicial432:
    """
    Importa stock inicial desde la plantilla generada en 4.3.1.

    Entrada:
    DATOS/stock_inicial_4_3_1.xlsx

    Salida:
    DATOS/db/stock_inicial.json
    """

    HOJA = "Stock inicial"

    def importar_desde_excel(
        self,
        ruta_excel: str,
        ruta_destino: str = "DATOS/db/stock_inicial.json",
    ) -> ResultadoImportacionStockInicial:
        filas = self._leer_excel(ruta_excel)
        return self.importar_filas(filas, ruta_destino)

    def importar_filas(
        self,
        filas: List[Dict[str, Any]],
        ruta_destino: str = "DATOS/db/stock_inicial.json",
    ) -> ResultadoImportacionStockInicial:
        ruta = Path(ruta_destino)
        ruta.parent.mkdir(parents=True, exist_ok=True)

        existentes = self._leer_json_lista(ruta)
        index = {item.get("codigo"): item for item in existentes if item.get("codigo")}

        importados = 0
        actualizados = 0
        ignorados = 0
        errores = 0
        mensajes: List[str] = []

        for numero, fila in enumerate(filas, start=2):
            codigo = self._texto(fila.get("Codigo"))
            articulo = self._texto(fila.get("Articulo"))
            unidad = self._texto(fila.get("Unidad"))
            stock_actual = self._numero(fila.get("Stock actual"))
            stock_minimo = self._numero(fila.get("Stock mínimo"))

            # Si la fila no tiene stock ni unidad, se considera no rellenada.
            if not unidad and stock_actual is None and stock_minimo is None:
                ignorados += 1
                continue

            if not codigo:
                errores += 1
                mensajes.append(f"Fila {numero}: falta código.")
                continue

            if not articulo:
                errores += 1
                mensajes.append(f"Fila {numero}: falta artículo.")
                continue

            if not unidad:
                errores += 1
                mensajes.append(f"Fila {numero}: falta unidad.")
                continue

            if stock_actual is None:
                errores += 1
                mensajes.append(f"Fila {numero}: stock actual inválido o vacío.")
                continue

            if stock_minimo is None:
                stock_minimo = 0.0

            if stock_actual < 0:
                errores += 1
                mensajes.append(f"Fila {numero}: stock actual negativo.")
                continue

            if stock_minimo < 0:
                errores += 1
                mensajes.append(f"Fila {numero}: stock mínimo negativo.")
                continue

            registro = RegistroStockInicial(
                codigo=codigo,
                articulo=articulo,
                proveedor=self._texto(fila.get("Proveedor")),
                familia=self._texto(fila.get("Familia")),
                unidad=unidad,
                stock_actual=stock_actual,
                stock_minimo=stock_minimo,
                ubicacion=self._texto(fila.get("Ubicación")),
                observaciones_stock=self._texto(fila.get("Observaciones stock")),
                fecha_importacion=datetime.now().isoformat(timespec="seconds"),
            )

            datos = asdict(registro)

            if codigo in index:
                index[codigo].update(datos)
                actualizados += 1
            else:
                index[codigo] = datos
                importados += 1

        ordenados = sorted(index.values(), key=lambda item: item.get("codigo", ""))
        ruta.write_text(json.dumps(ordenados, ensure_ascii=False, indent=2), encoding="utf-8")

        mensajes.append(
            f"Importación stock inicial completada. Importados: {importados}. "
            f"Actualizados: {actualizados}. Ignorados: {ignorados}. Errores: {errores}."
        )

        return ResultadoImportacionStockInicial(
            total_filas=len(filas),
            importados=importados,
            actualizados=actualizados,
            ignorados=ignorados,
            errores=errores,
            ruta_destino=str(ruta),
            mensajes=mensajes,
        )

    def _leer_excel(self, ruta_excel: str) -> List[Dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para leer la plantilla de stock.") from exc

        ruta = Path(ruta_excel)
        if not ruta.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta_excel}")

        wb = load_workbook(ruta, data_only=True)

        if self.HOJA not in wb.sheetnames:
            raise ValueError(f"No existe la hoja '{self.HOJA}'. Hojas disponibles: {wb.sheetnames}")

        ws = wb[self.HOJA]
        headers = [cell.value for cell in ws[1]]
        filas: List[Dict[str, Any]] = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            filas.append({str(headers[i]).strip(): row[i] if i < len(row) else None for i in range(len(headers))})

        return filas

    def _leer_json_lista(self, ruta: Path) -> List[Dict[str, Any]]:
        if not ruta.exists():
            return []
        try:
            contenido = json.loads(ruta.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return []
        return contenido if isinstance(contenido, list) else []

    def _texto(self, valor: Any) -> Optional[str]:
        if valor is None:
            return None
        texto = str(valor).strip()
        return texto if texto else None

    def _numero(self, valor: Any) -> Optional[float]:
        if valor is None or str(valor).strip() == "":
            return None

        if isinstance(valor, (int, float)):
            return float(valor)

        limpio = str(valor).replace(" ", "").replace("€", "")
        if "," in limpio and "." in limpio:
            limpio = limpio.replace(".", "").replace(",", ".")
        else:
            limpio = limpio.replace(",", ".")

        try:
            return float(limpio)
        except ValueError:
            return None


__all__ = [
    "ImportadorStockInicial432",
    "RegistroStockInicial",
    "ResultadoImportacionStockInicial",
]
