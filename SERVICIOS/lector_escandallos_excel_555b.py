from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from MODELOS.importacion_escandallos_excel_555b import (
    HojaEscandallosDetectada,
    ResultadoVistaPreviaExcel555B,
)
from SERVICIOS.perfil_mapeo_escandallos_555b import sugerir_mapeo


class LectorEscandallosExcel555B:
    """Lectura y diagnóstico de Excel en modo estrictamente solo lectura.

    Esta primera parte no importa escandallos ni escribe en la base de Host AI.
    Detecta hojas, cabeceras probables, columnas y propone un mapeo canónico.
    """

    EXTENSIONES_PERMITIDAS = {".xlsx", ".xlsm"}

    def analizar(self, ruta_archivo: str | Path, filas_preview: int = 8) -> dict[str, Any]:
        ruta = Path(ruta_archivo)
        resultado = ResultadoVistaPreviaExcel555B(archivo=str(ruta))

        if not ruta.exists():
            resultado.errores.append(f"No se encuentra el archivo: {ruta}")
            return resultado.to_dict()
        if ruta.suffix.lower() not in self.EXTENSIONES_PERMITIDAS:
            resultado.errores.append(
                f"Formato no compatible en esta fase: {ruta.suffix or 'sin extensión'}. "
                "Use .xlsx o .xlsm."
            )
            return resultado.to_dict()

        try:
            libro = load_workbook(ruta, read_only=True, data_only=True)
        except Exception as exc:
            resultado.errores.append(f"No se pudo abrir el Excel: {exc}")
            return resultado.to_dict()

        try:
            for hoja in libro.worksheets:
                detectada = self._analizar_hoja(hoja, filas_preview=max(1, filas_preview))
                resultado.hojas.append(detectada)
            if not resultado.hojas:
                resultado.avisos.append("El libro no contiene hojas visibles analizables.")
            elif resultado.hojas_candidatas == 0:
                resultado.avisos.append(
                    "No se ha detectado ninguna hoja con los cuatro campos mínimos: "
                    "receta, ingrediente, cantidad y unidad."
                )
        finally:
            libro.close()

        return resultado.to_dict()

    def _analizar_hoja(self, hoja, filas_preview: int) -> HojaEscandallosDetectada:
        filas = list(hoja.iter_rows(min_row=1, max_row=min(hoja.max_row or 1, 40), values_only=True))
        fila_cabecera = self._detectar_fila_cabecera(filas)
        valores_cabecera = filas[fila_cabecera - 1] if filas else tuple()
        columnas = self._normalizar_cabeceras(valores_cabecera)
        mapeo, confianza, faltantes = sugerir_mapeo(columnas)

        vista_previa: list[dict[str, Any]] = []
        filas_con_datos = 0
        inicio_datos = fila_cabecera + 1
        for numero_fila, valores in enumerate(
            hoja.iter_rows(min_row=inicio_datos, values_only=True), start=inicio_datos
        ):
            if not any(valor not in (None, "") for valor in valores):
                continue
            filas_con_datos += 1
            if len(vista_previa) < filas_preview:
                registro = {
                    columnas[indice]: valores[indice] if indice < len(valores) else None
                    for indice in range(len(columnas))
                }
                registro["__fila_excel__"] = numero_fila
                vista_previa.append(registro)

        return HojaEscandallosDetectada(
            nombre=hoja.title,
            fila_cabecera=fila_cabecera,
            columnas=columnas,
            filas_con_datos=filas_con_datos,
            vista_previa=vista_previa,
            mapeo_sugerido=mapeo,
            campos_faltantes=faltantes,
            confianza=confianza,
            candidata_escandallos=not faltantes and filas_con_datos > 0,
        )

    @staticmethod
    def _detectar_fila_cabecera(filas: list[tuple[Any, ...]]) -> int:
        mejor_fila = 1
        mejor_puntuacion = -1.0
        for indice, fila in enumerate(filas[:20], start=1):
            valores = [str(valor).strip() for valor in fila if valor not in (None, "")]
            if not valores:
                continue
            _, confianza, faltantes = sugerir_mapeo(valores)
            puntuacion = confianza + ((4 - len(faltantes)) * 15) + min(len(valores), 12)
            if puntuacion > mejor_puntuacion:
                mejor_puntuacion = puntuacion
                mejor_fila = indice
        return mejor_fila

    @staticmethod
    def _normalizar_cabeceras(valores: tuple[Any, ...]) -> list[str]:
        columnas: list[str] = []
        usados: dict[str, int] = {}
        for indice, valor in enumerate(valores, start=1):
            base = str(valor).strip() if valor not in (None, "") else f"COLUMNA_{indice}"
            contador = usados.get(base, 0) + 1
            usados[base] = contador
            columnas.append(base if contador == 1 else f"{base}_{contador}")
        while columnas and columnas[-1].startswith("COLUMNA_"):
            columnas.pop()
        return columnas
