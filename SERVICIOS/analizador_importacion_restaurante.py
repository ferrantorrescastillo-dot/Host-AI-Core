from __future__ import annotations

import base64
import csv
import io
import json
import re
import tempfile
import unicodedata
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from SERVICIOS.import_ambiguity_resolver import ImportAmbiguityResolver, resolve_regions_once


def _norm(value: Any) -> str:
    text = str(value or "").lower()
    if any(marker in text for marker in ("Ã", "Â", "â")):
        try:
            text = text.encode("latin-1").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    text = unicodedata.normalize("NFKD", text)
    return " ".join(re.sub(r"[^a-z0-9]+", " ", "".join(
        char for char in text if not unicodedata.combining(char)
    )).split())


FIELD_ALIASES = {
    "codigo": {"codigo", "codigo articulo", "referencia"},
    "nombre": {"nombre", "producto", "articulo", "nombre producto", "descripcion"},
    "familia": {"familia", "categoria", "grupo"},
    "receta": {"receta", "elaboracion", "plato", "nombre receta"},
    "ingrediente": {"ingrediente", "componente", "materia prima"},
    "cantidad": {"cantidad", "qty", "cant", "racion", "bruto", "neto", "kg total"},
    "unidad": {"unidad", "ud", "uom"},
    "procedimiento": {"procedimiento", "preparacion", "pasos", "metodo"},
    "rendimiento": {"rendimiento", "raciones", "porciones"},
    "proveedor": {"proveedor", "supplier"},
    "formato": {"formato", "formato compra", "envase"},
    "precio_compra": {"precio compra", "coste compra", "ultimo precio"},
    "precio_ambiguo": {"precio", "coste", "un", "precio unitario", "coste unitario"},
    "pvp": {"pvp", "precio venta", "precio carta"},
}


@dataclass
class ParsedTable:
    source: str
    sheet: str
    rows: list[dict[str, Any]]
    mapping: list[dict[str, Any]]
    kind: str
    region_id: str = "REGION-1"
    start_row: int = 1
    end_row: int = 1
    header_row: int | None = 1
    context_title: str = ""
    dimensions: str = ""
    nonempty_cells: int = 0
    merged_cells: list[str] | None = None
    confidence: float = 0.0
    reason: str = ""


class RestaurantImportAnalysisError(ValueError):
    pass


class RestaurantDataImportAnalyzer:
    """Analiza exportaciones tabulares en memoria; no conoce repositorios ni escribe datos."""

    ALLOWED = {".xlsx", ".csv", ".tsv", ".json", ".txt", ".md"}
    MAX_FILE_BYTES = 10 * 1024 * 1024
    MAX_TOTAL_BYTES = 30 * 1024 * 1024

    def __init__(self, ambiguity_resolver: ImportAmbiguityResolver | None = None) -> None:
        self.ambiguity_resolver = ambiguity_resolver

    def analyze(self, payload: dict[str, Any]) -> dict[str, Any]:
        inputs = list(payload.get("archivos") or [])
        pasted = str(payload.get("texto_pegado") or "").strip()
        if not inputs and pasted:
            inputs = [{"nombre": "datos-pegados.txt", "tipo_mime": "text/plain", "texto": pasted}]
        if not inputs:
            raise RestaurantImportAnalysisError("Añade al menos un archivo o pega datos.")
        if sum(self._size(item) for item in inputs) > self.MAX_TOTAL_BYTES:
            raise RestaurantImportAnalysisError("La sesión supera el límite total de 30 MB.")

        tables: list[ParsedTable] = []
        files: list[dict[str, Any]] = []
        for item in inputs:
            name = Path(str(item.get("nombre") or "datos.txt")).name
            content, text = self._content(item)
            if len(content) > self.MAX_FILE_BYTES:
                raise RestaurantImportAnalysisError(f"{name} supera el límite de 10 MB.")
            parsed = self._parse(name, content, text)
            tables.extend(parsed)
            files.append({
                "nombre": name, "tipo": Path(name).suffix.lower().lstrip("."),
                "tamano": len(content) or len(text.encode("utf-8")),
                "estado": "ANALIZADO", "hojas": len({table.sheet for table in parsed}),
                "filas": sum(len(table.rows) for table in parsed),
            })

        recipes = self._recipes(tables)
        menus = self._menus(tables)
        articles = self._named_rows(tables, "ARTICULOS")
        suppliers = self._named_rows(tables, "PROVEEDORES")
        normalized_articles = {_norm(item["nombre"]) for item in articles if item.get("nombre")}
        possible_duplicates = sorted({
            f"{recipe['nombre']}: {_norm(ingredient.get('nombre_original'))}"
            for recipe in recipes
            for ingredient in recipe["ingredientes_estructurados"]
            if sum(
                _norm(other.get("nombre_original")) == _norm(ingredient.get("nombre_original"))
                for other in recipe["ingredientes_estructurados"]
            ) > 1
        })
        recipe_names = {_norm(item.get("nombre")): item for item in recipes}
        same_name = sorted(normalized_articles.intersection(recipe_names))
        subdependencies = []
        for recipe in recipes:
            for ingredient in recipe["ingredientes_estructurados"]:
                target = recipe_names.get(_norm(ingredient.get("nombre_original")))
                if target and _norm(target.get("nombre")) != _norm(recipe.get("nombre")):
                    subdependencies.append({
                        "elaboracion": recipe["nombre"], "depende_de": target["nombre"],
                        "tipo_propuesto": "SUBELABORACION",
                    })
        sub_names = {_norm(item["depende_de"]) for item in subdependencies}
        for recipe in recipes:
            recipe["tipo_entidad_propuesto"] = (
                "SUBELABORACION" if _norm(recipe.get("nombre")) in sub_names
                else "ELABORACION_INTERNA"
            )
        price_doubts = [
            {"archivo": table.source, "hoja": table.sheet, "columna": item["columna"],
             "motivo": "Precio ambiguo: debe elegirse precio de compra, PVP o coste derivado."}
            for table in tables for item in table.mapping if item["destino"] == "precio_ambiguo"
        ]
        relation_doubts = [{
            "nombre": name, "tipo": "REQUIERE_REVISION",
            "motivo": "Mismo nombre en artículos y recetas no demuestra que sean la misma entidad.",
        } for name in same_name]
        mapped = [item for table in tables for item in table.mapping]
        ambiguous_regions = [self._ambiguity_region(table) for table in tables
                             if table.kind in {"DESCONOCIDA", "DESCONOCIDO"} or any(
                                 item["destino"] in {"precio_ambiguo", "sin_mapear"}
                                 for item in table.mapping)]
        ai_proposals: list[dict[str, Any]] = []
        ai_calls = 0
        if payload.get("resolver_ambiguedades_ia") and self.ambiguity_resolver:
            ai_proposals, ai_calls = resolve_regions_once(ambiguous_regions, self.ambiguity_resolver)
        unique_ambiguous_concepts = len({
            item.get("layout_signature")
            for item in ai_proposals
            if item.get("layout_signature")
        }) if ai_proposals else len({
            json.dumps({
                "headers": region.get("headers") or [],
                "opciones": region.get("opciones_permitidas") or [],
                "duda": region.get("duda") or "",
            }, ensure_ascii=False, sort_keys=True)
            for region in ambiguous_regions
        })
        reused_ai_responses = max(0, len(ai_proposals) - ai_calls)
        return {
            "archivos": files,
            "hojas": [{"archivo": t.source, "nombre": t.sheet, "region": t.region_id,
                        "tipo_propuesto": t.kind, "filas": len(t.rows), "mapping": t.mapping,
                        "fila_inicial": t.start_row, "fila_final": t.end_row,
                        "fila_encabezado": t.header_row, "titulo_contexto": t.context_title,
                        "dimensiones": t.dimensions, "celdas_no_vacias": t.nonempty_cells,
                        "merged_cells": list(t.merged_cells or []), "confianza": t.confidence,
                        "motivo": t.reason} for t in tables],
            "resumen": {
                "archivos_analizados": len(files),
                "hojas_analizadas": len({(t.source, t.sheet) for t in tables}),
                "filas_analizadas": sum(len(t.rows) for t in tables),
                "posibles_articulos": len(articles), "posibles_recetas": len(recipes),
                "posibles_menus": len(menus),
                "posibles_subelaboraciones": len({x["depende_de"] for x in subdependencies}),
                "posibles_productos_vendibles": sum(
                    1 for t in tables for row in t.rows if self._value(row, t.mapping, "pvp") not in {None, ""}
                ),
                "proveedores": len(suppliers), "relaciones_detectadas": len(subdependencies),
                "duplicados_posibles": len(possible_duplicates),
                "ambiguedades": len(price_doubts) + len(relation_doubts),
                "articulos_sin_coste": sum(
                    1 for item in articles if not item.get("precio_compra")
                ),
            },
            "mapping": mapped,
            "recetas": recipes,
            "menus": menus,
            "articulos": articles,
            "proveedores": suppliers,
            "relaciones": subdependencies,
            "duplicados": possible_duplicates,
            "dudas": {"clasificacion": relation_doubts, "precio": price_doubts,
                       "duplicados": possible_duplicates, "relaciones": relation_doubts},
            "regiones_ambiguas": ambiguous_regions,
            "propuestas_ia": ai_proposals,
            "coste_ia": {
                "usada": bool(ai_calls),
                "apariciones": len(ambiguous_regions),
                "conceptos_unicos": unique_ambiguous_concepts,
                "resueltos_sin_ia": max(0, len(tables) - len(ambiguous_regions)),
                "conceptos_enviados": ai_calls,
                "ambiguedades_enviadas": ai_calls,
                "llamadas": ai_calls,
                "respuestas_reutilizadas": reused_ai_responses,
                "layouts_reutilizados": reused_ai_responses,
                "ahorro_llamadas_deduplicacion": reused_ai_responses,
                "proveedor": None,
                "modelo": None,
                "tokens_reportados": None,
                "coste_reportado": None,
                "total": 0,
            },
            "solo_previsualizacion": True,
            "datos_operativos_modificados": False,
        }

    @staticmethod
    def _ambiguity_region(table: ParsedTable) -> dict[str, Any]:
        headers = list(table.rows[0].keys()) if table.rows else [
            item["columna"] for item in table.mapping
        ]
        return {
            "archivo": table.source, "hoja": table.sheet, "region_id": table.region_id,
            "rango": f"filas {table.start_row}:{table.end_row}",
            "titulo_contextual": table.context_title, "headers": headers,
            "filas": table.rows[:8], "interpretacion_determinista": table.mapping,
            "opciones_permitidas": sorted({value for value in FIELD_ALIASES}),
            "duda": "Clasificar la región y resolver únicamente el mapping ambiguo.",
        }

    def _parse(self, name: str, content: bytes, text: str) -> list[ParsedTable]:
        suffix = Path(name).suffix.lower()
        if suffix not in self.ALLOWED:
            if suffix == ".xls":
                raise RestaurantImportAnalysisError("XLS no se ejecuta ni interpreta; conviértelo a XLSX.")
            raise RestaurantImportAnalysisError(f"Formato no admitido: {suffix or 'sin extensión'}.")
        if suffix == ".xlsx":
            return self._xlsx(name, content)
        if suffix == ".json":
            return self._json(name, text or content.decode("utf-8-sig"))
        decoded = text or content.decode("utf-8-sig")
        delimiter = "\t" if suffix == ".tsv" else self._delimiter(decoded)
        return [self._table(name, "Datos", list(csv.DictReader(io.StringIO(decoded), delimiter=delimiter)))]

    def _xlsx(self, name: str, content: bytes) -> list[ParsedTable]:
        from openpyxl import load_workbook
        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp:
            temp.write(content)
            path = Path(temp.name)
        try:
            workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
            result = []
            for sheet in workbook.worksheets:
                values = [[value for value in row] for row in sheet.iter_rows(values_only=True)]
                result.extend(self._matrix_regions(
                    name, sheet.title, values, sheet.calculate_dimension(),
                    [str(item) for item in sheet.merged_cells.ranges],
                ))
            workbook.close()
            return result
        finally:
            path.unlink(missing_ok=True)

    def _json(self, name: str, text: str) -> list[ParsedTable]:
        value = json.loads(text)
        groups = value if isinstance(value, dict) else {Path(name).stem or "Datos": value}
        result = []
        for key, rows in groups.items():
            if isinstance(rows, dict):
                rows = [rows]
            if isinstance(rows, list) and all(isinstance(row, dict) for row in rows):
                result.append(self._table(name, str(key), rows))
        if not result:
            raise RestaurantImportAnalysisError("JSON sin colecciones de objetos importables.")
        return result

    def _matrix(self, source: str, sheet: str, values: list[list[Any]]) -> ParsedTable:
        nonempty = [row for row in values if any(value not in {None, ""} for value in row)]
        if not nonempty:
            return self._table(source, sheet, [])
        headers = [str(value or f"columna_{index + 1}").strip() for index, value in enumerate(nonempty[0])]
        rows = [{headers[index]: value for index, value in enumerate(row) if index < len(headers)}
                for row in nonempty[1:]]
        return self._table(source, sheet, rows)

    def _matrix_regions(
        self, source: str, sheet: str, values: list[list[Any]], dimensions: str,
        merged_cells: list[str],
    ) -> list[ParsedTable]:
        nonempty_cells = sum(value not in {None, ""} for row in values for value in row)
        single_column_table = any(word in _norm(sheet) for word in ("proveedor", "supplier"))
        candidates = [
            index for index, row in enumerate(values)
            if self._header_score(row) >= 2
            or (single_column_table and self._header_score(row) >= 1)
        ]
        if not candidates:
            vertical = self._vertical_region(source, sheet, values)
            if vertical:
                vertical.dimensions = dimensions
                vertical.nonempty_cells = nonempty_cells
                vertical.merged_cells = merged_cells
                return [vertical]
            unknown = self._matrix(source, sheet, values)
            menu_titles = [
                str(value).strip()
                for row in values[:5] for value in row
                if self._valid_entity_name(value) and _norm(value).startswith("menu ")
            ]
            menu_sections = sum(
                _norm(value) in {"aperitivo", "primero", "pescado", "carne", "postre", "bodega"}
                for row in values for value in row if value not in {None, ""}
            )
            complex_menu_document = len(menu_titles) >= 2 and menu_sections >= 3
            unknown.kind = "DOCUMENTACION" if complex_menu_document else "DESCONOCIDO"
            unknown.header_row = None
            unknown.dimensions = dimensions
            unknown.nonempty_cells = nonempty_cells
            unknown.merged_cells = merged_cells
            unknown.confidence = 0.2
            unknown.reason = "No se encontró un encabezado con evidencia estructural suficiente."
            if complex_menu_document:
                unknown.reason = (
                    "Documento multicolumna con varias identidades de menu; requiere desglose explicito."
                )
            return [unknown]
        regions: list[ParsedTable] = []
        for offset, header_index in enumerate(candidates):
            limit = candidates[offset + 1] if offset + 1 < len(candidates) else len(values)
            end = self._region_end(values, header_index + 1, limit)
            headers = [
                str(value or f"columna_{index + 1}").strip()
                for index, value in enumerate(values[header_index])
            ]
            rows = []
            for physical_index, row in enumerate(values[header_index + 1:end], header_index + 2):
                if not any(value not in {None, ""} for value in row):
                    continue
                parsed_row = {
                    headers[index]: value for index, value in enumerate(row)
                    if index < len(headers)
                }
                parsed_row["_fila"] = physical_index
                rows.append(parsed_row)
            table = self._table(source, sheet, rows)
            table.region_id = f"REGION-{offset + 1}"
            table.start_row = header_index + 1
            table.end_row = max(header_index + 1, end)
            table.header_row = header_index + 1
            table.context_title = self._context_title(values, header_index)
            destinations = {item["destino"] for item in table.mapping}
            structural_kind, structural_reason = self._structural_kind(
                sheet, values, header_index, table
            )
            if structural_kind:
                table.kind = structural_kind
            elif table.context_title and table.kind != "MENUS" and {"nombre", "cantidad"}.issubset(destinations):
                table.kind = "RECETAS"
            table.dimensions = dimensions
            table.nonempty_cells = nonempty_cells
            table.merged_cells = merged_cells
            table.confidence = min(0.98, 0.55 + self._header_score(values[header_index]) * 0.1)
            table.reason = "Encabezado desplazado detectado por aliases y patrón de columnas."
            if structural_reason:
                table.reason = structural_reason
            regions.append(table)
        return regions

    @classmethod
    def _structural_kind(
        cls, sheet: str, values: list[list[Any]], header_index: int, table: ParsedTable,
    ) -> tuple[str | None, str]:
        """Distingue fichas de receta, menus operativos y plantillas por estructura."""
        sheet_text = _norm(sheet)
        nearby = " ".join(
            _norm(value)
            for row in values[max(0, header_index - 8):header_index + 1]
            for value in row if value not in {None, ""}
        )
        header_text = " ".join(
            _norm(value) for value in values[header_index] if value not in {None, ""}
        )
        context = _norm(table.context_title)
        if "plantilla" in sheet_text:
            return "DOCUMENTACION", "Plantilla documental; no representa una entidad operativa."
        destinations = {item["destino"] for item in table.mapping}
        if "nombre" in destinations and destinations.intersection(
            {"codigo", "familia", "proveedor", "formato"}
        ):
            return "ARTICULOS", "Catalogo estructurado por identidad y atributos de articulo."
        recipe_markers = (
            "ficha tecnica plato" in nearby
            or sum(marker in header_text for marker in ("articulo", "bruto", "neto", "desp", "sucio")) >= 3
        )
        if recipe_markers:
            return "RECETAS", "Ficha tecnica con identidad de elaboracion y matriz de ingredientes."
        menu_title = context.startswith("menu ") or any(
            _norm(value).startswith("menu ")
            for row in values[:min(header_index, 5)]
            for value in row if cls._valid_entity_name(value)
        )
        menu_matrix = sum(
            marker in nearby for marker in ("p v p", "food cost", "euros racion", "precio kg")
        ) >= 2
        if menu_title and menu_matrix:
            return "MENUS", "Menu operativo con identidad en cabecera y matriz de componentes/coste."
        return None, ""

    def _vertical_region(self, source: str, sheet: str, values: list[list[Any]]) -> ParsedTable | None:
        ingredient_rows = []
        first_data_index = len(values)
        for index, row in enumerate(values):
            cells = [value for value in row if value not in {None, ""}]
            numeric = next((value for value in cells if isinstance(value, (int, float))), None)
            names = [str(value).strip() for value in cells if self._valid_entity_name(value)]
            if numeric is not None and names:
                first_data_index = min(first_data_index, index)
                ingredient_rows.append({
                    "Ingrediente": names[0], "Cantidad": numeric,
                    "Unidad": next((str(v) for v in cells if _norm(v) in {"kg", "g", "l", "ml", "ud"}), ""),
                    "_fila": index + 1,
                })
        title = self._context_title(values, first_data_index)
        if not title or len(ingredient_rows) < 2:
            return None
        table = self._table(source, sheet, ingredient_rows)
        table.kind = "RECETAS"
        table.header_row = None
        table.context_title = title
        structural_kind, structural_reason = self._structural_kind(
            sheet, values, first_data_index, table
        )
        if structural_kind:
            table.kind = structural_kind
        table.start_row = ingredient_rows[0]["_fila"]
        table.end_row = ingredient_rows[-1]["_fila"]
        table.confidence = min(0.9, 0.55 + len(ingredient_rows) * 0.03)
        table.reason = "Ficha vertical: título contextual seguido de nombres y cantidades."
        if structural_reason:
            table.reason = structural_reason
        return table

    @staticmethod
    def _header_score(row: list[Any]) -> int:
        destinations = set()
        for value in row:
            normalized = _norm(value)
            for field, aliases in FIELD_ALIASES.items():
                if normalized in aliases:
                    destinations.add(field)
        return len(destinations)

    @staticmethod
    def _region_end(values: list[list[Any]], start: int, limit: int) -> int:
        empty_run = 0
        for index in range(start, limit):
            if any(value not in {None, ""} for value in values[index]):
                empty_run = 0
            else:
                empty_run += 1
                if empty_run >= 2:
                    return index - 1
        return limit

    @classmethod
    def _context_title(cls, values: list[list[Any]], before: int) -> str:
        for index in range(min(before, len(values)) - 1, max(-1, before - 8), -1):
            cells = [value for value in values[index] if value not in {None, ""}]
            for position, value in enumerate(cells[:-1]):
                if _norm(value) in {"articulo", "receta", "plato", "elaboracion", "nombre receta"}:
                    candidate = cells[position + 1]
                    if cls._valid_entity_name(candidate):
                        return str(candidate).strip()
            texts = [str(value).strip() for value in cells if cls._valid_entity_name(value)]
            if len(texts) == 1 and cls._header_score(cells) == 0:
                return texts[0]
        return ""

    @staticmethod
    def _valid_entity_name(value: Any) -> bool:
        text = str(value or "").strip()
        if not text or not any(char.isalpha() for char in text):
            return False
        normalized = _norm(text)
        structural = {alias for aliases in FIELD_ALIASES.values() for alias in aliases}
        if normalized in structural:
            return False
        return not bool(re.search(r"\b(total|subtotal|coste total|food cost|margen)\b", normalized))

    def _table(self, source: str, sheet: str, rows: list[dict[str, Any]]) -> ParsedTable:
        headers = [header for header in (list(rows[0]) if rows else []) if not str(header).startswith("_")]
        mapping = []
        for header in headers:
            normalized = _norm(header)
            targets = (
                ["precio_ambiguo"]
                if "€" in str(header) and normalized in {"un", "racion"}
                else [field for field, aliases in FIELD_ALIASES.items() if normalized in aliases]
            )
            destination = targets[0] if len(targets) == 1 else "sin_mapear"
            mapping.append({"archivo": source, "hoja": sheet, "columna": header,
                            "destino": destination, "confianza": 1.0 if targets else 0.0,
                            "requiere_revision": destination in {"precio_ambiguo", "sin_mapear"}})
        sheet_signal = _norm(sheet)
        signal = _norm(sheet + " " + " ".join(headers))
        kind = "DATOS"
        for candidate, words in (("MENUS", ("menu",)),
                                 ("PROVEEDORES", ("proveedor", "supplier")),
                                 ("ARTICULOS", ("articulo", "producto")),
                                 ("RECETAS", ("receta", "elaboracion", "escandallo")),
                                 ("PRECIOS", ("precio", "pvp", "coste"))):
            if any(word in sheet_signal for word in words):
                kind = candidate
                break
        if kind == "DATOS":
            destinations = {item["destino"] for item in mapping}
            if "receta" in destinations or "ingrediente" in destinations:
                kind = "RECETAS"
            elif "proveedor" in destinations and "nombre" not in destinations:
                kind = "PROVEEDORES"
            elif "nombre" in destinations and destinations.intersection({"proveedor", "familia", "codigo", "formato"}):
                kind = "ARTICULOS"
            else:
                kind = "DESCONOCIDO"
        return ParsedTable(source, sheet, rows, mapping, kind)

    def _recipes(self, tables: list[ParsedTable]) -> list[dict[str, Any]]:
        grouped: dict[str, dict[str, Any]] = {}
        seen_physical_origins: set[tuple[str, str, int, str]] = set()
        for table in tables:
            if table.kind != "RECETAS" and not (
                table.kind == "MENUS"
                and any(item["destino"] == "receta" for item in table.mapping)
            ):
                continue
            for row in table.rows:
                recipe_name = (
                    self._value(row, table.mapping, "receta")
                    or table.context_title
                    or self._value(row, table.mapping, "nombre")
                )
                if _norm(recipe_name) == "tapa":
                    # En los escandallos físicos, TAPA es un rótulo de sección
                    # y no una elaboración identificable. Los nombres específicos
                    # (p. ej. "Tapa de anchoa") sí se conservan.
                    continue
                ingredient = self._value(row, table.mapping, "ingrediente")
                if not ingredient and table.context_title:
                    ingredient = self._value(row, table.mapping, "nombre")
                if not self._valid_entity_name(recipe_name):
                    continue
                key = "|".join((table.source, table.sheet, table.region_id, _norm(recipe_name)))
                recipe = grouped.setdefault(key, {
                    "id_origen": f"REC-IMPORT-{len(grouped)+1:03d}", "nombre": str(recipe_name),
                    "ingredientes_estructurados": [], "pasos": [], "bloques_origen": [],
                    "numero_raciones": self._value(row, table.mapping, "rendimiento"),
                    "trazabilidad": {"archivo": table.source, "hoja": table.sheet,
                                     "region": table.region_id, "fila_inicial": table.start_row,
                                     "fila_final": table.end_row,
                                     "columnas_origen": [item["columna"] for item in table.mapping]},
                })
                block_origin = f"{table.source}:{table.sheet}:{table.region_id}"
                if block_origin not in recipe["bloques_origen"]:
                    recipe["bloques_origen"].append(block_origin)
                if self._valid_entity_name(ingredient):
                    physical_row = int(row.get("_fila") or 0)
                    ingredient_column = next(
                        (item["columna"] for item in table.mapping if item["destino"] in {"ingrediente", "nombre"}),
                        "ingrediente",
                    )
                    origin_key = (table.source, table.sheet, physical_row, ingredient_column)
                    if physical_row and origin_key in seen_physical_origins:
                        continue
                    if physical_row:
                        seen_physical_origins.add(origin_key)
                    recipe["ingredientes_estructurados"].append({
                        "nombre_original": str(ingredient),
                        "cantidad_texto": str(self._value(row, table.mapping, "cantidad") or ""),
                        "unidad": self._value(row, table.mapping, "unidad"),
                        "trazabilidad": {
                            "archivo": table.source, "hoja": table.sheet,
                            "region": table.region_id, "fila": physical_row or None,
                            "columna": ingredient_column,
                            "rango": f"{ingredient_column}:{physical_row}" if physical_row else ingredient_column,
                        },
                    })
                procedure = self._value(row, table.mapping, "procedimiento")
                if procedure and str(procedure) not in recipe["pasos"]:
                    recipe["pasos"].append(str(procedure))
        return [
            recipe for recipe in grouped.values()
            if recipe["ingredientes_estructurados"] or recipe["pasos"]
        ]

    def _menus(self, tables: list[ParsedTable]) -> list[dict[str, Any]]:
        """Proyecta la hoja MENU como contenedor, conservando sus recetas separadas."""
        grouped: dict[tuple[str, str], dict[str, Any]] = {}
        for table in tables:
            if table.kind != "MENUS":
                continue
            key = (table.source, table.sheet)
            menu = grouped.setdefault(key, {
                "nombre": table.sheet.strip(), "tipo": "MENU", "componentes": [],
                "origen": {"source_filename": table.source, "sheet": table.sheet},
            })
            seen = {_norm(item.get("nombre")) for item in menu["componentes"]}
            for row in table.rows:
                recipe_name = (
                    self._value(row, table.mapping, "receta")
                    or self._value(row, table.mapping, "ingrediente")
                    or self._value(row, table.mapping, "nombre")
                )
                normalized = _norm(recipe_name)
                if (
                    not self._valid_entity_name(recipe_name)
                    or normalized in {_norm(table.sheet), _norm(menu["nombre"])}
                    or normalized in seen
                ):
                    continue
                seen.add(normalized)
                menu["componentes"].append({
                    "nombre": str(recipe_name).strip(), "seccion": "Otros",
                    "cantidad_origen": 1,
                    "origen": {"archivo": table.source, "hoja": table.sheet,
                               "region": table.region_id, "fila": row.get("_fila")},
                })
        return [item for item in grouped.values() if item["componentes"]]

    def _named_rows(self, tables: list[ParsedTable], kind: str) -> list[dict[str, Any]]:
        unique: dict[str, dict[str, Any]] = {}
        for table in tables:
            if table.kind != kind:
                continue
            for row in table.rows:
                name = self._value(row, table.mapping, "nombre") or self._value(row, table.mapping, "proveedor")
                if self._valid_entity_name(name):
                    unique.setdefault(_norm(name), {
                        "nombre": str(name), "proveedor": self._value(row, table.mapping, "proveedor"),
                        "formato": self._value(row, table.mapping, "formato"),
                        "precio_compra": self._value(row, table.mapping, "precio_compra"),
                        "pvp": self._value(row, table.mapping, "pvp"),
                        "origen": f"{table.source}:{table.sheet}:{table.region_id}",
                        "trazabilidad": {"archivo": table.source, "hoja": table.sheet,
                                         "region": table.region_id, "fila_inicial": table.start_row,
                                         "fila_final": table.end_row,
                                         "columnas_origen": [item["columna"] for item in table.mapping]},
                    })
        return list(unique.values())

    @staticmethod
    def _value(row: dict[str, Any], mapping: list[dict[str, Any]], field: str) -> Any:
        match = next((item for item in mapping if item["destino"] == field), None)
        return row.get(match["columna"]) if match else None

    @staticmethod
    def _delimiter(text: str) -> str:
        if "\t" in text.splitlines()[0] if text.splitlines() else False:
            return "\t"
        try:
            return csv.Sniffer().sniff(text[:4096], delimiters=",;\t|").delimiter
        except csv.Error:
            return ","

    @staticmethod
    def _content(item: dict[str, Any]) -> tuple[bytes, str]:
        text = str(item.get("texto") or "")
        raw = str(item.get("contenido_base64") or "")
        try:
            return (base64.b64decode(raw, validate=True) if raw else text.encode("utf-8")), text
        except (ValueError, TypeError) as exc:
            raise RestaurantImportAnalysisError("Contenido base64 no válido.") from exc

    @classmethod
    def _size(cls, item: dict[str, Any]) -> int:
        content, _text = cls._content(item)
        return len(content)


__all__ = ["RestaurantDataImportAnalyzer", "RestaurantImportAnalysisError"]
