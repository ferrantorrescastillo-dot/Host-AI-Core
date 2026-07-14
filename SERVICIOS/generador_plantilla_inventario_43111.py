from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List
import json


@dataclass
class ResultadoPlantillaInventario:
    total_registros_stock: int
    ruta_destino: str
    estado: str
    mensaje: str


class GeneradorPlantillaInventario43111:
    """
    Genera una plantilla Excel para hacer inventario físico.

    Fuente:
    DATOS/db/stock_inicial.json

    Salida:
    DATOS/inventario_4_3_11_1.xlsx
    """

    COLUMNAS = [
        "Codigo",
        "Articulo",
        "Unidad",
        "Stock sistema",
        "Stock contado",
        "Diferencia",
        "Ubicación",
        "Proveedor",
        "Familia",
        "Observaciones inventario",
    ]

    def __init__(self, ruta_stock: str = "DATOS/db/stock_inicial.json") -> None:
        self.ruta_stock = Path(ruta_stock)

    def generar_excel(self, ruta_destino: str = "DATOS/inventario_4_3_11_1.xlsx") -> ResultadoPlantillaInventario:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para generar la plantilla de inventario.") from exc

        stock = self._leer_stock()
        ruta = Path(ruta_destino)
        ruta.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Inventario"

        ws.append(self.COLUMNAS)

        for idx, item in enumerate(stock, start=2):
            ws.append([
                item.get("codigo"),
                item.get("articulo"),
                item.get("unidad"),
                item.get("stock_actual"),
                "",
                f"=E{idx}-D{idx}",
                item.get("ubicacion"),
                item.get("proveedor"),
                item.get("familia"),
                "",
            ])

        header_fill = PatternFill("solid", fgColor="92400E")
        header_font = Font(color="FFFFFF", bold=True)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        widths = {
            "A": 14,
            "B": 48,
            "C": 12,
            "D": 16,
            "E": 16,
            "F": 14,
            "G": 22,
            "H": 24,
            "I": 24,
            "J": 38,
        }

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # Proteger visualmente columnas calculadas/sistema con color suave
        system_fill = PatternFill("solid", fgColor="F3F4F6")
        for row in ws.iter_rows(min_row=2, max_row=len(stock) + 1):
            for cell in [row[0], row[1], row[2], row[3], row[5], row[6], row[7], row[8]]:
                cell.fill = system_fill

        resumen = wb.create_sheet("Resumen")
        resumen.append(["HOST AI 4.3.11.1 - Plantilla Inventario", ""])
        resumen.append(["Registros de stock incluidos", len(stock)])
        resumen.append(["Estado", "Plantilla generada"])
        resumen.append(["Instrucción", "Rellenar solo la columna Stock contado y Observaciones inventario si hace falta."])
        resumen.append(["Importante", "No modificar Codigo ni Stock sistema."])

        for cell in resumen[1]:
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
            cell.font = Font(color="FFFFFF", bold=True)

        resumen.column_dimensions["A"].width = 38
        resumen.column_dimensions["B"].width = 82

        wb.save(ruta)

        return ResultadoPlantillaInventario(
            total_registros_stock=len(stock),
            ruta_destino=str(ruta),
            estado="ok" if stock else "sin_stock",
            mensaje="Plantilla de inventario generada correctamente." if stock else "No hay stock cargado para inventariar.",
        )

    def _leer_stock(self) -> List[Dict[str, Any]]:
        if not self.ruta_stock.exists():
            return []
        try:
            contenido = json.loads(self.ruta_stock.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return []
        return contenido if isinstance(contenido, list) else []


__all__ = ["GeneradorPlantillaInventario43111", "ResultadoPlantillaInventario"]
