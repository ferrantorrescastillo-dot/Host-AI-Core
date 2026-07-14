from __future__ import annotations

from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional
from datetime import datetime
import json


@dataclass
class ArticuloHostAI:
    codigo: str
    nombre: str
    observaciones: Optional[str]
    proveedor: Optional[str]
    familia: Optional[str]
    precio: Optional[float]
    activo: bool = True
    origen: str = "excel"
    fecha_importacion: str = ""


@dataclass
class ResultadoImportacionArticulos:
    total_procesados: int
    nuevos: int
    actualizados: int
    errores: int
    ruta_destino: str
    mensajes: List[str]


class ImportadorArticulosRestaurante414:
    """
    Importa artículos reales a DATOS/db/articulos.json.

    Reglas:
    - Usa Codigo Host AI si existe la hoja "Codigos 4.1.3".
    - Si no existe, usa Codigo de "Listado de Artículos".
    - No duplica artículos: actualiza por codigo.
    - No modifica el Excel original.
    """

    def importar_desde_excel(
        self,
        ruta_excel: str,
        ruta_db: str = "DATOS/db/articulos.json",
        hoja_articulos: str = "Listado de Artículos",
        hoja_codigos: str = "Codigos 4.1.3",
    ) -> ResultadoImportacionArticulos:
        filas = self._leer_excel_con_codigos(ruta_excel, hoja_articulos, hoja_codigos)
        return self.importar_filas(filas, ruta_db)

    def importar_filas(self, filas: Iterable[Dict[str, Any]], ruta_db: str = "DATOS/db/articulos.json") -> ResultadoImportacionArticulos:
        ruta = Path(ruta_db)
        ruta.parent.mkdir(parents=True, exist_ok=True)

        articulos_actuales = self._leer_db(ruta)
        index = {articulo["codigo"]: articulo for articulo in articulos_actuales if articulo.get("codigo")}

        nuevos = 0
        actualizados = 0
        errores = 0
        mensajes: List[str] = []

        for numero, fila in enumerate(filas, start=1):
            try:
                articulo = self._crear_articulo(fila)
            except ValueError as exc:
                errores += 1
                mensajes.append(f"Fila {numero}: {exc}")
                continue

            datos = asdict(articulo)

            if articulo.codigo in index:
                index[articulo.codigo].update(datos)
                actualizados += 1
            else:
                index[articulo.codigo] = datos
                nuevos += 1

        articulos_ordenados = sorted(index.values(), key=lambda item: item.get("codigo", ""))
        ruta.write_text(json.dumps(articulos_ordenados, ensure_ascii=False, indent=2), encoding="utf-8")

        mensajes.append(f"Importación completada. Nuevos: {nuevos}. Actualizados: {actualizados}. Errores: {errores}.")

        return ResultadoImportacionArticulos(
            total_procesados=nuevos + actualizados + errores,
            nuevos=nuevos,
            actualizados=actualizados,
            errores=errores,
            ruta_destino=str(ruta),
            mensajes=mensajes,
        )

    def _leer_excel_con_codigos(self, ruta_excel: str, hoja_articulos: str, hoja_codigos: str) -> List[Dict[str, Any]]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ImportError("Necesitas instalar openpyxl para leer archivos .xlsx.") from exc

        ruta = Path(ruta_excel)
        if not ruta.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta_excel}")

        wb = load_workbook(ruta, data_only=True)

        if hoja_articulos not in wb.sheetnames:
            raise ValueError(f"No existe la hoja '{hoja_articulos}'. Hojas disponibles: {wb.sheetnames}")

        filas_articulos = self._leer_hoja(wb[hoja_articulos])

        if hoja_codigos in wb.sheetnames:
            filas_codigos = self._leer_hoja(wb[hoja_codigos])
            return self._combinar_articulos_y_codigos(filas_articulos, filas_codigos)

        return filas_articulos

    def _leer_hoja(self, ws) -> List[Dict[str, Any]]:
        headers = [cell.value for cell in ws[1]]
        filas: List[Dict[str, Any]] = []

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            filas.append({str(headers[i]).strip(): row[i] if i < len(row) else None for i in range(len(headers))})

        return filas

    def _combinar_articulos_y_codigos(self, articulos: List[Dict[str, Any]], codigos: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        combinadas: List[Dict[str, Any]] = []

        for indice, fila_articulo in enumerate(articulos):
            fila = dict(fila_articulo)

            if indice < len(codigos):
                codigo_host_ai = codigos[indice].get("Codigo Host AI") or codigos[indice].get("Codigo")
                if codigo_host_ai:
                    fila["Codigo"] = codigo_host_ai

            combinadas.append(fila)

        return combinadas

    def _crear_articulo(self, fila: Dict[str, Any]) -> ArticuloHostAI:
        codigo = self._texto(fila.get("Codigo", fila.get("codigo", "")))
        nombre = self._texto(fila.get("Articulo", fila.get("articulo", "")))

        if not codigo:
            raise ValueError("Artículo sin código.")
        if not nombre:
            raise ValueError("Artículo sin nombre.")

        return ArticuloHostAI(
            codigo=codigo,
            nombre=nombre,
            observaciones=self._texto(fila.get("Observaciones", fila.get("observaciones", ""))),
            proveedor=self._texto(fila.get("Proveedor", fila.get("proveedor", ""))),
            familia=self._texto(fila.get("Familia", fila.get("familia", ""))),
            precio=self._precio(fila.get("Precio", fila.get("precio", None))),
            fecha_importacion=datetime.now().isoformat(timespec="seconds"),
        )

    def _leer_db(self, ruta: Path) -> List[Dict[str, Any]]:
        if not ruta.exists():
            return []

        try:
            contenido = json.loads(ruta.read_text(encoding="utf-8"))
            if isinstance(contenido, list):
                return contenido
            return []
        except json.JSONDecodeError:
            return []

    def _texto(self, valor: Any) -> Optional[str]:
        if valor is None:
            return None
        texto = str(valor).strip()
        return texto if texto else None

    def _precio(self, valor: Any) -> Optional[float]:
        if valor is None or str(valor).strip() == "":
            return None
        if isinstance(valor, (int, float)):
            return float(valor)

        limpio = str(valor).replace("€", "").replace(" ", "").strip()
        if "," in limpio and "." in limpio:
            limpio = limpio.replace(".", "").replace(",", ".")
        else:
            limpio = limpio.replace(",", ".")

        try:
            return float(limpio)
        except ValueError:
            return None


__all__ = [
    "ImportadorArticulosRestaurante414",
    "ArticuloHostAI",
    "ResultadoImportacionArticulos",
]
