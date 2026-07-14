from __future__ import annotations

from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any
import re
import unicodedata

from openpyxl import load_workbook


def _norm(valor: Any) -> str:
    texto = str(valor or "").strip().lower()
    texto = "".join(
        c for c in unicodedata.normalize("NFD", texto)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"[^a-z0-9]+", " ", texto).strip()


def _numero(valor: Any) -> float | None:
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = str(valor).strip().replace("€", "").replace("%", "").replace(" ", "")
    if texto.count(",") == 1 and texto.count(".") == 0:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except (TypeError, ValueError):
        return None


TIPOS_FILA = {
    "TITULO_MENU", "CABECERA_COLUMNAS", "SECCION", "PLATO",
    "ARTICULO_DIRECTO", "COMPLEMENTO", "RESUMEN_ECONOMICO",
    "FILA_VACIA", "RUIDO",
}

SECCIONES_EXACTAS = {
    "aperitivo": "Cóctel",
    "aperitivos": "Cóctel",
    "coctel": "Cóctel",
    "cocktail": "Cóctel",
    "entrante": "Entrante",
    "entrantes": "Entrante",
    "primero": "Primero",
    "primer plato": "Primero",
    "segundo": "Segundo",
    "segundo plato": "Segundo",
    "principal": "Principal",
    "plato principal": "Principal",
    "postre": "Postre",
    "postres": "Postre",
    "recena": "Recena",
    "infantil": "Infantil",
    "menu infantil": "Infantil",
}

CABECERAS = {
    "pvp", "p v p", "neto", "costo", "coste", "food cost",
    "gr o kg", "gr kg", "cantidad", "precio kg", "precio unitario",
    "euros racion", "euro racion", "unidad", "observaciones",
}

ETIQUETAS_RESUMEN = {
    "escandallo", "benefici", "beneficio", "margen", "coste total",
    "costo total", "precio venta", "precio menu", "food cost",
}

PALABRAS_ARTICULO_DIRECTO = {
    "pan", "agua", "vino", "cafe", "cafes", "refresco", "cerveza",
    "hielo", "servilleta", "mantel", "cubierto",
}


def _texto_principal(valores: list[Any]) -> tuple[int | None, str]:
    for idx, valor in enumerate(valores):
        if isinstance(valor, str) and valor.strip():
            return idx, valor.strip()
    return None, ""


def _es_titulo_menu(normal: str) -> bool:
    return bool(normal and normal.startswith("menu") and len(normal.split()) <= 6)


def _es_cabecera(valores: list[Any]) -> bool:
    textos = [_norm(v) for v in valores if isinstance(v, str) and str(v).strip()]
    coincidencias = sum(1 for t in textos if t in CABECERAS)
    return coincidencias >= 2 or (coincidencias >= 1 and len(textos) >= 3)


def _es_complemento(normal: str) -> bool:
    tokens = set(normal.split())
    grupos = {"bodega", "pan", "vino", "cafes", "cafe", "bebidas"}
    return len(tokens & grupos) >= 2 or normal in {"bodega", "bebidas", "pan y cafes"}


def _es_seccion_descriptiva(normal: str, valores: list[Any]) -> bool:
    numeros = [_numero(v) for v in valores[1:] if _numero(v) is not None]
    if numeros:
        return False
    return normal.startswith(("coctel ", "cocktail ", "aperitivo ")) and len(normal.split()) >= 3


def _es_articulo_directo(normal: str, seccion_actual: str, cantidad: float | None, precio_unitario: float | None) -> bool:
    primero = normal.split()[0] if normal else ""
    if seccion_actual == "Complementos" and primero in PALABRAS_ARTICULO_DIRECTO:
        return True
    if primero in PALABRAS_ARTICULO_DIRECTO and cantidad is not None and precio_unitario is not None:
        return cantidad != 1.0
    return False


@dataclass
class FilaMenuI131:
    fila: int
    tipo: str
    texto: str = ""
    seccion: str = ""
    cantidad: float | None = None
    precio_unitario: float | None = None
    coste_racion: float | None = None
    valores: list[Any] = field(default_factory=list)
    motivo: str = ""

    def a_dict(self) -> dict[str, Any]:
        dato = asdict(self)
        if dato["tipo"] not in TIPOS_FILA:
            raise ValueError(f"Tipo de fila no válido: {dato['tipo']}")
        return dato


@dataclass
class MenuLimpioI131:
    hoja: str
    nombre: str
    filas: list[FilaMenuI131]
    platos: list[dict[str, Any]]
    articulos_directos: list[dict[str, Any]]
    complementos: list[dict[str, Any]]
    secciones: list[dict[str, Any]]
    economico: dict[str, float | None]
    avisos: list[str] = field(default_factory=list)

    def a_dict(self) -> dict[str, Any]:
        return {
            "hoja": self.hoja,
            "nombre": self.nombre,
            "filas": [f.a_dict() for f in self.filas],
            "platos": self.platos,
            "articulos_directos": self.articulos_directos,
            "complementos": self.complementos,
            "secciones": self.secciones,
            "economico": self.economico,
            "avisos": self.avisos,
        }


class DetectorLimpioMenusI131:
    """I1.3.1: clasifica filas y genera una vista previa limpia.

    No vincula recetas, no importa menús y no escribe en la base de datos.
    """

    @staticmethod
    def _nombre_menu(ws) -> str:
        for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
            for valor in fila:
                if isinstance(valor, str) and _es_titulo_menu(_norm(valor)):
                    return valor.strip()
        return ws.title.strip()

    @staticmethod
    def _mapas_columnas(ws) -> tuple[dict[str, int], dict[str, int]]:
        detalle: dict[str, int] = {}
        resumen: dict[str, int] = {}
        for fila in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 15), values_only=True):
            valores = list(fila)
            textos = [_norm(v) if isinstance(v, str) else "" for v in valores]
            for idx, texto in enumerate(textos):
                if texto in {"gr o kg", "gr kg", "cantidad"}:
                    detalle["cantidad"] = idx
                elif texto in {"precio kg", "precio unitario"}:
                    detalle["precio_unitario"] = idx
                elif texto in {"euros racion", "euro racion"}:
                    detalle["coste_racion"] = idx
                elif texto in {"pvp", "p v p"}:
                    resumen["precio_venta"] = idx
                elif texto == "neto":
                    resumen["venta_neta"] = idx
                elif texto in {"costo", "coste"}:
                    resumen["coste_total"] = idx
                elif texto == "food cost":
                    resumen["food_cost"] = idx
        return detalle, resumen

    @staticmethod
    def _valor_columna(valores: list[Any], mapa: dict[str, int], clave: str) -> float | None:
        idx = mapa.get(clave)
        return _numero(valores[idx]) if idx is not None and idx < len(valores) else None

    def _detectar_hoja(self, ws) -> dict[str, Any] | None:
        nombre = self._nombre_menu(ws)
        detalle, resumen_cols = self._mapas_columnas(ws)
        filas: list[FilaMenuI131] = []
        platos: list[dict[str, Any]] = []
        articulos: list[dict[str, Any]] = []
        complementos: list[dict[str, Any]] = []
        secciones: list[dict[str, Any]] = []
        economico = {
            "precio_venta": None,
            "venta_neta": None,
            "coste_total": None,
            "food_cost_pct": None,
            "beneficio": None,
        }
        seccion_actual = "Sin sección"

        for numero_fila, fila in enumerate(ws.iter_rows(values_only=True), 1):
            valores = list(fila)
            idx_texto, texto = _texto_principal(valores)
            normal = _norm(texto)

            if not any(v not in (None, "") for v in valores):
                continue

            if _es_titulo_menu(normal):
                filas.append(FilaMenuI131(numero_fila, "TITULO_MENU", texto, valores=valores, motivo="Título del menú"))
                continue

            if normal in SECCIONES_EXACTAS:
                seccion_actual = SECCIONES_EXACTAS[normal]
                dato = {"fila": numero_fila, "nombre": texto, "seccion": seccion_actual}
                secciones.append(dato)
                filas.append(FilaMenuI131(numero_fila, "SECCION", texto, seccion_actual, valores=valores, motivo="Sección reconocida"))
                continue

            if _es_seccion_descriptiva(normal, valores):
                seccion_actual = "Cóctel"
                dato = {"fila": numero_fila, "nombre": texto, "seccion": seccion_actual}
                secciones.append(dato)
                filas.append(FilaMenuI131(numero_fila, "SECCION", texto, seccion_actual, valores=valores, motivo="Subsección descriptiva de cóctel"))
                continue

            if _es_cabecera(valores):
                filas.append(FilaMenuI131(numero_fila, "CABECERA_COLUMNAS", texto, valores=valores, motivo="Cabeceras económicas u operativas"))
                continue

            if _es_complemento(normal):
                seccion_actual = "Complementos"
                dato = {"fila": numero_fila, "nombre": texto, "seccion": seccion_actual}
                complementos.append(dato)
                filas.append(FilaMenuI131(numero_fila, "COMPLEMENTO", texto, seccion_actual, valores=valores, motivo="Bloque de complemento del menú"))
                continue

            if normal in ETIQUETAS_RESUMEN or any(_norm(v) in ETIQUETAS_RESUMEN for v in valores if isinstance(v, str)):
                if normal == "escandallo":
                    economico["precio_venta"] = self._valor_columna(valores, resumen_cols, "precio_venta")
                    economico["venta_neta"] = self._valor_columna(valores, resumen_cols, "venta_neta")
                    economico["coste_total"] = self._valor_columna(valores, resumen_cols, "coste_total")
                    food = self._valor_columna(valores, resumen_cols, "food_cost")
                    economico["food_cost_pct"] = food * 100 if food is not None and abs(food) <= 1 else food
                etiqueta_beneficio = any(_norm(v) in {"benefici", "beneficio"} for v in valores if isinstance(v, str))
                if etiqueta_beneficio:
                    numeros = [_numero(v) for v in valores if _numero(v) is not None]
                    economico["beneficio"] = numeros[0] if numeros else None
                filas.append(FilaMenuI131(numero_fila, "RESUMEN_ECONOMICO", texto or "Beneficio", valores=valores, motivo="Fila de resumen económico"))
                continue

            cantidad = self._valor_columna(valores, detalle, "cantidad")
            precio_unitario = self._valor_columna(valores, detalle, "precio_unitario")
            coste_racion = self._valor_columna(valores, detalle, "coste_racion")

            if texto and _es_articulo_directo(normal, seccion_actual, cantidad, precio_unitario):
                dato = {
                    "fila": numero_fila, "nombre": texto, "seccion": seccion_actual,
                    "cantidad": cantidad, "precio_unitario": precio_unitario,
                    "coste_racion": coste_racion,
                }
                articulos.append(dato)
                filas.append(FilaMenuI131(numero_fila, "ARTICULO_DIRECTO", texto, seccion_actual, cantidad, precio_unitario, coste_racion, valores, "Artículo directo, no exige receta"))
                continue

            if texto and coste_racion is not None:
                dato = {
                    "fila": numero_fila, "nombre": texto, "seccion": seccion_actual,
                    "cantidad": cantidad, "precio_unitario": precio_unitario,
                    "coste_racion": coste_racion,
                }
                platos.append(dato)
                filas.append(FilaMenuI131(numero_fila, "PLATO", texto, seccion_actual, cantidad, precio_unitario, coste_racion, valores, "Fila culinaria con coste por ración"))
                continue

            filas.append(FilaMenuI131(numero_fila, "RUIDO", texto, seccion_actual, valores=valores, motivo="No cumple reglas de plato, artículo, sección o resumen"))

        if len(platos) < 2:
            return None

        avisos: list[str] = []
        suma_costes = sum(float(p["coste_racion"] or 0) for p in platos) + sum(float(a["coste_racion"] or 0) for a in articulos)
        if economico["coste_total"] is not None and abs(suma_costes - float(economico["coste_total"])) > 0.02:
            avisos.append(
                f"La suma detectada ({suma_costes:.4f}) no coincide con el coste total ({economico['coste_total']:.4f})."
            )

        return MenuLimpioI131(
            hoja=ws.title,
            nombre=nombre,
            filas=filas,
            platos=platos,
            articulos_directos=articulos,
            complementos=complementos,
            secciones=secciones,
            economico=economico,
            avisos=avisos,
        ).a_dict()

    def preparar(self, ruta_excel: str | Path, hojas=None) -> dict[str, Any]:
        ruta = Path(ruta_excel)
        if not ruta.exists():
            raise FileNotFoundError(f"No existe el archivo Excel: {ruta}")
        wb = load_workbook(ruta, data_only=True, read_only=True)
        seleccion = list(hojas or [])
        if seleccion:
            inexistentes = [h for h in seleccion if h not in wb.sheetnames]
            if inexistentes:
                wb.close()
                raise ValueError(f"No existen las hojas: {', '.join(inexistentes)}")
        else:
            seleccion = [
                h for h in wb.sheetnames
                if "menu" in _norm(h) and not _norm(h).startswith("mp ")
                and "materia prima" not in _norm(h)
            ]
        menus = []
        for hoja in seleccion:
            detectado = self._detectar_hoja(wb[hoja])
            if detectado:
                menus.append(detectado)
        wb.close()

        conteo_tipos = {tipo: 0 for tipo in TIPOS_FILA}
        for menu in menus:
            for fila in menu["filas"]:
                conteo_tipos[fila["tipo"]] += 1
        return {
            "ok": True,
            "ruta": str(ruta),
            "modo": "vista_previa_sin_escritura_sin_vinculacion",
            "menus": menus,
            "resumen": {
                "hojas_analizadas": len(seleccion),
                "menus_detectados": len(menus),
                "platos": sum(len(m["platos"]) for m in menus),
                "articulos_directos": sum(len(m["articulos_directos"]) for m in menus),
                "complementos": sum(len(m["complementos"]) for m in menus),
                "secciones": sum(len(m["secciones"]) for m in menus),
                "tipos_fila": conteo_tipos,
            },
        }


def formatear_previa_i131(previa: dict[str, Any]) -> str:
    r = previa["resumen"]
    lineas = [
        "I1.3.1 — DETECTOR LIMPIO DE MENÚS",
        "=" * 78,
        f"Hojas: {r['hojas_analizadas']} | Menús: {r['menus_detectados']} | Secciones: {r['secciones']}",
        f"Platos: {r['platos']} | Artículos directos: {r['articulos_directos']} | Complementos: {r['complementos']}",
    ]
    for indice, menu in enumerate(previa.get("menus", []), 1):
        lineas.append(f"{indice}. {menu['nombre']} | Hoja: {menu['hoja']}")
        for fila in menu["filas"]:
            if fila["tipo"] in {"FILA_VACIA", "RUIDO", "CABECERA_COLUMNAS", "TITULO_MENU", "RESUMEN_ECONOMICO"}:
                continue
            prefijo = {
                "SECCION": "[SECCIÓN]",
                "PLATO": "- [PLATO]",
                "ARTICULO_DIRECTO": "- [ARTÍCULO]",
                "COMPLEMENTO": "[COMPLEMENTO]",
            }.get(fila["tipo"], f"[{fila['tipo']}]")
            coste = f" | coste/ración {fila['coste_racion']:.4f} €" if fila.get("coste_racion") is not None else ""
            lineas.append(f"   {prefijo} {fila.get('texto') or '-'}{coste}")
        eco = menu["economico"]
        lineas.append(
            "   Económico: "
            f"venta={eco.get('precio_venta')} | neto={eco.get('venta_neta')} | "
            f"coste={eco.get('coste_total')} | food cost={eco.get('food_cost_pct')}% | "
            f"beneficio={eco.get('beneficio')}"
        )
        for aviso in menu.get("avisos", []):
            lineas.append(f"   AVISO: {aviso}")
    lineas += [
        "=" * 78,
        "VISTA PREVIA LIMPIA: no se vinculó ninguna receta y no se modificó la base de datos.",
    ]
    return "\n".join(lineas)
