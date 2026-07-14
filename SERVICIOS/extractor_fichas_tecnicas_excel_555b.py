from __future__ import annotations

import math
import re
import unicodedata
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from MODELOS.extraccion_fichas_555b import (
    FichaTecnicaExtraida555B,
    IngredienteExtraido555B,
    ResultadoExtraccionFichas555B,
)


def _norm(valor: object) -> str:
    texto = str(valor or "").strip().lower()
    texto = "".join(c for c in unicodedata.normalize("NFD", texto) if unicodedata.category(c) != "Mn")
    return " ".join(re.sub(r"\s+", " ", texto).split())


def _numero(valor: object) -> float | None:
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        numero = float(valor)
        return numero if math.isfinite(numero) else None
    texto = str(valor).strip().replace("€", "").replace(" ", "")
    if not texto:
        return None
    # Acepta coma decimal sin romper números con punto decimal.
    if texto.count(",") == 1 and texto.count(".") == 0:
        texto = texto.replace(",", ".")
    try:
        numero = float(texto)
        return numero if math.isfinite(numero) else None
    except ValueError:
        return None


class ExtractorFichasTecnicasExcel555B:
    """Extrae fichas a una vista previa normalizada sin escribir en Host AI."""

    MARCADORES_FICHA = ("ficha tecnica plato", "ficha tecnica")
    CABECERAS_INGREDIENTE = ("articulo", "ingrediente", "producto")
    UNIDADES = {
        "kg": "kg", "kilo": "kg", "kilos": "kg", "gr": "g", "g": "g",
        "l": "l", "lt": "l", "litro": "l", "litros": "l", "ml": "ml",
        "unid": "u", "unidad": "u", "unidades": "u", "ud": "u", "uds": "u",
        "pax": "pax", "racion": "ración", "raciones": "ración",
        "tapa": "u", "elaborados": "u",
    }

    def extraer(self, ruta_archivo: str | Path) -> dict[str, Any]:
        ruta = Path(ruta_archivo)
        resultado = ResultadoExtraccionFichas555B(archivo=str(ruta))
        if not ruta.exists():
            resultado.errores.append(f"No se encuentra el archivo: {ruta}")
            return resultado.to_dict()
        if ruta.suffix.lower() not in {".xlsx", ".xlsm"}:
            resultado.errores.append("Esta fase admite únicamente archivos .xlsx o .xlsm.")
            return resultado.to_dict()

        try:
            libro = load_workbook(ruta, read_only=False, data_only=True)
        except Exception as exc:
            resultado.errores.append(f"No se pudo abrir el Excel: {exc}")
            return resultado.to_dict()

        try:
            for hoja in libro.worksheets:
                marcadores = self._marcadores(hoja)
                for indice, fila_inicio in enumerate(marcadores):
                    fila_fin = marcadores[indice + 1] - 1 if indice + 1 < len(marcadores) else hoja.max_row
                    ficha = self._extraer_bloque(hoja, fila_inicio, fila_fin)
                    resultado.fichas.append(ficha)
        finally:
            libro.close()

        if not resultado.fichas:
            resultado.avisos.append("No se han localizado bloques con marcador 'FICHA TÉCNICA'.")
        return resultado.to_dict()

    def _marcadores(self, hoja) -> list[int]:
        """Localiza los inicios de ficha en una única pasada.

        El extractor carga el libro en modo normal y solo lectura lógica. Esto
        evita el coste extremo de acceso aleatorio de las hojas ``read_only``
        de openpyxl, que podía aparentar que la terminal estaba bloqueada.
        """
        encontrados: list[int] = []
        for n, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
            if any(
                any(marcador in _norm(valor) for marcador in self.MARCADORES_FICHA)
                for valor in fila
                if valor not in (None, "")
            ):
                encontrados.append(n)
        return encontrados

    def _extraer_bloque(self, hoja, inicio: int, fin: int) -> FichaTecnicaExtraida555B:
        limite_busqueda = min(fin, inicio + 12)
        nombre, fila_nombre = self._buscar_nombre(hoja, inicio, limite_busqueda)
        cabecera = self._buscar_cabecera_ingredientes(hoja, inicio, min(fin, inicio + 18))
        rendimiento, unidad_rendimiento = self._buscar_rendimiento(hoja, inicio, cabecera or limite_busqueda)
        ingredientes = self._extraer_ingredientes(hoja, cabecera, fin) if cabecera else []

        avisos: list[str] = []
        errores: list[str] = []
        if not nombre:
            errores.append("No se ha podido identificar el nombre de la receta.")
            nombre = f"Ficha sin nombre (fila {inicio})"
        if rendimiento is None or rendimiento <= 0:
            avisos.append("Rendimiento no localizado o no válido.")
        if not unidad_rendimiento:
            avisos.append("Unidad de rendimiento no localizada.")
        if not cabecera:
            errores.append("No se ha localizado la cabecera de ingredientes.")
        if not ingredientes:
            errores.append("No se han extraído ingredientes con nombre y cantidad.")

        puntuacion = 25.0
        if fila_nombre: puntuacion += 25.0
        if rendimiento and rendimiento > 0: puntuacion += 15.0
        if unidad_rendimiento: puntuacion += 10.0
        if cabecera: puntuacion += 10.0
        if ingredientes: puntuacion += min(15.0, 3.0 + len(ingredientes) * 1.5)
        confianza = round(min(100.0, puntuacion), 2)

        return FichaTecnicaExtraida555B(
            hoja=hoja.title,
            fila_inicio=inicio,
            fila_fin=fin,
            nombre=nombre,
            rendimiento=rendimiento,
            unidad_rendimiento=unidad_rendimiento,
            ingredientes=ingredientes,
            confianza=confianza,
            avisos=avisos,
            errores=errores,
        )

    def _buscar_nombre(self, hoja, inicio: int, fin: int) -> tuple[str, int | None]:
        for n in range(inicio + 1, fin + 1):
            valores = [hoja.cell(n, c).value for c in range(1, min(hoja.max_column, 20) + 1)]
            for i, valor in enumerate(valores):
                if _norm(valor) in self.CABECERAS_INGREDIENTE:
                    # La cabecera de nombre suele ser ARTÍCULO | Nombre del plato.
                    for candidato in valores[i + 1:]:
                        texto = str(candidato or "").strip()
                        if texto and _norm(texto) not in self.CABECERAS_INGREDIENTE and _numero(candidato) is None:
                            return texto, n
        return "", None

    def _buscar_cabecera_ingredientes(self, hoja, inicio: int, fin: int) -> int | None:
        for n in range(inicio + 1, fin + 1):
            valores = [hoja.cell(n, c).value for c in range(1, min(hoja.max_column, 20) + 1)]
            normales = [_norm(v) for v in valores]
            tiene_articulo = any(v in self.CABECERAS_INGREDIENTE for v in normales)
            tiene_cantidad = any(v in {"kg", "gr", "g", "cantidad", "peso", "unid", "unidad"} for v in normales)
            tiene_coste = any("/un" in v or "racion" in v or "bruto" in v for v in normales)
            if tiene_articulo and (tiene_cantidad or tiene_coste):
                return n
        return None

    def _buscar_rendimiento(self, hoja, inicio: int, fin: int) -> tuple[float | None, str | None]:
        for n in range(inicio + 1, fin + 1):
            fila = [hoja.cell(n, c).value for c in range(1, min(hoja.max_column, 12) + 1)]
            for i, valor in enumerate(fila):
                normal = _norm(valor).rstrip(".")
                if normal in self.UNIDADES:
                    for candidato in fila[i + 1:i + 4]:
                        numero = _numero(candidato)
                        if numero is not None and numero > 0:
                            return numero, self.UNIDADES[normal]
        return None, None

    def _extraer_ingredientes(self, hoja, cabecera: int | None, fin: int) -> list[IngredienteExtraido555B]:
        if cabecera is None:
            return []
        columnas = self._columnas_tabla(hoja, cabecera)
        col_nombre = columnas.get("nombre")
        col_cantidad = columnas.get("cantidad")
        col_precio = columnas.get("precio")
        col_coste = columnas.get("coste")
        if not col_nombre or not col_cantidad:
            return []

        ingredientes: list[IngredienteExtraido555B] = []
        vacias = 0
        for n in range(cabecera + 1, fin + 1):
            nombre_val = hoja.cell(n, col_nombre).value
            cantidad_val = hoja.cell(n, col_cantidad).value
            nombre = str(nombre_val or "").strip()
            cantidad = _numero(cantidad_val)
            if not nombre and cantidad is None:
                vacias += 1
                if vacias >= 3 and ingredientes:
                    break
                continue
            vacias = 0
            if not nombre:
                continue
            normal = _norm(nombre)
            if any(m in normal for m in self.MARCADORES_FICHA) or normal in self.CABECERAS_INGREDIENTE:
                break
            if cantidad is None:
                continue
            avisos: list[str] = []
            if cantidad < 0:
                avisos.append("Cantidad negativa.")
            unidad = self._unidad_desde_cabecera(hoja.cell(cabecera, col_cantidad).value)
            ingredientes.append(IngredienteExtraido555B(
                nombre=nombre,
                cantidad=cantidad,
                unidad=unidad,
                fila_origen=n,
                precio_unitario=_numero(hoja.cell(n, col_precio).value) if col_precio else None,
                coste_racion=_numero(hoja.cell(n, col_coste).value) if col_coste else None,
                avisos=avisos,
            ))
        return ingredientes

    def _columnas_tabla(self, hoja, fila: int) -> dict[str, int]:
        resultado: dict[str, int] = {}
        for c in range(1, min(hoja.max_column, 30) + 1):
            normal = _norm(hoja.cell(fila, c).value)
            if normal in self.CABECERAS_INGREDIENTE and "nombre" not in resultado:
                resultado["nombre"] = c
            elif normal in {"kg", "gr", "g", "cantidad", "peso", "unid", "unidad"} and "cantidad" not in resultado:
                resultado["cantidad"] = c
            elif ("€/un" in normal or "precio" in normal) and "precio" not in resultado:
                resultado["precio"] = c
            elif "racion" in normal and ("€" in str(hoja.cell(fila, c).value) or "coste" in normal) and "coste" not in resultado:
                resultado["coste"] = c
        return resultado

    def _unidad_desde_cabecera(self, valor: object) -> str:
        normal = _norm(valor).rstrip(".")
        return self.UNIDADES.get(normal, normal or "unidad_origen")
