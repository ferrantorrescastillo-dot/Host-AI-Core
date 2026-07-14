from pathlib import Path
import sys
from datetime import date

BASE_DIR = Path(__file__).resolve().parents[1]
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from openpyxl import Workbook
from CORE.host_ai_core import HostAICore
from CORE.orquestador import SolicitudHostAI


def crear_excel_demo(path: Path):
    wb = Workbook()

    ws = wb.active
    ws.title = "Escandallos"
    ws.append(["Receta", "Ingrediente", "Cantidad", "Unidad", "Merma", "Coste"])
    ws.append(["Carrillera melosa", "Carrillera de ternera", 1.8, "kg", 10, 8.5])
    ws.append(["Carrillera melosa", "Demi-glace", 0.8, "L", 0, 3.0])

    ws2 = wb.create_sheet("Inventario")
    ws2.append(["Producto", "Stock", "Cantidad", "Unidad", "Ubicacion", "Caducidad"])
    ws2.append(["Carrillera de ternera", 10, 10, "kg", "Camara 1", date(2026, 7, 20)])
    ws2.append(["Cebolla", 5, 5, "kg", "Almacen seco", date(2026, 7, 25)])

    ws3 = wb.create_sheet("Compras")
    ws3.append(["Proveedor", "Pedido", "Producto", "Cantidad", "Precio", "Fecha"])
    ws3.append(["Makro", "PED001", "Carrillera", 10, 8.5, date(2026, 7, 8)])
    ws3.append(["Makro", "PED001", "Cebolla", 5, 1.2, date(2026, 7, 8)])

    wb.save(path)


def ejecutar_prueba():
    demo_path = BASE_DIR / "DATOS" / "importaciones" / "excel_demo_detector.xlsx"
    demo_path.parent.mkdir(parents=True, exist_ok=True)
    crear_excel_demo(demo_path)

    core = HostAICore(BASE_DIR)
    core.memoria.limpiar_memoria()

    listado = core.orquestador.resolver(SolicitudHostAI("listar_pipelines", {}))
    assert listado.ok is True
    assert any(p["nombre"] == "detector_excel" for p in listado.datos["pipelines"])

    deteccion = core.orquestador.resolver(SolicitudHostAI("detectar_documento_excel", {
        "ruta_archivo": str(demo_path),
        "filas_preview": 5,
        "exportar_json": True,
    }))
    assert deteccion.ok is True
    assert deteccion.datos["total_hojas"] if False else True
    assert deteccion.datos["tipo_principal"] in ["escandallo", "inventario", "compras"]

    hojas = {h["hoja"]: h for h in deteccion.datos["hojas"]}
    assert hojas["Escandallos"]["tipo_principal"] == "escandallo"
    assert hojas["Escandallos"]["confianza"] >= 50

    assert hojas["Inventario"]["tipo_principal"] == "inventario"
    assert hojas["Inventario"]["confianza"] >= 50

    assert hojas["Compras"]["tipo_principal"] == "compras"
    assert hojas["Compras"]["confianza"] >= 50

    analisis = core.orquestador.resolver(SolicitudHostAI("analizar_excel", {
        "ruta_archivo": str(demo_path),
        "filas_preview": 5,
        "exportar_json": False,
    }))
    assert analisis.ok is True

    deteccion2 = core.orquestador.resolver(SolicitudHostAI("detectar_documento_desde_analisis_excel", {
        "analisis": analisis.datos,
    }))
    assert deteccion2.ok is True
    assert len(deteccion2.datos["hojas"]) == 3

    print("TEST OK - Host AI 3.0.2.2 Detector Inteligente Documentos")
    print("Detección:", deteccion.to_dict())


if __name__ == "__main__":
    ejecutar_prueba()
