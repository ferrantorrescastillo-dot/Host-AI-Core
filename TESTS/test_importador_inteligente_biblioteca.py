from __future__ import annotations

import base64
import json
import logging
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from API.app import HostAIPlatformAPI
from API.facade.catalog_public_facade import CatalogPublicFacade
from API.http_server import create_app
from MODELOS.importador_inteligente_biblioteca import DocumentType
from SERVICIOS.importador_inteligente_biblioteca import (
    ImportDocumentService,
    RuleBasedDocumentClassifier,
)
from SERVICIOS.borrador_importacion_biblioteca import (
    CulinaryQuantityParser,
    IngredientTextNormalizer,
)


RECIPE_TEXT = """Receta: Salsa verde
Rendimiento: 4 raciones
Ingredientes:
Perejil: 0.1 kg
Aceite: 0.2 l
Procedimiento: Triturar todos los ingredientes.
"""


def _payload(name: str = "salsa-verde.txt", text: str = RECIPE_TEXT) -> dict:
    return {
        "nombre": name,
        "tipo_mime": "text/plain",
        "contenido_base64": base64.b64encode(text.encode("utf-8")).decode("ascii"),
        "texto": text,
    }


def _word_payload(
    name: str = "receta-cocido.docx",
    *,
    recipes: int = 1,
    table: bool = False,
    unstructured: bool = False,
    empty: bool = False,
    image: bool = False,
    mixed: bool = False,
) -> dict:
    document = Document()
    if empty:
        pass
    elif unstructured:
        document.add_heading("Notas de cocina", level=1)
        document.add_paragraph("Documento general con observaciones internas del equipo.")
    else:
        for index in range(1, recipes + 1):
            document.add_heading(f"Receta {index}", level=1)
            document.add_paragraph("Rendimiento: 4 raciones")
            document.add_heading("Ingredientes", level=2)
            if table and index == 1:
                ingredient_table = document.add_table(rows=1, cols=3)
                ingredient_table.rows[0].cells[0].text = "Ingrediente"
                ingredient_table.rows[0].cells[1].text = "Cantidad"
                ingredient_table.rows[0].cells[2].text = "Unidad"
                row = ingredient_table.add_row().cells
                row[0].text, row[1].text, row[2].text = "Patata", "2", "kg"
            else:
                document.add_paragraph("1 kg Patata", style="List Bullet")
                document.add_paragraph("250 g Mantequilla", style="List Bullet")
            document.add_heading("Procedimiento", level=2)
            document.add_paragraph("Cocer a 90 ºC y triturar.")
        if mixed:
            document.add_paragraph("Menú: servicio de verano")
        if image:
            image_bytes = base64.b64decode(
                "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
            )
            document.add_picture(BytesIO(image_bytes))
    stream = BytesIO()
    document.save(stream)
    return {
        "nombre": name,
        "tipo_mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "contenido_base64": base64.b64encode(stream.getvalue()).decode("ascii"),
    }


def test_clasificador_determinista_cubre_tipos_principales() -> None:
    classifier = RuleBasedDocumentClassifier()
    cases = {
        "receta con ingredientes y procedimiento": DocumentType.RECETA,
        "escandallo coste unitario": DocumentType.ESCANDALLO,
        "ficha tecnica appcc alergenos": DocumentType.FICHA_TECNICA,
        "menu entrante primer plato": DocumentType.MENU,
        "factura base imponible": DocumentType.FACTURA,
        "albaran entrega mercancia": DocumentType.ALBARAN,
    }
    for text, expected in cases.items():
        result, confidence = classifier.classify(filename="documento.txt", text=text)
        assert result == expected
        assert 0 <= confidence.value <= 1
        assert confidence.reason


def test_importacion_genera_solo_propuestas_y_no_conserva_documento(tmp_path: Path) -> None:
    service = ImportDocumentService(tmp_path)
    result = service.import_document(_payload())

    assert result["ok"] is True
    preview = result["importacion"]
    document = preview["documento"]
    assert document["clasificacion"]["tipo"] == "RECETA"
    assert document["contenido_almacenado"] is False
    assert preview["solo_previsualizacion"] is True
    assert preview["confirmacion_disponible"] is True
    assert preview["propuestas"]
    assert all(item["estado"] == "PENDIENTE_REVISION" for item in preview["propuestas"])
    assert all(item["persistida"] is False for item in preview["propuestas"])
    assert RECIPE_TEXT not in json.dumps(preview, ensure_ascii=False)

    import_id = document["id"]
    assert service.get_import(import_id)["importacion"] == preview
    proposals = service.get_proposals(import_id)
    assert proposals["total"] == len(preview["propuestas"])
    assert proposals["solo_previsualizacion"] is True


def test_importacion_rechaza_formato_y_controla_id_ausente(tmp_path: Path) -> None:
    service = ImportDocumentService(tmp_path)
    invalid = service.import_document(_payload("programa.exe"))
    assert invalid["error"]["code"] == "formato_no_soportado"
    assert invalid["error"]["status"] == 400
    assert service.get_import("NO-EXISTE")["error"]["status"] == 404
    assert service.get_proposals("NO-EXISTE")["error"]["status"] == 404


def test_http_importacion_biblioteca_post_get_y_propuestas(tmp_path: Path) -> None:
    client = TestClient(create_app(HostAIPlatformAPI(base_dir=tmp_path)))
    created = client.post("/api/v1/biblioteca/importaciones", json=_payload())
    assert created.status_code == 200
    body = created.json()
    assert body["ok"] is True
    assert body["modo_seguro"] is True
    assert body["datos_reales_modificados"] is False
    import_id = body["importacion"]["documento"]["id"]

    detail = client.get(f"/api/v1/biblioteca/importaciones/{import_id}")
    assert detail.status_code == 200
    assert detail.json()["importacion"]["documento"]["id"] == import_id

    proposals = client.get(f"/api/v1/biblioteca/importaciones/{import_id}/propuestas")
    assert proposals.status_code == 200
    assert proposals.json()["total"] >= 1
    assert proposals.json()["datos_reales_modificados"] is False

    missing = client.get("/api/v1/biblioteca/importaciones/NO-EXISTE")
    assert missing.status_code == 404


def test_http_word_genera_sesion_clasificacion_y_propuestas_sin_escribir(
    tmp_path: Path,
) -> None:
    client = TestClient(create_app(HostAIPlatformAPI(base_dir=tmp_path)))

    response = client.post("/api/v1/biblioteca/importaciones", json=_word_payload())

    assert response.status_code == 200
    body = response.json()
    assert body["ok"] is True
    assert body["datos_reales_modificados"] is False
    session = body["importacion"]
    assert session["documento"]["origen"] == "WORD"
    assert session["documento"]["clasificacion"]["tipo"] == "RECETA"
    assert session["documento"]["contenido_almacenado"] is False
    assert session["propuestas"]
    assert session["resumen"]["recetas_detectadas"] == 1
    assert session["resumen"]["ingredientes_detectados"] == 2
    assert not any(
        "Lector Word pendiente" in warning
        for warning in session["documento"]["advertencias"]
    )
    assert session["solo_previsualizacion"] is True
    assert session["confirmacion_disponible"] is True


def test_word_con_dos_recetas_tabla_listas_y_propuestas_separadas(
    tmp_path: Path,
) -> None:
    result = ImportDocumentService(tmp_path).import_document(
        _word_payload(recipes=2, table=True)
    )

    assert result["ok"] is True
    session = result["importacion"]
    assert session["documento"]["clasificacion"]["tipo"] == "RECETA"
    assert session["resumen"]["recetas_detectadas"] == 2
    assert session["resumen"]["ingredientes_detectados"] == 3
    recipes = [
        item for item in session["documento"]["entidades"] if item["kind"] == "RECETA"
    ]
    assert [item["name"] for item in recipes] == ["Receta 1", "Receta 2"]
    assert any(item["fields"]["pasos"] for item in recipes)
    create_recipes = [
        item for item in session["propuestas"] if item["tipo"] == "CREAR_RECETA"
    ]
    assert len(create_recipes) == 2
    assert all(item["bloques_origen"] for item in create_recipes)
    assert all(item["persistida"] is False for item in session["propuestas"])


def test_word_relaciona_exactos_detecta_dudosos_nuevos_y_duplicado(
    tmp_path: Path,
) -> None:
    data_dir = tmp_path / "DATOS" / "db"
    data_dir.mkdir(parents=True)
    (data_dir / "articulos.json").write_text(json.dumps([
        {"codigo": "M.P-PAT", "nombre": "Patata", "unidad": "kg", "precio": 1},
        {"codigo": "M.P-MANT-1", "nombre": "Mantequilla dulce", "unidad": "kg", "precio": 8},
        {"codigo": "M.P-MANT-2", "nombre": "Mantequilla salada", "unidad": "kg", "precio": 9},
    ]), encoding="utf-8")
    (data_dir / "biblioteca_recetas_601.json").write_text(json.dumps({
        "version_modelo": "6.0.1",
        "recetas": [{
            "id": "REC-1", "nombre": "Receta 1", "codigo": "REC-1",
            "familia": "", "tipo": "", "numero_raciones": 4,
            "ingredientes": ["Patata"], "cantidades": ["1 kg"],
            "elaboracion": "Cocer.", "tiempo_elaboracion": "",
            "conservacion": "", "alergenos": [], "observaciones": "",
            "fotografia": "", "estado": "OPERATIVA", "version": 1,
            "creado_en": "", "actualizado_en": "",
        }],
    }), encoding="utf-8")

    result = ImportDocumentService(tmp_path).import_document(_word_payload())

    assert result["ok"] is True
    session = result["importacion"]
    ingredients = [
        item["fields"]
        for item in session["documento"]["entidades"]
        if item["kind"] == "INGREDIENTE"
    ]
    assert ingredients[0]["estado_relacion"] == "relacionado"
    assert ingredients[0]["articulo_id"] == "M.P-PAT"
    assert ingredients[1]["estado_relacion"] == "coincidencia_dudosa"
    assert session["resumen"]["ingredientes_relacionados"] == 1
    assert session["resumen"]["coincidencias_dudosas"] == 1
    assert any(item["tipo"] == "ACTUALIZAR_RECETA" for item in session["propuestas"])
    assert any(item["tipo"] == "REVISAR_COINCIDENCIA" for item in session["propuestas"])


def test_word_corrupto_vacio_y_sin_marcadores_no_provocan_500(tmp_path: Path) -> None:
    service = ImportDocumentService(tmp_path)
    corrupt = service.import_document({
        "nombre": "corrupto.docx",
        "tipo_mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "contenido_base64": base64.b64encode(b"no es docx").decode("ascii"),
    })
    assert corrupt["error"]["code"] == "documento_corrupto"
    assert corrupt["error"]["status"] == 400

    empty = service.import_document(_word_payload(empty=True))
    assert empty["error"]["code"] == "documento_vacio"
    assert empty["error"]["status"] == 400

    unstructured = service.import_document(
        _word_payload(name="notas-cocina.docx", unstructured=True)
    )
    assert unstructured["ok"] is True
    assert unstructured["importacion"]["documento"]["clasificacion"]["tipo"] == "DOCUMENTACION"
    assert unstructured["importacion"]["resumen"]["recetas_detectadas"] == 0


def test_word_mixto_e_imagen_informa_limitacion_sin_perder_recetas(
    tmp_path: Path,
) -> None:
    result = ImportDocumentService(tmp_path).import_document(
        _word_payload(mixed=True, image=True)
    )

    assert result["ok"] is True
    session = result["importacion"]
    assert session["documento"]["clasificacion"]["tipo"] == "MIXTO"
    assert session["resumen"]["recetas_detectadas"] == 1
    assert any(
        "imágenes" in warning for warning in session["documento"]["advertencias"]
    )


def test_python_docx_disponible_y_previsualizacion_no_escribe_datos(
    tmp_path: Path,
) -> None:
    import docx

    assert docx.__version__ == "1.2.0"
    service = ImportDocumentService(tmp_path)
    before = {
        path.relative_to(tmp_path): path.read_bytes()
        for path in (tmp_path / "DATOS").rglob("*")
        if path.is_file()
    }

    result = service.import_document(_word_payload(recipes=2))

    after = {
        path.relative_to(tmp_path): path.read_bytes()
        for path in (tmp_path / "DATOS").rglob("*")
        if path.is_file()
    }
    assert result["ok"] is True
    # La fase final persiste únicamente conocimiento estructurado y auditoría,
    # nunca el DOCX ni cambios en las fuentes maestras.
    domain_files = {
        Path("DATOS/db/articulos.json"),
        Path("DATOS/db/biblioteca_recetas_601.json"),
        Path("DATOS/db/biblioteca_escandallos_601.json"),
    }
    assert {path: before.get(path) for path in domain_files} == {
        path: after.get(path) for path in domain_files
    }
    store = (tmp_path / "DATOS/db/biblioteca_importaciones_web.json").read_text(encoding="utf-8")
    assert "UEsDB" not in store
    assert all(not item["persistida"] for item in result["importacion"]["propuestas"])


def test_borrador_normaliza_cantidades_y_observaciones_sin_corregir_ambiguedad() -> None:
    parser = CulinaryQuantityParser()
    assert parser.parse("1/2")[0] == 0.5
    assert parser.parse("½")[0] == 0.5
    assert parser.parse("2,3")[0] == 2.3
    assert parser.parse("2 3")[0] is None
    assert parser.parse("2 3")[1][0].code == "CANTIDAD_AMBIGUA"
    assert parser.parse("2-3")[0] is None

    normalizer = IngredientTextNormalizer()
    assert normalizer.split("azúcar para un kilo de miso") == (
        "azúcar", "para un kilo de miso"
    )
    assert normalizer.split("hueso jamón(previamente blanqueado 3 veces)") == (
        "hueso de jamón", "previamente blanqueado 3 veces"
    )


def test_borrador_editable_jerarquia_ingrediente_y_version_optimista(
    tmp_path: Path,
) -> None:
    service = ImportDocumentService(tmp_path)
    created = service.import_document(_word_payload(recipes=2))
    import_id = created["importacion"]["documento"]["id"]
    draft = service.get_draft(import_id)["borrador"]
    assert len(draft["recipes"]) == 2
    first, second = draft["recipes"]
    ingredient = second["ingredients"][0]

    updated = service.update_draft(import_id, {
        "draft_version": draft["version"],
        "recipes": [
            {**first, "title": "Elaboración principal", "entity_type": "PRINCIPAL"},
            {
                **second,
                "entity_type": "SUBELABORACION",
                "parent_recipe_id": first["id"],
                "proposed_action": "CREAR_SUBELABORACION",
                "ingredients": [{
                    **ingredient,
                    "quantity_raw": "2 3",
                    "unit_raw": "l",
                    "name_raw": "nata",
                    "observations": "para terminar la salsa",
                    "relation_status": "SIN_RELACIONAR",
                    "article_id": None,
                }],
            },
        ],
    })
    assert updated["ok"] is True
    saved = updated["borrador"]
    assert saved["version"] == 2
    assert saved["recipes"][0]["title"] == "Elaboración principal"
    assert saved["recipes"][1]["parent_recipe_id"] == first["id"]
    assert saved["recipes"][1]["ingredients"][0]["quantity"] is None
    assert saved["recipes"][1]["ingredients"][0]["validation_errors"][0]["code"] == "CANTIDAD_AMBIGUA"

    conflict = service.update_draft(import_id, {
        "draft_version": 1,
        "recipes": saved["recipes"],
    })
    assert conflict["error"]["status"] == 409
    assert conflict["error"]["code"] == "draft_version_conflict"


def test_borrador_articulos_exactos_candidatos_y_sin_candidato(tmp_path: Path) -> None:
    db = tmp_path / "DATOS" / "db"
    db.mkdir(parents=True)
    (db / "articulos.json").write_text(json.dumps([
        {"codigo": "ART-PAT", "nombre": "Patata", "unidad": "kg", "precio": 1.2},
        {"codigo": "ART-NATA", "nombre": "Nata culinaria", "unidad": "l", "precio": 4.5},
        {"codigo": "ART-MANT", "nombre": "Mantequilla", "unidad": "kg", "precio": 8},
    ]), encoding="utf-8")
    service = ImportDocumentService(tmp_path)
    created = service.import_document(_word_payload())
    ingredients = created["importacion"]["borrador"]["recipes"][0]["ingredients"]
    patata, mantequilla = ingredients
    assert patata["relation_status"] == "COINCIDENCIA_EXACTA_PROPUESTA"
    assert patata["article_id"] == "ART-PAT"
    assert patata["article_candidates"][0]["motivo"]
    assert mantequilla["article_id"] == "ART-MANT"

    candidate = service.drafts.articles.find("nata", "l")
    assert candidate["status"] == "REVISAR_COINCIDENCIA"
    assert candidate["candidates"][0]["articulo_id"] == "ART-NATA"
    assert service.drafts.articles.find("ingrediente inexistente", "kg")["status"] == "SIN_RELACIONAR"


def test_borrador_permite_anadir_varios_eliminar_y_revalidar_ingredientes(
    tmp_path: Path,
) -> None:
    service = ImportDocumentService(tmp_path)
    created = service.import_document(_word_payload(recipes=2))
    import_id = created["importacion"]["documento"]["id"]
    draft = service.get_draft(import_id)["borrador"]
    empty_recipe = draft["recipes"][0]
    empty_recipe["ingredients"] = []

    with_invalid = service.update_draft(import_id, {
        "draft_version": draft["version"],
        "recipes": [{
            **empty_recipe,
            "ingredients": [
                {
                    "id": "NEW-1", "quantity_raw": "cantidad desconocida",
                    "unit_raw": "kg", "name_raw": "Naranja",
                    "observations": "Sin piel", "relation_status": "SIN_RELACIONAR",
                    "article_id": None,
                },
                {
                    "id": "NEW-2", "quantity_raw": "0,5",
                    "unit_raw": "l", "name_raw": "Fondo",
                    "observations": "", "relation_status": "SIN_RELACIONAR",
                    "article_id": None,
                },
            ],
        }],
    })
    assert with_invalid["ok"] is True
    saved = with_invalid["borrador"]
    ingredients = saved["recipes"][0]["ingredients"]
    assert len(ingredients) == 2
    assert all(not item["id"].startswith("NEW-") for item in ingredients)
    assert ingredients[0]["validation_errors"][0]["code"] == "CANTIDAD_INVALIDA"
    assert not any(
        issue["code"] == "INGREDIENTES_VACIOS"
        for issue in saved["recipes"][0]["validation_errors"]
    )

    corrected = service.update_draft(import_id, {
        "draft_version": saved["version"],
        "recipes": [{
            **saved["recipes"][0],
            "ingredients": [{
                **ingredients[0],
                "quantity_raw": "1,25",
            }],
        }],
    })
    assert corrected["ok"] is True
    final_recipe = corrected["borrador"]["recipes"][0]
    assert len(final_recipe["ingredients"]) == 1
    assert final_recipe["ingredients"][0]["quantity"] == 1.25
    assert final_recipe["ingredients"][0]["validation_errors"] == []


def test_http_borrador_get_patch_conflicto_y_sin_escritura(tmp_path: Path) -> None:
    client = TestClient(create_app(HostAIPlatformAPI(base_dir=tmp_path)))
    created = client.post("/api/v1/biblioteca/importaciones", json=_word_payload()).json()
    import_id = created["importacion"]["documento"]["id"]
    draft_response = client.get(
        f"/api/v1/biblioteca/importaciones/{import_id}/borrador"
    )
    assert draft_response.status_code == 200
    draft = draft_response.json()["borrador"]
    recipe = {**draft["recipes"][0], "entity_type": "DESCARTAR", "proposed_action": "IGNORAR"}

    saved = client.patch(
        f"/api/v1/biblioteca/importaciones/{import_id}/borrador",
        json={"draft_version": draft["version"], "recipes": [recipe]},
    )
    assert saved.status_code == 200
    assert saved.json()["borrador"]["recipes"][0]["entity_type"] == "DESCARTAR"
    assert saved.json()["datos_reales_modificados"] is False

    stale = client.patch(
        f"/api/v1/biblioteca/importaciones/{import_id}/borrador",
        json={"draft_version": draft["version"], "recipes": [recipe]},
    )
    assert stale.status_code == 409
    assert stale.json()["error"]["code"] == "draft_version_conflict"


def test_fachada_registra_y_expone_excepcion_real_solo_en_desarrollo(
    tmp_path: Path,
    monkeypatch,
    caplog,
) -> None:
    class FailingImportService:
        @staticmethod
        def import_document(_payload: dict) -> dict:
            raise RuntimeError("fallo reproducible del importador")

    facade = CatalogPublicFacade(base_dir=tmp_path)
    facade._biblioteca_import_service = FailingImportService()
    monkeypatch.setenv("HOST_AI_API_ENV", "development")

    with caplog.at_level(logging.ERROR):
        result = facade.crear_importacion_biblioteca(_word_payload())

    assert result["ok"] is False
    assert result["error"]["code"] == "library_import_failed"
    assert result["error"]["message"] == "RuntimeError: fallo reproducible del importador"
    assert "Error al crear una importación de Biblioteca." in caplog.text
