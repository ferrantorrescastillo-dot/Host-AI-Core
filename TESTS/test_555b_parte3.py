from __future__ import annotations

import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from openpyxl import Workbook

from SERVICIOS.extractor_fichas_tecnicas_excel_555b import ExtractorFichasTecnicasExcel555B
from SERVICIOS.normalizador_fichas_escandallo_555b import NormalizadorFichasEscandallo555B


def crear_excel(ruta: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "M.P PRUEBA"
    ws["C4"] = "FICHA TÉCNICA PLATO"
    ws["C7"] = "ARTÍCULO"
    ws["D7"] = "Paella de prueba"
    ws["C9"] = "PAX"
    ws["D9"] = 10
    ws["B11"] = "ARTICULO"
    ws["C11"] = "KG"
    ws["D11"] = "€/UN"
    ws["B12"] = "Arroz bomba"
    ws["C12"] = 1.0
    ws["D12"] = 2.15
    ws["B13"] = "Caldo de pescado"
    ws["C13"] = 5.0
    wb.save(ruta)


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        ruta = Path(tmp) / "escandallos.xlsx"
        crear_excel(ruta)
        resultado = ExtractorFichasTecnicasExcel555B().extraer(ruta)
        assert resultado["resumen"]["fichas_detectadas"] == 1
        assert resultado["resumen"]["fichas_validas_para_importar"] == 1
        ficha = resultado["fichas"][0]
        assert ficha["nombre"] == "Paella de prueba"
        assert ficha["rendimiento"] == 10
        assert ficha["unidad_rendimiento"] == "pax"
        assert len(ficha["ingredientes"]) == 2
        assert ficha["ingredientes"][0]["nombre"] == "Arroz bomba"
        assert ficha["ingredientes"][0]["cantidad"] == 1.0
        canonico = NormalizadorFichasEscandallo555B().normalizar_ficha(ficha)
        assert canonico["estado_importacion"] == "VISTA_PREVIA"
        assert canonico["ingredientes"][0]["origen"]["fila"] == 12
        assert resultado["datos_reales_modificados"] is False
        assert resultado["escandallos_importados"] == 0
    print("TEST OK 5.5.5B PARTE 3 - Extracción y normalización de fichas técnicas")


if __name__ == "__main__":
    main()
