from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import sys
import argparse
import base64
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import uvicorn
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from API.app import HostAIPlatformAPI
from API.http_server import create_app
from SERVICIOS.biblioteca_recetas_601 import RepositorioBibliotecaRecetas601
from SERVICIOS.importador_inteligente_biblioteca import ImportDocumentService
from SERVICIOS.catalog_crud_write_service import CatalogCrudWriteService
from SERVICIOS.host_ai_authorized_execution_context import AuthorizedExecutionContext
from SERVICIOS.precio_referencia_web_service import PrecioReferenciaWebService
from SERVICIOS.precio_referencias_import_service import PrecioReferenciasImportService
from SERVICIOS.recipe_completion_exchange_service import RecipeCompletionExchangeService
from SERVICIOS.receta_documentacion_batch_service import RecetaDocumentacionBatchService
from SERVICIOS.receta_documentacion_write_service import RecetaDocumentacionWriteService


def _runtime_dir() -> Path:
    raw = str(os.environ.get("HOST_AI_E2E_BASE_DIR") or "").strip()
    if not raw:
        raise RuntimeError("HOST_AI_E2E_BASE_DIR es obligatorio.")
    runtime = Path(raw).resolve()
    allowed_parent = (PROJECT_ROOT / ".test-runs").resolve()
    if runtime.parent != allowed_parent or runtime.name not in {
        "fase1-e2e", "fase1-articles-e2e", "fase1-closing-e2e", "fase1-closing-smoke",
        "fase1-mass-e2e", "fase1-mass-smoke", "fase1-download-e2e",
        "fase1-download-clean-e2e", "fase1-contract-e2e",
        "fase1-rehydration-e2e", "fase1-persistence-e2e", "fase1-master-contract-e2e",
        "fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke",
    }:
        raise RuntimeError(f"Directorio E2E no autorizado: {runtime}")
    return runtime


MASS_RECIPE_NAMES = (
    "Ensaladilla", "Crema de calabaza", "Gazpacho de tomate", "Salmorejo cordobés",
    "Salsa romesco", "Paella de alcachofas", "Agua de jamaica", "Hummus de garbanzos",
    "Croquetas de setas", "Tortilla de patata", "Bacalao al pilpil", "Pollo al ajillo",
    "Lentejas estofadas", "Crema de puerros", "Pisto de verduras", "Albóndigas en salsa",
    "Merluza en salsa verde", "Arroz caldoso de marisco", "Fideuá de verduras",
    "Canelones de espinacas", "Sopa de cebolla", "Ensalada de quinoa",
    "Patatas a la riojana", "Marmitako de bonito", "Consomé de ave", "Salsa bechamel",
    "Alioli de ajo asado", "Vinagreta de mostaza", "Puré de patata", "Risotto de setas",
    "Cuscús de verduras", "Curry de garbanzos", "Estofado de ternera", "Lubina al horno",
    "Pulpo a la gallega", "Calamares encebollados", "Flan de vainilla", "Arroz con leche",
    "Tarta de queso", "Compota de manzana", "Bizcocho de limón", "Crema catalana",
    "Helado de chocolate", "Sorbete de mango", "Pan de focaccia", "Masa de pizza",
    "Fondo oscuro", "Fumet de pescado", "Aceite de hierbas", "Cebolla encurtida",
    "Verduras asadas", "Fruta marinada",
)


def _fixture_recipe_name(index: int) -> str:
    if index >= len(MASS_RECIPE_NAMES):
        raise AssertionError("El fixture E2E necesita un nombre humano explícito por receta.")
    return MASS_RECIPE_NAMES[index]


def _seed_recipes(runtime: Path, count: int = 30) -> list[str]:
    assert count <= len(MASS_RECIPE_NAMES)
    recipe_ids = [f"REC601-{index + 1:06d}" for index in range(count)]
    now = datetime.now(timezone.utc).isoformat()
    recipes = [{
        "id": recipe_id,
        "codigo": recipe_id,
        "nombre": _fixture_recipe_name(index),
        "familia": "fixture_e2e",
        "tipo": "ELABORACION",
        "ingredientes": (
            ["Ingrediente controlado", "Ingrediente sin precio"] if count >= 50 and index == 5
            else ["Flor de hibiscus", "Limones", "Azúcar"] if count >= 50 and index == 6
            else ["Ingrediente controlado"]
        ),
        "cantidades": (
            ["1 kg", "0.5 kg"] if count >= 50 and index == 5
            else ["0.5", "1", "1"] if count >= 50 and index == 6
            else ["1 kg"]
        ),
        "descripcion": "",
        "elaboracion": "",
        "observaciones": "",
        "alergenos": [],
        "estado": "PENDIENTE_DE_COMPLETAR",
        "version": 1,
        "creado_en": now,
        "actualizado_en": now,
    } for index, recipe_id in enumerate(recipe_ids)]
    path = runtime / "DATOS/db/biblioteca_recetas_601.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({
        "version_modelo": "6.0.1",
        "actualizado_en": now,
        "recetas": recipes,
        "fichas_tecnicas": recipes,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return recipe_ids


def _seed_articles(runtime: Path) -> None:
    path = runtime / "DATOS/db/articulos.json"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps([{
        "codigo": "ART-INGREDIENTE-CONTROLADO",
        "nombre": "Ingrediente controlado",
        "unidad_base": "kg", "unidad_compra": "kg", "precio": 1.0,
        "proveedor": "Proveedor fixture", "estado": "ACTIVO", "activo": True,
    }, {
        "codigo": "ART-INGREDIENTE-SIN-PRECIO", "nombre": "Ingrediente sin precio",
        "unidad_base": "kg", "unidad_compra": "kg", "precio": None,
        "proveedor": "", "estado": "PENDIENTE_DE_COMPLETAR", "activo": True,
        "precios_referencia": [{
            "tipo": "PRECIO_REFERENCIA_IMPORTADA", "origen": "REFERENCIA_EXTERNA",
            "precio_normalizado": 3.5, "unidad_normalizada": "kg",
            "tienda_referencia": "Proveedor externo fixture",
            "url": "https://example.test/referencia", "consultado_en": "2026-09-05",
            "confianza": 0.8, "autoridad": "REFERENCIA_NO_REAL",
        }],
    }, {
        "codigo": "ART-FLOR-HIBISCUS", "nombre": "Flor de hibiscus",
        "unidad_base": "kg", "unidad_compra": "kg", "precio": 4.0,
        "proveedor": "Proveedor fixture", "estado": "ACTIVO", "activo": True,
    }, {
        "codigo": "ART-LIMONES", "nombre": "Limones",
        "unidad_base": "kg", "unidad_compra": "kg", "precio": 2.0,
        "proveedor": "Proveedor fixture", "estado": "ACTIVO", "activo": True,
    }, {
        "codigo": "ART-AZUCAR", "nombre": "Azúcar",
        "unidad_base": "kg", "unidad_compra": "kg", "precio": 1.8,
        "proveedor": "Proveedor fixture", "estado": "ACTIVO", "activo": True,
    }], ensure_ascii=False, indent=2), encoding="utf-8")


def _prepare_economic_sentinel(runtime: Path) -> None:
    """Remove real-price authority so the XLSX references drive the sentinel."""
    path = runtime / "DATOS/db/articulos.json"
    articles = json.loads(path.read_text(encoding="utf-8"))
    sentinel_names = {"Flor de hibiscus", "Limones", "Azúcar"}
    for article in articles:
        if str(article.get("nombre") or "") in sentinel_names:
            article["precio"] = None
            article["proveedor"] = ""
    path.write_text(json.dumps(articles, ensure_ascii=False, indent=2), encoding="utf-8")


def _complete_economic_references(workbook, recipe_id: str) -> None:
    references = workbook["PRECIOS_REFERENCIA"]
    headers = {cell.value: cell.column for cell in references[1]}
    fixtures = {
        "Flor de hibiscus": ("Hibisco flor a granel 1 kg", "Herbolínea", 15.9),
        "Limones": ("Limones malla 1 kg", "Alcampo", 2.79),
        "Azúcar": ("Azúcar blanco 1 kg", "Alcampo", 0.89),
    }
    completed: set[str] = set()
    for row in range(2, references.max_row + 1):
        name = str(references.cell(row, headers["nombre_canonico"]).value or "")
        fixture = fixtures.get(name)
        if not fixture:
            continue
        product, supplier, price = fixture
        values = {
            "producto_encontrado": product, "comercio_fuente": supplier,
            "precio_observado": price, "moneda": "EUR", "formato_envase": "1 kg",
            "cantidad_envase": 1, "unidad_envase": "kg", "precio_normalizado": price,
            "unidad_precio_normalizado": "kg",
            "url_fuente": f"https://example.test/{name.casefold().replace(' ', '-')}",
            "fecha_consulta": "2026-09-05",
            "observacion_equivalencia": "Referencia externa controlada para el sentinel E2E",
            "confianza": 0.9, "price_basis": "IVA_INCLUIDO",
            "procedencia": "REFERENCIA_EXTERNA", "estado_referencia": "REFERENCIA_PROPUESTA",
        }
        for field, value in values.items():
            references.cell(row, headers[field]).value = value
        completed.add(name)
    assert completed == set(fixtures)
    assert recipe_id


def _seed_import(runtime: Path, recipe_ids: list[str]) -> str:
    package = {
        "schema": "hostai.import.package",
        "version": "0.1",
        "metadata": {"source": "fixture-boronat-equivalente", "generator": "fase1-e2e-post-fix"},
        "recipes": [{
            "source_id": f"BORONAT-E2E-{index + 1:02d}",
            "name": _fixture_recipe_name(index),
            "ingredients": ([
                {"name": "Flor de hibiscus", "quantity": 0.5, "unit": None, "observed": "Flor de hibiscus 0.5"},
                {"name": "Limones", "quantity": 1, "unit": None, "observed": "Limones 1"},
                {"name": "Azúcar", "quantity": 1, "unit": None, "observed": "Azúcar 1"},
            ] if len(recipe_ids) >= 50 and index == 6 else [{
                "name": "Ingrediente controlado", "quantity": 1, "unit": "kg",
                "observed": "Ingrediente controlado",
            }] + ([{
                "name": "Ingrediente sin precio", "quantity": 0.5, "unit": "kg",
                "observed": "Ingrediente sin precio",
            }] if len(recipe_ids) >= 50 and index == 5 else [])),
            "procedure": [],
            "occurrences": [{"sheet": f"RECETA-{index + 1:02d}", "row": 2}],
        } for index, _recipe_id in enumerate(recipe_ids)],
        "articles": [{
            "source_code": "BORONAT-ART-E2E-1", "name": "Ingrediente controlado", "unit": "kg",
            "interpretation": {"classification": "ARTICULO_COMPRADO_CANDIDATE"},
        }],
        "suppliers": [], "menus": [], "relations": [], "ambiguities": [], "variant_groups": [],
    }
    service = ImportDocumentService(runtime, persistent_sessions=True)
    result = service.import_document({
        "hostai_import_package": package,
    })
    assert result["ok"] is True
    session = result["importacion"]
    assert session["schema_version"] == 2
    assert session["resumen"]["recetas_detectadas"] == len(recipe_ids)
    assert session["resolucion_identidad"]["recetas"]["ya_canonicas"] == len(recipe_ids)
    decisions = [{
        "id": item["id"],
        "proposed_action": "REUTILIZAR_EXISTENTE",
        "identity_decision": "MISMA_RECETA",
        "selected_canonical_recipe_id": item["duplicate_candidates"][0]["id"],
    } for item in session["borrador"]["recipes"]]
    updated = service.update_draft(session["documento"]["id"], {
        "draft_version": session["borrador"]["version"],
        "recipes": decisions,
    })
    assert updated["ok"] is True
    source_names = [_fixture_recipe_name(index) for index in range(len(recipe_ids))]
    assert [item.get("title") for item in updated["borrador"]["recipes"]] == source_names
    selected = {
        str(item.get("selected_canonical_recipe_id") or "")
        for item in updated["borrador"]["recipes"]
    }
    assert set(recipe_ids).issubset(selected)
    draft = updated["borrador"]
    recipes = json.loads(json.dumps(draft["recipes"]))
    for index in range(2):
        recipes[index]["ingredients"].append({
            "id": f"NEW-E2E-ESTRAGON-{index + 1}",
            "name_raw": "Estragón nuevo E2E",
            "quantity_raw": "0.05" if index == 0 else "50",
            "unit_raw": "kg" if index == 0 else "g",
            "observations": "Ingrediente nuevo controlado para E2E",
            "relation_status": "SIN_RELACIONAR",
            "article_id": None,
        })
    enriched = service.update_draft(session["documento"]["id"], {
        "draft_version": draft["draft_version"],
        "recipes": recipes,
        "variant_decisions": draft.get("variant_decisions") or [],
        "article_decisions": draft.get("article_decisions") or [],
        "menu_decisions": draft.get("menu_decisions") or [],
    })
    assert enriched["ok"] is True
    assert sum(
        ingredient.get("name_raw") == "Estragón nuevo E2E"
        for recipe in enriched["borrador"]["recipes"]
        for ingredient in recipe.get("ingredients") or []
    ) == 2
    return str(session["documento"]["id"])


def _canonical_recipe_id(candidate: dict) -> str:
    value = str(candidate.get("recipe_id") or candidate.get("receta_id") or candidate.get("id") or "")
    return value if value.startswith("REC601-") and value[7:].isdigit() else ""


def _seed_real_boronat_import(runtime: Path) -> tuple[str, list[str], dict[str, int]]:
    source_db = PROJECT_ROOT / "DATOS/db"
    target_db = runtime / "DATOS/db"
    shutil.copytree(source_db, target_db)
    for workflow in (
        "biblioteca_importaciones_web.json",
        "biblioteca_completado_recetas_batches.json",
        "biblioteca_precio_referencias_importaciones.json",
    ):
        (target_db / workflow).unlink(missing_ok=True)

    excel = PROJECT_ROOT / "Documentos/Escandallos Boronat.xlsx"
    service = ImportDocumentService(runtime, persistent_sessions=True)
    imported = service.import_document({
        "archivos": [{
            "nombre": excel.name,
            "tipo_mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "contenido_base64": base64.b64encode(excel.read_bytes()).decode("ascii"),
        }],
        "excluir_ap_antiguos": True,
    })
    assert imported["ok"] is True
    session = imported["importacion"]
    decisions = []
    for recipe in session["borrador"]["recipes"]:
        candidates = list(dict.fromkeys(
            candidate_id for candidate_id in (
                _canonical_recipe_id(item) for item in recipe.get("duplicate_candidates") or []
            ) if candidate_id
        ))
        selected = str(recipe.get("selected_canonical_recipe_id") or "")
        if not selected and len(candidates) == 1:
            selected = candidates[0]
        if selected:
            decisions.append({
                "id": recipe["id"], "proposed_action": "REUTILIZAR_EXISTENTE",
                "identity_decision": "MISMA_RECETA",
                "selected_canonical_recipe_id": selected,
            })
    updated = service.update_draft(session["documento"]["id"], {
        "draft_version": session["borrador"]["version"], "recipes": decisions,
    })
    assert updated["ok"] is True
    selected_ids = list(dict.fromkeys(
        str(item.get("selected_canonical_recipe_id") or "")
        for item in updated["borrador"]["recipes"]
        if str(item.get("selected_canonical_recipe_id") or "")
    ))
    assert len(selected_ids) >= 30

    draft = updated["borrador"]
    recipes = json.loads(json.dumps(draft["recipes"]))
    selected_recipes = [item for item in recipes if item.get("selected_canonical_recipe_id")]
    for index, recipe in enumerate(selected_recipes[:2]):
        recipe["ingredients"].append({
            "id": f"NEW-E2E-ESTRAGON-{index + 1}",
            "name_raw": "Estragón nuevo E2E",
            "quantity_raw": "0.05" if index == 0 else "50",
            "unit_raw": "kg" if index == 0 else "g",
            "observations": "Ingrediente nuevo controlado para E2E",
            "relation_status": "SIN_RELACIONAR", "article_id": None,
        })
    enriched = service.update_draft(session["documento"]["id"], {
        "draft_version": draft["draft_version"], "recipes": recipes,
        "variant_decisions": draft.get("variant_decisions") or [],
        "article_decisions": draft.get("article_decisions") or [],
        "menu_decisions": draft.get("menu_decisions") or [],
    })
    assert enriched["ok"] is True
    entities = session.get("documento", {}).get("entidades") or []
    analysis = session.get("analisis_restaurante") or {}
    metrics = {
        "detected_recipes": int((session.get("resolucion_identidad", {}).get("recetas") or {}).get("total") or 0),
        "menus": sum(item.get("kind") == "MENU" for item in entities),
        "ap_excluded": len(analysis.get("exclusiones_sesion") or []),
    }
    assert metrics["menus"] > 0 and metrics["ap_excluded"] > 0
    return str(session["documento"]["id"]), selected_ids, metrics


def _seed_external_batch(runtime: Path, recipe_ids: list[str], import_id: str) -> dict[str, object]:
    repository = RepositorioBibliotecaRecetas601(runtime)
    write_service = RecetaDocumentacionWriteService(runtime, repository=repository)
    batch_service = RecetaDocumentacionBatchService(
        runtime, repository=repository, recipe_service=write_service,
    )
    exchange = RecipeCompletionExchangeService(
        runtime, repository=repository, recipe_service=write_service,
        batch_service=batch_service,
    )
    exported = exchange.export(
        recipe_ids=recipe_ids, scope="IMPORTACION", import_id=import_id,
    )
    included_ids = list(exported["recipe_ids"])
    assert included_ids
    workbook = load_workbook(BytesIO(__import__("base64").b64decode(exported["contenido_base64"])))
    assert dict(workbook["METADATA"].values)["version"] == "0.3"
    assert "PROMPT_IA" in workbook.sheetnames
    assert "Procesa TODAS las recetas" in str(workbook["PROMPT_IA"].cell(2, 3).value or "")
    assert "PRECIOS_REFERENCIA" in workbook.sheetnames
    assert "SCHEMA_PRECIOS" in workbook.sheetnames
    assert "INSTRUCCIONES" in workbook.sheetnames
    assert any(
        "NO_APLICA" in str(cell.value or "")
        for row in workbook["INSTRUCCIONES"].iter_rows() for cell in row
    )
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    recipes_by_id = {
        str(item.get("id") or item.get("codigo") or ""): item
        for item in repository.listar()
    }
    expected_names = {
        recipe_id: str((recipes_by_id.get(recipe_id) or {}).get("nombre") or "").strip()
        for recipe_id in included_ids
    }
    xlsx_names = {
        str(sheet.cell(row, headers["recipe_id"]).value or ""): str(
            sheet.cell(row, headers["nombre"]).value or ""
        ).strip()
        for row in range(2, sheet.max_row + 1)
    }
    assert all(expected_names.values())
    assert all(not re.fullmatch(r"Receta\s+\d+", name) for name in expected_names.values())
    assert xlsx_names == expected_names
    agua_recipe_id = next((
        recipe_id for recipe_id in included_ids
        if str((recipes_by_id.get(recipe_id) or {}).get("nombre") or "").strip().casefold()
        == "agua de jamaica"
    ), "")
    if runtime.name in {"fase1-closing-e2e", "fase1-closing-smoke"}:
        assert agua_recipe_id, "El fixture real debe incluir Agua de jamaica."
    agua_export_context: dict[str, object] = {}
    if agua_recipe_id:
        agua_row = included_ids.index(agua_recipe_id) + 2
        agua_export_context = {
            "import_id": sheet.cell(agua_row, headers["import_id"]).value,
            "contexto_documental": json.loads(sheet.cell(agua_row, headers["contexto_documental"]).value or "{}"),
            "ingredientes": json.loads(sheet.cell(agua_row, headers["ingredientes_contexto"]).value or "[]"),
            "articulos": json.loads(sheet.cell(agua_row, headers["articulos_relacionados"]).value or "[]"),
            "menus": json.loads(sheet.cell(agua_row, headers["menus_contexto"]).value or "[]"),
            "servicio": json.loads(sheet.cell(agua_row, headers["contexto_servicio"]).value or "[]"),
            "datos_documentales": json.loads(sheet.cell(agua_row, headers["datos_documentales"]).value or "{}"),
            "datos_calculados": json.loads(sheet.cell(agua_row, headers["datos_calculados"]).value or "{}"),
        }
        assert agua_export_context["import_id"] == import_id
        assert agua_export_context["ingredientes"]
        if runtime.name in {"fase1-closing-e2e", "fase1-closing-smoke"}:
            assert agua_export_context["menus"] and agua_export_context["servicio"]
        assert all("nombre_original" in item for item in agua_export_context["ingredientes"])
    critical_values = {
        "vida_util_refrigerado": "24 horas", "vida_util_congelado": "30 dias",
        "puede_congelarse": "true",
        "puede_refrigerarse": "true", "regeneracion": "Regenerar y validar",
        "tiempo_activo": "30 minutos", "tiempo_pasivo": "60 minutos",
        "tiempo_total": "90 minutos", "numero_raciones": 10,
    }
    for index, recipe_id in enumerate(included_ids):
        row = index + 2
        if index != 0 and index <= 23:
            sheet.cell(row, headers["descripcion_propuesto"]).value = f"Descripcion {recipe_id}"
        sheet.cell(row, headers["elaboracion_propuesto"]).value = f"Elaboracion {recipe_id}"
        sheet.cell(row, headers["observaciones_propuesto"]).value = f"Observaciones {recipe_id}"
        critical_count = 9 if index < 17 else 8
        for field, value in list(critical_values.items())[:critical_count]:
            sheet.cell(row, headers[f"{field}_propuesto"]).value = value
        if runtime.name in {
            "fase1-mass-e2e", "fase1-mass-smoke", "fase1-contract-e2e",
            "fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke",
        }:
            recipe = recipes_by_id[recipe_id]
            structured = []
            for ingredient_index, ingredient_name in enumerate(recipe.get("ingredientes") or []):
                quantity = 0.5 if ingredient_name == "Ingrediente sin precio" else 1
                structured.append({
                    "line_id": f"{recipe_id}-ING-{ingredient_index + 1}",
                    "nombre_original": ingredient_name, "name_raw": ingredient_name,
                    "cantidad": quantity, "unidad": "kg",
                    "cantidad_normalizada": quantity, "unidad_normalizada": "kg",
                    "procedencia_propuesta": {"origen": "IA_PROPUESTA", "confianza": 0.9},
                })
            operational = {
                "tipo_elaboracion": "ELABORACION", "rendimiento": 10,
                "unidad_rendimiento": "raciones", "numero_raciones": 10,
                "cantidad_por_racion": {"cantidad": 250, "unidad": "g"},
                "ingredientes_estructurados": structured,
                "tiempo_activo": "20 minutos", "tiempo_pasivo": "10 minutos",
                "tiempo_total": "30 minutos", "produccion_maxima": 20,
                "unidad_tanda": "raciones", "rendimiento_por_tanda": 20,
                "personal_recomendado": {"personas": 1, "rol": "cocinero"},
                "recursos_necesarios": ["mesa de preparación", "recipiente"],
                "puede_refrigerarse": False, "puede_congelarse": False,
                "conservacion": "Servicio inmediato; propuesta pendiente de revisión.",
                "regeneracion": "NO_APLICA",
            }
            if index == 1:
                operational["puede_congelarse"] = "quizas"
            if index == 2:
                operational["unidad_tanda"] = (
                    "PENDIENTE_IMPOSIBLE_DE_ESTIMAR: falta capacidad del equipo"
                )
            for field, value in operational.items():
                sheet.cell(row, headers[f"{field}_propuesto"]).value = (
                    json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
                )
            sheet.cell(row, headers["metadatos_propuestas"]).value = json.dumps({
                field: {
                    "origen": "IA_PROPUESTA", "confianza": 0.5 if index == 4 else 0.9,
                    "motivo": "Propuesta mecánica E2E; no acredita calidad culinaria.",
                    "modelo": "MOCK_GPT_MASS_E2E",
                } for field in operational
            }, ensure_ascii=False)
            if index == 0 and runtime.name not in {
                "fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke",
            }:
                sheet.cell(row, headers["recipe_fingerprint"]).value = "stale-e2e"
        if recipe_id == agua_recipe_id:
            agua = recipes_by_id[recipe_id]
            structured = []
            for ingredient in list(agua.get("ingredientes_estructurados") or []):
                structured.append({
                    **dict(ingredient), "unidad": "kg",
                    "cantidad_normalizada": float(ingredient.get("cantidad") or 0),
                    "unidad_normalizada": "kg", "dato_provisional": True,
                })
            if not structured:
                for ingredient_index, ingredient_name in enumerate(agua.get("ingredientes") or []):
                    quantity_text = str((agua.get("cantidades") or [])[ingredient_index] or "1")
                    quantity = float(quantity_text.split()[0].replace(",", "."))
                    structured.append({
                        "line_id": f"{recipe_id}-ING-{ingredient_index + 1}",
                        "nombre_original": ingredient_name, "name_raw": ingredient_name,
                        "cantidad": quantity, "unidad": "kg", "cantidad_normalizada": quantity,
                        "unidad_normalizada": "kg", "dato_provisional": True,
                        "procedencia_propuesta": {"origen": "IA_PROPUESTA", "confianza": 0.9},
                    })
            operational = {
                "categoria": "Bebida", "tipo_elaboracion": "BEBIDA",
                "descripcion": "Infusion fria de flor de jamaica para servicio de menu.",
                "rendimiento": 10,
                "unidad_rendimiento": "raciones", "numero_raciones": 10,
                "cantidad_por_racion": {"valor": 250, "unidad": "ml"},
                "tiempo_activo": "25 minutos", "tiempo_pasivo": "65 minutos",
                "tiempo_total": "90 minutos", "intervencion_activa": True,
                "produccion_maxima": 20, "unidad_tanda": "raciones",
                "rendimiento_por_tanda": 20,
                "personal_recomendado": {"personas": 1, "rol": "cocinero"},
                "recursos_necesarios": ["Olla", "Colador", "Camara frigorifica"],
                "puede_refrigerarse": True, "vida_util_refrigerado": "48 horas",
                "puede_congelarse": False, "vida_util_congelado": "NO_APLICA",
                "conservacion": "Conservar refrigerada; propuesta pendiente de validacion sanitaria.",
                "tiempo_descongelacion": "NO_APLICA",
                "regeneracion": "NO_APLICA",
                "ingredientes_estructurados": structured,
            }
            if runtime.name in {"fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke"}:
                operational["ingredientes_estructurados"].append({
                    "line_id": f"{recipe_id}-AGUA-PENDIENTE",
                    "nombre_original": "Agua", "name_raw": "Agua",
                    "cantidad": 8, "unidad": "l", "cantidad_normalizada": 8,
                    "unidad_normalizada": "l", "estado_relacion": "CANDIDATO_NUEVO",
                    "dato_provisional": True,
                    "procedencia_propuesta": {"origen": "IA_PROPUESTA", "confianza": 0.9},
                })
            for field, value in operational.items():
                sheet.cell(row, headers[f"{field}_propuesto"]).value = (
                    json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value
                )
            sheet.cell(row, headers["metadatos_propuestas"]).value = json.dumps({
                field: {
                    "origen": "IA_PROPUESTA", "confianza": 0.9,
                    "motivo": "Propuesta E2E basada en menú, cantidades documentales y artículos relacionados.",
                    "modelo": "MOCK_E2E",
                }
                for field in operational
            }, ensure_ascii=False)
    if runtime.name in {"fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke"}:
        _complete_economic_references(workbook, agua_recipe_id)
    physical_path = runtime / "fixture-completado-externo.xlsx"
    workbook.save(physical_path)
    encoded = __import__("base64").b64encode(physical_path.read_bytes()).decode("ascii")
    imported = exchange.import_package(
        filename=physical_path.name,
        content_base64=encoded,
        expected_recipe_ids=included_ids,
        scope="IMPORTACION",
        import_id=import_id,
        source="CHATGPT",
    )
    validation = imported["validacion"]
    batch_names = {
        str(item.get("recipe_id") or ""): str(item.get("nombre") or "").strip()
        for item in imported["batch"]["resultados"]
    }
    assert batch_names == expected_names
    expected_safe = sum(
        len(item.get("datos_propuestos_seguros_masivo") or {})
        for item in imported["batch"]["resultados"]
    )
    expected_grouped = sum(
        len(item.get("datos_operativos_agrupables") or {})
        for item in imported["batch"]["resultados"]
    )
    expected_critical = sum(
        len(item.get("datos_requieren_revision_individual") or {})
        for item in imported["batch"]["resultados"]
    )
    agua_ready = False
    if agua_recipe_id:
        agua_result = next(
            item for item in imported["batch"]["resultados"]
            if item["recipe_id"] == agua_recipe_id
        )
        agua_ready = bool((agua_result.get("completitud") or {}).get("production_ready_provisional"))
        assert agua_ready is True
    assert validation["filas_recibidas"] == len(included_ids)
    if runtime.name in {
        "fase1-mass-e2e", "fase1-mass-smoke", "fase1-contract-e2e",
        "fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke",
    }:
        assert validation["filas_recibidas"] >= 50
        expected_rejected = 0 if runtime.name in {
            "fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke",
        } else 1
        assert validation["filas_rechazadas"] == expected_rejected
        if expected_rejected:
            assert any("RECETA_CAMBIO_DESDE_EXPORTACION" in row["errores"] for row in validation["filas"])
        assert any("puede_congelarse:invalid_boolean" in row["errores"] for row in validation["filas"])
        assert any(row.get("campos_imposibles_estimar") for row in validation["filas"])
    else:
        assert validation["filas_utiles"] == len(included_ids)
        assert validation["filas_requieren_revision"] == len(included_ids)
        assert validation["filas_rechazadas"] == 0
    assert validation["campos"]["utiles"] == expected_safe
    expected_impossible = sum(
        len(row.get("campos_imposibles_estimar") or []) for row in validation["filas"]
    )
    assert validation["campos"]["requieren_revision"] == expected_critical + expected_impossible
    assert imported["batch"]["progreso"]["propuestas"] == expected_safe + expected_grouped + expected_critical
    return {
        "batch_id": str(imported["batch"]["batch_id"]),
        "safe_count": int(validation["campos"]["utiles"]),
        "grouped_count": expected_grouped,
        "critical_count": expected_critical,
        "review_count": int(validation["campos"]["requieren_revision"]),
        "proposal_count": int(imported["batch"]["progreso"]["propuestas"]),
        "recipe_count": len(included_ids),
        "agua_recipe_id": agua_recipe_id,
        "xlsx_version": "0.3",
        "export_context_complete": bool(
            agua_export_context.get("import_id")
            and agua_export_context.get("ingredientes")
            and agua_export_context.get("menus")
            and agua_export_context.get("servicio")
        ),
        "production_ready_provisional": int(imported["batch"]["resumen_masivo"]["production_ready_provisional"]),
        "mass_errors": int(imported["batch"]["resumen_masivo"]["errores"]),
        "mass_impossible": int(imported["batch"]["resumen_masivo"]["imposibles_estimar"]),
        "mass_no_aplica": int(imported["batch"]["resumen_masivo"]["no_aplica"]),
        "agua_production_ready_provisional": agua_ready,
        "recipe_names": expected_names,
        "source_file": physical_path.name,
        "source_sha256": hashlib.sha256(physical_path.read_bytes()).hexdigest().upper(),
        "name_audit": {
            "canonical_matches_xlsx": xlsx_names == expected_names,
            "xlsx_matches_batch": batch_names == xlsx_names,
            "human_names": len(expected_names),
        },
    }


def prepare() -> dict[str, object]:
    runtime = _runtime_dir()
    if runtime.exists():
        shutil.rmtree(runtime)
    runtime.mkdir(parents=True)
    if runtime.name in {"fase1-closing-e2e", "fase1-closing-smoke"}:
        import_id, recipe_ids, metrics = _seed_real_boronat_import(runtime)
    else:
        _seed_articles(runtime)
        recipe_count = 52 if runtime.name in {
            "fase1-mass-e2e", "fase1-mass-smoke", "fase1-download-e2e",
            "fase1-download-clean-e2e", "fase1-contract-e2e", "fase1-master-contract-e2e",
            "fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke",
        } else 30
        recipe_ids = _seed_recipes(runtime, recipe_count)
        import_id = _seed_import(runtime, recipe_ids)
        if runtime.name in {"fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke"}:
            _prepare_economic_sentinel(runtime)
        metrics = {"detected_recipes": recipe_count, "menus": 0, "ap_excluded": 0}
    if runtime.name == "fase1-download-clean-e2e":
        state = {
            "import_id": import_id,
            "batch_id": None,
            "recipe_count": len(recipe_ids),
            "recipe_names": {
                recipe_id: _fixture_recipe_name(index)
                for index, recipe_id in enumerate(recipe_ids)
            },
            "xlsx_version": "0.3",
            "has_external_batch": False,
            **metrics,
        }
    else:
        batch = _seed_external_batch(runtime, recipe_ids, import_id)
        state = {
            "import_id": import_id,
            "batch_id": batch["batch_id"],
            "recipe_count": batch["recipe_count"],
            "safe_count": batch["safe_count"],
            "grouped_count": batch["grouped_count"],
            "critical_count": batch["critical_count"],
            "review_count": batch["review_count"],
            "proposal_count": batch["proposal_count"],
            "agua_recipe_id": batch["agua_recipe_id"],
            "xlsx_version": batch["xlsx_version"],
            "export_context_complete": batch["export_context_complete"],
            "production_ready_provisional": batch["production_ready_provisional"],
            "mass_errors": batch["mass_errors"],
            "mass_impossible": batch["mass_impossible"],
            "mass_no_aplica": batch["mass_no_aplica"],
            "agua_production_ready_provisional": batch["agua_production_ready_provisional"],
            "recipe_names": batch["recipe_names"],
            "source_file": batch["source_file"],
            "source_sha256": batch["source_sha256"],
            "name_audit": batch["name_audit"],
            "has_external_batch": True,
            **metrics,
        }
        if runtime.name in {"fase1-economic-exceptions-e2e", "fase1-economic-exceptions-smoke"}:
            protected = ("biblioteca_recetas_601.json", "articulos.json")
            state["protected_hashes"] = {
                filename: hashlib.sha256((runtime / "DATOS/db" / filename).read_bytes()).hexdigest().upper()
                for filename in protected
            }
    (runtime / "e2e-state.json").write_text(json.dumps(state, indent=2), encoding="utf-8")
    return state


def serve() -> None:
    runtime = _runtime_dir()
    if not (runtime / "e2e-state.json").exists():
        raise RuntimeError("El runtime E2E no está preparado.")
    app = create_app(HostAIPlatformAPI(base_dir=runtime))
    port = int(os.environ.get("HOST_AI_E2E_BACKEND_PORT") or 8011)
    uvicorn.run(app, host="127.0.0.1", port=port, log_level="warning")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--prepare-only", action="store_true")
    parser.add_argument("--serve-only", action="store_true")
    args = parser.parse_args()
    if args.prepare_only:
        prepare()
    elif args.serve_only:
        serve()
    else:
        prepare()
        serve()
