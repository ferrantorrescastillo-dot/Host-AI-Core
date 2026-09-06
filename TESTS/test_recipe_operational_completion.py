from __future__ import annotations

import base64
import json
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from SERVICIOS.biblioteca_culinaria_read_service import BibliotecaCulinariaReadService
from SERVICIOS.biblioteca_recetas_601 import RepositorioBibliotecaRecetas601
from SERVICIOS.host_ai_authorized_execution_context import AuthorizedExecutionContext
from SERVICIOS.recipe_completion_exchange_service import RecipeCompletionExchangeService
from SERVICIOS.receta_documentacion_batch_service import RecetaDocumentacionBatchService
from SERVICIOS.receta_documentacion_write_service import (
    RecetaDocumentacionError,
    RecetaDocumentacionWriteService,
    classify_recipe_proposals,
)


def _context() -> AuthorizedExecutionContext:
    return AuthorizedExecutionContext(
        "REQ-OPERATIONAL", "CHEF", "LOCAL", ("chef",), frozenset({"recetas:write"}),
    )


def _fixture(base: Path) -> Path:
    db = base / "DATOS" / "db"
    db.mkdir(parents=True)
    recipe = {
        "id": "REC601-000777", "codigo": "AGUA-JAMAICA", "nombre": "Agua de jamaica",
        "familia": "Bebidas", "tipo": "BEBIDA", "descripcion": "Bebida del menú de verano.",
        "numero_raciones": 0, "rendimiento": "", "unidad_rendimiento": "",
        "ingredientes": ["Flor de hibiscus", "Limón", "Azúcar"],
        "cantidades": ["0.5", "1", "1"],
        "ingredientes_estructurados": [
            {"nombre_original": "Flor de hibiscus", "cantidad": 0.5, "unidad": None},
            {"nombre_original": "Limón", "cantidad": 1, "unidad": None},
            {"nombre_original": "Azúcar", "cantidad": 1, "unidad": None},
        ],
        "elaboracion": "Infusionar, colar, endulzar y enfriar.",
        "tiempo_activo": "15 minutos", "tiempo_pasivo": "45 minutos", "tiempo_total": "",
        "conservacion": "", "alergenos": [], "observaciones": "", "fotografia": "",
        "estado": "PENDIENTE_DE_COMPLETAR", "version": 1,
        "creado_en": "2026-09-01T10:00:00", "actualizado_en": "2026-09-01T10:00:00",
        "completitud": {"campos_obligatorios_pendientes": ["Rendimiento", "Documentación"]},
        "procedencia_campos": {"descripcion": {"tipo": "DOCUMENTO"}},
    }
    recipe_path = db / "biblioteca_recetas_601.json"
    recipe_path.write_text(json.dumps({"version_modelo": "6.0.1", "recetas": [recipe]}, ensure_ascii=False), encoding="utf-8")
    (db / "biblioteca_escandallos_601.json").write_text(json.dumps({
        "version_modelo": "6.0.1", "escandallos": [{
            "id": "ESC601-000777", "codigo": "ESC-AGUA-JAMAICA", "nombre": "Agua de jamaica",
            "receta_asociada": {"id": recipe["id"], "codigo": recipe["codigo"]},
            "estado": "CON_INCIDENCIAS", "lineas": [], "otros_costes": 0,
        }], "historial_calculos": [], "simulaciones": [],
    }, ensure_ascii=False), encoding="utf-8")
    (db / "articulos.json").write_text(json.dumps([
        {"codigo": "ART-HIBISCUS", "nombre": "Flor de hibiscus", "precio": 10, "unidad": "kg"},
        {"codigo": "ART-LIMON", "nombre": "Limón", "precio": 2, "unidad": "u"},
        {"codigo": "ART-AZUCAR", "nombre": "Azúcar", "precio": 1, "unidad": "kg"},
    ], ensure_ascii=False), encoding="utf-8")
    (db / "menus.json").write_text(json.dumps([{
        "menu_id": "MENU601-000001", "codigo": "MENU-VERANO", "nombre": "Menú verano",
        "tipo": "Servicio de terraza", "estado": "BORRADOR", "modelo_biblioteca": "MENU_601",
        "comensales_recomendado": 10,
        "composicion": {"bebidas": [{
            "tipo_referencia": "RECETA", "referencia": recipe["id"], "cantidad": 10,
        }]},
    }], ensure_ascii=False), encoding="utf-8")
    (db / "proveedores.json").write_text("[]", encoding="utf-8")
    (db / "compras_producto_proveedor.json").write_text("[]", encoding="utf-8")
    invoices = base / "DATOS" / "facturas"
    invoices.mkdir(parents=True)
    (invoices / "historico_precios.json").write_text('{"registros":[]}', encoding="utf-8")
    return recipe_path


def _proposals() -> dict[str, object]:
    return {
        "descripcion": "No debe sustituir el texto documental.",
        "tiempo_total": {"valor": "60 minutos", "origen": "CALCULADO", "confianza": 1},
        "rendimiento": {"valor": 10, "confianza": 0.9, "motivo": "El menú prevé diez servicios."},
        "unidad_rendimiento": "raciones", "numero_raciones": 10,
        "cantidad_por_racion": {"valor": 250, "unidad": "ml"},
        "produccion_maxima": 20, "personal_recomendado": {"personas": 1, "rol": "cocinero"},
        "conservacion": "Mantener refrigerada hasta revisión del responsable.",
        "tiempo_descongelacion": {
            "estado": "NO_APLICA", "confianza": 0.98,
            "motivo": "La bebida se conserva refrigerada y no se congela.",
        },
        "ingredientes_estructurados": [
            {"nombre_original": "Flor de hibiscus", "cantidad": 0.5, "unidad": "kg", "confianza": 0.85},
            {"nombre_original": "Limón", "cantidad": 1, "unidad": "u", "confianza": 0.9},
            {"nombre_original": "Azúcar", "cantidad": 1, "unidad": "kg", "conversion": {"factor": 1}},
        ],
    }


def test_agua_jamaica_contexto_precedencia_propuestas_y_cero_write(tmp_path: Path) -> None:
    recipe_path = _fixture(tmp_path)
    calls: list[dict[str, object]] = []

    class Generator:
        def generate(self, *, recipe, missing_fields, allowed_fields):
            calls.append({"recipe": recipe, "missing_fields": missing_fields, "allowed_fields": allowed_fields})
            return _proposals()

    before = recipe_path.read_bytes()
    service = RecetaDocumentacionWriteService(tmp_path, generator=Generator())
    result = service.proposal(recipe_id="REC601-000777", proposed={})

    assert calls and calls[0]["recipe"]["_contexto_completado"]["menus"][0]["tipo"] == "Servicio de terraza"
    assert calls[0]["recipe"]["_contexto_completado"]["menus"][0]["pax"] == 10
    assert "descripcion" not in result["datos_propuestos_ia"]
    assert result["datos_propuestos_ia"]["tiempo_total"] == "60 minutos"
    assert result["metadatos_propuestas"]["tiempo_total"]["origen"] == "CALCULADO"
    assert result["datos_propuestos_ia"]["ingredientes_estructurados"][0]["unidad"] == "kg"
    assert result["datos_propuestos_ia"]["ingredientes_estructurados"][0]["cantidad_normalizada"] == 0.5
    safe, critical = classify_recipe_proposals(result["datos_propuestos_ia"])
    assert set(safe) <= {"descripcion", "elaboracion", "observaciones"}
    assert {"rendimiento", "numero_raciones", "cantidad_por_racion", "produccion_maxima", "personal_recomendado", "ingredientes_estructurados"} <= set(critical)
    assert result["completitud"]["propuesta"]["porcentaje"] > result["completitud"]["documental"]["porcentaje"]
    assert result["datos_reales_modificados"] is False
    assert recipe_path.read_bytes() == before


def test_referencia_externa_costea_candidato_reutilizado_sin_crear_articulo(tmp_path: Path) -> None:
    recipe_path = _fixture(tmp_path)
    db = tmp_path / "DATOS" / "db"
    recipes_payload = json.loads(recipe_path.read_text(encoding="utf-8"))
    second = dict(recipes_payload["recetas"][0])
    second.update({"id": "REC601-000778", "codigo": "AGUA-MENTA", "nombre": "Agua de menta"})
    recipes_payload["recetas"].append(second)
    recipe_path.write_text(json.dumps(recipes_payload, ensure_ascii=False), encoding="utf-8")
    articles_path = db / "articulos.json"
    articles = json.loads(articles_path.read_text(encoding="utf-8"))
    for article in articles:
        article.pop("precio", None)
    articles_path.write_text(json.dumps(articles, ensure_ascii=False), encoding="utf-8")
    before_recipes = recipe_path.read_bytes()
    before_articles = articles_path.read_bytes()

    proposals = _proposals()
    proposals["ingredientes_estructurados"] = [
        *proposals["ingredientes_estructurados"],
        {
            "line_id": "MENTA-PROPUESTA", "nombre_original": "Menta",
            "cantidad": 0.1, "unidad": "kg", "cantidad_normalizada": 0.1,
            "unidad_normalizada": "kg", "estado_relacion": "CANDIDATO_NUEVO",
            "dato_provisional": True,
        },
    ]
    reference = {
        "article_key": "CANDIDATO_NUEVO|menta|kg", "article_id": "",
        "nombre_canonico": "Menta", "producto": "Menta fresca 1 kg",
        "tienda_referencia": "Comercio de prueba", "precio_comercial": 8.0,
        "moneda": "EUR", "cantidad_formato": 1.0, "unidad_formato": "kg",
        "precio_normalizado": 8.0, "unidad_normalizada": "kg",
        "url": "https://example.test/menta", "consultado_en": "2026-09-05",
        "confianza": 0.8, "price_basis": "IVA_INCLUIDO",
        "origen": "REFERENCIA_EXTERNA", "autoridad": "REFERENCIA_NO_REAL",
        "estado_revision": "REQUIERE_REVISION_HUMANA",
    }
    read = BibliotecaCulinariaReadService(tmp_path)
    projections = [
        read.proyectar_provisional(recipe_id, propuestas=proposals, referencias_precio=[reference])
        for recipe_id in ("REC601-000777", "REC601-000778")
    ]

    assert [projection["escandallo"]["estado_coste"] for projection in projections] == ["PARCIAL", "PARCIAL"]
    assert [projection["escandallo"]["precios_referencia"] for projection in projections] == [1, 1]
    assert all(projection["escandallo"]["coste_provisional"] is True for projection in projections)
    for projection in projections:
        candidate = next(line for line in projection["escandallo"]["lineas"] if line["nombre_original"] == "Menta")
        assert candidate["coste_linea"] == 0.8
        assert candidate["clasificacion_precio"] == "REFERENCIA"
        assert candidate["referencia_precio"]["autoridad"] == "REFERENCIA_NO_REAL"
    assert recipe_path.read_bytes() == before_recipes
    assert articles_path.read_bytes() == before_articles


def test_precio_real_prevalece_sobre_referencia_externa_provisional(tmp_path: Path) -> None:
    recipe_path = _fixture(tmp_path)
    before = recipe_path.read_bytes()
    reference = {
        "article_key": "ART-HIBISCUS", "article_id": "ART-HIBISCUS",
        "nombre_canonico": "Flor de hibiscus", "producto": "Referencia cara",
        "tienda_referencia": "Comercio externo", "precio_comercial": 999.0,
        "cantidad_formato": 1.0, "unidad_formato": "kg",
        "precio_normalizado": 999.0, "unidad_normalizada": "kg",
        "origen": "REFERENCIA_EXTERNA", "autoridad": "REFERENCIA_NO_REAL",
        "estado_revision": "REQUIERE_REVISION_HUMANA",
    }

    projection = BibliotecaCulinariaReadService(tmp_path).proyectar_provisional(
        "REC601-000777", propuestas=_proposals(), referencias_precio=[reference],
    )

    hibiscus = next(
        line for line in projection["escandallo"]["lineas"]
        if line["nombre_original"] == "Flor de hibiscus"
    )
    assert hibiscus["coste_unitario"] == 10.0
    assert hibiscus["coste_linea"] == 5.0
    assert hibiscus["clasificacion_precio"] == "CONFIRMADO"
    assert hibiscus["origen_precio"] != "referencia_externa"
    assert recipe_path.read_bytes() == before


def test_preview_batch_rehidratado_conserva_escandallo_completo_de_referencias_externas(
    tmp_path: Path,
) -> None:
    recipe_path = _fixture(tmp_path)
    articles_path = tmp_path / "DATOS" / "db" / "articulos.json"
    articles = json.loads(articles_path.read_text(encoding="utf-8"))
    for article in articles:
        article.pop("precio", None)
    articles_path.write_text(json.dumps(articles, ensure_ascii=False), encoding="utf-8")
    before_recipe = recipe_path.read_bytes()
    before_articles = articles_path.read_bytes()

    proposals = _proposals()
    proposals["ingredientes_estructurados"] = [
        {"line_id": "HIBISCUS", "nombre_original": "Flor de hibiscus", "cantidad": 0.5, "unidad": "kg"},
        {"line_id": "LIMONES", "nombre_original": "Limón", "cantidad": 1, "unidad": "kg"},
        {"line_id": "AZUCAR", "nombre_original": "Azúcar", "cantidad": 1, "unidad": "kg"},
        {"line_id": "AGUA", "nombre_original": "Agua", "cantidad": 8, "unidad": "l", "estado_relacion": "CANDIDATO_NUEVO"},
    ]
    references = [
        {
            "article_key": "ART-HIBISCUS", "article_id": "ART-HIBISCUS",
            "nombre_canonico": "Flor de hibiscus", "producto": "Hibisco flor a granel 1 kg",
            "tienda_referencia": "Herbolínea", "precio_normalizado": 15.9,
            "unidad_normalizada": "kg", "precio_comercial": 15.9, "unidad_formato": "kg",
            "origen": "REFERENCIA_EXTERNA", "autoridad": "REFERENCIA_NO_REAL",
        },
        {
            "article_key": "CANDIDATO_NUEVO|limones|kg", "article_id": "",
            "nombre_canonico": "Limón", "producto": "Limones malla 1 kg",
            "tienda_referencia": "Alcampo", "precio_normalizado": 2.79,
            "unidad_normalizada": "kg", "precio_comercial": 2.79, "unidad_formato": "kg",
            "origen": "REFERENCIA_EXTERNA", "autoridad": "REFERENCIA_NO_REAL",
        },
        {
            "article_key": "ART-AZUCAR", "article_id": "ART-AZUCAR",
            "nombre_canonico": "Azúcar", "producto": "Azúcar blanco 1 kg",
            "tienda_referencia": "Alcampo", "precio_normalizado": 0.89,
            "unidad_normalizada": "kg", "precio_comercial": 0.89, "unidad_formato": "kg",
            "origen": "REFERENCIA_EXTERNA", "autoridad": "REFERENCIA_NO_REAL",
        },
    ]
    repository = RepositorioBibliotecaRecetas601(tmp_path)
    write = RecetaDocumentacionWriteService(tmp_path, repository=repository)
    result = write.proposal(recipe_id="REC601-000777", proposed=proposals)
    batch_service = RecetaDocumentacionBatchService(
        tmp_path, repository=repository, recipe_service=write,
    )
    started = batch_service.start_external(
        proposal_results=[{
            "recipe_id": "REC601-000777", "proposal_result": result,
            "proposal_source": {"fuente": "ARCHIVO_EXTERNO"},
        }],
        validation={"filas": [{"recipe_id": "REC601-000777", "estado": "UTIL"}]},
        price_references=references,
    )
    batch_id = started["batch_id"]
    grouped = started["resultados"][0]["datos_operativos_agrupables"]
    batch_service.select(
        batch_id, {}, grouped_selections={"REC601-000777": grouped},
    )
    preview = batch_service.preview(batch_id, context=_context())

    projection = preview["preview"]["items"][0]["proyeccion_provisional"]
    costing = projection["escandallo"]
    assert costing["estado_coste"] == "PARCIAL"
    assert costing["coste_total_parcial"] == 11.63
    assert costing["precios_referencia"] == 3
    assert costing["completitud_coste_porcentaje"] == 75.0
    assert len(costing["lineas"]) == 4
    assert preview["preview"]["cambios_a_aplicar"] == len(grouped)
    assert "ingredientes_estructurados" not in preview["preview"]["items"][0]["cambios"]

    restarted = RecetaDocumentacionBatchService(
        tmp_path, repository=repository, recipe_service=write,
    ).get(batch_id)
    restarted_costing = restarted["preview"]["items"][0]["proyeccion_provisional"]["escandallo"]
    assert restarted_costing["estado_coste"] == "PARCIAL"
    assert restarted_costing["coste_total_parcial"] == 11.63
    assert restarted_costing["precios_referencia"] == 3
    assert restarted["preview"]["items"][0]["proyeccion_provisional"]["datos_reales_modificados"] is False
    assert recipe_path.read_bytes() == before_recipe
    assert articles_path.read_bytes() == before_articles


def test_agua_jamaica_escandallo_y_ficha_provisionales_rechazo_y_confirmacion(tmp_path: Path) -> None:
    recipe_path = _fixture(tmp_path)
    service = RecetaDocumentacionWriteService(tmp_path)
    result = service.proposal(recipe_id="REC601-000777", proposed=_proposals())
    proposals = result["datos_propuestos_ia"]
    read = BibliotecaCulinariaReadService(tmp_path)
    before = recipe_path.read_bytes()

    projection = read.proyectar_provisional(
        "REC601-000777", propuestas=proposals,
        metadatos_propuestas=result["metadatos_propuestas"], completitud=result["completitud"],
    )
    rejected = read.proyectar_provisional("REC601-000777", propuestas={})

    assert projection["datos_reales_modificados"] is False and recipe_path.read_bytes() == before
    assert projection["ficha_tecnica"]["estado"] == "PROVISIONAL"
    assert projection["ficha_tecnica"]["estados_campos_operativos"]["tiempo_descongelacion"]["estado"] == "NO_APLICA"
    assert projection["ficha_tecnica"]["rendimiento"] == 10
    assert projection["escandallo"]["estado_coste"] == "PROVISIONAL"
    assert [line["coste_linea"] for line in projection["escandallo"]["lineas"]] == [5.0, 2.0, 1.0]
    assert projection["escandallo"]["coste_total"] == 8.0
    assert projection["escandallo"]["coste_por_racion"] == 0.8
    assert projection["escandallo"]["lineas_datos_propuestos"] == 3
    assert rejected["escandallo"]["coste_total"] is None

    selected = {
        key: proposals[key] for key in (
            "ingredientes_estructurados", "rendimiento", "unidad_rendimiento",
            "numero_raciones", "cantidad_por_racion", "produccion_maxima", "personal_recomendado",
            "tiempo_descongelacion",
        )
    }
    metadata = {key: result["metadatos_propuestas"][key] for key in selected}
    preview = service.preview(
        recipe_id="REC601-000777", selected=selected, overwrite_fields=[], context=_context(),
        proposal_metadata_by_field=metadata,
    )
    assert preview["datos_reales_modificados"] is False and recipe_path.read_bytes() == before
    confirmed = service.confirm(
        recipe_id="REC601-000777", selected=selected, overwrite_fields=[],
        preview_token=preview["preview_token"], context=_context(),
        proposal_metadata_by_field=metadata,
    )
    repeated = service.confirm(
        recipe_id="REC601-000777", selected=selected, overwrite_fields=[],
        preview_token=preview["preview_token"], context=_context(),
        proposal_metadata_by_field=metadata,
    )
    reread = RepositorioBibliotecaRecetas601(tmp_path).obtener("REC601-000777")
    canonical = read.detalle("REC601-000777")["elaboracion"]
    assert confirmed["lectura_posterior_verificada"] is True
    assert repeated["idempotente"] is True
    assert reread["ingredientes_estructurados"][0]["unidad"] == "kg"
    assert reread.get("tiempo_descongelacion") in (None, "")
    assert reread["estados_campos_operativos"]["tiempo_descongelacion"]["estado"] == "NO_APLICA"
    assert canonical["ficha_tecnica"]["estados_campos_operativos"]["tiempo_descongelacion"]["estado"] == "NO_APLICA"
    assert canonical["escandallo"]["coste_total"] == 8.0
    assert canonical["escandallo"]["estado_coste"] == "DISPONIBLE"


def test_stale_y_round_trip_xlsx_fisico_versionado(tmp_path: Path) -> None:
    _fixture(tmp_path)
    repository = RepositorioBibliotecaRecetas601(tmp_path)
    service = RecetaDocumentacionWriteService(tmp_path, repository=repository)
    exchange = RecipeCompletionExchangeService(tmp_path, repository=repository, recipe_service=service)
    exported = exchange.export(recipe_ids=["REC601-000777"], import_id="IMPORT-AGUA")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    assert dict(workbook["METADATA"].values)["version"] == "0.3"
    assert "PROMPT_IA" in workbook.sheetnames
    assert "Procesa TODAS las recetas" in str(workbook["PROMPT_IA"].cell(2, 3).value or "")
    assert "INSTRUCCIONES" in workbook.sheetnames
    instructions = " ".join(str(cell.value or "") for row in workbook["INSTRUCCIONES"] for cell in row)
    assert "NO_APLICA" in instructions
    sheet = workbook["RECETAS"]
    headers = [cell.value for cell in sheet[1]]
    row = {name: index + 1 for index, name in enumerate(headers)}
    exported_row = {name: sheet.cell(2, column).value for name, column in row.items()}
    assert exported_row["import_id"] == "IMPORT-AGUA"
    assert json.loads(exported_row["contexto_servicio"])[0]["pax"] == 10
    assert json.loads(exported_row["contexto_servicio"])[0]["servicio"] == "Servicio de terraza"
    assert len(json.loads(exported_row["ingredientes_contexto"])) == 3
    sheet.cell(2, row["rendimiento_propuesto"], 10)
    sheet.cell(2, row["unidad_rendimiento_propuesto"], "raciones")
    sheet.cell(2, row["numero_raciones_propuesto"], 10)
    sheet.cell(2, row["cantidad_por_racion_propuesto"], json.dumps({"valor": 250, "unidad": "ml"}))
    sheet.cell(2, row["produccion_maxima_propuesto"], 20)
    sheet.cell(2, row["personal_recomendado_propuesto"], json.dumps({"personas": 1, "rol": "cocinero"}))
    sheet.cell(2, row["tiempo_descongelacion_propuesto"], "NO_APLICA")
    sheet.cell(2, row["ingredientes_estructurados_propuesto"], json.dumps(_proposals()["ingredientes_estructurados"], ensure_ascii=False))
    sheet.cell(2, row["metadatos_propuestas"], json.dumps({
        "rendimiento": {"origen": "CONTEXTO_INTERNO", "confianza": 0.9, "motivo": "Contexto de menú"},
        "tiempo_descongelacion": {
            "origen": "IA_PROPUESTA", "confianza": 0.9,
            "motivo": "La bebida no se congela en el proceso propuesto.",
        },
    }, ensure_ascii=False))
    sheet.cell(2, row["origen_propuesta"], "CHATGPT")
    output = BytesIO(); workbook.save(output)
    imported = exchange.import_package(
        filename="agua-completada.xlsx", content_base64=base64.b64encode(output.getvalue()).decode(),
        expected_recipe_ids=["REC601-000777"], import_id="IMPORT-AGUA",
    )
    item = imported["batch"]["resultados"][0]
    assert imported["contrato"] == {
        "format": "HOSTAI_RECIPE_COMPLETION_PACKAGE", "version": "0.3", "version_importada": "0.3",
    }
    assert item["datos_operativos_agrupables"]["rendimiento"] == 10
    assert item["metadatos_propuestas"]["rendimiento"]["confianza"] == 0.9
    assert item["metadatos_propuestas"]["rendimiento"]["origen"] == "CONTEXTO_INTERNO"
    assert item["datos_requieren_revision_individual"]["tiempo_descongelacion"] == {"estado": "NO_APLICA"}
    assert item["metadatos_propuestas"]["tiempo_descongelacion"]["estado_campo"] == "NO_APLICA"
    assert imported["datos_reales_modificados"] is False

    proposal = service.proposal(recipe_id="REC601-000777", proposed={"rendimiento": 10})
    preview = service.preview(
        recipe_id="REC601-000777", selected=proposal["datos_propuestos_ia"],
        overwrite_fields=[], context=_context(), proposal_metadata_by_field=proposal["metadatos_propuestas"],
    )
    recipe_path = tmp_path / "DATOS" / "db" / "biblioteca_recetas_601.json"
    concurrent = json.loads(recipe_path.read_text(encoding="utf-8"))
    concurrent["recetas"][0]["actualizado_en"] = "2026-09-02T12:00:00"
    recipe_path.write_text(json.dumps(concurrent, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(RecetaDocumentacionError, match="vista previa"):
        service.confirm(
            recipe_id="REC601-000777", selected=proposal["datos_propuestos_ia"],
            overwrite_fields=[], preview_token=preview["preview_token"], context=_context(),
            proposal_metadata_by_field=proposal["metadatos_propuestas"],
        )


def test_confirmar_campo_seguro_no_exige_rendimiento_ajeno_en_borrador(tmp_path: Path) -> None:
    _fixture(tmp_path)
    service = RecetaDocumentacionWriteService(tmp_path)
    selected = {"observaciones": "Texto culinario revisado."}
    preview = service.preview(
        recipe_id="REC601-000777", selected=selected, overwrite_fields=[], context=_context(),
    )

    confirmed = service.confirm(
        recipe_id="REC601-000777", selected=selected, overwrite_fields=[],
        preview_token=preview["preview_token"], context=_context(),
    )

    assert confirmed["ok"] is True
    assert confirmed["lectura_posterior_verificada"] is True
    assert confirmed["receta"]["numero_raciones"] == 0
