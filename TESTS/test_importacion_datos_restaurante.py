from __future__ import annotations

import base64
import io
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook

from SERVICIOS.analizador_importacion_restaurante import ParsedTable, RestaurantDataImportAnalyzer
from SERVICIOS.importador_inteligente_biblioteca import ImportDocumentService
from SERVICIOS.borrador_importacion_biblioteca import normalize_text
from SERVICIOS.repositorio_productos_maestro_601 import RepositorioProductosMaestro601
from SERVICIOS.biblioteca_recetas_601 import RepositorioBibliotecaRecetas601
from SERVICIOS.motor_escritura_segura_i1342 import MotorEscrituraSeguraI1342
from SERVICIOS.biblioteca_culinaria_read_service import BibliotecaCulinariaReadService
from SERVICIOS.catalog_crud_write_service import CatalogCrudWriteService
from SERVICIOS.host_ai_authorized_execution_context import AuthorizedExecutionContext
from SERVICIOS.produccion_inteligente_workflow import IntelligentProductionWorkflow


def _xlsx() -> bytes:
    workbook = Workbook()
    articles = workbook.active
    articles.title = "Artículos"
    articles.append(["Producto", "Proveedor", "Formato", "Precio compra"])
    articles.append(["Ceviche de corvina", "Pescados", "kg", 12.5])
    recipes = workbook.create_sheet("Recetas")
    recipes.append(["Receta", "Ingrediente", "Cantidad", "Unidad", "Procedimiento", "Rendimiento"])
    recipes.append(["Ceviche de corvina", "Corvina", 1, "kg", "Cortar y mezclar", 4])
    recipes.append(["Salsa cítrica", "Ceviche de corvina", 0.2, "kg", "Mezclar", 4])
    suppliers = workbook.create_sheet("Proveedores")
    suppliers.append(["Proveedor"])
    suppliers.append(["Pescados"])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _file(name: str, content: bytes) -> dict:
    return {"nombre": name, "contenido_base64": base64.b64encode(content).decode("ascii")}


def _complete_xlsx() -> bytes:
    workbook = Workbook()
    suppliers = workbook.active
    suppliers.title = "Proveedores"
    suppliers.append(["Proveedor"])
    suppliers.append(["Makro"])
    suppliers.append(["Pescados Sur"])
    articles = workbook.create_sheet("Artículos")
    articles.append(["Producto", "Proveedor", "Formato"])
    articles.append(["Nata 35%", "Makro", "L"])
    articles.append(["Sal fina", "Makro", "kg"])
    articles.append(["Corvina", "Pescados Sur", "kg"])
    recipes = workbook.create_sheet("Recetas")
    recipes.append(["Receta", "Ingrediente", "Cantidad", "Unidad", "Procedimiento", "Rendimiento"])
    recipes.append(["Fondo de pescado", "Corvina", 1, "kg", "Cocer", 4])
    recipes.append(["Ceviche de corvina", "Corvina", 1, "kg", "Cortar", 4])
    recipes.append(["Ceviche de corvina", "Sal fina", 0.01, "kg", "Mezclar", 4])
    recipes.append(["Salsa ceviche", "Ceviche de corvina", 0.2, "kg", "Triturar", 4])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _realistic_structures_xlsx() -> bytes:
    workbook = Workbook()
    articles = workbook.active
    articles.title = "Listado de Artículos"
    articles.merge_cells("A1:F1")
    articles["A1"] = "CATÁLOGO GENERAL"
    articles.append([])
    articles.append(["FICHA TÉCNICA PLATO"])
    articles.append(["Código", "Artículo", "Familia", "Proveedor", "Formato", "Precio"])
    articles.append(["A-1", "Nata 35%", "Lácteos", "Makro", "L", 5.99])
    articles.append(["A-2", "Sal fina", "Secos", "Makro", "kg", 0.11])
    articles.append([None, "TOTAL", None, None, None, 6.10])
    blocks = workbook.create_sheet("Escandallos")
    for name in ("Receta A", "Receta B", "Receta C"):
        blocks.append([name])
        blocks.append(["Ingrediente", "Cantidad", "Unidad"])
        blocks.append(["Corvina", 1, "kg"])
        blocks.append(["Sal fina", 0.1, "kg"])
        blocks.append(["SUBTOTAL", 5.45, None])
        blocks.append([])
        blocks.append([])
    vertical = workbook.create_sheet("FICHA TECNICA VERTICAL")
    vertical.merge_cells("A1:C1")
    vertical["A1"] = "FICHA TÉCNICA PLATO"
    vertical.append(["Crema de verduras"])
    vertical.append([])
    for index in range(10):
        vertical.append([f"Ingrediente {index + 1}", index + 1, "g"])
    irrelevant = workbook.create_sheet("Auxiliar")
    irrelevant.append(["Notas internas"])
    irrelevant.append([0.11, 0.12, 5.99])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _library_match_xlsx(*, incompatible_ceviche: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recetas"
    sheet.append(["Receta", "Ingrediente", "Cantidad", "Unidad", "Procedimiento", "Rendimiento"])
    fixtures = [
        ("Ceviche de corvina", "Chocolate" if incompatible_ceviche else "Corvina", "Cortar", 4),
        ("Pico de gallo", "Tomate", "Picar", 4),
        ("Salsa naranja", "Naranja", "Reducir", 4),
        ("Receta realmente nueva", "Patata", "Cocer", 4),
    ]
    for name, ingredient, procedure, servings in fixtures:
        sheet.append([name, ingredient, 1, "kg", procedure, servings])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _incomplete_recipe_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recetas históricas"
    sheet.append(["Receta", "Ingrediente", "Cantidad", "Unidad"])
    for index in range(8):
        sheet.append(["Ensaladilla", f"Ingrediente {index + 1}", index + 1, "g"])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _menu_multisheet_xlsx() -> bytes:
    workbook = Workbook()
    raw = workbook.active
    raw.title = "M.P RECETAS"
    raw.append(["Receta", "Ingrediente", "Cantidad", "Unidad", "Procedimiento"])
    raw.append(["Patatas bravas", "Patata", 1, "kg", "Freír"])
    menu = workbook.create_sheet("MENU BODA 31-1")
    menu.append(["Receta", "Ingrediente", "Cantidad", "Unidad", "Procedimiento"])
    menu.append(["Patatas Bravas.", "Patata", 1, "kg", "Freír"])
    menu.append(["Salsa romesco", "Tomate", 1, "kg", "Triturar"])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _articles_with_unregistered_supplier_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Artículos"
    sheet.append(["Producto", "Proveedor", "Formato"])
    sheet.append(["Patata", "Proveedor no identificado", "kg"])
    sheet.append(["Tomate", "Proveedor no identificado", "kg"])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _ap_mp_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Artículos"
    sheet.append(["Producto", "Formato"])
    sheet.append(["A.P MINI GOFRE CON SOBRASADA Y MIEL.", "ud"])
    sheet.append(["M.P HARINA", "kg"])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _three_recipe_blocks_xlsx() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Escandallos"
    for name, count in (("RECETA A", 5), ("RECETA B", 7), ("RECETA C", 4)):
        sheet.append([name])
        sheet.append(["Ingrediente", "Cantidad", "Unidad"])
        for index in range(count):
            sheet.append([f"{name} ingrediente {index + 1}", index + 1, "g"])
        sheet.append([])
        sheet.append([])
    target = io.BytesIO()
    workbook.save(target)
    return target.getvalue()


def _seed_library_recipes(base_dir: Path) -> dict[str, dict]:
    repository = RepositorioBibliotecaRecetas601(base_dir)
    created: dict[str, dict] = {}
    for name, ingredient, procedure in (
        ("Ceviche de corvina", "Corvina", "Procedimiento humano del ceviche"),
        ("Pico de gallo", "Tomate", "Procedimiento humano del pico"),
        ("Salsa naranja", "Naranja", "Procedimiento humano de la salsa"),
    ):
        result = repository.crear_ficha_tecnica({
            "nombre": name, "ingredientes": [ingredient], "cantidades": ["1 kg"],
            "ingredientes_estructurados": [{"nombre_original": ingredient, "cantidad": 1, "unidad": "kg"}],
            "elaboracion": procedure, "numero_raciones": 4, "tipo": "PRINCIPAL",
            "descripcion": f"Descripción humana de {name}",
        })
        assert result["ok"] is True
        created[name] = result["receta"]
    return created


def _read(path: Path):
    if not path.exists():
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def test_xlsx_reads_every_sheet_and_crosses_entities_without_write():
    result = RestaurantDataImportAnalyzer().analyze({"archivos": [_file("restaurante.xlsx", _xlsx())]})
    assert result["resumen"]["hojas_analizadas"] == 3
    assert result["resumen"]["posibles_recetas"] == 2
    assert result["resumen"]["posibles_articulos"] == 1
    assert result["resumen"]["proveedores"] == 1
    assert result["resumen"]["posibles_subelaboraciones"] == 1
    assert result["dudas"]["relaciones"][0]["tipo"] == "REQUIERE_REVISION"
    assert result["datos_operativos_modificados"] is False
    assert result["coste_ia"]["usada"] is False


def test_csv_and_xlsx_are_one_session_with_deterministic_mapping():
    csv_content = "Producto,Proveedor,Precio\nNata,SARDA,2.50\n".encode()
    result = RestaurantDataImportAnalyzer().analyze({
        "archivos": [_file("articulos.csv", csv_content), _file("recetas.xlsx", _xlsx())]
    })
    assert result["resumen"]["archivos_analizados"] == 2
    assert result["resumen"]["hojas_analizadas"] == 4
    assert any(item["destino"] == "precio_ambiguo" for item in result["mapping"])
    assert result["dudas"]["precio"]


def test_tsv_json_and_pasted_data_are_supported_without_ai():
    analyzer = RestaurantDataImportAnalyzer()
    tsv = analyzer.analyze({"archivos": [_file("proveedores.tsv", b"Proveedor\nSARDA\n")]})
    json_result = analyzer.analyze({"archivos": [_file("articulos.json", b'[{"Producto":"Azucar"}]')]})
    pasted = analyzer.analyze({"texto_pegado": "Producto\tProveedor\nAzucar\tSARDA"})
    assert tsv["resumen"]["proveedores"] == 1
    assert json_result["resumen"]["posibles_articulos"] == 1
    assert pasted["resumen"]["archivos_analizados"] == 1
    assert all(not item["coste_ia"]["usada"] for item in (tsv, json_result, pasted))


def test_preview_and_confirm_persist_catalog_then_recipes_without_touching_stock(tmp_path: Path):
    stock_paths = [
        "DATOS/db/stock.json", "DATOS/db/stock_lotes.json", "DATOS/db/stock_movimientos.json"
    ]
    for rel in stock_paths:
        path = tmp_path / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps({"fixture": rel}), encoding="utf-8")
    before_stock = {rel: _read(tmp_path / rel) for rel in stock_paths}
    service = ImportDocumentService(tmp_path)
    response = service.import_document({"archivos": [_file("restaurante.xlsx", _complete_xlsx())]})
    assert response["ok"] is True
    session = response["importacion"]
    preview = session["preview_global"]
    assert {
        key: preview["contadores"][key]
        for key in (
            "proveedores_nuevos", "proveedores_reutilizados",
            "articulos_nuevos", "articulos_reutilizados",
            "articulos_requieren_revision", "recetas_elaboraciones",
            "relaciones", "pendientes",
        )
    } == {
            "proveedores_nuevos": 2, "proveedores_reutilizados": 0,
            "articulos_nuevos": 3, "articulos_reutilizados": 0,
            "articulos_requieren_revision": 0,
            "recetas_elaboraciones": 3, "relaciones": 3, "pendientes": 0,
    }
    assert RepositorioProductosMaestro601(tmp_path).listar_proveedores() == []
    assert RepositorioProductosMaestro601(tmp_path).listar_productos() == []
    assert before_stock == {rel: _read(tmp_path / rel) for rel in stock_paths}

    result = service.confirm(session["documento"]["id"], {
        "draft_version": 1, "usuario": "test", "confirmacion": "CONFIRMAR",
    })
    assert result["ok"] is True
    repository = RepositorioProductosMaestro601(tmp_path)
    assert {item["nombre"] for item in repository.listar_proveedores()} == {"Makro", "Pescados Sur"}
    articles = repository.listar_productos()
    assert {item["nombre"] for item in articles} == {"Nata 35%", "Sal fina", "Corvina"}
    article_ids = {item["nombre"]: item["codigo"] for item in articles}
    recipes = _read(tmp_path / "DATOS/db/biblioteca_recetas_601.json")["recetas"]
    assert len(recipes) == 3
    ceviche = next(item for item in recipes if item["nombre"] == "Ceviche de corvina")
    assert {item["article_id"] for item in ceviche["ingredientes_estructurados"]} == {
        article_ids["Corvina"], article_ids["Sal fina"],
    }
    salsa = next(item for item in recipes if item["nombre"] == "Salsa ceviche")
    assert salsa["ingredientes_estructurados"][0]["elaboracion_id"] == ceviche["id"]
    escandallos = _read(tmp_path / "DATOS/db/biblioteca_escandallos_601.json")["escandallos"]
    linked_ids = {line.get("producto_codigo") for esc in escandallos for line in esc["lineas"]}
    assert article_ids["Corvina"] in linked_ids
    assert article_ids["Sal fina"] in linked_ids
    assert before_stock == {rel: _read(tmp_path / rel) for rel in stock_paths}

    second = service.confirm(session["documento"]["id"], {
        "draft_version": 1, "usuario": "test", "confirmacion": "CONFIRMAR",
    })
    assert second["ok"] is False
    assert len(repository.listar_proveedores()) == 2
    assert len(repository.listar_productos()) == 3


def test_existing_catalog_is_reused_and_ceviche_collision_is_excluded_from_safe_confirm(tmp_path: Path):
    repository = RepositorioProductosMaestro601(tmp_path)
    makro = repository.crear_proveedor({"nombre": "Makro"})
    salt = repository.crear_producto({"nombre": "Sal fina", "unidad_base": "kg"})
    service = ImportDocumentService(tmp_path)
    response = service.import_document({"archivos": [_file("restaurante.xlsx", _complete_xlsx())]})
    preview = response["importacion"]["preview_global"]
    assert preview["proveedores"]["reutilizar"][0]["proveedor_id"] == makro["codigo"]
    assert preview["articulos"]["reutilizar"][0]["article_id"] == salt["codigo"]

    collision = service.import_document({"archivos": [_file("colision.xlsx", _xlsx())]})
    assert collision["importacion"]["preview_global"]["contadores"]["pendientes"] > 0
    confirmed = service.confirm(collision["importacion"]["documento"]["id"], {
        "draft_version": 1, "usuario": "test", "confirmacion": "CONFIRMAR",
    })
    assert confirmed["ok"] is True
    assert all(item["nombre"] != "Ceviche de corvina" for item in repository.listar_productos())
    assert collision["importacion"]["preview_global"]["contadores"]["articulos_requieren_revision"] == 1


def test_menu_sheet_is_container_and_keeps_component_recipes_separate(tmp_path: Path):
    service = ImportDocumentService(tmp_path)
    response = service.import_document({"archivos": [_file("menus.xlsx", _menu_multisheet_xlsx())]})

    assert response["ok"] is True
    session = response["importacion"]
    analysis = session["analisis_restaurante"]
    assert any(item["nombre"] == "MENU BODA 31-1" and item["tipo"] == "MENU" for item in analysis["menus"])
    assert {normalize_text(item["nombre"]) for item in analysis["recetas"]} >= {"patatas bravas", "salsa romesco"}
    menu = session["preview_global"]["menus"]["crear"][0]
    assert menu["nombre"] == "MENU BODA 31-1"
    assert {line["nombre"] for line in menu["lineas"]} == {"Patatas Bravas.", "Salsa romesco"}
    assert all(line["tipo_referencia"] == "RECETA" and line["futura"] for line in menu["lineas"])
    assert session["preview_global"]["contadores"]["menus_recibidos"] == 1
    assert response["datos_operativos_modificados"] is False


def test_session_option_excludes_legacy_ap_but_keeps_mp_without_derived_pending(tmp_path: Path):
    service = ImportDocumentService(tmp_path)
    response = service.import_document({
        "archivos": [_file("catalogo.xlsx", _ap_mp_xlsx())],
        "excluir_ap_antiguos": True,
    })

    session = response["importacion"]
    catalog = session["borrador"]["catalogo"]
    assert [item["nombre"] for item in catalog["articulos"]] == ["M.P HARINA"]
    assert catalog["relaciones"] == []
    assert session["analisis_restaurante"]["exclusiones_sesion"] == [{
        "nombre": "A.P MINI GOFRE CON SOBRASADA Y MIEL.",
        "tipo_origen": "ARTICULOS", "accion": "IGNORAR",
        "motivo": "Exclusión A.P confirmada para esta sesión.",
        "origen": "catalogo.xlsx:Artículos:REGION-1",
    }]
    assert session["analisis_restaurante"]["opciones_sesion"] == {"excluir_ap_antiguos": True}
    assert response["datos_operativos_modificados"] is False


def test_supplier_text_without_detected_supplier_does_not_create_artificial_relations(tmp_path: Path):
    session = ImportDocumentService(tmp_path).import_document({
        "archivos": [_file("sin-proveedores.xlsx", _articles_with_unregistered_supplier_xlsx())]
    })["importacion"]

    assert session["analisis_restaurante"]["proveedores"] == []
    assert session["borrador"]["catalogo"]["relaciones"] == []
    counts = session["preview_global"]["contadores"]
    assert counts["relaciones"] == 0
    assert counts["relaciones_pendientes"] == 0
    assert counts["pendientes_desglose"]["relacion"] == 0


def test_visible_recipe_proposals_follow_create_reuse_and_pending_actions(tmp_path: Path):
    _seed_library_recipes(tmp_path)
    session = ImportDocumentService(tmp_path).import_document({
        "archivos": [_file("recetas.xlsx", _library_match_xlsx(incompatible_ceviche=True))]
    })["importacion"]

    actions = {item["title"]: item["proposed_action"] for item in session["borrador"]["recipes"]}
    visible = {
        ((item["datos_propuestos"]["entidad"])["name"]): item["titulo"]
        for item in session["propuestas"]
        if item["datos_propuestos"]["entidad"]["kind"] == "RECETA"
    }
    assert actions["Pico de gallo"] == "REUTILIZAR_EXISTENTE"
    assert visible["Pico de gallo"].startswith("Reutilizar receta existente:")
    assert actions["Receta realmente nueva"] == "CREAR_RECETA"
    assert visible["Receta realmente nueva"].startswith("Crear receta nueva:")
    assert actions["Ceviche de corvina"] == "REQUIERE_REVISION"
    assert visible["Ceviche de corvina"].startswith("Pendiente de decisión de identidad:")


def test_existing_library_recipes_are_checked_reused_and_not_overwritten(tmp_path: Path):
    existing = _seed_library_recipes(tmp_path)
    service = ImportDocumentService(tmp_path)
    response = service.import_document({
        "archivos": [_file("recetas.xlsx", _library_match_xlsx())]
    })
    assert response["ok"] is True
    session = response["importacion"]
    drafts = {item["title"]: item for item in session["borrador"]["recipes"]}
    assert {drafts[name]["proposed_action"] for name in existing} == {"REUTILIZAR_EXISTENTE"}
    assert drafts["Receta realmente nueva"]["proposed_action"] == "CREAR_RECETA"
    assert len(session["preview_global"]["elaboraciones"]["reutilizar_existente"]) == 3
    assert len(session["preview_global"]["elaboraciones"]["crear_nueva"]) == 1
    assert len(RepositorioBibliotecaRecetas601(tmp_path).listar()) == 3  # ANALYZE/PREVIEW: 0 WRITE

    confirmed = service.confirm(session["documento"]["id"], {
        "draft_version": 1, "usuario": "test", "confirmacion": "CONFIRMAR",
    })
    assert confirmed["ok"] is True
    stored = RepositorioBibliotecaRecetas601(tmp_path).listar()
    assert len(stored) == 4
    assert sum(action["tipo"] == "REUTILIZAR_RECETA" for action in confirmed["resultado"]["acciones"]) == 3
    assert sum(action["tipo"] == "CREAR_RECETA" for action in confirmed["resultado"]["acciones"]) == 1
    after_ceviche = next(item for item in stored if item["nombre"] == "Ceviche de corvina")
    assert after_ceviche["id"] == existing["Ceviche de corvina"]["id"]
    assert after_ceviche["elaboracion"] == "Procedimiento humano del ceviche"
    assert after_ceviche["descripcion"] == "Descripción humana de Ceviche de corvina"

    reimported = ImportDocumentService(tmp_path).import_document({
        "archivos": [_file("recetas.xlsx", _library_match_xlsx())]
    })["importacion"]
    assert {item["proposed_action"] for item in reimported["borrador"]["recipes"]} == {
        "REUTILIZAR_EXISTENTE"
    }


def test_same_recipe_name_without_compatible_structure_requires_review(tmp_path: Path):
    _seed_library_recipes(tmp_path)
    session = ImportDocumentService(tmp_path).import_document({
        "archivos": [_file("recetas.xlsx", _library_match_xlsx(incompatible_ceviche=True))]
    })["importacion"]
    ceviche = next(item for item in session["borrador"]["recipes"] if item["title"] == "Ceviche de corvina")
    assert ceviche["proposed_action"] == "REQUIERE_REVISION"
    assert ceviche["duplicate_candidates"][0]["estructura_compatible"] is False
    assert session["confirmacion_disponible"] is False


def test_incomplete_historical_recipe_is_importable_without_inventing_fields(tmp_path: Path):
    service = ImportDocumentService(tmp_path)
    session = service.import_document({
        "archivos": [_file("historico.xlsx", _incomplete_recipe_xlsx())]
    })["importacion"]
    recipe = session["borrador"]["recipes"][0]
    assert recipe["procedure"] == []
    assert recipe["servings"] is None
    assert recipe["yield_value"] is None
    warning_codes = {item["code"] for item in session["borrador"]["validation"]["warnings"]}
    assert {"PROCEDIMIENTO_PENDIENTE", "RENDIMIENTO_PENDIENTE"}.issubset(warning_codes)
    assert not {
        "PROCEDIMIENTO_PENDIENTE", "RENDIMIENTO_PENDIENTE"
    }.intersection(item["code"] for item in session["borrador"]["validation"]["blocking_errors"])
    assert RepositorioBibliotecaRecetas601(tmp_path).listar() == []  # ANALYZE/PREVIEW: 0 WRITE

    confirmed = service.confirm(session["documento"]["id"], {
        "draft_version": 1, "usuario": "test", "confirmacion": "CONFIRMAR",
    })
    assert confirmed["ok"] is True
    stored = RepositorioBibliotecaRecetas601(tmp_path).listar()[0]
    assert stored["estado"] == "PENDIENTE_DE_COMPLETAR"
    assert stored["elaboracion"] == ""
    assert stored["numero_raciones"] == 0
    assert stored["campos_pendientes_importacion"] == ["procedimiento", "rendimiento"]


def test_structural_overlap_deduplicates_by_physical_origin_but_real_repeat_remains():
    analyzer = RestaurantDataImportAnalyzer()
    mapping = [
        {"columna": "Receta", "destino": "receta"},
        {"columna": "Ingrediente", "destino": "ingrediente"},
        {"columna": "Cantidad", "destino": "cantidad"},
        {"columna": "Unidad", "destino": "unidad"},
    ]
    shared = {"Receta": "RECETA A", "Ingrediente": "Tomate", "Cantidad": 100, "Unidad": "g", "_fila": 10}
    tables = [
        ParsedTable("fixture.xlsx", "Hoja", [shared], mapping, "RECETAS", region_id="REGION-1"),
        ParsedTable("fixture.xlsx", "Hoja", [shared], mapping, "RECETAS", region_id="REGION-2"),
        ParsedTable("fixture.xlsx", "Hoja", [
            {"Receta": "RECETA B", "Ingrediente": "Tomate", "Cantidad": 100, "Unidad": "g", "_fila": 20},
            {"Receta": "RECETA B", "Ingrediente": "Tomate", "Cantidad": 50, "Unidad": "g", "_fila": 21},
        ], mapping, "RECETAS", region_id="REGION-3"),
    ]
    recipes = analyzer._recipes(tables)
    recipe_a = next(item for item in recipes if item["nombre"] == "RECETA A")
    recipe_b = next(item for item in recipes if item["nombre"] == "RECETA B")
    assert len(recipe_a["ingredientes_estructurados"]) == 1
    assert len(recipe_b["ingredientes_estructurados"]) == 2
    assert {item["trazabilidad"]["fila"] for item in recipe_b["ingredientes_estructurados"]} == {20, 21}


def test_generic_tapa_heading_is_not_a_recipe_but_specific_name_is_preserved():
    analyzer = RestaurantDataImportAnalyzer()
    mapping = [{"columna": "ARTICULO", "destino": "nombre"}]
    tables = [
        ParsedTable(
            "fixture.xlsx", "M.P Menú finde", [{"ARTICULO": "Tomate", "_fila": 10}],
            mapping, "RECETAS", region_id="REGION-1", context_title="TAPA",
        ),
        ParsedTable(
            "fixture.xlsx", "M.P Menú finde", [{"ARTICULO": "Anchoa", "_fila": 20}],
            mapping, "RECETAS", region_id="REGION-2", context_title="Tapa de anchoa",
        ),
    ]
    recipes = analyzer._recipes(tables)
    assert [item["nombre"] for item in recipes] == ["Tapa de anchoa"]


def test_recipe_blocks_end_before_the_next_recipe():
    result = RestaurantDataImportAnalyzer().analyze({
        "archivos": [_file("bloques.xlsx", _three_recipe_blocks_xlsx())]
    })
    counts = {item["nombre"]: len(item["ingredientes_estructurados"]) for item in result["recetas"]}
    assert counts == {"RECETA A": 5, "RECETA B": 7, "RECETA C": 4}


def test_new_recipe_has_one_canonical_creation_proposal(tmp_path: Path):
    session = ImportDocumentService(tmp_path).import_document({
        "archivos": [_file("historico.xlsx", _incomplete_recipe_xlsx())]
    })["importacion"]
    recipe_proposals = [
        item for item in session["propuestas"]
        if item["tipo"] in {"CREAR_RECETA", "CREAR_ELABORACION"}
    ]
    assert len(recipe_proposals) == 1
    assert recipe_proposals[0]["tipo"] == "CREAR_RECETA"
    assert recipe_proposals[0]["titulo"] == "Crear receta nueva: Ensaladilla"


def test_unknown_imported_yield_is_null_publicly_pending_and_not_used_in_calculations(tmp_path: Path):
    repository = RepositorioBibliotecaRecetas601(tmp_path)
    created = repository.crear_incompleta_desde_importacion({
        "nombre": "Ensaladilla histórica", "ingredientes": ["Patata"],
        "cantidades": ["1 kg"], "ingredientes_estructurados": [{
            "nombre_original": "Patata", "cantidad": 1, "unidad": "kg",
        }], "elaboracion": "", "numero_raciones": None,
    })["receta"]
    assert created["numero_raciones"] == 0
    assert "rendimiento" in created["campos_pendientes_importacion"]

    detail = BibliotecaCulinariaReadService(tmp_path).detalle(created["id"])["elaboracion"]
    assert detail["rendimiento"] is None
    assert detail["raciones"] is None
    assert detail["receta"]["rendimiento"] is None
    assert detail["receta"]["raciones"] is None
    assert detail["coste_por_racion"] is None
    assert detail["motivo_coste_no_disponible"] == "Pendiente de rendimiento."
    assert "Rendimiento" in detail["pendientes"]
    assert "Rendimiento" in detail["ficha_tecnica"]["campos_pendientes"]

    needs, unknown = IntelligentProductionWorkflow._needs([{
        "id": created["id"], "nombre": created["nombre"], "rendimiento": detail["rendimiento"],
        "ingredientes": [{"nombre_original": "Patata", "cantidad": 1, "unidad": "kg"}],
    }], 20)
    assert needs == []
    assert unknown == ["rendimiento o pax de Ensaladilla histórica"]


def test_confirmed_yield_removes_import_pending_and_becomes_public(tmp_path: Path):
    repository = RepositorioBibliotecaRecetas601(tmp_path)
    created = repository.crear_incompleta_desde_importacion({
        "nombre": "Ensaladilla histórica", "ingredientes": ["Patata"],
        "cantidades": ["1 kg"], "elaboracion": "", "numero_raciones": None,
    })["receta"]
    service = CatalogCrudWriteService(tmp_path)
    preview_context = AuthorizedExecutionContext(
        request_id="REQ-YIELD-1", user_id="USR-1", tenant_id="TEN-1",
        roles=("chef",), scopes=frozenset({"recetas:preview"}),
    )
    write_context = AuthorizedExecutionContext(
        request_id="REQ-YIELD-1", user_id="USR-1", tenant_id="TEN-1",
        roles=("chef",), scopes=frozenset({"recetas:write"}),
    )
    preview = service.preview(
        domain="RECETA", operation="MODIFICAR", entity_id=created["id"], session_id="S1",
        context=preview_context, payload={"numero_raciones": 20},
    )
    assert preview["datos_reales_modificados"] is False
    assert repository.obtener(created["id"])["numero_raciones"] == 0
    confirmed = service.confirm(
        preview_token=preview["preview_token"], session_id="S1", context=write_context,
    )
    assert confirmed["registro"]["numero_raciones"] == 20
    assert "rendimiento" not in confirmed["registro"]["campos_pendientes_importacion"]
    detail = BibliotecaCulinariaReadService(tmp_path).detalle(created["id"])["elaboracion"]
    assert detail["rendimiento"] == 20
    assert detail["raciones"] == 20
    assert "Rendimiento" not in detail["pendientes"]


def test_controlled_transaction_failure_leaves_no_partial_catalog(tmp_path: Path, monkeypatch):
    service = ImportDocumentService(tmp_path)
    response = service.import_document({"archivos": [_file("restaurante.xlsx", _complete_xlsx())]})
    original = MotorEscrituraSeguraI1342.ejecutar
    calls = 0

    def fail_domain_commit(self, cambios, *args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 1:
            return SimpleNamespace(estado="ROLLBACK", error="fallo controlado", transaccion_id="TEST")
        return original(self, cambios, *args, **kwargs)

    monkeypatch.setattr(MotorEscrituraSeguraI1342, "ejecutar", fail_domain_commit)
    result = service.confirm(response["importacion"]["documento"]["id"], {
        "draft_version": 1, "usuario": "test", "confirmacion": "CONFIRMAR",
    })
    assert result["ok"] is False
    repository = RepositorioProductosMaestro601(tmp_path)
    assert repository.listar_proveedores() == []
    assert repository.listar_productos() == []
    recipes = _read(tmp_path / "DATOS/db/biblioteca_recetas_601.json")
    assert recipes is None or recipes.get("recetas") == []


def test_realistic_excel_detects_shifted_headers_regions_and_filters_structural_values():
    result = RestaurantDataImportAnalyzer().analyze({
        "archivos": [_file("realista.xlsx", _realistic_structures_xlsx())]
    })
    names = {item["nombre"] for item in result["articulos"]}
    assert names == {"Nata 35%", "Sal fina"}
    assert not names.intersection({"0.11", "0.12", "5.99", "ARTÍCULO", "FAMILIA", "TOTAL", "FICHA TÉCNICA PLATO"})
    article_region = next(item for item in result["hojas"] if item["nombre"] == "Listado de Artículos")
    assert article_region["fila_encabezado"] == 4
    assert article_region["merged_cells"] == ["A1:F1"]
    block_regions = [item for item in result["hojas"] if item["nombre"] == "Escandallos"]
    assert len(block_regions) == 3
    assert {item["titulo_contexto"] for item in block_regions} == {"Receta A", "Receta B", "Receta C"}
    recipes = {item["nombre"]: item for item in result["recetas"]}
    assert {"Receta A", "Receta B", "Receta C", "Crema de verduras"}.issubset(recipes)
    assert len(recipes["Crema de verduras"]["ingredientes_estructurados"]) == 10
    assert all(item.get("trazabilidad", {}).get("region") for item in recipes.values())


def test_menu_classification_uses_structure_and_excludes_recipe_sheets_and_templates():
    workbook = Workbook()
    recipe = workbook.active
    recipe.title = "M.P MENU EVENTO"
    recipe.append(["FICHA TÃ‰CNICA PLATO"])
    recipe.append(["ARTÃCULO", "Salsa romesco"])
    recipe.append(["ARTICULO", "KG", "BRUTO", "NETO", "%DESP", "SUCIO"])
    recipe.append(["Tomate", 1, 1, 1, 0, 1])
    menu = workbook.create_sheet("MENU EVENTO")
    menu.append(["MENU EVENTO 2026"])
    menu.append(["APERITIVO", "P.V.P", "NETO", "COSTO", "FOOD COST"])
    menu.append(["Salsa romesco", 1, 1, 1, 1])
    menu.append(["Agua", 1, 1, 1, 1])
    template = workbook.create_sheet("PLANTILLA COSTE MENU")
    template.append(["MENU MODELO"])
    template.append(["APERITIVO", "P.V.P", "NETO", "COSTO", "FOOD COST"])
    template.append(["Linea ejemplo", 1, 1, 1, 1])
    template.append(["Otra linea", 1, 1, 1, 1])
    target = io.BytesIO()
    workbook.save(target)

    result = RestaurantDataImportAnalyzer().analyze({
        "archivos": [_file("estructuras.xlsx", target.getvalue())]
    })
    kinds = {(item["nombre"], item["tipo_propuesto"]) for item in result["hojas"]}
    assert ("M.P MENU EVENTO", "RECETAS") in kinds
    assert ("MENU EVENTO", "MENUS") in kinds
    assert ("PLANTILLA COSTE MENU", "DOCUMENTACION") in kinds
    assert [item["nombre"] for item in result["menus"]] == ["MENU EVENTO"]
