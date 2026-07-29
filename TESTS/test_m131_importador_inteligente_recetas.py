from pathlib import Path
import json,tempfile
import pytest
from openpyxl import Workbook
from SERVICIOS.importador_inteligente_recetas_m131 import ImportadorInteligenteRecetasM131,clasificar_prefijo_articulo

def preparar(p):
    a=p/'articulos.json';a.write_text(json.dumps([
      {'codigo':'M.P-PAT','nombre':'Patata','unidad':'kg','precio':1},
      {'codigo':'M.P-MANT','nombre':'Mantequilla','unidad':'kg','precio':8},
      {'codigo':'A.P-BRIOX','nombre':'A.P Briox de curri','unidad':'ud','precio':3},
    ]),encoding='utf-8');return a

def texto(): return 'Parmentier de patata\nRendimiento: 10 raciones\nIngredientes\n2 kg Patata\n0.2 kg Mantequilla\nElaboración\nCocer y triturar.'

def test_prefijos():
 assert clasificar_prefijo_articulo({'codigo':'M.P-1','nombre':'Patata'})=='MATERIA_PRIMA'
 assert clasificar_prefijo_articulo({'codigo':'A.P-1','nombre':'Briox'})=='APERITIVO'

def test_texto_completo():
 with tempfile.TemporaryDirectory() as d:
  imp=ImportadorInteligenteRecetasM131(preparar(Path(d)));b=imp.desde_texto(texto())
  assert b.nombre=='Parmentier de patata' and b.rendimiento==10 and len(b.ingredientes)==2 and not b.bloqueos

def test_ap_no_se_propone_como_ingrediente():
 with tempfile.TemporaryDirectory() as d:
  imp=ImportadorInteligenteRecetasM131(preparar(Path(d)));b=imp.desde_texto('X\nRendimiento: 1\nIngredientes\n1 ud Briox de curri')
  assert any('Artículo no resuelto' in x for x in b.bloqueos)

def test_sin_cantidad_bloquea():
 with tempfile.TemporaryDirectory() as d:
  imp=ImportadorInteligenteRecetasM131(preparar(Path(d)));b=imp.desde_texto('X\nRendimiento: 1\nIngredientes\nPatata')
  assert b.bloqueos

def test_payload_solo_sin_bloqueos():
 with tempfile.TemporaryDirectory() as d:
  imp=ImportadorInteligenteRecetasM131(preparar(Path(d)));b=imp.desde_texto(texto());p=imp.payload_creacion(b)
  assert p['raciones_base']==10 and p['ingredientes'][0]['articulo_id']=='M.P-PAT'

def test_docx():
 docx = pytest.importorskip('docx', reason='python-docx es opcional y solo necesario para importar .docx')
 Document = docx.Document
 with tempfile.TemporaryDirectory() as d:
  p=Path(d);a=preparar(p);f=p/'r.docx';doc=Document();[doc.add_paragraph(x) for x in texto().splitlines()];doc.save(f)
  b=ImportadorInteligenteRecetasM131(a).desde_archivo(f);assert b.origen_tipo=='WORD' and not b.bloqueos

def test_excel():
 with tempfile.TemporaryDirectory() as d:
  p=Path(d);a=preparar(p);f=p/'r.xlsx';wb=Workbook();ws=wb.active
  for i,x in enumerate(texto().splitlines(),1):ws.cell(i,1,x)
  wb.save(f);b=ImportadorInteligenteRecetasM131(a).desde_archivo(f);assert b.origen_tipo=='EXCEL' and not b.bloqueos

def test_articulo_inexistente_bloquea_m132():
 with tempfile.TemporaryDirectory() as d:
  imp=ImportadorInteligenteRecetasM131(preparar(Path(d)));b=imp.desde_texto('X\nRendimiento: 1\nIngredientes\n1 kg Nuez moscada')
  assert any('Artículo no resuelto' in x for x in b.bloqueos)
