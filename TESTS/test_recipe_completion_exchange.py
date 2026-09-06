import base64
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from fastapi.testclient import TestClient

from API.app import HostAIPlatformAPI
from API.http_server import create_app
from SERVICIOS.host_ai_authorized_execution_context import AuthorizedExecutionContext
from SERVICIOS.confirmacion_importacion_biblioteca import ImportSessionRepository
from SERVICIOS.recipe_completion_exchange_service import (
    RecipeCompletionExchangeError,
    RecipeCompletionExchangeService,
)
from SERVICIOS.recipe_completion_contract import (
    FIELD_CONTRACTS,
    PRICE_REFERENCE_COLUMNS,
    PRODUCTION_READINESS_BASE_FIELDS,
)
from SERVICIOS.receta_documentacion_batch_service import RecetaDocumentacionBatchService
from SERVICIOS.receta_documentacion_write_service import RecetaDocumentacionWriteService
from SERVICIOS.receta_documentacion_write_service import RecetaDocumentacionError


def _context():
    return AuthorizedExecutionContext("REQ", "CHEF", "LOCAL", ("chef",), frozenset({"recetas:write"}))


class FakeRepository:
    def __init__(self):
        self.items = {
            "R1": {
                "id": "R1", "codigo": "REC-1", "nombre": "Salsa uno", "version": 1,
                "ingredientes": ["Tomate"], "cantidades": ["1 kg"], "descripcion": "",
                "elaboracion": "", "alergenos": [],
                "completitud": {"campos_obligatorios_pendientes": ["Descripcion", "Elaboracion paso a paso"]},
            },
            "R2": {
                "id": "R2", "codigo": "REC-2", "nombre": "Salsa dos", "version": 1,
                "ingredientes": ["Cebolla"], "cantidades": ["500 g"], "descripcion": "Ya existe",
                "elaboracion": "", "alergenos": [],
                "completitud": {"campos_obligatorios_pendientes": ["Elaboracion paso a paso"]},
            },
            "R3": {
                "id": "R3", "codigo": "REC-3", "nombre": "Completa", "version": 1,
                "descripcion": "Lista", "elaboracion": "Lista",
                "completitud": {"campos_obligatorios_pendientes": []},
            },
        }
        self.edits = []

    def listar(self):
        return list(self.items.values())

    def obtener(self, recipe_id):
        return deepcopy(self.items.get(recipe_id))

    def editar(self, recipe_id, changes):
        self.edits.append((recipe_id, deepcopy(changes)))
        self.items[recipe_id].update(deepcopy(changes))
        return {"ok": True, "receta": deepcopy(self.items[recipe_id])}


def _service(tmp_path: Path):
    repository = FakeRepository()
    write = RecetaDocumentacionWriteService(tmp_path, repository=repository)
    batch = RecetaDocumentacionBatchService(tmp_path, repository=repository, recipe_service=write)
    return RecipeCompletionExchangeService(
        tmp_path, repository=repository, recipe_service=write, batch_service=batch,
    ), repository, batch


def _workbook(exported, changes=None):
    raw = base64.b64decode(exported["contenido_base64"])
    workbook = load_workbook(BytesIO(raw), data_only=False)
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for recipe_id, field, value in changes or []:
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, headers["recipe_id"]).value == recipe_id:
                sheet.cell(row, headers[f"{field}_propuesto"]).value = value
    buffer = BytesIO(); workbook.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode("ascii")


def _physical_workbook(tmp_path: Path, exported, changes=None, *, without_dimensions=False):
    path = tmp_path / "round-trip-fisico.xlsx"
    path.write_bytes(base64.b64decode(exported["contenido_base64"]))
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for recipe_id, field, value in changes or []:
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row, headers["recipe_id"]).value == recipe_id:
                sheet.cell(row, headers[f"{field}_propuesto"]).value = value
    workbook.save(path)
    if without_dimensions:
        source = BytesIO(path.read_bytes()); target = BytesIO()
        with ZipFile(source) as incoming, ZipFile(target, "w", ZIP_DEFLATED) as outgoing:
            for entry in incoming.infolist():
                content = incoming.read(entry.filename)
                if entry.filename.startswith("xl/worksheets/") and b"recipe_id" in content:
                    content = re.sub(br"<dimension[^>]*/>", b"", content, count=1)
                outgoing.writestr(entry, content)
        path.write_bytes(target.getvalue())
    return path, base64.b64encode(path.read_bytes()).decode("ascii")


def test_exportacion_versionada_respeta_frontera_estado_y_no_escribe(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    before = deepcopy(repository.items)
    result = service.export(recipe_ids=["R2", "R1", "R2", "R3", "NO"], scope="IMPORTACION", import_id="IMP-1")
    workbook = load_workbook(BytesIO(base64.b64decode(result["contenido_base64"])), read_only=True, data_only=False)

    assert result["contrato"] == {"format": service.FORMAT, "version": service.VERSION}
    assert result["recipe_ids"] == ["R2", "R1"]
    assert {item["motivo"] for item in result["excluidas"]} == {"YA_COMPLETA", "RECETA_NO_EXISTE"}
    assert workbook.sheetnames == [
        "METADATA", "PROMPT_IA", "INSTRUCCIONES", "SCHEMA",
        "ARTICULOS_PENDIENTES", "PRECIOS_REFERENCIA", "SCHEMA_PRECIOS", "RECETAS",
    ]
    assert repository.items == before and repository.edits == []


def test_xlsx_03_expone_campos_estructurales_production_ready_y_conserva_compatibilidad(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-PROD")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])), read_only=True)
    headers = {cell.value for cell in workbook["RECETAS"][1]}
    expected = {
        "tipo_elaboracion_propuesto", "rendimiento_neto_propuesto", "merma_propuesto",
        "tiempo_preparacion_propuesto", "tiempo_coccion_propuesto", "tiempo_reposo_propuesto",
        "tiempo_enfriamiento_propuesto", "unidad_tanda_propuesto",
        "rendimiento_por_tanda_propuesto", "limitacion_tanda_propuesto",
        "intervencion_activa_propuesto", "estacion_zona_propuesto", "cuello_botella_propuesto",
    }
    assert expected.issubset(headers)
    schema_rows = list(workbook["SCHEMA"].iter_rows(values_only=True))
    schema = {row[0]: dict(zip(schema_rows[0], row)) for row in schema_rows[1:]}
    assert set(schema) == set(FIELD_CONTRACTS)
    assert len(schema) == 36
    assert all(row["que_es"] and row["por_que_host_ai_lo_necesita"] and row["como_estimarlo"] for row in schema.values())
    assert {field for field, row in schema.items() if row["participa_readiness"] == "ALWAYS"} == set(PRODUCTION_READINESS_BASE_FIELDS)
    assert schema["ingredientes_estructurados"]["tipo_esperado"] == "json"
    assert "CANDIDATO_NUEVO" in schema["ingredientes_estructurados"]["shape_json"]
    assert schema["regeneracion"]["admite_no_aplica"] is True
    assert json.loads(schema["regeneracion"]["ejemplo_no_aplica"])["estado"] == "NO_APLICA"
    instructions = dict(workbook["INSTRUCCIONES"].iter_rows(min_row=2, values_only=True))
    assert "estado_campo" in instructions["EJEMPLO_METADATOS"]
    assert "campos_pendientes" in instructions["SOLO_PENDIENTES"]
    prompt_rows = list(workbook["PROMPT_IA"].iter_rows(min_row=2, values_only=True))
    assert len(prompt_rows) == 1
    assert "Procesa TODAS las recetas" in prompt_rows[0][2]
    assert "VACÍO ES EL ÚLTIMO RECURSO" in prompt_rows[0][2]
    assert "REFERENCIA_EXTERNA" in prompt_rows[0][2]
    assert tuple(cell.value for cell in workbook["PRECIOS_REFERENCIA"][1]) == PRICE_REFERENCE_COLUMNS
    assert {row[0] for row in workbook["SCHEMA_PRECIOS"].iter_rows(min_row=2, values_only=True)} == set(PRICE_REFERENCE_COLUMNS)
    assert service.SUPPORTED_VERSIONS == {"0.1", "0.2", "0.3"}
    assert repository.edits == []


def test_referencias_precio_externas_se_validan_sin_escribir_precio_real(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    before = deepcopy(repository.items)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-PRICE")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    workbook["ARTICULOS_PENDIENTES"].append([
        "CANDIDATO-TOMATE", "", "Tomate", "", '["R1"]',
        "PRECIO_REFERENCIA_PENDIENTE", False, "Buscar referencia",
    ])
    references = workbook["PRECIOS_REFERENCIA"]
    headers = {cell.value: cell.column for cell in references[1]}
    values = {
        "article_key": "CANDIDATO-TOMATE", "article_id": "", "nombre_canonico": "Tomate",
        "unidad_base": "kg", "recetas_json": '["R1"]',
        "producto_encontrado": "Tomate rama 1 kg", "comercio_fuente": "Comercio de prueba",
        "precio_observado": 2.5, "moneda": "EUR", "formato_envase": "Bolsa 1 kg",
        "cantidad_envase": 1, "unidad_envase": "kg", "precio_normalizado": 2.5,
        "unidad_precio_normalizado": "kg", "url_fuente": "https://example.test/tomate",
        "fecha_consulta": "2026-09-05", "observacion_equivalencia": "Producto comparable",
        "confianza": 0.8, "price_basis": "IVA_INCLUIDO", "procedencia": "REFERENCIA_EXTERNA",
        "estado_referencia": "REFERENCIA_PROPUESTA",
    }
    for field, value in values.items():
        references.cell(2, headers[field]).value = value
    output = BytesIO(); workbook.save(output)

    imported = service.import_package(
        filename="precio-externo.xlsx", content_base64=base64.b64encode(output.getvalue()).decode("ascii"),
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-PRICE",
    )

    price = imported["validacion"]["referencias_precio"]
    assert price["referencias_utiles"] == 1
    assert price["referencias_consolidadas"] == 1
    assert price["rechazadas"] == 0
    assert price["filas"][0]["referencia"]["origen"] == "REFERENCIA_EXTERNA"
    assert price["filas"][0]["referencia"]["autoridad"] == "REFERENCIA_NO_REAL"
    assert price["datos_reales_modificados"] is False
    assert imported["datos_reales_modificados"] is False
    assert imported["batch"]["referencias_precio_externas"][0]["nombre_canonico"] == "Tomate"
    restored = RecetaDocumentacionBatchService(
        tmp_path, repository=repository, recipe_service=service.recipe_service,
    ).get(imported["batch"]["batch_id"])
    assert restored["referencias_precio_externas"] == imported["batch"]["referencias_precio_externas"]
    assert repository.items == before and repository.edits == []


def test_referencia_precio_invalida_o_duplicada_se_rechaza_sin_write(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-PRICE-BAD")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    workbook["ARTICULOS_PENDIENTES"].append([
        "CANDIDATO-TOMATE", "", "Tomate", "kg", '["R1"]',
        "PRECIO_REFERENCIA_PENDIENTE", False, "Buscar referencia",
    ])
    references = workbook["PRECIOS_REFERENCIA"]
    headers = {cell.value: cell.column for cell in references[1]}
    base = {
        "article_key": "CANDIDATO-TOMATE", "article_id": "", "nombre_canonico": "Tomate",
        "unidad_base": "kg", "recetas_json": '["R1"]',
    }
    bad = {
        **base, "producto_encontrado": "Tomate", "comercio_fuente": "Comercio",
        "precio_observado": 2, "cantidad_envase": 1, "unidad_envase": "kg",
        "precio_normalizado": float("nan"), "unidad_precio_normalizado": "kg",
        "url_fuente": "javascript:alert(1)", "fecha_consulta": "hoy", "confianza": 2,
        "price_basis": "REAL", "procedencia": "REAL", "estado_referencia": "REFERENCIA_PROPUESTA",
    }
    for field, value in bad.items(): references.cell(2, headers[field]).value = value
    references.append([bad.get(cell.value, "") for cell in references[1]])
    output = BytesIO(); workbook.save(output)

    imported = service.import_package(
        filename="precio-invalido.xlsx", content_base64=base64.b64encode(output.getvalue()).decode("ascii"),
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-PRICE-BAD",
    )
    price = imported["validacion"]["referencias_precio"]
    assert price["referencias_utiles"] == 0
    assert price["rechazadas"] == 2
    assert price["duplicadas"] == 1
    assert repository.edits == []


def test_candidato_nuevo_del_completado_puede_aportar_referencia_en_el_mismo_xlsx(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    before = deepcopy(repository.items)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-NEW-CANDIDATE")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    recipes = workbook["RECETAS"]
    recipe_headers = {cell.value: cell.column for cell in recipes[1]}
    ingredients = [
        {"line_id": "TOMATE", "nombre_original": "Tomate", "cantidad": 1, "unidad": "kg", "estado_relacion": "DOCUMENTAL"},
        {"line_id": "MENTA", "nombre_original": "Menta", "cantidad": 0.1, "unidad": "kg", "estado_relacion": "CANDIDATO_NUEVO"},
    ]
    recipes.cell(2, recipe_headers["ingredientes_estructurados_propuesto"]).value = json.dumps(ingredients)
    references = workbook["PRECIOS_REFERENCIA"]
    price_headers = [cell.value for cell in references[1]]
    values = {
        "article_key": "CANDIDATO_NUEVO|menta|kg", "article_id": "",
        "nombre_canonico": "Menta", "unidad_base": "kg", "recetas_json": '["R1"]',
        "producto_encontrado": "Menta fresca 1 kg", "comercio_fuente": "Comercio",
        "precio_observado": 8, "moneda": "EUR", "formato_envase": "Bolsa 1 kg",
        "cantidad_envase": 1, "unidad_envase": "kg", "precio_normalizado": 8,
        "unidad_precio_normalizado": "kg", "url_fuente": "https://example.test/menta",
        "fecha_consulta": "2026-09-05", "observacion_equivalencia": "Menta fresca comparable",
        "confianza": 0.8, "price_basis": "IVA_INCLUIDO", "procedencia": "REFERENCIA_EXTERNA",
        "estado_referencia": "REFERENCIA_PROPUESTA",
    }
    references.append([values.get(header, "") for header in price_headers])
    output = BytesIO(); workbook.save(output)

    imported = service.import_package(
        filename="candidato-con-precio.xlsx", content_base64=base64.b64encode(output.getvalue()).decode(),
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-NEW-CANDIDATE",
    )

    assert imported["validacion"]["referencias_precio"]["referencias_utiles"] == 1
    assert imported["batch"]["referencias_precio_externas"][0]["article_key"] == "CANDIDATO_NUEVO|menta|kg"
    assert imported["batch"]["resultados"][0]["datos_requieren_revision_individual"]["ingredientes_estructurados"][1]["estado_relacion"] == "CANDIDATO_NUEVO"
    assert imported["datos_reales_modificados"] is False
    assert repository.items == before and repository.edits == []


def test_consolidacion_precio_reutiliza_nombre_normalizado_solo_con_unidad_compatible(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    contexts = {
        "R1": {
            "articulos_relacionados": [{"article_id": "ART-AZUCAR", "nombre": "Azúcar", "unidad_base": "kg", "precio": None}],
            "ingredientes_contexto": [
                {"nombre_original": "AZUCAR", "unidad": "g"},
                {"nombre_original": "Azúcar", "unidad": "l"},
            ],
        },
        "R2": {
            "articulos_relacionados": [],
            "ingredientes_contexto": [{"nombre_original": "azúcar", "unidad": "ml"}],
        },
    }
    rows = service._pending_price_articles([], contexts)
    assert len(rows) == 1
    article = next(row for row in rows if row["article_key"] == "ART-AZUCAR")
    assert article["unidad_base"] == "kg"


def test_no_aplica_sin_motivo_se_rechaza_por_campo_sin_abortar_la_fila(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-NO-APLICA")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["tiempo_descongelacion_propuesto"]).value = "NO_APLICA"
    sheet.cell(2, headers["elaboracion_propuesto"]).value = "Elaboración segura conservada."
    output = BytesIO()
    workbook.save(output)

    imported = service.import_package(
        filename="sin-motivo.xlsx",
        content_base64=base64.b64encode(output.getvalue()).decode("ascii"),
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-NO-APLICA",
    )

    row = imported["validacion"]["filas"][0]
    item = imported["batch"]["resultados"][0]
    assert "tiempo_descongelacion:missing_no_aplica_reason" in row["errores"]
    assert item["datos_propuestos_seguros_masivo"] == {
        "elaboracion": "Elaboración segura conservada."
    }
    assert "tiempo_descongelacion" not in item["datos_requieren_revision_individual"]
    assert repository.edits == []


def test_no_aplica_estructurado_y_metadata_masiva_se_normalizan_sin_scalar_canonico(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-NO-APLICA-OK")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["regeneracion_propuesto"]).value = json.dumps({
        "estado": "NO_APLICA", "origen": "IA_PROPUESTA", "confianza": 0.85,
        "motivo": "La preparación se sirve fría.", "fuente": "Contexto del XLSX", "modelo": "GPT",
    })
    metadata = {
        field: {
            "origen": "IA_PROPUESTA", "confianza": 0.7,
            "motivo": "Motivo culinario provisional suficientemente concreto.",
            "fuente": "Nombre, ingredientes y contexto del XLSX", "modelo": "GPT",
            "detalle": "x" * 40,
        }
        for field in sorted(service.recipe_service.FIELDS)
    }
    assert len(json.dumps(metadata)) > service.MAX_TEXT_LENGTH
    sheet.cell(2, headers["metadatos_propuestas"]).value = json.dumps(metadata)
    output = BytesIO(); workbook.save(output)

    imported = service.import_package(
        filename="no-aplica-estructurado.xlsx",
        content_base64=base64.b64encode(output.getvalue()).decode(),
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-NO-APLICA-OK",
    )

    item = imported["batch"]["resultados"][0]
    assert item["datos_requieren_revision_individual"]["regeneracion"] == {"estado": "NO_APLICA"}
    assert item["metadatos_propuestas"]["regeneracion"]["estado_campo"] == "NO_APLICA"
    assert "metadatos_propuestas:value_too_long" not in imported["validacion"]["filas"][0]["errores"]
    assert item["completitud"]["resolucion_campos"]["regeneracion"]["estado"] == "NO_APLICA"
    assert repository.edits == []


def test_ingredientes_documentales_se_preservan_y_nuevo_queda_como_candidato(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    repository.items["R1"]["ingredientes_estructurados"] = [{
        "nombre_original": "Tomate", "articulo_id": "ART-TOMATE",
    }]
    proposed = [{
        "line_id": "L-ORIGINAL", "nombre_original": "Tomate", "cantidad": 1, "unidad": "kg",
        "procedencia_propuesta": {"origen": "IA_PROPUESTA", "confianza": 0.9},
    }, {
        "line_id": "L-NUEVA", "nombre_original": "Albahaca", "cantidad": 20, "unidad": "g",
        "procedencia_propuesta": {"origen": "IA_PROPUESTA", "confianza": 0.6},
    }]

    result = service.recipe_service.proposal(
        recipe_id="R1", proposed={"ingredientes_estructurados": proposed},
    )

    lines = result["datos_propuestos_ia"]["ingredientes_estructurados"]
    assert [line["nombre_original"] for line in lines] == ["Tomate", "Albahaca"]
    assert lines[0]["articulo_id"] == "ART-TOMATE"
    assert lines[1]["estado_relacion"] == "CANDIDATO_NUEVO"
    assert "articulo_id" not in lines[1]
    assert lines[1]["cantidad_normalizada"] == pytest.approx(0.02)
    assert result["completitud"]["resolucion_campos"]["ingredientes_estructurados"]["estado"] == "RESUELTO_IA"
    assert repository.edits == []


def test_ingrediente_documental_omitido_se_rechaza_con_razon_accionable(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    result = service.recipe_service.proposal(
        recipe_id="R1", proposed={"ingredientes_estructurados": [{
            "nombre_original": "Albahaca", "cantidad": 20, "unidad": "g",
        }]},
    )
    assert result["datos_propuestos_ia"] == {}
    assert result["campos_descartados"] == [{
        "key": "ingredientes_estructurados", "reason": "ingredient_original_missing",
    }]


def test_ingrediente_candidato_duplicado_se_rechaza_sin_crear_articulo(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    proposed = [
        {"nombre_original": "Tomate", "cantidad": 1, "unidad": "kg"},
        {"nombre_original": "tomáte", "cantidad": 0.5, "unidad": "kg", "estado_relacion": "CANDIDATO_NUEVO"},
    ]
    result = service.recipe_service.proposal(
        recipe_id="R1", proposed={"ingredientes_estructurados": proposed},
    )
    assert result["datos_propuestos_ia"] == {}
    assert result["campos_descartados"] == [{
        "key": "ingredientes_estructurados", "reason": "duplicate_ingredient_candidate",
    }]
    assert repository.edits == []


@pytest.mark.parametrize(("mutation", "reason"), [
    ({"articulo_id": "ART-INVENTADO"}, "ingredient_article_link_unauthorized"),
    ({"line_id": "LINEA-CAMBIADA"}, "ingredient_line_id_changed"),
])
def test_propuesta_ingrediente_no_puede_inventar_enlace_ni_cambiar_line_id(
    tmp_path: Path, mutation: dict[str, str], reason: str,
):
    service, repository, _ = _service(tmp_path)
    repository.items["R1"]["ingredientes_estructurados"] = [{
        "line_id": "LINEA-DOCUMENTAL", "nombre_original": "Tomate",
        "articulo_id": "ART-TOMATE",
    }]
    result = service.recipe_service.proposal(
        recipe_id="R1", proposed={"ingredientes_estructurados": [{
            "line_id": "LINEA-DOCUMENTAL", "nombre_original": "Tomate",
            "cantidad": 1, "unidad": "kg", **mutation,
        }]},
    )
    assert result["datos_propuestos_ia"] == {}
    assert result["campos_descartados"] == [{
        "key": "ingredientes_estructurados", "reason": reason,
    }]
    assert repository.edits == []


def test_line_id_duplicado_en_ingredientes_se_rechaza(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    result = service.recipe_service.proposal(
        recipe_id="R1", proposed={"ingredientes_estructurados": [{
            "line_id": "LINEA-1", "nombre_original": "Tomate", "cantidad": 1, "unidad": "kg",
        }, {
            "line_id": "LINEA-1", "nombre_original": "Albahaca", "cantidad": 20, "unidad": "g",
        }]},
    )
    assert result["datos_propuestos_ia"] == {}
    assert result["campos_descartados"] == [{
        "key": "ingredientes_estructurados", "reason": "duplicate_ingredient_line_id",
    }]
    assert repository.edits == []


@pytest.mark.parametrize(("field", "value", "code"), [
    ("tipo_elaboracion", "TIPO_INVENTADO", "invalid_enum"),
    ("unidad_rendimiento", "Texto culinario provisional", "placeholder_value"),
    ("unidad_tanda", "sacos", "invalid_unit"),
    ("tiempo_activo", "pronto", "invalid_duration"),
    ("tiempo_activo", "-30 minutos", "invalid_duration"),
    ("tiempo_activo", "30 minutos aproximadamente", "invalid_duration"),
    ("rendimiento", "NaN", "invalid_number"),
    ("produccion_maxima", "Infinity", "invalid_number"),
    ("puede_refrigerarse", "quizá", "invalid_boolean"),
    ("cantidad_por_racion", '{"cantidad":"mucho"}', "invalid_structured_shape"),
    ("personal_recomendado", "1", "invalid_structured_shape"),
    ("recursos_necesarios", "{}", "invalid_structured_shape"),
    ("alergenos", "gluten, leche", "invalid_allergen_list"),
    ("descripcion", "NO_APLICA", "no_aplica_no_permitido"),
])
def test_validacion_tipada_rechaza_basura_y_no_la_cuenta_para_readiness(
    tmp_path: Path, field: str, value: object, code: str,
):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-TIPOS")
    content = _workbook(exported, [("R1", field, value)])
    before = deepcopy(repository.items)
    imported = service.import_package(
        filename="tipos-invalidos.xlsx", content_base64=content,
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-TIPOS",
    )
    row = imported["validacion"]["filas"][0]
    item = imported["batch"]["resultados"][0]
    assert f"{field}:{code}" in row["errores"]
    assert field not in item.get("datos_propuestos_ia", {})
    assert item["completitud"]["production_ready_provisional"] is False
    assert repository.items == before and repository.edits == []


def test_alergenos_exigen_fundamento_y_fuente_sin_inventar_un_default(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-ALERGENOS")
    content = _workbook(exported, [("R1", "alergenos", '["gluten"]')])
    result = service.import_package(
        filename="alergenos-sin-fuente.xlsx", content_base64=content,
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-ALERGENOS",
    )
    assert "alergenos:insufficient_allergen_provenance" in result["validacion"]["filas"][0]["errores"]
    assert result["batch"]["resultados"][0]["datos_propuestos_ia"] == {}
    assert repository.edits == []


def test_contradicciones_objetivas_quedan_fuera_y_bloquean_readiness(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    result = service.recipe_service.proposal(recipe_id="R1", proposed={
        "tiempo_activo": "60 minutos", "tiempo_total": "30 minutos",
        "intervencion_activa": False,
        "puede_congelarse": True,
        "vida_util_congelado": {"estado": "NO_APLICA", "motivo": "No se congela."},
        "tiempo_descongelacion": {"estado": "NO_APLICA", "motivo": "No se congela."},
        "rendimiento": 10, "unidad_rendimiento": "raciones", "numero_raciones": 8,
        "produccion_maxima": 10, "rendimiento_por_tanda": 20,
    })
    discarded = {item["key"]: item["reason"] for item in result["campos_descartados"]}
    assert discarded["intervencion_activa"] == "contradiction_active_intervention"
    assert discarded["tiempo_total"] == "contradiction_total_below_active"
    assert discarded["vida_util_congelado"] == "contradiction_vida_util_congelado_no_aplica"
    assert discarded["tiempo_descongelacion"] == "contradiction_tiempo_descongelacion_no_aplica"
    assert discarded["rendimiento"] == "contradiction_yield_portions"
    assert discarded["rendimiento_por_tanda"] == "contradiction_batch_above_maximum"
    assert result["completitud"]["production_ready_provisional"] is False
    assert repository.edits == []


def test_cantidad_por_racion_y_no_aplica_activo_no_pueden_forzar_readiness(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    result = service.recipe_service.proposal(recipe_id="R1", proposed={
        "tiempo_activo": "30 minutos",
        "intervencion_activa": {"estado": "NO_APLICA"},
        "rendimiento": 2,
        "unidad_rendimiento": "kg",
        "numero_raciones": 10,
        "cantidad_por_racion": {"cantidad": 100, "unidad": "g"},
    })
    discarded = {item["key"]: item["reason"] for item in result["campos_descartados"]}
    assert discarded["intervencion_activa"] == "contradiction_active_intervention"
    assert discarded["cantidad_por_racion"] == "contradiction_portion_quantity_yield"
    assert result["completitud"]["production_ready_provisional"] is False
    assert repository.edits == []


def test_campo_texto_no_acepta_objeto_serializado_como_dato(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    result = service.recipe_service.proposal(
        recipe_id="R1", proposed={"descripcion": {"texto": "Crema"}},
    )
    assert result["datos_propuestos_ia"] == {}
    assert result["campos_descartados"] == [{"key": "descripcion", "reason": "invalid_string"}]
    assert repository.edits == []


def test_orden_de_filas_y_doble_upload_no_contaminan_ni_escriben(tmp_path: Path):
    service, repository, batch_service = _service(tmp_path)
    exported = service.export(recipe_ids=["R1", "R2"], scope="IMPORTACION", import_id="IMP-ORDER")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    first = [cell.value for cell in sheet[2]]
    second = [cell.value for cell in sheet[3]]
    for column, value in enumerate(second, start=1):
        sheet.cell(2, column).value = value
    for column, value in enumerate(first, start=1):
        sheet.cell(3, column).value = value
    sheet.cell(2, headers["elaboracion_propuesto"]).value = "Pochar la cebolla lentamente."
    sheet.cell(3, headers["elaboracion_propuesto"]).value = "Triturar el tomate hasta textura homogénea."
    buffer = BytesIO(); workbook.save(buffer)
    content = base64.b64encode(buffer.getvalue()).decode()
    before = deepcopy(repository.items)
    first_import = service.import_package(
        filename="reordenado.xlsx", content_base64=content,
        expected_recipe_ids=["R1", "R2"], scope="IMPORTACION", import_id="IMP-ORDER",
    )
    second_import = service.import_package(
        filename="reordenado.xlsx", content_base64=content,
        expected_recipe_ids=["R1", "R2"], scope="IMPORTACION", import_id="IMP-ORDER",
    )
    assert [item["recipe_id"] for item in first_import["batch"]["resultados"]] == ["R2", "R1"]
    assert second_import["batch"]["batch_id"] != first_import["batch"]["batch_id"]
    assert batch_service.active_external("IMP-ORDER")["batch"]["batch_id"] == second_import["batch"]["batch_id"]
    assert repository.items == before and repository.edits == []


def test_exportacion_vacia_es_valida_y_no_amplia_a_toda_biblioteca(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    result = service.export(recipe_ids=[], scope="IMPORTACION")
    workbook = load_workbook(BytesIO(base64.b64decode(result["contenido_base64"])), read_only=True)
    assert result["recetas_exportadas"] == 0
    assert workbook["RECETAS"].max_row == 1


def test_reimportacion_converge_en_batch_preview_confirmacion_y_procedencia(tmp_path: Path):
    service, repository, batch_service = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-1")
    content = _workbook(exported, [
        ("R1", "elaboracion", "Pochar y triturar."),
        ("R1", "numero_raciones", 8),
        ("R1", "personal_recomendado", '{"personas": 1, "rol": "cocinero"}'),
    ])
    before = deepcopy(repository.items)

    imported = service.import_package(
        filename=exported["filename"], content_base64=content,
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-1", source="CHATGPT",
    )
    batch = imported["batch"]
    assert repository.items == before and repository.edits == []
    assert batch["modo_generacion"] == "ARCHIVO_EXTERNO"
    assert batch["resultados"][0]["datos_propuestos_seguros_masivo"] == {"elaboracion": "Pochar y triturar."}
    assert batch["resultados"][0]["datos_operativos_agrupables"] == {
        "numero_raciones": 8,
        "personal_recomendado": {"personas": 1, "rol": "cocinero"},
    }
    assert batch["resultados"][0]["datos_requieren_revision_individual"] == {}
    assert imported["validacion"]["filas_utiles"] == 1
    assert imported["validacion"]["filas_requieren_revision"] == 0
    assert imported["validacion"]["campos"]["operativos_agrupables"] == 2
    assert imported["validacion"]["filas"][0]["estado"] == "UTIL"

    batch_service.select(
        batch["batch_id"], {"R1": {"elaboracion": "Pochar y triturar."}},
        {"R1": {"numero_raciones": 8}},
    )
    preview = batch_service.preview(batch["batch_id"], context=_context())
    assert repository.items == before and preview["preview"]["recetas_afectadas"] == 1
    confirmed = batch_service.confirm(
        batch["batch_id"], fingerprint=preview["preview"]["fingerprint"], context=_context(),
    )
    repeated = batch_service.confirm(
        batch["batch_id"], fingerprint=preview["preview"]["fingerprint"], context=_context(),
    )
    assert confirmed["ok"] is True and repeated["idempotente"] is True
    assert repository.items["R1"]["procedencia_campos"]["elaboracion"]["fuente_propuesta"] == "ARCHIVO_EXTERNO"
    assert repository.items["R1"]["procedencia_campos"]["elaboracion"]["origen_externo"] == "CHATGPT"


def test_reimportacion_mixta_aisla_filas_y_campos_invalidos_sin_escribir(tmp_path: Path):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1", "R2"], scope="IMPORTACION")
    raw = base64.b64decode(_workbook(exported, [
        ("R1", "elaboracion", "Pochar."), ("R1", "puede_congelarse", "quizas"),
        ("R2", "elaboracion", "Triturar."),
    ]))
    workbook = load_workbook(BytesIO(raw)); sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(3, headers["recipe_fingerprint"]).value = "obsoleta"
    sheet.cell(1, sheet.max_column + 1).value = "precio_propuesto"
    sheet.cell(2, sheet.max_column).value = "999"
    buffer = BytesIO(); workbook.save(buffer)

    result = service.import_package(
        filename="completado.xlsx", content_base64=base64.b64encode(buffer.getvalue()).decode(),
        expected_recipe_ids=["R1", "R2"], scope="IMPORTACION",
    )
    rows = result["validacion"]["filas"]
    assert rows[0]["estado"] == "UTIL_Y_REQUIERE_REVISION" and "puede_congelarse:invalid_boolean" in rows[0]["errores"]
    assert rows[1]["estado"] == "RECHAZADA" and "RECETA_CAMBIO_DESDE_EXPORTACION" in rows[1]["errores"]
    assert result["validacion"]["columnas_propuesta_desconocidas"] == ["precio_propuesto"]
    assert result["batch"]["recipe_ids"] == ["R1", "R2"] and repository.edits == []
    rejected = next(item for item in result["batch"]["resultados"] if item["recipe_id"] == "R2")
    assert rejected["estado"] == "ERROR_VALIDACION"


def test_id_duplicado_rechaza_todas_sus_filas(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="BIBLIOTECA")
    workbook = load_workbook(BytesIO(base64.b64decode(_workbook(exported, [("R1", "elaboracion", "Uno")]))))
    sheet = workbook["RECETAS"]
    sheet.append([cell.value for cell in sheet[2]])
    buffer = BytesIO(); workbook.save(buffer)
    result = service.import_package(filename="duplicado.xlsx", content_base64=base64.b64encode(buffer.getvalue()).decode())
    assert [row["errores"] for row in result["validacion"]["filas"]] == [["ID_DUPLICADO"], ["ID_DUPLICADO"]]


def test_contenido_critico_embebido_no_llega_a_seleccion_masiva(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"])
    content = _workbook(exported, [("R1", "descripcion", "Vida util 5 dias en nevera."), ("R1", "elaboracion", "Triturar.")])
    result = service.import_package(filename="externo.xlsx", content_base64=content, expected_recipe_ids=["R1"])
    item = result["batch"]["resultados"][0]
    assert item["datos_propuestos_seguros_masivo"] == {"elaboracion": "Triturar."}
    assert item["propuestas_bloqueadas_revision"][0]["campo"] == "descripcion"


def test_ingredientes_en_descripcion_no_son_falso_positivo_pero_declaracion_alergenica_si(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"])
    content = _workbook(exported, [
        ("R1", "observaciones", "Revisar cantidades, alérgenos y parámetros sanitarios antes de confirmar este lote de forma masiva."),
        ("R1", "descripcion", "Crema con leche y huevo."),
    ])
    result = service.import_package(filename="externo.xlsx", content_base64=content, expected_recipe_ids=["R1"])
    item = result["batch"]["resultados"][0]
    assert item["datos_propuestos_seguros_masivo"]["observaciones"].startswith("Revisar cantidades")
    assert item["datos_propuestos_seguros_masivo"]["descripcion"] == "Crema con leche y huevo."
    assert item["propuestas_bloqueadas_revision"] == []

    explicit = _workbook(exported, [
        ("R1", "descripcion", "Contiene alérgenos: leche y huevo."),
    ])
    blocked = service.import_package(
        filename="externo-explicito.xlsx", content_base64=explicit,
        expected_recipe_ids=["R1"],
    )["batch"]["resultados"][0]
    assert [entry["campo"] for entry in blocked["propuestas_bloqueadas_revision"]] == ["descripcion"]


def test_cambio_posterior_a_reimportacion_invalida_preview(tmp_path: Path):
    service, repository, batch_service = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"])
    result = service.import_package(
        filename="externo.xlsx", content_base64=_workbook(exported, [("R1", "elaboracion", "Triturar.")]),
        expected_recipe_ids=["R1"],
    )
    batch = result["batch"]
    batch_service.select(batch["batch_id"], {"R1": {"elaboracion": "Triturar."}})
    repository.items["R1"]["version"] = 2
    with pytest.raises(RecetaDocumentacionError) as error:
        batch_service.preview(batch["batch_id"], context=_context())
    assert error.value.code == "stale_external_recipe"


@pytest.mark.parametrize(
    ("filename", "content", "code"),
    [
        ("../externo.xlsx", base64.b64encode(b"x").decode(), "invalid_completion_filename"),
        ("externo.txt", base64.b64encode(b"x").decode(), "invalid_completion_filename"),
        ("externo.xlsx", "%%%", "invalid_completion_base64"),
        ("externo.xlsx", base64.b64encode(b"not-xlsx").decode(), "invalid_completion_xlsx"),
    ],
)
def test_entradas_maliciosas_o_corruptas_se_rechazan(tmp_path: Path, filename, content, code):
    service, _, _ = _service(tmp_path)
    with pytest.raises(RecipeCompletionExchangeError) as error:
        service.import_package(filename=filename, content_base64=content)
    assert error.value.code == code


def test_formula_se_rechaza_sin_ejecutarla(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"])
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["elaboracion_propuesto"]).value = "=HYPERLINK(\"http://evil\")"
    buffer = BytesIO(); workbook.save(buffer)
    with pytest.raises(RecipeCompletionExchangeError) as error:
        service.import_package(filename="formula.xlsx", content_base64=base64.b64encode(buffer.getvalue()).decode())
    assert error.value.code == "unsafe_formula"


def test_zip_bomb_contexto_y_receta_ajena_se_rechazan(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-1")

    service.MAX_UNCOMPRESSED_BYTES = 1
    with pytest.raises(RecipeCompletionExchangeError) as archive_error:
        service.import_package(
            filename=exported["filename"], content_base64=exported["contenido_base64"],
            expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-1",
        )
    assert archive_error.value.code == "unsafe_completion_archive"

    service.MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
    with pytest.raises(RecipeCompletionExchangeError) as import_error:
        service.import_package(
            filename=exported["filename"], content_base64=exported["contenido_base64"],
            expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-AJENA",
        )
    assert import_error.value.code == "completion_context_mismatch"
    with pytest.raises(RecipeCompletionExchangeError) as scope_error:
        service.import_package(
            filename=exported["filename"], content_base64=exported["contenido_base64"],
            expected_recipe_ids=["R1"], scope="BIBLIOTECA", import_id="IMP-1",
        )
    assert scope_error.value.code == "completion_scope_mismatch"

    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["recipe_id"]).value = "R2"
    buffer = BytesIO(); workbook.save(buffer)
    foreign = service.import_package(
        filename="ajena.xlsx", content_base64=base64.b64encode(buffer.getvalue()).decode(),
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-1",
    )
    assert foreign["validacion"]["filas"][0]["estado"] == "RECHAZADA"
    assert "FUERA_DE_LA_SELECCION_EXPORTADA" in foreign["validacion"]["filas"][0]["errores"]


def test_version_incompatible_se_rechaza(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"])
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    metadata = workbook["METADATA"]
    for row in metadata.iter_rows():
        if row[0].value == "version": row[1].value = "99"
    buffer = BytesIO(); workbook.save(buffer)
    with pytest.raises(RecipeCompletionExchangeError) as error:
        service.import_package(filename="version.xlsx", content_base64=base64.b64encode(buffer.getvalue()).decode())
    assert error.value.code == "unsupported_completion_version"


@pytest.mark.parametrize("legacy_version", ["0.1", "0.2"])
def test_versiones_anteriores_siguen_siendo_importables_tras_exportar_03(tmp_path: Path, legacy_version: str):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"])
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
    for row in workbook["METADATA"].iter_rows():
        if row[0].value == "version":
            row[1].value = legacy_version
    headers = {cell.value: cell.column for cell in workbook["RECETAS"][1]}
    workbook["RECETAS"].cell(2, headers["elaboracion_propuesto"]).value = "Triturar y colar."
    buffer = BytesIO(); workbook.save(buffer)

    result = service.import_package(
        filename="legacy-01.xlsx", content_base64=base64.b64encode(buffer.getvalue()).decode(),
        expected_recipe_ids=["R1"],
    )

    assert result["contrato"]["version"] == "0.3"
    assert result["contrato"]["version_importada"] == legacy_version
    assert result["batch"]["resultados"][0]["datos_propuestos_seguros_masivo"] == {
        "elaboracion": "Triturar y colar.",
    }
    assert result["datos_reales_modificados"] is False


def test_endpoints_exportan_reimportan_y_comparten_el_mismo_batch(tmp_path: Path):
    service, _, batch_service = _service(tmp_path)
    platform = HostAIPlatformAPI(base_dir=tmp_path)
    platform.facade._recipe_completion_exchange_service = service
    platform.facade._recipe_docs_batch_service = batch_service
    client = TestClient(create_app(platform))

    exported_response = client.post(
        "/api/v1/biblioteca/recetas/completado-externo/exportar",
        json={"recipe_ids": ["R1"], "scope": "IMPORTACION", "import_id": "IMP-1"},
    )
    exported = exported_response.json()
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])), read_only=True)
    headers = {cell.value: cell.column for cell in workbook["RECETAS"][1]}
    assert workbook["RECETAS"].cell(2, headers["recipe_id"]).value == "R1"
    assert workbook["RECETAS"].cell(2, headers["nombre"]).value == "Salsa uno"
    workbook.close()
    imported_response = client.post(
        "/api/v1/biblioteca/recetas/completado-externo/importar",
        json={
            "filename": exported["filename"],
            "contenido_base64": _workbook(exported, [("R1", "elaboracion", "Triturar.")]),
            "recipe_ids": ["R1"], "scope": "IMPORTACION", "import_id": "IMP-1",
            "origen_propuesta": "CHATGPT",
        },
    )
    imported = imported_response.json()
    status = client.get(
        f"/api/v1/biblioteca/recetas/completado-ia/{imported['batch']['batch_id']}/estado"
    )

    assert exported_response.status_code == 200 and exported["recipe_ids"] == ["R1"]
    assert imported_response.status_code == 200 and imported["datos_reales_modificados"] is False
    assert status.status_code == 200 and status.json()["modo_generacion"] == "ARCHIVO_EXTERNO"
    assert imported["batch"]["resultados"][0]["nombre"] == "Salsa uno"
    assert status.json()["resultados"][0]["nombre"] == "Salsa uno"


@pytest.mark.parametrize("field", ["descripcion", "elaboracion", "observaciones"])
def test_round_trip_fisico_cada_campo_seguro_llega_como_util(tmp_path: Path, field: str):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-REAL")
    path, content = _physical_workbook(
        tmp_path, exported, [("R1", field, f"Propuesta culinaria física para {field}.")],
    )
    before = deepcopy(repository.items)
    imported = service.import_package(
        filename=path.name, content_base64=content, expected_recipe_ids=["R1"],
        scope="IMPORTACION", import_id="IMP-REAL", source="CHATGPT",
    )
    row = imported["validacion"]["filas"][0]
    trace = row["traza_campos"][0]
    assert imported["validacion"]["filas_utiles"] == 1
    assert imported["validacion"]["campos"]["utiles"] == 1
    assert row["campos_utiles"] == [field] and row["estado"] == "UTIL"
    assert trace == {
        "columna": f"{field}_propuesto", "campo": field,
        "valor_bruto": f"Propuesta culinaria física para {field}.", "tipo_bruto": "str",
        "estado": "ACEPTADO_SELECCION_MASIVA", "politica": "SELECCION_MASIVA", "motivos": [],
        "valor_normalizado": f"Propuesta culinaria física para {field}.",
    }
    assert imported["batch"]["resultados"][0]["datos_propuestos_seguros_masivo"][field]
    assert repository.items == before and repository.edits == []


def test_round_trip_fisico_tres_campos_safe_preview_y_cero_write(tmp_path: Path):
    service, repository, batch_service = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-REAL")
    changes = [
        ("R1", "descripcion", "Descripción culinaria propuesta de prueba"),
        ("R1", "elaboracion", "Elaboración culinaria propuesta de prueba"),
        ("R1", "observaciones", "Observación culinaria propuesta de prueba"),
    ]
    path, content = _physical_workbook(tmp_path, exported, changes)
    before = deepcopy(repository.items)
    imported = service.import_package(
        filename=path.name, content_base64=content, expected_recipe_ids=["R1"],
        scope="IMPORTACION", import_id="IMP-REAL", source="CHATGPT",
    )
    batch = imported["batch"]
    assert imported["validacion"]["filas_utiles"] == 1
    assert imported["validacion"]["campos"]["utiles"] == 3
    assert imported["validacion"]["campos"]["requieren_revision"] == 0
    assert imported["validacion"]["filas"][0]["campos_utiles"] == ["descripcion", "elaboracion", "observaciones"]
    assert imported["validacion"]["filas"][0]["origen_propuesta"] == "CHATGPT"
    assert batch["fuentes_propuesta"]["R1"] == {"fuente": "ARCHIVO_EXTERNO", "origen_externo": "CHATGPT"}
    assert batch["archivo_externo"] == imported["archivo"]
    assert batch["archivo_externo"]["nombre"] == path.name
    assert len(batch["archivo_externo"]["sha256"]) == 64
    assert RecetaDocumentacionBatchService(
        tmp_path, repository=repository, recipe_service=service.recipe_service,
    ).get(batch["batch_id"])["archivo_externo"] == imported["archivo"]
    batch_service.select(batch["batch_id"], {"R1": batch["resultados"][0]["datos_propuestos_seguros_masivo"]})
    preview = batch_service.preview(batch["batch_id"], context=_context())
    assert preview["preview"]["cambios_a_aplicar"] == 3
    assert repository.items == before and repository.edits == []


@pytest.mark.parametrize(
    ("field", "value", "policy"),
    [
        ("alergenos", '["gluten"]', "INDIVIDUAL"),
        ("vida_util_refrigerado", "24 horas", "INDIVIDUAL"),
        ("tiempo_activo", "30 minutos", "AGRUPADA"),
        ("numero_raciones", 8, "AGRUPADA"),
    ],
)
def test_round_trip_campo_critico_llega_a_politica_de_revision(tmp_path: Path, field, value, policy):
    service, repository, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-CRIT")
    path, content = _physical_workbook(tmp_path, exported, [("R1", field, value)])
    if field == "alergenos":
        workbook = load_workbook(path)
        sheet = workbook["RECETAS"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        sheet.cell(2, headers["metadatos_propuestas"]).value = json.dumps({
            "alergenos": {"motivo": "Ingrediente con gluten documentado.", "fuente": "Ingredientes del XLSX"},
        })
        workbook.save(path)
        content = base64.b64encode(path.read_bytes()).decode()
    before = deepcopy(repository.items)
    imported = service.import_package(
        filename=path.name, content_base64=content, expected_recipe_ids=["R1"],
        scope="IMPORTACION", import_id="IMP-CRIT", source="CHATGPT",
    )
    row = imported["validacion"]["filas"][0]
    item = imported["batch"]["resultados"][0]
    assert imported["validacion"]["filas_utiles"] == 0
    assert imported["validacion"]["filas_requieren_revision"] == (1 if policy == "INDIVIDUAL" else 0)
    assert imported["validacion"]["campos"]["requieren_revision"] == (1 if policy == "INDIVIDUAL" else 0)
    assert row["estado"] == ("REQUIERE_REVISION" if policy == "INDIVIDUAL" else "OPERATIVA_AGRUPABLE")
    if policy == "INDIVIDUAL":
        assert row["campos_requieren_revision_individual"] == [field]
        assert field in item["datos_requieren_revision_individual"]
    else:
        assert row["campos_operativos_agrupables"] == [field]
        assert field in item["datos_operativos_agrupables"]
    assert item["datos_propuestos_seguros_masivo"] == {}
    assert repository.items == before and repository.edits == []


def test_mezcla_safe_y_critico_conserva_ambos_canales_y_fila_superpuesta(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-MIX")
    path, content = _physical_workbook(tmp_path, exported, [
        ("R1", "descripcion", "Descripción culinaria sin datos críticos."),
        ("R1", "alergenos", '["gluten"]'),
    ])
    workbook = load_workbook(path)
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sheet.cell(2, headers["metadatos_propuestas"]).value = json.dumps({
        "alergenos": {"motivo": "Ingrediente con gluten documentado.", "fuente": "Ingredientes del XLSX"},
    })
    workbook.save(path)
    content = base64.b64encode(path.read_bytes()).decode()
    imported = service.import_package(
        filename=path.name, content_base64=content, expected_recipe_ids=["R1"],
        scope="IMPORTACION", import_id="IMP-MIX", source="CHATGPT",
    )
    row = imported["validacion"]["filas"][0]
    item = imported["batch"]["resultados"][0]
    assert row["estado"] == "UTIL_Y_REQUIERE_REVISION"
    assert imported["validacion"]["filas_utiles"] == 1
    assert imported["validacion"]["filas_requieren_revision"] == 1
    assert imported["validacion"]["campos"]["utiles"] == 1
    assert imported["validacion"]["campos"]["requieren_revision"] == 1
    assert item["datos_propuestos_seguros_masivo"] == {"descripcion": "Descripción culinaria sin datos críticos."}
    assert item["datos_requieren_revision_individual"] == {"alergenos": ["gluten"]}


def test_batch_externo_declara_errores_si_acepta_una_fila_con_otro_campo_invalido(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-MIX-ERROR")
    content = _workbook(exported, [
        ("R1", "elaboracion", "Triturar hasta textura homogénea."),
        ("R1", "unidad_tanda", "sacos"),
    ])
    imported = service.import_package(
        filename="mezcla-error.xlsx", content_base64=content,
        expected_recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-MIX-ERROR",
    )
    assert imported["batch"]["estado"] == "PROPUESTAS_LISTAS_CON_ERRORES"
    assert imported["batch"]["progreso"]["propuestas"] == 1
    assert imported["validacion"]["campos"]["rechazados"] == 1


def test_xlsx_externo_sin_dimension_calculada_sigue_siendo_legible(tmp_path: Path):
    service, _, _ = _service(tmp_path)
    exported = service.export(recipe_ids=["R1"], scope="IMPORTACION", import_id="IMP-NODIM")
    path, content = _physical_workbook(
        tmp_path, exported, [("R1", "elaboracion", "Elaboración desde editor sin dimensión.")],
        without_dimensions=True,
    )
    assert load_workbook(path, read_only=True)["RECETAS"].max_row is None
    imported = service.import_package(
        filename=path.name, content_base64=content, expected_recipe_ids=["R1"],
        scope="IMPORTACION", import_id="IMP-NODIM", source="CHATGPT",
    )
    assert imported["validacion"]["filas_utiles"] == 1
    assert imported["batch"]["progreso"]["total"] == 1


def test_contrato_publico_xlsx_permite_completar_52_sin_constantes_privadas(tmp_path: Path):
    repository = FakeRepository()
    template = deepcopy(repository.items["R1"])
    repository.items = {}
    recipe_ids = []
    for index in range(52):
        recipe_id = f"PUBLIC-{index + 1:03d}"
        recipe_ids.append(recipe_id)
        repository.items[recipe_id] = {
            **deepcopy(template), "id": recipe_id, "codigo": recipe_id,
            "nombre": "Agua de jamaica" if index == 0 else f"Preparación pública {index + 1}",
            "ingredientes": ["Ingrediente documental"],
            "cantidades": [""], "ingredientes_estructurados": [{"nombre_original": "Ingrediente documental"}],
            "completitud": {"campos_obligatorios_pendientes": ["Ficha operativa"]},
        }
    write = RecetaDocumentacionWriteService(tmp_path, repository=repository)
    batch_service = RecetaDocumentacionBatchService(tmp_path, repository=repository, recipe_service=write)
    service = RecipeCompletionExchangeService(
        tmp_path, repository=repository, recipe_service=write, batch_service=batch_service,
    )
    exported = service.export(recipe_ids=recipe_ids, scope="IMPORTACION", import_id="IMP-PUBLIC-52")
    workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))

    # Este completador conoce exclusivamente las hojas públicas del XLSX.
    schema_rows = list(workbook["SCHEMA"].iter_rows(values_only=True))
    public_schema = {row[0]: dict(zip(schema_rows[0], row)) for row in schema_rows[1:]}
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row_number in range(2, sheet.max_row + 1):
        pending = json.loads(sheet.cell(row_number, headers["campos_pendientes_claves"]).value)
        ingredients = json.loads(sheet.cell(row_number, headers["ingredientes"]).value)
        metadata = {}
        for field in pending:
            spec = public_schema[field]
            if field == "alergenos":
                # Sin fundamento específico, el completador externo deja la excepción al humano.
                continue
            value = spec["ejemplo_valido"]
            if field == "puede_congelarse":
                value = False
            elif field == "intervencion_activa":
                value = True
            elif field in {"tiempo_descongelacion", "vida_util_congelado"}:
                value = spec["ejemplo_no_aplica"]
            elif field == "regeneracion" and row_number == 2:
                value = spec["ejemplo_no_aplica"]
            elif field == "ingredientes_estructurados":
                value = [{
                    "line_id": f"{row_number}-DOC", "nombre_original": name, "name_raw": name,
                    "cantidad": 1, "unidad": "kg", "cantidad_normalizada": 1,
                    "unidad_normalizada": "kg", "dato_provisional": True,
                    "procedencia_propuesta": {"origen": "IA_PROPUESTA", "confianza": 0.75, "motivo": "Estimación de contrato público"},
                } for name in ingredients]
                if row_number == 2:
                    value.append({
                        "line_id": "AGUA-CANDIDATA", "nombre_original": "Agua de proceso",
                        "cantidad": 0.5, "unidad": "l", "dato_provisional": True,
                        "estado_relacion": "CANDIDATO_NUEVO",
                    })
                value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
            if row_number == 3 and field == "unidad_rendimiento":
                value = "Texto culinario provisional"
            if row_number == 4 and field == "tipo_elaboracion":
                value = "TIPO_INVENTADO"
            sheet.cell(row_number, headers[f"{field}_propuesto"]).value = value
            metadata[field] = {
                "origen": "IA_PROPUESTA", "confianza": 0.75,
                "motivo": "Propuesta creada siguiendo exclusivamente el contrato público.",
                "fuente": "Hojas públicas del XLSX", "modelo": "COMPLETADOR_CONTRATO",
            }
            if field in {"tiempo_descongelacion", "vida_util_congelado"} or (
                field == "regeneracion" and row_number == 2
            ):
                metadata[field]["estado_campo"] = "NO_APLICA"
        sheet.cell(row_number, headers["metadatos_propuestas"]).value = json.dumps(
            metadata, ensure_ascii=False, separators=(",", ":"),
        )
        sheet.cell(row_number, headers["origen_propuesta"]).value = "CHATGPT"
    output = BytesIO(); workbook.save(output)

    imported = service.import_package(
        filename="contrato-publico-52.xlsx",
        content_base64=base64.b64encode(output.getvalue()).decode(),
        expected_recipe_ids=recipe_ids, scope="IMPORTACION", import_id="IMP-PUBLIC-52",
        source="CHATGPT",
    )

    results = imported["batch"]["resultados"]
    assert imported["validacion"]["filas_recibidas"] == 52
    assert imported["validacion"]["campos"]["rechazados"] == 2
    assert imported["validacion"]["campos"]["bloqueados_criticos"] == 0
    assert sum(item["completitud"]["production_ready_provisional"] for item in results) == 50
    assert all(not item["completitud"]["production_ready_confirmed"] for item in results)
    detailed = [
        message
        for row in imported["validacion"]["filas"]
        for rejection in row["rechazos_detallados"]
        for message in rejection["mensajes"]
    ]
    assert any("comodín" in message for message in detailed)
    assert any("enum_permitidos_json" in message for message in detailed)
    sentinel = next(item for item in results if item["nombre"] == "Agua de jamaica")
    assert sentinel["completitud"]["no_aplica"] == [
        "regeneracion", "tiempo_descongelacion", "vida_util_congelado",
    ]
    assert sentinel["datos_requieren_revision_individual"]["ingredientes_estructurados"][1]["estado_relacion"] == "CANDIDATO_NUEVO"
    assert sentinel["datos_requieren_revision_individual"]["ingredientes_estructurados"][1]["nombre_original"] == "Agua de proceso"
    assert "alergenos" not in sentinel["datos_propuestos_ia"]
    assert all(
        "Texto culinario provisional" not in str(value)
        for item in results for value in item.get("datos_propuestos_ia", {}).values()
    )
    assert repository.edits == []


def test_round_trip_treinta_filas_safe_mas_criticas_no_colapsa_a_cero(tmp_path: Path):
    repository = FakeRepository()
    template = deepcopy(repository.items["R1"])
    repository.items = {}
    recipe_ids = []
    for index in range(30):
        recipe_id = f"REC601-{index + 1:06d}"
        recipe_ids.append(recipe_id)
        repository.items[recipe_id] = {
            **deepcopy(template), "id": recipe_id, "codigo": recipe_id,
            "nombre": "Ensaladilla" if index == 0 else f"Receta {index + 1}",
        }
    write = RecetaDocumentacionWriteService(tmp_path, repository=repository)
    batch_service = RecetaDocumentacionBatchService(tmp_path, repository=repository, recipe_service=write)
    service = RecipeCompletionExchangeService(
        tmp_path, repository=repository, recipe_service=write, batch_service=batch_service,
    )
    exported = service.export(recipe_ids=recipe_ids, scope="IMPORTACION", import_id="IMP-30")
    critical_values = {
        "alergenos": '["gluten"]',
        "vida_util_refrigerado": "24 horas",
        "vida_util_congelado": "30 dias",
        "puede_congelarse": "true",
        "puede_refrigerarse": "true",
        "regeneracion": "Regenerar y validar",
        "tiempo_activo": "30 minutos",
        "tiempo_pasivo": "60 minutos",
        "tiempo_total": "90 minutos",
    }
    changes = []
    for index, recipe_id in enumerate(recipe_ids):
        # 23 descripciones seguras, pero Ensaladilla conserva exactamente dos
        # campos seguros para reproducir el caso real 2 + 9 críticos.
        if index != 0 and index <= 23:
            changes.append((recipe_id, "descripcion", f"Descripcion culinaria propuesta {recipe_id}."))
        changes.extend([
            (recipe_id, "elaboracion", f"Elaboracion culinaria propuesta {recipe_id}."),
            (recipe_id, "observaciones", f"Observacion culinaria propuesta {recipe_id}."),
        ])
        critical_count = 9 if index < 17 else 8
        changes.extend(
            (recipe_id, field, value)
            for field, value in list(critical_values.items())[:critical_count]
        )
    path, content = _physical_workbook(tmp_path, exported, changes)
    workbook = load_workbook(path)
    sheet = workbook["RECETAS"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row_number, headers["metadatos_propuestas"]).value = json.dumps({
            "alergenos": {"motivo": "Declaración basada en ingredientes.", "fuente": "Ingredientes del XLSX"},
        })
    workbook.save(path)
    content = base64.b64encode(path.read_bytes()).decode()
    before = deepcopy(repository.items)
    imported = service.import_package(
        filename=path.name, content_base64=content, expected_recipe_ids=recipe_ids,
        scope="IMPORTACION", import_id="IMP-30", source="CHATGPT",
    )
    validation = imported["validacion"]
    assert validation["filas_recibidas"] == 30
    assert validation["filas_utiles"] == 30
    assert validation["filas_requieren_revision"] == 30
    assert validation["filas_rechazadas"] == 0
    assert validation["campos"]["utiles"] == 83
    assert validation["campos"]["operativos_agrupables"] == 77
    assert validation["campos"]["requieren_revision"] == 180
    assert imported["batch"]["progreso"]["total"] == 30
    assert imported["batch"]["progreso"]["propuestas"] == 340
    assert all(row["estado"] == "UTIL_Y_REQUIERE_REVISION" for row in validation["filas"])
    assert repository.items == before and repository.edits == []

    batch_id = imported["batch"]["batch_id"]
    all_safe = {
        item["recipe_id"]: item["datos_propuestos_seguros_masivo"]
        for item in imported["batch"]["resultados"]
    }
    selected_all = batch_service.select(batch_id, all_safe, {})
    preview_all = batch_service.preview(batch_id, context=_context())
    assert sum(len(fields) for fields in selected_all["selecciones"].values()) == 83
    assert preview_all["preview"]["recetas_afectadas"] == 30
    assert preview_all["preview"]["cambios_a_aplicar"] == 83
    assert selected_all["selecciones_individuales"] == {}

    ensaladilla_safe = next(
        item["datos_propuestos_seguros_masivo"]
        for item in imported["batch"]["resultados"]
        if item["recipe_id"] == recipe_ids[0]
    )
    assert len(ensaladilla_safe) == 2
    batch_service.select(batch_id, {recipe_ids[0]: ensaladilla_safe}, {})
    preview_two = batch_service.preview(batch_id, context=_context())
    assert preview_two["preview"]["recetas_afectadas"] == 1
    assert preview_two["preview"]["cambios_a_aplicar"] == 2

    batch_service.select(
        batch_id, {recipe_ids[0]: ensaladilla_safe},
        {recipe_ids[0]: {"alergenos": ["gluten"]}},
    )
    preview_three = batch_service.preview(batch_id, context=_context())
    assert preview_three["preview"]["cambios_a_aplicar"] == 3

    batch_service.select(batch_id, {recipe_ids[0]: ensaladilla_safe}, {})
    preview_two_again = batch_service.preview(batch_id, context=_context())
    assert preview_two_again["preview"]["cambios_a_aplicar"] == 2
    confirmed = batch_service.confirm(
        batch_id,
        fingerprint=preview_two_again["preview"]["fingerprint"],
        context=_context(),
    )
    assert confirmed["estado"] == "COMPLETADO"

    restarted_write = RecetaDocumentacionWriteService(tmp_path, repository=repository)
    restarted_batch = RecetaDocumentacionBatchService(
        tmp_path, repository=repository, recipe_service=restarted_write,
    )
    restored_response = restarted_batch.active_external("IMP-30")
    restored = restored_response["batch"]
    assert restored["batch_id"] == batch_id
    assert restored["estado"] == "COMPLETADO"
    assert restored["preview"]["cambios_a_aplicar"] == 2
    assert restored_response["datos_reales_modificados"] is False

    ImportSessionRepository(tmp_path).save_all({
        "IMP-30": {
            "id": "IMP-30",
            "estado": "PENDIENTE_REVISION",
            "documento": {"nombre_archivo": "recetas_30.xlsx"},
            "propuestas": [],
        }
    })
    restarted_api = TestClient(create_app(HostAIPlatformAPI(base_dir=tmp_path)))
    detail = restarted_api.get("/api/v1/biblioteca/importaciones/IMP-30")
    assert detail.status_code == 200
    assert detail.json()["importacion"]["completado_recetas_activo"]["batch_id"] == batch_id
    assert detail.json()["importacion"]["completado_recetas_activo"]["preview"]["cambios_a_aplicar"] == 2

    active = restarted_api.get(
        "/api/v1/biblioteca/importaciones/IMP-30/completado-recetas-activo"
    )
    assert active.status_code == 200
    assert active.json()["batch"]["batch_id"] == batch_id
    assert active.json()["batch"]["preview"]["cambios_a_aplicar"] == 2
    assert active.json()["datos_reales_modificados"] is False
