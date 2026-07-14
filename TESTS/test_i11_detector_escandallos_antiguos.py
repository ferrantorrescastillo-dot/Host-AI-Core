from pathlib import Path

from openpyxl import Workbook

from SERVICIOS.detector_escandallos_antiguos_i11 import DetectorEscandallosAntiguosI11


def _crear_excel_legacy(ruta: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "MP boda 31-1"
    filas = [
        [None, None, None, "BODA PRUEBA"],
        [None, None, None, "FICHA TÉCNICA PLATO"],
        [], [],
        [None, None, None, "ARTÍCULO", "Ceviche de corvina"],
        [None, None, None, "TAPA"],
        [None, None, None, "unid.", 58, "Coste TOTAL", None, 34.4878, "Coste pax", None, 0.5946],
        [],
        [None, None, "ARTICULO", "KG", "€/UN", "BRUTO", "NETO", "%NETO", "%DESP", "DESP", "SUCIO", "€ / RACION"],
        [None, None, "Corvina fileteada", 7, 4.2, 1, 1, 1, 0, 0, 7, 29.4],
        [None, None, "Limones", 2, 1.5, 1, 1, 1, 0, 0, 2, 3.0],
        [None, None, "Ají limo", 0.5, 1660, 1, 1, 1, 0, 0, 0.5, 830],
        [], [],
        [None, None, None, "BODA PRUEBA"],
        [None, None, None, "FICHA TÉCNICA PLATO"],
        [], [],
        [None, None, None, "ARTÍCULO", "Taco de carrillera"],
        [None, None, None, "TAPA", "ELABORADOS."],
        [None, None, None, "unid.", 40, "Coste TOTAL", None, 20.0, "Coste pax", None, 0.5],
        [],
        [None, None, "ARTICULO", "KG", "€/UN", "BRUTO", "NETO", "%NETO", "%DESP", "DESP", "SUCIO", "€ / RACION"],
        [None, None, "Carrillera cocinada", 2, 8.5, 1, 1, 1, 0, 0, 2, 17],
        [None, None, "Tortilla maíz", 40, 0.075, 1, 1, 1, 0, 0, 40, 3],
    ]
    for row in filas:
        ws.append(row)
    wb.create_sheet("MENU BODA 31-1").append(["MENÚ RESUMEN SIN FICHAS"])
    wb.save(ruta)


def test_detecta_dos_bloques_y_anomalia(tmp_path):
    ruta = tmp_path / "escandallos.xlsx"
    _crear_excel_legacy(ruta)
    resultado = DetectorEscandallosAntiguosI11().analizar(ruta, hojas=["MP boda 31-1"])
    assert resultado["ok"] is True
    assert resultado["modo"] == "solo_deteccion_no_importa"
    assert resultado["resumen"]["escandallos_detectados"] == 2
    assert resultado["resumen"]["ingredientes_detectados"] == 5
    primero = resultado["escandallos"][0]
    assert primero["nombre"] == "Ceviche de corvina"
    assert primero["rendimiento_cantidad"] == 58
    assert len(primero["ingredientes"]) == 3
    aji = next(i for i in primero["ingredientes"] if i["nombre"] == "Ají limo")
    assert any("anómalo" in a for a in aji["avisos"])
    assert Path(resultado["archivo_json"]).exists()


def test_no_detecta_hoja_menu_sin_fichas(tmp_path):
    ruta = tmp_path / "escandallos.xlsx"
    _crear_excel_legacy(ruta)
    resultado = DetectorEscandallosAntiguosI11().analizar(ruta, hojas=["MENU BODA 31-1"], exportar_json=False)
    assert resultado["resumen"]["escandallos_detectados"] == 0


def test_detecta_varias_recetas_con_un_solo_marcador_ficha(tmp_path):
    ruta = tmp_path / "boda_realista.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "MP boda 31-1"
    filas = [
        [None, "Boda 31-1"],
        [None, "FICHA TÉCNICA PLATO"],
        [], [],
        [None, "ARTÍCULO", "Ceviche de corvina"],
        [None, "TAPA", "ELABORADOS"],
        [None, "unid.", 58, "Coste TOTAL", None, 34.4878, "Coste UNID:", None, 0.5946],
        [],
        ["ARTICULO", "KG", "€/UN", "BRUTO", "NETO", "%NETO", "%DESP", "DESP", "SUCIO", "€ / RACION"],
        ["Corvina fileteada", 7, 3.2, 1, 1, 1, 0, 0, 7, 22.4],
        ["Limones extra.", 2, 1.38, 1, 1, 1, 0, 0, 2, 2.76],
        [], [], [],
        [None, "ARTÍCULO", "Taco de carrillera"],
        [None, "TAPA", "ELABORADOS"],
        [None, "unid.", 58, "Coste TOTAL", None, 55.0, "Coste UNID:", None, 0.9482],
        [],
        ["ARTICULO", "KG", "€/UN", "BRUTO", "NETO", "%NETO", "%DESP", "DESP", "SUCIO", "€ / RACION"],
        ["Galtes porc", 9, 5.35, 1, 1, 1, 0, 0, 9, 48.15],
        ["Cebolla blanca.", 1, 4.5, 1, 1, 1, 0, 0, 1, 4.5],
        [], [],
        [None, "ARTÍCULO", "Guacamole i totopos"],
        [None, "TAPA", "ELABORADOS"],
        [None, "unid.", 58, "Coste TOTAL", None, 48.55, "Coste UNID:", None, 0.8371],
        [],
        ["ARTICULO", "KG", "€/UN", "BRUTO", "NETO", "%NETO", "%DESP", "DESP", "SUCIO", "€ / RACION"],
        ["Aguacate", 8, 5.65, 1, 1, 1, 0, 0, 8, 45.2],
        ["Nachos 750g", 0.5, 0.82, 1, 1, 1, 0, 0, 0.5, 0.41],
    ]
    for row in filas:
        ws.append(row)
    wb.save(ruta)

    resultado = DetectorEscandallosAntiguosI11().analizar(
        ruta, hojas=["MP boda 31-1"], exportar_json=False
    )
    assert resultado["resumen"]["escandallos_detectados"] == 3
    assert [e["nombre"] for e in resultado["escandallos"]] == [
        "Ceviche de corvina", "Taco de carrillera", "Guacamole i totopos"
    ]
    assert [len(e["ingredientes"]) for e in resultado["escandallos"]] == [2, 2, 2]
    assert resultado["escandallos"][0]["coste_total_declarado"] == 34.4878
    assert resultado["escandallos"][1]["rendimiento_cantidad"] == 58


def test_no_confunde_cabecera_articulo_con_inicio_receta(tmp_path):
    ruta = tmp_path / "cabeceras.xlsx"
    _crear_excel_legacy(ruta)
    resultado = DetectorEscandallosAntiguosI11().analizar(
        ruta, hojas=["MP boda 31-1"], exportar_json=False
    )
    nombres = [e["nombre"] for e in resultado["escandallos"]]
    assert nombres == ["Ceviche de corvina", "Taco de carrillera"]
