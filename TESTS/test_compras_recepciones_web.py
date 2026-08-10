import json
import base64
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient

from API.app import HostAIPlatformAPI
from API.http_server import create_app
from MOTORES.motor_compras import MotorCompras
from SERVICIOS.base_datos_local import BaseDatosLocal


def _write(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")


def _seed(base: Path, *, article_id: str = "ART-1", unit: str = "kg") -> tuple[TestClient, str]:
    _write(base / "DATOS/db/articulos.json", [{"codigo": "ART-1", "nombre": "Patata Monalisa", "unidad": "kg", "precio": 2}])
    motor = MotorCompras(BaseDatosLocal(base)); motor.crear_proveedor_manual("Proveedor A")
    order = motor.crear_pedidos_borrador_transaccional([{"proveedor": "Proveedor A", "lineas": [{"nombre": "Patata Monalisa", "articulo_id": article_id, "cantidad": 10, "unidad": unit, "precio_unitario": 2}]}])[0]
    motor.confirmar_borrador_pedido(order["id"], usuario="test")
    return TestClient(create_app(platform_api=HostAIPlatformAPI(base_dir=base))), order["id"]


def test_borrador_desde_pedido_no_modifica_stock_y_conserva_contrato(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    stock_path = tmp_path / "DATOS/db/stock_movimientos.json"; before = stock_path.read_bytes()
    response = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={})
    assert response.status_code == 201
    reception = response.json()["recepcion"]
    assert reception["estado"] == "BORRADOR" and reception["order_id"] == order_id
    assert reception["proveedor"] == "Proveedor A" and reception["document_type"] == "albaran"
    assert reception["lineas"][0]["ordered_quantity"] == 10
    assert reception["lineas"][0]["previously_received"] == 0
    assert reception["lineas"][0]["pending_quantity"] == 10
    assert stock_path.read_bytes() == before


def test_dashboard_expone_pedido_preparado_como_recepcionable(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    response = client.get("/api/v1/dashboard")
    assert response.status_code == 200
    pedidos = response.json()["dashboard"]["modulos"]["compras"]["pedidos"]
    pedido = next(item for item in pedidos if item["id"] == order_id)
    assert pedido["estado"] == "preparado"
    assert pedido["lineas"][0]["nombre"] == "Patata Monalisa"


def test_confirmacion_atomica_rechaza_una_version_antigua_sin_tocar_stock(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    visible_line = {**reception["lineas"][0], "received_quantity": 1}
    response = client.post(f"/api/v1/compras/recepciones/{reception['id']}/confirmar", json={
        "confirmacion": "CONFIRMAR_RECEPCION", "actualizado_en": "VERSION-ANTIGUA", "lineas": [visible_line],
    })
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "stale_reception"
    current = client.get(f"/api/v1/compras/recepciones/{reception['id']}").json()["recepcion"]
    assert current["estado"] == "BORRADOR"
    assert current["lineas"][0]["received_quantity"] == 10
    assert json.loads((tmp_path / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8")) == []


def test_recepcion_parcial_y_segunda_recepcion_completan_sin_duplicar(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    first = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    line = {**first["lineas"][0], "received_quantity": 6, "received_price": 2.2, "lot": "L-1", "expiry": "2026-12-31", "location": "Cámara"}
    saved = client.patch(f"/api/v1/compras/recepciones/{first['id']}", json={"referencia": "ALB-1", "lineas": [line]}).json()["recepcion"]
    assert any(issue["code"] == "DIFERENCIA_CANTIDAD" for issue in saved["incidencias"])
    confirmed = client.post(f"/api/v1/compras/recepciones/{first['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"})
    assert confirmed.status_code == 200 and confirmed.json()["pedido"]["estado"] == "parcialmente_recibido"
    movement = confirmed.json()["movimientos"][0]
    assert movement["cantidad"] == 6 and movement["trazabilidad"]["reception_id"] == first["id"]
    assert movement["trazabilidad"] | {"lot": "L-1", "expiry": "2026-12-31", "location": "Cámara"} == movement["trazabilidad"]
    count = len(json.loads((tmp_path / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8")))
    repeated = client.post(f"/api/v1/compras/recepciones/{first['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"})
    assert repeated.json()["idempotente"] is True
    assert repeated.json()["stock_modificado"] is False
    assert len(json.loads((tmp_path / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8"))) == count
    second = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    assert second["lineas"][0]["previously_received"] == 6 and second["lineas"][0]["pending_quantity"] == 4
    done = client.post(f"/api/v1/compras/recepciones/{second['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"})
    assert done.json()["pedido"]["estado"] == "recibido"
    assert sum(item["cantidad"] for item in json.loads((tmp_path / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8"))) == 10


def test_recepcion_realista_010_mas_015_cierra_solo_al_completar(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    first = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    first_line = {**first["lineas"][0], "ordered_quantity": 0.25, "pending_quantity": 0.25, "received_quantity": 0.10}
    # El pedido temporal reproduce exactamente 0,25 kg sin tocar datos reales.
    orders_path = tmp_path / "DATOS/db/compras_pedidos.json"
    orders = json.loads(orders_path.read_text(encoding="utf-8"))
    orders[0]["lineas"][0]["cantidad"] = 0.25
    orders_path.write_text(json.dumps(orders), encoding="utf-8")
    # Reabrir la API hace que MotorCompras lea la cantidad contractual actualizada del fixture.
    client = TestClient(create_app(platform_api=HostAIPlatformAPI(base_dir=tmp_path)))
    first_line.update({"lot": "LOTE-VISIBLE", "expiry": "2026-12-31", "location": "Cámara 1", "received_price": 2.25, "observations": "Caja revisada"})
    result1 = client.post(f"/api/v1/compras/recepciones/{first['id']}/confirmar", json={
        "confirmacion": "CONFIRMAR_RECEPCION", "actualizado_en": first["actualizado_en"],
        "fecha": "2026-08-10", "referencia": "ALB-VISIBLE", "observaciones": "Recepción visible",
        "lineas": [first_line],
    }).json()
    assert result1["pedido"]["estado"] == "parcialmente_recibido"
    assert result1["movimientos"][0]["cantidad"] == 0.10
    assert result1["movimientos"][0]["trazabilidad"] | {
        "lot": "LOTE-VISIBLE", "expiry": "2026-12-31", "location": "Cámara 1"
    } == result1["movimientos"][0]["trazabilidad"]
    confirmed_line = result1["recepcion"]["lineas"][0]
    assert confirmed_line["received_price"] == 2.25
    assert confirmed_line["observations"] == "Caja revisada"
    assert result1["recepcion"]["referencia"] == "ALB-VISIBLE"
    second = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    assert second["lineas"][0]["previously_received"] == 0.10
    assert second["lineas"][0]["pending_quantity"] == 0.15
    assert second["lineas"][0]["received_quantity"] == 0.15
    result2 = client.post(f"/api/v1/compras/recepciones/{second['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"}).json()
    assert result2["pedido"]["estado"] == "recibido"
    quantities = [item["cantidad"] for item in json.loads((tmp_path / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8"))]
    assert quantities == [0.10, 0.15]
    repeated = client.post(f"/api/v1/compras/recepciones/{second['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"}).json()
    assert repeated["idempotente"] is True
    assert [item["cantidad"] for item in json.loads((tmp_path / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8"))] == quantities


def test_pedido_multilinea_sigue_parcial_si_una_linea_tiene_pendiente(tmp_path: Path) -> None:
    _write(tmp_path / "DATOS/db/articulos.json", [
        {"codigo": "ART-A", "nombre": "A", "unidad": "kg"},
        {"codigo": "ART-B", "nombre": "B", "unidad": "kg"},
    ])
    motor = MotorCompras(BaseDatosLocal(tmp_path)); motor.crear_proveedor_manual("Proveedor A")
    order = motor.crear_pedidos_borrador_transaccional([{"proveedor": "Proveedor A", "lineas": [
        {"nombre": "A", "articulo_id": "ART-A", "cantidad": 10, "unidad": "kg"},
        {"nombre": "B", "articulo_id": "ART-B", "cantidad": 5, "unidad": "kg"},
    ]}])[0]
    motor.confirmar_borrador_pedido(order["id"], usuario="test")
    client = TestClient(create_app(platform_api=HostAIPlatformAPI(base_dir=tmp_path)))
    reception = client.post(f"/api/v1/compras/pedidos/{order['id']}/recepciones", json={}).json()["recepcion"]
    lines = [{**line, "received_quantity": 10 if line["article_id"] == "ART-A" else 2} for line in reception["lineas"]]
    client.patch(f"/api/v1/compras/recepciones/{reception['id']}", json={"lineas": lines})
    confirmed = client.post(f"/api/v1/compras/recepciones/{reception['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"}).json()
    assert confirmed["pedido"]["estado"] == "parcialmente_recibido"
    next_reception = client.post(f"/api/v1/compras/pedidos/{order['id']}/recepciones", json={}).json()["recepcion"]
    assert len(next_reception["lineas"]) == 1
    assert next_reception["lineas"][0]["article_id"] == "ART-B"
    assert next_reception["lineas"][0]["previously_received"] == 2
    assert next_reception["lineas"][0]["pending_quantity"] == 3


def test_cantidad_distinta_registra_exactamente_y_no_actualiza_precio_maestro(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    line = {**reception["lineas"][0], "received_quantity": 9.8, "received_price": 3.5}
    client.patch(f"/api/v1/compras/recepciones/{reception['id']}", json={"lineas": [line]})
    result = client.post(f"/api/v1/compras/recepciones/{reception['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"}).json()
    assert result["movimientos"][0]["cantidad"] == 9.8
    assert json.loads((tmp_path / "DATOS/db/articulos.json").read_text(encoding="utf-8"))[0]["precio"] == 2


def test_unidad_incompatible_y_articulo_sin_relacionar_bloquean(tmp_path: Path) -> None:
    for suffix, unit, code in (("unit", "l", "UNIDAD_INCOMPATIBLE"), ("article", "kg", "ARTICULO_SIN_RELACIONAR")):
        base = tmp_path / suffix; client, order_id = _seed(base, unit=unit)
        reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
        if suffix == "article":
            line = {**reception["lineas"][0], "article_id": ""}
            reception = client.patch(f"/api/v1/compras/recepciones/{reception['id']}", json={"lineas": [line]}).json()["recepcion"]
        assert any(issue["code"] == code for issue in reception["incidencias"])
        response = client.post(f"/api/v1/compras/recepciones/{reception['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"})
        assert response.status_code == 400
        assert json.loads((base / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8")) == []


def test_unidad_canonicamente_convertible_se_registra_en_unidad_base(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path, unit="g")
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    assert reception["confirmable"] is True
    result = client.post(
        f"/api/v1/compras/recepciones/{reception['id']}/confirmar",
        json={"confirmacion": "CONFIRMAR_RECEPCION"},
    ).json()
    movement = result["movimientos"][0]
    assert movement["cantidad"] == 0.01
    assert movement["unidad"] == "kg"
    assert movement["trazabilidad"]["quantity_received"] == 10
    assert movement["trazabilidad"]["unit_received"] == "g"


def test_albaran_se_adjunta_sin_stock_y_permanece_trazable_al_confirmar(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    stock_path = tmp_path / "DATOS/db/stock_movimientos.json"
    stock_before = stock_path.read_bytes()
    payload = {
        "nombre": "ALB-PROVEEDOR.pdf", "tipo_mime": "application/pdf",
        "contenido_base64": base64.b64encode(b"%PDF-1.4 fake delivery note").decode("ascii"),
        "usuario": "test", "referencia": "ALB-123",
    }
    attached = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento", json=payload)
    assert attached.status_code == 201
    document = attached.json()["documento"]
    assert document["order_id"] == order_id and document["reception_id"] == reception["id"]
    assert document["proveedor"] == "Proveedor A" and document["estado"] == "PENDIENTE_REVISION"
    assert document["checksum_sha256"] and stock_path.read_bytes() == stock_before
    duplicate = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento", json=payload).json()
    assert duplicate["idempotente"] is True and duplicate["documento"]["document_id"] == document["document_id"]
    viewed = client.get(f"/api/v1/compras/recepciones/{reception['id']}/documento").json()["documento"]
    assert base64.b64decode(viewed["contenido_base64"]) == b"%PDF-1.4 fake delivery note"
    saved = client.patch(f"/api/v1/compras/recepciones/{reception['id']}", json={
        "referencia": "ALB-124", "fecha_albaran": "2026-08-10", "observaciones": "Revisado manualmente",
    }).json()["recepcion"]
    assert saved["documento"]["referencia"] == "ALB-124"
    assert saved["documento"]["fecha_albaran"] == "2026-08-10"
    assert stock_path.read_bytes() == stock_before
    confirmed = client.post(f"/api/v1/compras/recepciones/{reception['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"}).json()
    assert confirmed["recepcion"]["documento"]["document_id"] == document["document_id"]
    assert confirmed["movimientos"][0]["trazabilidad"]["document_id"] == document["document_id"]
    assert client.delete(f"/api/v1/compras/recepciones/{reception['id']}/documento").status_code == 409


def test_albaran_puede_quitarse_y_recepcion_sin_documento_sigue_operativa(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    stock_path = tmp_path / "DATOS/db/stock_movimientos.json"; before = stock_path.read_bytes()
    payload = {"nombre": "albaran.xlsx", "tipo_mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
               "contenido_base64": base64.b64encode(b"temporary spreadsheet").decode("ascii")}
    client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento", json=payload)
    removed = client.delete(f"/api/v1/compras/recepciones/{reception['id']}/documento").json()
    assert removed["documento_eliminado"] is True and removed["recepcion"]["documento"] is None
    assert before == stock_path.read_bytes()
    confirmed = client.post(f"/api/v1/compras/recepciones/{reception['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"})
    assert confirmed.status_code == 200 and confirmed.json()["stock_modificado"] is True


def test_albaran_rechaza_formato_y_tamano_sin_modificar_stock(tmp_path: Path, monkeypatch) -> None:
    from SERVICIOS.importador_inteligente_biblioteca import ImportDocumentService
    monkeypatch.setattr(ImportDocumentService, "MAX_BYTES", 4)
    client, order_id = _seed(tmp_path)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    stock_path = tmp_path / "DATOS/db/stock_movimientos.json"; before = stock_path.read_bytes()
    bad = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento", json={
        "nombre": "malware.exe", "contenido_base64": base64.b64encode(b"x").decode("ascii")})
    assert bad.status_code == 400
    large = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento", json={
        "nombre": "grande.pdf", "contenido_base64": base64.b64encode(b"xxxxx").decode("ascii")})
    assert large.status_code == 413 and stock_path.read_bytes() == before


def _attach_for_extraction(client: TestClient, reception_id: str, name: str = "albaran.jpg", content: bytes = b"fake image") -> None:
    response = client.post(f"/api/v1/compras/recepciones/{reception_id}/documento", json={
        "nombre": name, "tipo_mime": "image/jpeg" if name.endswith(".jpg") else "application/pdf",
        "contenido_base64": base64.b64encode(content).decode("ascii"),
    })
    assert response.status_code == 201


def test_extraccion_imagen_propone_linea_exacta_y_aplicar_no_toca_stock(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    articles = json.loads((tmp_path / "DATOS/db/articulos.json").read_text(encoding="utf-8"))
    articles[0]["catalogo_maestro"] = {"referencia_proveedor": "PAT-01"}
    (tmp_path / "DATOS/db/articulos.json").write_text(json.dumps(articles), encoding="utf-8")
    client = TestClient(create_app(platform_api=HostAIPlatformAPI(base_dir=tmp_path)))
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    _attach_for_extraction(client, reception["id"])
    stock_path = tmp_path / "DATOS/db/stock_movimientos.json"; before = stock_path.read_bytes()
    text = "Proveedor A\nPedido %s\nAlbaran ALB-77 10/08/2026\nRef PAT-01 Patata 6 kg 2,20 13,20 lote L-77 cad 31/12/2026\nTotal 13,20" % order_id
    analyzed = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento/analizar", json={"texto_ocr": text})
    assert analyzed.status_code == 200
    extraction = analyzed.json()["extraccion"]; line = extraction["lines"][0]
    assert analyzed.json()["recepcion"]["confirmable"] is False
    assert extraction["provider_match"]["status"] == "PROVEEDOR_COINCIDE"
    assert line["matched_article_id"] == "ART-1" and line["match_status"] == "MATCH_EXACTO"
    assert line["quantity"] == 6 and line["unit_price"] == 2.2 and line["line_total"] == 13.2
    assert line["price_variation_pct"] == 10
    assert line["lot"] == "L-77" and line["expiration_date"] == "31/12/2026"
    assert any(issue["code"] == "CANTIDAD_DISTINTA" for issue in line["issues"])
    assert any(issue["code"] == "PRECIO_DISTINTO" for issue in line["issues"])
    assert stock_path.read_bytes() == before
    applied = client.post(f"/api/v1/compras/recepciones/{reception['id']}/extraccion/aplicar", json={
        "extraction_id": extraction["extraction_id"], "lines": extraction["lines"],
    })
    assert applied.status_code == 200 and applied.json()["confirmada"] is False
    assert applied.json()["recepcion"]["confirmable"] is True
    assert applied.json()["recepcion"]["lineas"][0]["received_quantity"] == 6
    assert applied.json()["recepcion"]["lineas"][0]["received_price"] == 2.2
    assert stock_path.read_bytes() == before
    confirmed = client.post(f"/api/v1/compras/recepciones/{reception['id']}/confirmar", json={"confirmacion": "CONFIRMAR_RECEPCION"}).json()
    assert confirmed["movimientos"][0]["cantidad"] == 6
    assert confirmed["movimientos"][0]["trazabilidad"]["document_id"] == extraction["document_id"]
    assert confirmed["movimientos"][0]["trazabilidad"]["extraction_id"] == extraction["extraction_id"]


def test_extraccion_proveedor_incorrecto_y_unidad_incompatible_bloquean_aplicar(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    _attach_for_extraction(client, reception["id"])
    result = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento/analizar", json={
        "texto_ocr": "Proveedor totalmente distinto\nPatata Monalisa 2 l 2,00 4,00",
    }).json()["extraccion"]
    assert result["provider_match"]["status"] == "PROVEEDOR_NO_COINCIDE"
    assert any(issue["code"] == "CONVERSION_PENDIENTE" for issue in result["lines"][0]["issues"])
    applied = client.post(f"/api/v1/compras/recepciones/{reception['id']}/extraccion/aplicar", json={
        "extraction_id": result["extraction_id"], "lines": result["lines"],
    })
    assert applied.status_code == 400
    assert json.loads((tmp_path / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8")) == []


def test_extraccion_ambigua_desconocida_y_reanalisis_no_duplican(tmp_path: Path) -> None:
    _write(tmp_path / "DATOS/db/articulos.json", [
        {"codigo": "ART-A", "nombre": "Tomate rojo", "unidad": "kg"},
        {"codigo": "ART-B", "nombre": "Tomate rojo", "unidad": "kg"},
    ])
    motor = MotorCompras(BaseDatosLocal(tmp_path)); motor.crear_proveedor_manual("Proveedor A")
    order = motor.crear_pedidos_borrador_transaccional([{"proveedor": "Proveedor A", "lineas": [{"nombre": "Tomate rojo", "articulo_id": "ART-A", "cantidad": 2, "unidad": "kg", "precio_unitario": 1}]}])[0]
    motor.confirmar_borrador_pedido(order["id"], usuario="test")
    client = TestClient(create_app(platform_api=HostAIPlatformAPI(base_dir=tmp_path)))
    reception = client.post(f"/api/v1/compras/pedidos/{order['id']}/recepciones", json={}).json()["recepcion"]
    _attach_for_extraction(client, reception["id"])
    first = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento/analizar", json={"texto_ocr": "Proveedor A\nTomate rojo 2 kg 1,00 2,00"}).json()["extraccion"]
    assert first["lines"][0]["match_status"] == "AMBIGUO"
    second = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento/analizar", json={"texto_ocr": "Proveedor A\nProducto imposible xyz 2 kg 1,00 2,00"}).json()["extraccion"]
    assert second["extraction_id"] == first["extraction_id"]
    assert len(second["lines"]) == 1 and second["lines"][0]["match_status"] == "SIN_MATCH"


def test_pdf_sin_lector_ni_ocr_devuelve_bloqueo_honesto(tmp_path: Path) -> None:
    client, order_id = _seed(tmp_path)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    _attach_for_extraction(client, reception["id"], "albaran.pdf", b"%PDF-1.4 scanned")
    response = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento/analizar", json={})
    assert response.status_code == 422 and response.json()["error"]["code"] == "ocr_required"
    assert json.loads((tmp_path / "DATOS/db/stock_movimientos.json").read_text(encoding="utf-8")) == []


def test_pdf_con_texto_embebido_se_extrae_por_lector_canonico(tmp_path: Path) -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    client, order_id = _seed(tmp_path)
    writer = PdfWriter(); writer.add_blank_page(width=612, height=792); page = writer.pages[0]
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica"), NameObject("/Encoding"): NameObject("/WinAnsiEncoding")})
    resources = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject(); stream.set_data(f"BT /F1 12 Tf 72 720 Td (Proveedor A) Tj 0 -20 Td (Patata Monalisa 10 kg 2.00 20.00) Tj ET".encode("ascii"))
    page[NameObject("/Resources")] = resources; page[NameObject("/Contents")] = writer._add_object(stream)
    buffer = BytesIO(); writer.write(buffer)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    _attach_for_extraction(client, reception["id"], "estructurado.pdf", buffer.getvalue())
    analyzed = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento/analizar", json={})
    assert analyzed.status_code == 200
    extraction = analyzed.json()["extraccion"]
    assert extraction["method"] == "pypdf" and extraction["summary"]["detected_lines"] == 1
    assert extraction["lines"][0]["matched_article_id"] == "ART-1"


def test_excel_se_extrae_con_lector_nativo(tmp_path: Path) -> None:
    from openpyxl import Workbook
    client, order_id = _seed(tmp_path)
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["Proveedor A"]); sheet.append(["Patata Monalisa 10 kg 2,00 20,00"])
    buffer = BytesIO(); workbook.save(buffer)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    response = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento", json={
        "nombre": "albaran.xlsx", "tipo_mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "contenido_base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
    })
    assert response.status_code == 201
    analyzed = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento/analizar", json={}).json()["extraccion"]
    assert analyzed["method"] == "lector_nativo" and analyzed["summary"]["detected_lines"] == 1
    assert analyzed["lines"][0]["match_status"] == "MATCH_EXACTO"


def test_word_se_extrae_con_lector_nativo(tmp_path: Path) -> None:
    from docx import Document
    client, order_id = _seed(tmp_path)
    document = Document(); document.add_paragraph("Proveedor A"); document.add_paragraph("Patata Monalisa 10 kg 2,00 20,00")
    buffer = BytesIO(); document.save(buffer)
    reception = client.post(f"/api/v1/compras/pedidos/{order_id}/recepciones", json={}).json()["recepcion"]
    attached = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento", json={
        "nombre": "albaran.docx", "tipo_mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "contenido_base64": base64.b64encode(buffer.getvalue()).decode("ascii"),
    })
    assert attached.status_code == 201
    analyzed = client.post(f"/api/v1/compras/recepciones/{reception['id']}/documento/analizar", json={}).json()["extraccion"]
    assert analyzed["method"] == "lector_nativo" and analyzed["lines"][0]["matched_article_id"] == "ART-1"
