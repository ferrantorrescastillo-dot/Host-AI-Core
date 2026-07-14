from __future__ import annotations

import sys
import tempfile
from pathlib import Path

RAIZ = Path(__file__).resolve().parents[1]
if str(RAIZ) not in sys.path:
    sys.path.insert(0, str(RAIZ))

from openpyxl import Workbook

from SERVICIOS.clasificador_hojas_excel_555b import ClasificadorHojasExcel555B


def crear_excel(ruta: Path) -> None:
    libro = Workbook()
    articulos = libro.active
    articulos.title = "Listado de Artículos"
    articulos.append(["Código", "Artículo", "Proveedor", "Familia", "Precio"])
    articulos.append(["ART1", "Arroz bomba", "Makro", "Secos", 2.15])

    fichas = libro.create_sheet("M.P Menú paella 45€")
    fichas.append([None, None, "FICHA TÉCNICA PLATO"])
    fichas.append(["Paella de marisco"])
    fichas.append(["Ingrediente", "Cantidad", "Unidad", "Precio kg"])
    fichas.append(["Arroz bomba", 0.1, "kg", 2.15])

    menu = libro.create_sheet("Menú paella 45€")
    menu.append(["Aperitivo", "Plato principal", "Postre"])
    menu.append(["Croquetas", "Paella de marisco", "Tarta"])

    resumen = libro.create_sheet("Resumen 4.1.3")
    resumen.append(["Total artículos", 1])

    auxiliar = libro.create_sheet("Hoja 1")
    auxiliar["A1"] = None
    libro.save(ruta)


def main() -> None:
    with tempfile.TemporaryDirectory() as temporal:
        ruta = Path(temporal) / "estructura.xlsx"
        crear_excel(ruta)
        resultado = ClasificadorHojasExcel555B().analizar(ruta)
        assert not resultado["errores"], resultado["errores"]
        por_nombre = {hoja["nombre"]: hoja for hoja in resultado["hojas"]}
        assert por_nombre["Listado de Artículos"]["tipo"] == "ARTICULOS"
        assert por_nombre["M.P Menú paella 45€"]["tipo"] == "FICHAS_TECNICAS"
        assert por_nombre["Menú paella 45€"]["tipo"] == "MENUS"
        assert por_nombre["Resumen 4.1.3"]["tipo"] == "RESUMENES"
        assert por_nombre["Hoja 1"]["tipo"] == "AUXILIARES"
        assert por_nombre["M.P Menú paella 45€"]["bloques"]
        assert resultado["datos_reales_modificados"] is False
    print("TEST OK 5.5.5B PARTE 2 - Clasificación de hojas y detección de bloques")


if __name__ == "__main__":
    main()
