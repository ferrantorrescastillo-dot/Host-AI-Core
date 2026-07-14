from pathlib import Path

from openpyxl import Workbook

from SERVICIOS.detector_multihoja_menus_i13411 import DetectorMultihojaMenusI13411
from SERVICIOS.simulador_importacion_menus_i13411 import SimuladorImportacionMenusI13411


def _crear_libro(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "MENU BODA"
    ws.append(["MENU BODA", None, None, None])
    ws.append(["APERITIVO", "P.V.P", "NETO", "COSTO", "FOOD COST"])
    ws.append(["Ceviche", 1, 0.5, 0.5])
    ws.append(["Taco", 1, 0.7, 0.7])
    ws.append(["Postre", 1, 0.4, 0.4])

    ws = wb.create_sheet("COMUNION ANA")
    ws.append(["MENU COCTEL", None, None, None])
    ws.append(["FRIOS", "P.V.P", "NETO", "COSTO", "FOOD COST"])
    ws.append(["Tabla de quesos", 1, 1.0, 1.0])
    ws.append(["Croqueta", 1, 0.8, 0.8])
    ws.append(["Sorbete", 1, 0.3, 0.3])

    ws = wb.create_sheet("M.P MENU BODA")
    ws.append(["FICHA TÉCNICA PLATO"])
    ws.append(["ARTÍCULO", "KG", "€/UN", "BRUTO", "NETO"])

    ws = wb.create_sheet("PLANTILLA COSTE MENU")
    ws.append(["MENU PLANTILLA"])
    ws.append(["P.V.P", "NETO", "COSTO", "FOOD COST"])

    ws = wb.create_sheet("Listado de Artículos")
    ws.append(["ARTICULO", "PRECIO"])
    wb.save(path)


def test_detecta_nombre_y_contenido_sin_depender_solo_de_menu(tmp_path: Path):
    libro = tmp_path / "menus.xlsx"
    _crear_libro(libro)
    r = DetectorMultihojaMenusI13411().detectar(libro)
    assert r["hojas_confirmadas"] == ["MENU BODA", "COMUNION ANA"]
    assert "M.P MENU BODA" not in r["hojas_confirmadas"]
    assert "PLANTILLA COSTE MENU" not in r["hojas_confirmadas"]


def test_simulador_usa_todos_los_confirmados_si_no_hay_seleccion(tmp_path: Path, monkeypatch):
    libro = tmp_path / "menus.xlsx"
    _crear_libro(libro)
    sim = SimuladorImportacionMenusI13411(tmp_path)
    monkeypatch.setattr(sim.preimportador, "preparar", lambda ruta, hojas=None: {
        "menus": [{"nombre": h, "hoja": h, "secciones": [], "platos": [], "articulos_directos": [], "complementos": [], "economico": {}} for h in hojas],
        "bloqueos": [],
    })
    r = sim.simular(libro)
    assert r["seleccion"]["menus_detectados_en_archivo"] == 2
    assert r["seleccion"]["menus_seleccionados"] == 2
    assert r["seleccion"]["menus_a_crear"] == 2
    assert r["integridad"]["sin_escrituras"] is True


def test_permite_seleccionar_un_subconjunto(tmp_path: Path, monkeypatch):
    libro = tmp_path / "menus.xlsx"
    _crear_libro(libro)
    sim = SimuladorImportacionMenusI13411(tmp_path)
    monkeypatch.setattr(sim.preimportador, "preparar", lambda ruta, hojas=None: {
        "menus": [{"nombre": h, "hoja": h, "secciones": [], "platos": [], "articulos_directos": [], "complementos": [], "economico": {}} for h in hojas],
        "bloqueos": [],
    })
    r = sim.simular(libro, hojas=["COMUNION ANA"])
    assert r["seleccion"]["menus_seleccionados"] == 1
    assert r["seleccion"]["hojas_seleccionadas"] == ["COMUNION ANA"]


def test_rechaza_hoja_tecnica_seleccionada_manualmente(tmp_path: Path):
    libro = tmp_path / "menus.xlsx"
    _crear_libro(libro)
    sim = SimuladorImportacionMenusI13411(tmp_path)
    try:
        sim.simular(libro, hojas=["M.P MENU BODA"])
    except ValueError as exc:
        assert "no reconocidas como menú" in str(exc).lower()
    else:
        raise AssertionError("Debía rechazar una hoja M.P")
