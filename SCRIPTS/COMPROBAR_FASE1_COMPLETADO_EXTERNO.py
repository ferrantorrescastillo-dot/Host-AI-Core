from __future__ import annotations

import argparse
import base64
import hashlib
import json
import shutil
import sys
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from SERVICIOS.host_ai_authorized_execution_context import AuthorizedExecutionContext
from SERVICIOS.importador_inteligente_biblioteca import ImportDocumentService
from SERVICIOS.recipe_completion_exchange_service import RecipeCompletionExchangeService
from SERVICIOS.receta_documentacion_batch_service import RecetaDocumentacionBatchService
from SERVICIOS.receta_documentacion_write_service import RecetaDocumentacionError, RecetaDocumentacionWriteService


def hashes(path: Path) -> dict[str, str]:
    return {
        str(item.relative_to(path)): hashlib.sha256(item.read_bytes()).hexdigest()
        for item in path.rglob("*") if item.is_file()
    }


def canonical_id(candidate: dict[str, Any]) -> str:
    value = str(candidate.get("recipe_id") or candidate.get("receta_id") or candidate.get("id") or "")
    return value if value.startswith("REC601-") and value[7:].isdigit() else ""


def contextual_ids(session: dict[str, Any]) -> list[str]:
    result = []
    for recipe in session.get("borrador", {}).get("recipes") or []:
        incomplete = not any(str(step).strip() for step in recipe.get("procedure") or []) or not (
            float(recipe.get("servings") or recipe.get("yield_value") or 0) > 0
        )
        if not incomplete:
            continue
        selected = str(recipe.get("selected_canonical_recipe_id") or "").strip()
        if selected:
            result.append(selected)
            continue
        if recipe.get("proposed_action") != "REUTILIZAR_EXISTENTE":
            continue
        candidates = [canonical_id(item) for item in recipe.get("duplicate_candidates") or []]
        candidates = [item for item in candidates if item]
        if len(candidates) == 1:
            result.append(candidates[0])
    return list(dict.fromkeys(result))


class ControlledProposalService:
    def __init__(self, write_service: RecetaDocumentacionWriteService, fail_id: str) -> None:
        self.write_service = write_service
        self.repository = write_service.repository
        self.fail_id = fail_id
        self.failed = False
        self.calls: list[str] = []

    def proposal(self, *, recipe_id: str, proposed: dict[str, Any], session_id: str = "") -> dict[str, Any]:
        self.calls.append(recipe_id)
        if recipe_id == self.fail_id and not self.failed:
            self.failed = True
            raise RecetaDocumentacionError("ai_provider_timeout", "Timeout controlado del diagnostico.")
        return self.write_service.proposal(recipe_id=recipe_id, proposed={
            "elaboracion": f"Propuesta culinaria controlada para {recipe_id}.",
            "descripcion": "Rinde 4 raciones.",
            "alergenos": ["gluten"],
        })

    def preview(self, **kwargs):
        return self.write_service.preview(**kwargs)

    def confirm(self, **kwargs):
        return self.write_service.confirm(**kwargs)


def main() -> int:
    parser = argparse.ArgumentParser(description="Diagnostico aislado del completado de recetas de Fase 1.")
    parser.add_argument("excel", type=Path)
    args = parser.parse_args()
    excel = args.excel.resolve()
    if not excel.is_file():
        raise SystemExit(f"No existe: {excel}")

    original_before = hashes(ROOT / "DATOS")
    with tempfile.TemporaryDirectory(prefix="hostai-fase1-completion-") as holder:
        isolated = Path(holder)
        shutil.copytree(ROOT / "DATOS" / "db", isolated / "DATOS" / "db")
        isolated_before = hashes(isolated / "DATOS" / "db")
        imported = ImportDocumentService(isolated).import_document({
            "archivos": [{
                "nombre": excel.name,
                "tipo_mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "contenido_base64": base64.b64encode(excel.read_bytes()).decode("ascii"),
            }],
            "excluir_ap_antiguos": True,
        })
        if not imported.get("ok"):
            raise SystemExit(json.dumps(imported, ensure_ascii=False))
        session = imported["importacion"]
        ids = contextual_ids(session)
        preview_counts = session.get("preview_global", {}).get("contadores") or {}
        analysis = session.get("analisis_restaurante") or {}
        identity = session.get("resolucion_identidad") or {}

        write_service = RecetaDocumentacionWriteService(isolated)
        controlled = ControlledProposalService(write_service, ids[1] if len(ids) > 1 else ids[0])
        batch_service = RecetaDocumentacionBatchService(
            isolated, repository=write_service.repository, recipe_service=controlled,
        )
        started = batch_service.start(recipe_ids=ids)
        current = started
        while current["progreso"]["pendientes"]:
            current = batch_service.next(current["batch_id"])
        failed_ids = [item["recipe_id"] for item in current["resultados"] if item["estado"] == "ERROR_PROVIDER"]
        continued_after_failure = bool(failed_ids and current["progreso"]["analizadas"] == len(ids))
        if failed_ids:
            batch_service.retry(current["batch_id"], failed_ids[0])
            retried = batch_service.next(current["batch_id"])
        else:
            retried = current
        cancelled = batch_service.cancel(batch_service.start(recipe_ids=ids[:1])["batch_id"])

        exchange_batch = RecetaDocumentacionBatchService(
            isolated, repository=write_service.repository, recipe_service=write_service,
        )
        exchange = RecipeCompletionExchangeService(
            isolated, repository=write_service.repository, recipe_service=write_service,
            batch_service=exchange_batch,
        )
        exported = exchange.export(recipe_ids=ids, scope="IMPORTACION", import_id=session["documento"]["id"])
        workbook = load_workbook(BytesIO(base64.b64decode(exported["contenido_base64"])))
        sheet = workbook["RECETAS"]
        headers = {cell.value: cell.column for cell in sheet[1]}
        if sheet.max_row > 1:
            sheet.cell(2, headers["elaboracion_propuesto"]).value = "Elaboracion externa sintetica para revision."
            sheet.cell(2, headers["alergenos_propuesto"]).value = '["gluten"]'
        buffer = BytesIO(); workbook.save(buffer)
        reimported = exchange.import_package(
            filename=exported["filename"], content_base64=base64.b64encode(buffer.getvalue()).decode("ascii"),
            expected_recipe_ids=ids, scope="IMPORTACION", import_id=session["documento"]["id"], source="CHATGPT",
        )
        external_batch = reimported["batch"]
        external_preview = None
        if external_batch["resultados"]:
            first = external_batch["resultados"][0]
            exchange_batch.select(external_batch["batch_id"], {
                first["recipe_id"]: dict(first.get("datos_propuestos_seguros_masivo") or {})
            })
            context = AuthorizedExecutionContext(
                "DIAG-F1", "CHEF", "LOCAL", ("chef",), frozenset({"recetas:write"}),
            )
            external_preview = exchange_batch.preview(external_batch["batch_id"], context=context)["preview"]

        isolated_after = hashes(isolated / "DATOS" / "db")
        original_after = hashes(ROOT / "DATOS")
        changed_isolated = sorted(key for key in set(isolated_before) | set(isolated_after) if isolated_before.get(key) != isolated_after.get(key))
        workflow_state_files = {
            "biblioteca_importaciones_web.json",
            "biblioteca_completado_recetas_batches.json",
        }
        operational_changes = [key for key in changed_isolated if key not in workflow_state_files]
        entities = session.get("documento", {}).get("entidades") or []
        generic_recipe_names = [
            str(recipe.get("nombre") or "")
            for recipe in analysis.get("recetas") or []
            if str(recipe.get("nombre") or "").strip().casefold() == "tapa"
        ]
        result = {
            "excel": str(excel),
            "importacion_real_aislada": {
                "hojas": int((analysis.get("resumen") or {}).get("hojas_analizadas") or 0),
                "filas": int((analysis.get("resumen") or {}).get("filas_analizadas") or 0),
                "recetas": int((identity.get("recetas") or {}).get("total") or 0),
                "recetas_canonicas_aptas": len(ids),
                "articulos_efectivos": int((identity.get("articulos") or {}).get("total") or 0),
                "articulos_reutilizados": int(preview_counts.get("articulos_reutilizados") or 0),
                "articulos_revision": int(preview_counts.get("articulos_requieren_revision") or 0),
                "menus": sum(item.get("kind") == "MENU" and item.get("fields", {}).get("tipo") != "CONTEXT" for item in entities),
                "ap_excluidos": len(analysis.get("exclusiones_sesion") or []),
                "rotulos_genericos_como_receta": generic_recipe_names,
            },
            "batch_controlado": {
                "start_ok": started.get("ok"), "total": started["progreso"]["total"],
                "fallo_aislado": failed_ids, "continuo": continued_after_failure,
                "retry_resuelto": retried["progreso"]["fallidas"] == 0,
                "cancelado": cancelled["estado"] == "CANCELADO",
                "seguras": sum(len(item.get("datos_propuestos_seguros_masivo") or {}) for item in retried["resultados"]),
                "criticas_individuales": sum(len(item.get("datos_requieren_revision_individual") or {}) for item in retried["resultados"]),
                "criticas_embebidas_bloqueadas": sum(len(item.get("propuestas_bloqueadas_revision") or []) for item in retried["resultados"]),
            },
            "flujo_externo": {
                "exportadas": exported["recetas_exportadas"], "version": exported["contrato"]["version"],
                "filas_utiles": reimported["validacion"]["filas_utiles"],
                "filas_revision": reimported["validacion"]["filas_requieren_revision"],
                "preview_recetas": int((external_preview or {}).get("recetas_afectadas") or 0),
                "confirmacion_ejecutada": False,
            },
            "seguridad": {
                "cambios_operativos_en_copia": operational_changes,
                "datos_originales_equivalentes": original_before == original_after,
                "llamadas_ia_reales": 0, "coste_ia_usd": 0,
            },
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if not operational_changes and not generic_recipe_names and original_before == original_after else 2


if __name__ == "__main__":
    raise SystemExit(main())
